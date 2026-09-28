#!/usr/bin/env python3
"""Run a compact maintainer probe matrix for public capabilities with no smoke tier.

This is not a replacement for `portable_smoke_test.py --profile full`. It is a
thin diagnostic pass for entries whose registry smoke tier is `none`: optional
backends, platform-bound helpers, legacy lanes, and maintainer tools. Probes are
deliberately copy/synthetic/preflight oriented.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import os
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from _internal.datanalysis_bootstrap import ensure_datanalysis_runtime
from _internal.public_contract import build_tool_payload
from _internal.provenance_utils import public_path, standard_qa_payload
from _internal.runtime_common import configure_runtime


ENTRY_DIR = Path(__file__).resolve().parent
ENTRY_ROOT = ENTRY_DIR.parent
ROOT = (
    ENTRY_ROOT.parent / "scientific-data-analysis"
    if ENTRY_ROOT.name == "scientific-data-maintainer"
    else ENTRY_ROOT
)
SCRIPT_DIR = ROOT / "scripts"
PYTHON = Path(sys.executable)
SEVERITY_ORDER = {
    "FAIL": 0,
    "BLOCKED_CONTROLADO": 1,
    "WARNING": 2,
    "PASS": 3,
    "NO_PROBE": 4,
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", required=True, help="Directory for probe artifacts.")
    parser.add_argument("--registry", default=str(ROOT / "public_surface_registry.yaml"), help="Registry YAML to inspect.")
    parser.add_argument("--summary-json", help="Optional standard-envelope JSON summary path.")
    parser.add_argument("--include-maintainer-help", action="store_true", help="Also run cheap --help probes for maintainer-only entries.")
    parser.add_argument("--timeout-sec", type=int, default=90)
    return parser.parse_args()


def write_csv(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    fieldnames = list(rows[0].keys())
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def emit_payload_safely(payload: dict, summary_json: str | None = None) -> None:
    rendered = json.dumps(payload, indent=2, ensure_ascii=True) + "\n"
    if summary_json:
        target = Path(summary_json)
        try:
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_text(rendered, encoding="utf-8")
        except OSError as exc:
            raise RuntimeError(
                f"Could not write summary JSON to {public_path(target)}: {exc.__class__.__name__}: {exc}"
            ) from None
    print(rendered, end="")


def build_failure_payload(args: argparse.Namespace, exc: Exception) -> dict:
    output_dir = Path(args.output_dir).expanduser() if getattr(args, "output_dir", None) else None
    matrix_csv = output_dir / "capability_probe_matrix.csv" if output_dir else None
    run_details = output_dir / "capability_probe_runs.json" if output_dir else None
    return build_tool_payload(
        "capability_probe_matrix",
        status="fail",
        notes=[
            "Maintainer-only probe matrix for public capabilities with smoke_tier=none.",
            "The probe matrix failed before it could complete cleanly.",
        ],
        artifacts={
            "matrix_csv": matrix_csv,
            "run_details_json": run_details,
            "summary_json": getattr(args, "summary_json", None),
        },
        results={
            "entry_count": 0,
            "counts_by_probe_status": {"FAIL": 1},
            "rows": [],
            "error_type": exc.__class__.__name__,
            "error": str(exc),
        },
        qa=standard_qa_payload(
            status="fail",
            findings=[{"probe_status": "FAIL", "error_type": exc.__class__.__name__, "detail": str(exc)}],
            metrics={"entry_count": 0, "FAIL": 1},
        ),
        include_environment=True,
    )


def write_table(path: Path, rows: list[dict]) -> None:
    write_csv(path, rows)


def make_gaussian_fits(path: Path, filt: str, *, shift: float = 0.0, wcs: bool = True) -> None:
    import numpy as np
    from astropy.io import fits

    yy, xx = np.indices((48, 48))
    data = 50.0 + 1500.0 * np.exp(-(((yy - 24 - shift) ** 2) + ((xx - 24 + shift) ** 2)) / (2.0 * 2.0**2))
    header = fits.Header()
    header["OBJECT"] = "PROBE_RGB"
    header["FILTER"] = filt
    header["EXPTIME"] = 20.0
    header["DATE-OBS"] = "2026-01-01T00:00:00"
    header["RA"] = "10:00:00"
    header["DEC"] = "+20:00:00"
    header["OBSERVAT"] = "CAHA"
    if wcs:
        header["CTYPE1"] = "RA---TAN"
        header["CTYPE2"] = "DEC--TAN"
        header["CRVAL1"] = 150.0
        header["CRVAL2"] = 20.0
        header["CRPIX1"] = 24.0 - shift
        header["CRPIX2"] = 24.0 + shift
        header["CDELT1"] = -0.0002777778
        header["CDELT2"] = 0.0002777778
    path.parent.mkdir(parents=True, exist_ok=True)
    fits.PrimaryHDU(data.astype("float32"), header=header).writeto(path, overwrite=True)


def make_docx(path: Path) -> None:
    from docx import Document

    doc = Document()
    para = doc.add_paragraph("Normal ")
    run = para.add_run("MARKED")
    run.bold = True
    run.italic = True
    run.underline = True
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def make_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Probe deck"
    slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(7), Inches(1)).text = "Editable caption and short scientific note."
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)


def make_notebook(path: Path) -> None:
    import nbformat

    nb = nbformat.v4.new_notebook(
        cells=[
            nbformat.v4.new_markdown_cell("# Probe notebook\n\nQuestion: explain result."),
            nbformat.v4.new_code_cell("value = input('value: ')\nprint(value)"),
            nbformat.v4.new_markdown_cell("Answer:"),
        ],
        metadata={"kernelspec": {"name": "python3", "display_name": "Python 3", "language": "python"}},
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    nbformat.write(nb, path)


def make_pages_bundle(path: Path) -> None:
    from PIL import Image, ImageDraw

    preview = path.parent / "preview.jpg"
    preview.parent.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (320, 180), "white")
    draw = ImageDraw.Draw(image)
    draw.text((20, 80), "Synthetic Pages preview", fill="black")
    image.save(preview)
    with zipfile.ZipFile(path, "w") as bundle:
        bundle.write(preview, "preview.jpg")
        bundle.writestr("Metadata/Properties.plist", "<plist></plist>")
        bundle.writestr("Index/Document.iwa", b"synthetic iwa placeholder")


def make_rgb_png(path: Path, *, red_offset: int = 0) -> None:
    from PIL import Image
    import numpy as np

    yy, xx = np.indices((24, 32))
    red = np.clip(xx * 8 + red_offset, 0, 255).astype("uint8")
    green = np.clip(yy * 10, 0, 255).astype("uint8")
    blue = np.full_like(red, 70, dtype="uint8")
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.fromarray(np.dstack([red, green, blue]), mode="RGB").save(path)


def make_fixtures(root: Path) -> dict[str, Path]:
    fixtures: dict[str, Path] = {}
    write_table(root / "sources.csv", [{"id": "A", "x": 10, "y": 20, "ra": 10.0, "dec": 20.0}, {"id": "B", "x": 30, "y": 40, "ra": 10.001, "dec": 20.001}])
    write_table(
        root / "timeseries.csv",
        [{"date": f"2026-{month:02d}-01", "value": float(month + (month % 3) * 0.2)} for month in range(1, 13)],
    )
    write_table(root / "apt.tbl", [{"Number": 1, "Image": "image.fits", "X": 10, "Y": 20, "SourceIntensity": 1000.0}])
    write_table(
        root / "bad_fxcor.csv",
        [{"case_id": "pwand_rv", "image_name": "npwand_n3.fits", "template_image": "nhd161096_n3.fits"}],
    )
    write_table(
        root / "ccf.csv",
        [
            {
                "velocity_kms": v,
                "ccf_value": 0.1
                + 0.7 * math.exp(-0.5 * ((v + 36) / 8) ** 2)
                + 0.6 * math.exp(-0.5 * ((v - 32) / 9) ** 2),
            }
            for v in range(-80, 81, 4)
        ],
    )
    make_docx(root / "marked.docx")
    make_pptx(root / "deck.pptx")
    make_notebook(root / "notebook.ipynb")
    make_pages_bundle(root / "synthetic.pages")
    make_rgb_png(root / "previous_rgb.png", red_offset=0)
    make_rgb_png(root / "new_rgb.png", red_offset=20)
    for filt, shift in [("B", -1.0), ("V", 0.0), ("R", 1.0)]:
        make_gaussian_fits(root / "rgb" / f"probe_{filt}.fits", filt, shift=shift, wcs=True)

    legacy_root = root / "legacy"
    fits_dir = legacy_root / "fits_p1"
    source = ROOT / "examples" / "science" / "legacy_spectroscopy_mini" / "mini_multispec.fits"
    for name in [
        "npwand_n3.fits",
        "nhd161096_n3.fits",
        "nhd166620_n2.fits",
        "nrej1101_foces02_n2.fits",
        "nhd100696_n3.fits",
        "nhd97004_n4.fits",
        "nhd92588_n2.fits",
    ]:
        fits_dir.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, fits_dir / name)

    istar = root / "istarmod"
    istar.mkdir(parents=True, exist_ok=True)
    (istar / "sample.sm").write_text("! synthetic iSTARMOD script placeholder\n", encoding="utf-8")
    (istar / "rvvalues.dat").write_text("cache\n", encoding="utf-8")

    fixtures.update(
        {
            "sources": root / "sources.csv",
            "timeseries": root / "timeseries.csv",
            "apt_tbl": root / "apt.tbl",
            "bad_fxcor": root / "bad_fxcor.csv",
            "ccf": root / "ccf.csv",
            "docx": root / "marked.docx",
            "pptx": root / "deck.pptx",
            "notebook": root / "notebook.ipynb",
            "pages": root / "synthetic.pages",
            "previous_rgb_png": root / "previous_rgb.png",
            "new_rgb_png": root / "new_rgb.png",
            "rgb": root / "rgb",
            "legacy": legacy_root,
            "legacy_fits": fits_dir,
            "istar": istar,
        }
    )
    return fixtures


def command(script: str, *args) -> list[str]:
    return [str(PYTHON), str(SCRIPT_DIR / script), *[str(arg) for arg in args]]


def run_command(name: str, cmd: list[str], cwd: Path, timeout: int, summary_path: Path | None = None) -> dict:
    env = os.environ.copy()
    env.setdefault("PYTHONDONTWRITEBYTECODE", "1")
    try:
        completed = subprocess.run(cmd, cwd=cwd, text=True, capture_output=True, check=False, timeout=timeout, env=env)
        timed_out = False
    except subprocess.TimeoutExpired as exc:
        completed = None
        timed_out = True
        stdout = exc.stdout if isinstance(exc.stdout, str) else ""
        stderr = exc.stderr if isinstance(exc.stderr, str) else ""
    else:
        stdout = completed.stdout
        stderr = completed.stderr

    payload = None
    if summary_path and summary_path.exists():
        try:
            payload = json.loads(summary_path.read_text(encoding="utf-8"))
        except Exception:
            payload = None
    return {
        "name": name,
        "command": cmd,
        "returncode": None if timed_out else completed.returncode,
        "timed_out": timed_out,
        "stdout_tail": stdout[-1200:],
        "stderr_tail": stderr[-1200:],
        "summary_json": str(summary_path) if summary_path else None,
        "payload_status": payload.get("status") if isinstance(payload, dict) else None,
        "payload": payload,
    }


def classify(runs: list[dict], *, allow_blocked: bool = False, prefer_warning_for_blocked: bool = False) -> str:
    if any(item["timed_out"] for item in runs):
        return "FAIL"
    blocked = False
    warning = False
    for item in runs:
        payload_status = item.get("payload_status")
        if item["returncode"] not in (0, None):
            if payload_status == "blocked" and allow_blocked:
                blocked = True
                continue
            return "FAIL"
        if payload_status == "blocked":
            if allow_blocked:
                blocked = True
            else:
                return "FAIL"
        elif payload_status == "warning":
            warning = True
        elif payload_status == "fail":
            return "FAIL"
    if blocked and prefer_warning_for_blocked and any(item["returncode"] == 0 for item in runs):
        return "WARNING"
    if blocked:
        return "BLOCKED_CONTROLADO"
    if warning:
        return "WARNING"
    return "PASS"


def make_probe_commands(entry_id: str, fixtures: dict[str, Path], out: Path) -> tuple[list[tuple[str, list[str], Path | None]], dict]:
    probe_dir = out / entry_id.replace(".", "__")
    probe_dir.mkdir(parents=True, exist_ok=True)
    meta = {"allow_blocked": False, "prefer_warning_for_blocked": False, "evidence": ""}

    def s(name: str) -> Path:
        return probe_dir / f"{name}.json"

    mapping: dict[str, tuple[list[tuple[str, list[str], Path | None]], dict]] = {
        "datanalysis_env.status": ([(entry_id, command("datanalysis_env.py", "status"), None)], {"evidence": "status probe"}),
        "datanalysis_healthcheck": ([(entry_id, command("datanalysis_healthcheck.py", "--skip-notebook-exec", "--summary-json", s("healthcheck")), s("healthcheck"))], {"evidence": "dependency healthcheck without notebook execution"}),
        "companion_route_check": ([(entry_id, command("companion_route_check.py", "--task", "revisar PDF científico", "--file", "report.pdf", "--summary-json", s("companion")), s("companion"))], {"evidence": "advisory routing probe"}),
        "external_astro_tools_preflight": ([(entry_id, command("external_astro_tools_preflight.py", "--probe", "--summary-json", s("external")), s("external"))], {"evidence": "external optional backend preflight", "allow_blocked": False}),
        "fits_rgb_batch": ([(entry_id, command("fits_rgb_batch.py", "--input-root", fixtures["rgb"], "--output-dir", probe_dir / "rgb_out", "--no-stacks", "--summary-json", s("rgb")), s("rgb"))], {"evidence": "synthetic RGB batch"}),
        "rgb_visual_fits_export": (
            [
                (
                    entry_id,
                    command(
                        "rgb_visual_fits_export.py",
                        fixtures["new_rgb_png"],
                        probe_dir / "visual_rgb.fits",
                        "--previous-png",
                        fixtures["previous_rgb_png"],
                        "--comparison-png",
                        probe_dir / "before_after.png",
                        "--manifest-json",
                        probe_dir / "visual_manifest.json",
                        "--summary-json",
                        s("rgb_visual"),
                        "--object-name",
                        "PROBE_RGB",
                        "--reason",
                        "capability probe visual export",
                    ),
                    s("rgb_visual"),
                )
            ],
            {"evidence": "display-only RGB FITS export with comparison and manifest"},
        ),
        "stilts_workbench": ([(entry_id, command("stilts_workbench.py", "preflight", "--summary-json", s("stilts")), s("stilts"))], {"evidence": "STILTS optional backend preflight", "allow_blocked": True}),
        "radial_velocity_workbench.validate-manifest": (
            [
                ("rv_scaffold", command("radial_velocity_workbench.py", "scaffold-session", probe_dir / "rv_session", "--system-name", "Probe RV", "--planet-count", "1", "--overwrite"), None),
                (entry_id, command("radial_velocity_workbench.py", "validate-manifest", probe_dir / "rv_session" / "session_manifest.json", "--summary-json", s("rv_validate")), s("rv_validate")),
            ],
            {"evidence": "synthetic RV session manifest"},
        ),
        "legacy_spectroscopy_envcheck": ([(entry_id, command("legacy_spectroscopy_envcheck.py", fixtures["legacy"], "--summary-json", s("legacy_env")), s("legacy_env"))], {"evidence": "synthetic legacy practice preflight"}),
        "fxcor_iraf_workbench.prepare-session": ([(entry_id, command("fxcor_iraf_workbench.py", "prepare-session", fixtures["legacy"], "--output-dir", probe_dir / "fxcor_workspace", "--summary-json", s("fxcor_prepare")), s("fxcor_prepare"))], {"evidence": "synthetic fxcor workspace"}),
        "fxcor_iraf_workbench.run-auto": (
            [
                ("fxcor_prepare_for_run", command("fxcor_iraf_workbench.py", "prepare-session", fixtures["legacy"], "--output-dir", probe_dir / "fxcor_workspace", "--summary-json", s("fxcor_prepare")), s("fxcor_prepare")),
                (entry_id, command("fxcor_iraf_workbench.py", "run-auto", probe_dir / "fxcor_workspace", "--summary-json", s("fxcor_run")), s("fxcor_run")),
            ],
            {"evidence": "fxcor auto run blocks cleanly when IRAF output is unavailable", "allow_blocked": True},
        ),
        "legacy_rv_coursework_workbench.analyze": ([(entry_id, command("legacy_rv_coursework_workbench.py", "analyze", fixtures["bad_fxcor"], "--fits-root", fixtures["legacy_fits"], "--output-dir", probe_dir / "legacy_rv", "--summary-json", s("legacy_rv")), s("legacy_rv"))], {"evidence": "input validation blocked case", "allow_blocked": True}),
        "sb2_double_gaussian_workbench.fit": ([(entry_id, command("sb2_double_gaussian_workbench.py", "fit", fixtures["ccf"], "--output-dir", probe_dir / "sb2", "--summary-json", s("sb2")), s("sb2"))], {"evidence": "synthetic SB2 CCF"}),
        "legacy_external_reference_check": ([(entry_id, command("legacy_external_reference_check.py", "check", "--rv-summary", probe_dir / "missing_rv.json", "--output-dir", probe_dir / "external_ref", "--summary-json", s("external_ref")), s("external_ref"))], {"evidence": "reference check with missing measurements"}),
        "istarmod_workbench.inspect-tree": ([(entry_id, command("istarmod_workbench.py", "inspect-tree", fixtures["istar"], "--summary-json", s("istar_inspect")), s("istar_inspect"))], {"evidence": "synthetic iSTARMOD tree"}),
        "istarmod_workbench.prepare-copy": ([(entry_id, command("istarmod_workbench.py", "prepare-copy", fixtures["istar"], "--output-dir", probe_dir / "istar_copy", "--summary-json", s("istar_copy")), s("istar_copy"))], {"evidence": "synthetic iSTARMOD copy"}),
        "legacy_spectroscopy_report_builder.scaffold": ([(entry_id, command("legacy_spectroscopy_report_builder.py", "scaffold", probe_dir / "legacy_report", "--summary-json", s("legacy_report")), s("legacy_report"))], {"evidence": "legacy report scaffold"}),
        "legacy_spectroscopy_report_builder.populate": (
            [
                ("legacy_report_scaffold", command("legacy_spectroscopy_report_builder.py", "scaffold", probe_dir / "legacy_report", "--summary-json", s("legacy_report_scaffold")), s("legacy_report_scaffold")),
                (entry_id, command("legacy_spectroscopy_report_builder.py", "populate", probe_dir / "legacy_report", "--summary-json", s("legacy_report_populate")), s("legacy_report_populate")),
            ],
            {"evidence": "populate a fresh synthetic legacy report scaffold without scientific input data"},
        ),
        "apt_workbench": (
            [
                ("apt_preflight", command("apt_workbench.py", "preflight", "--summary-json", s("apt_preflight")), s("apt_preflight")),
                ("apt_source_list", command("apt_workbench.py", "prepare-source-list", fixtures["sources"], probe_dir / "apt_sources.txt", "--x-col", "x", "--y-col", "y", "--id-col", "id", "--summary-json", s("apt_source")), s("apt_source")),
                ("apt_parse", command("apt_workbench.py", "parse-results", fixtures["apt_tbl"], probe_dir / "apt_parsed.csv", "--summary-json", s("apt_parse")), s("apt_parse")),
            ],
            {"evidence": "APT preflight plus source-list and .tbl parse probes", "allow_blocked": True, "prefer_warning_for_blocked": True},
        ),
        "teareduce_router": ([(entry_id, command("teareduce_router.py", "--intent", "decidir ruta de reducción CCD docente", "--summary-json", s("teareduce")), s("teareduce"))], {"evidence": "TEAREDUCE routing decision"}),
        "iwork_workbench": ([(entry_id, command("iwork_workbench.py", fixtures["pages"], "--output-dir", probe_dir / "iwork", "--output-json", s("iwork")), s("iwork"))], {"evidence": "synthetic iWork bundle inspection"}),
        "office_roundtrip.docx-style-inventory": ([(entry_id, command("office_roundtrip.py", "docx-style-inventory", fixtures["docx"], "--require-bold", "--require-italic", "--require-underline", "--summary-json", s("docx_inventory")), s("docx_inventory"))], {"evidence": "styled DOCX inventory"}),
        "office_roundtrip.docx-styled-replace": ([(entry_id, command("office_roundtrip.py", "docx-styled-replace", fixtures["docx"], probe_dir / "replaced.docx", "--find", "MARKED", "--replace", "REPLACED", "--require-bold", "--require-italic", "--require-underline", "--summary-json", s("docx_replace")), s("docx_replace"))], {"evidence": "styled DOCX replacement on copy"}),
        "keynote_export": ([(entry_id, command("keynote_export.py", "--preflight-only", fixtures["pptx"], probe_dir / "deck.pdf", "--summary-json", s("keynote")), s("keynote"))], {"evidence": "Keynote preflight only", "allow_blocked": True}),
        "duckdb_workbench": ([(entry_id, command("duckdb_workbench.py", fixtures["sources"], "--sql", "select count(*) as n from source0", "--output", probe_dir / "duckdb.csv", "--summary-json", s("duckdb")), s("duckdb"))], {"evidence": "synthetic DuckDB query"}),
        "coursework_notebook_fidelity_check": ([(entry_id, command("coursework_notebook_fidelity_check.py", fixtures["notebook"], "--summary-json", s("fidelity")), s("fidelity"))], {"evidence": "synthetic professor-notebook fidelity scan"}),
        "notebook_branch_compare": ([(entry_id, command("notebook_branch_compare.py", fixtures["rgb"].parent, "--output-dir", probe_dir / "branches", "--summary-json", s("branches")), s("branches"))], {"evidence": "synthetic branch/product inventory"}),
        "timeseries_forecasting_workbench": ([(entry_id, command("timeseries_forecasting_workbench.py", "--output-dir", probe_dir / "forecast", "--data-path", fixtures["timeseries"], "--date-column", "date", "--value-column", "value", "--test-horizon", "1", "--summary-json", s("forecast"), "--overwrite"), s("forecast"))], {"evidence": "forecasting notebook scaffold"}),
        "physical_qa": ([(entry_id, command("physical_qa.py", fixtures["sources"], "--summary-json", s("physical_qa")), s("physical_qa"))], {"evidence": "synthetic table physical QA"}),
        "sync_public_surface_docs": ([(entry_id, command("sync_public_surface_docs.py", "--check", "--summary-json", s("sync")), s("sync"))], {"evidence": "snapshot drift check"}),
        "skill_surface_audit": ([(entry_id, command("skill_surface_audit.py", "--skip-hygiene", "--summary-json", s("surface")), s("surface"))], {"evidence": "surface audit without hygiene residue check"}),
    }

    maintainer_help = {
        "capability_probe_matrix": "capability_probe_matrix.py",
        "external_astro_tools_local_validation": "external_astro_tools_local_validation.py",
        "portable_smoke_test": "portable_smoke_test.py",
        "validate_skill_samples": "validate_skill_samples.py",
    }
    if entry_id in maintainer_help:
        mapping[entry_id] = (
            [(entry_id, command(maintainer_help[entry_id], "--help"), None)],
            {"evidence": "cheap maintainer CLI help probe"},
        )

    if entry_id == "legacy_external_reference_check":
        (probe_dir / "missing_rv.json").write_text(json.dumps({"tool": "legacy_rv_coursework_workbench.analyze", "status": "blocked", "results": {}}), encoding="utf-8")

    return mapping.get(entry_id, ([], {"evidence": "No safe probe configured."}))


def run_matrix(args: argparse.Namespace) -> tuple[dict, int]:
    from _internal.public_surface_registry import load_public_surface_registry

    configure_runtime("capability_probe_matrix")
    output_dir = Path(args.output_dir).expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    fixture_root = output_dir / "_fixtures"

    with tempfile.TemporaryDirectory(prefix="sda_capability_probe_") as tmp_raw:
        tmp_root = Path(tmp_raw)
        fixtures = make_fixtures(tmp_root / "fixtures")
        entries = load_public_surface_registry(args.registry)
        selected = [entry for entry in entries if entry.get("smoke_tier") == "none"]
        if not args.include_maintainer_help:
            selected = [entry for entry in selected if entry.get("kind") != "maintainer_only"]

        rows = []
        run_details = {}
        for entry in selected:
            entry_id = entry["id"]
            command_specs, meta = make_probe_commands(entry_id, fixtures, output_dir)
            runs = []
            for name, cmd, summary_path in command_specs:
                runs.append(run_command(name, cmd, ROOT, args.timeout_sec, summary_path))
            if command_specs:
                probe_status = classify(
                    runs,
                    allow_blocked=bool(meta.get("allow_blocked")),
                    prefer_warning_for_blocked=bool(meta.get("prefer_warning_for_blocked")),
                )
            else:
                probe_status = "NO_PROBE"
            row = {
                "id": entry_id,
                "label": entry.get("label", entry_id),
                "visible_block": entry.get("visible_block"),
                "kind": entry.get("kind"),
                "support_level": entry.get("support_level"),
                "probe_status": probe_status,
                "command_count": len(runs),
                "evidence": meta.get("evidence", ""),
            }
            rows.append(row)
            run_details[entry_id] = runs

        rows = sorted(rows, key=lambda item: (SEVERITY_ORDER.get(item["probe_status"], 99), item["visible_block"], item["id"]))
        write_csv(output_dir / "capability_probe_matrix.csv", rows)
        (output_dir / "capability_probe_runs.json").write_text(json.dumps(run_details, indent=2, ensure_ascii=True) + "\n", encoding="utf-8")
        if fixture_root.exists():
            shutil.rmtree(fixture_root)

    counts = {}
    for row in rows:
        counts[row["probe_status"]] = counts.get(row["probe_status"], 0) + 1
    overall_status = "fail" if counts.get("FAIL") else ("warning" if counts.get("WARNING") or counts.get("BLOCKED_CONTROLADO") or counts.get("NO_PROBE") else "ok")
    findings = [
        row
        for row in rows
        if row["probe_status"] in {"FAIL", "WARNING", "BLOCKED_CONTROLADO", "NO_PROBE"}
    ]
    payload = build_tool_payload(
        "capability_probe_matrix",
        status=overall_status,
        notes=[
            "Maintainer-only probe matrix for public capabilities with smoke_tier=none.",
            "Optional external or platform-bound backends may report BLOCKED_CONTROLADO without indicating a core skill failure.",
            "Synthetic fixtures are temporary; the persisted artifacts are the matrix CSV and run JSON.",
        ],
        artifacts={
            "matrix_csv": output_dir / "capability_probe_matrix.csv",
            "run_details_json": output_dir / "capability_probe_runs.json",
            "summary_json": args.summary_json,
        },
        results={
            "entry_count": len(rows),
            "counts_by_probe_status": counts,
            "rows": rows,
        },
        qa=standard_qa_payload(status=overall_status, findings=findings, metrics={"entry_count": len(rows), **counts}),
        include_environment=True,
    )
    return payload, 1 if counts.get("FAIL") else 0


def main() -> int:
    args = parse_args()
    try:
        ensure_datanalysis_runtime("capability_probe_matrix", argv=sys.argv[1:], strict=True)
    except SystemExit as exc:
        if isinstance(exc.code, str):
            payload = build_failure_payload(args, RuntimeError(exc.code))
            try:
                emit_payload_safely(payload, args.summary_json)
            except Exception as emit_exc:
                payload["notes"].append(str(emit_exc))
                payload["results"]["summary_json_error"] = str(emit_exc)
                emit_payload_safely(payload, None)
            return 1
        raise
    try:
        payload, exit_code = run_matrix(args)
        emit_payload_safely(payload, args.summary_json)
        return exit_code
    except Exception as exc:
        payload = build_failure_payload(args, exc)
        try:
            emit_payload_safely(payload, args.summary_json)
        except Exception as emit_exc:
            payload["notes"].append(str(emit_exc))
            payload["results"]["summary_json_error"] = str(emit_exc)
            emit_payload_safely(payload, None)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
