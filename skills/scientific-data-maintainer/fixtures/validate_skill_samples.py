#!/usr/bin/env python3
"""Run non-destructive smoke tests for the scientific-data-analysis skill."""

import argparse
import bz2
import contextlib
import gzip
import io
import json
import os
import lzma
import subprocess
import sys
import tarfile
import zipfile
from collections import defaultdict
from pathlib import Path

os.environ.setdefault("ARROW_USER_SIMD_LEVEL", "NONE")
os.environ.setdefault("PYTHONDONTWRITEBYTECODE", "1")

from _internal.provenance_utils import public_path, sanitize_payload, standard_tool_payload, write_manifest
from _internal.runtime_common import configure_runtime, suppress_fd_output
from _internal.tabular_io import read_table_any
from maintainer_path_safety import PathSafetyError, ensure_safe_paths, operand


SUPPORTED_EXTS = {
    ".fits",
    ".fit",
    ".fts",
    ".fz",
    ".csv",
    ".tsv",
    ".ecsv",
    ".tbl",
    ".dat",
    ".txt",
    ".text",
    ".ipynb",
    ".pdf",
    ".docx",
    ".docm",
    ".rtf",
    ".pptx",
    ".pptm",
    ".xlsx",
    ".xlsm",
    ".xls",
    ".xlsb",
    ".doc",
    ".pages",
    ".numbers",
    ".key",
    ".odt",
    ".ods",
    ".odp",
    ".epub",
    ".reg",
    ".pref",
    ".jmars",
    ".tex",
    ".bib",
    ".html",
    ".htm",
    ".png",
    ".jpg",
    ".jpeg",
    ".json",
    ".jsonl",
    ".ndjson",
    ".yaml",
    ".yml",
    ".toml",
    ".npy",
    ".npz",
    ".h5",
    ".hdf5",
    ".mat",
    ".sqlite",
    ".sqlite3",
    ".db",
    ".nc",
    ".nc4",
    ".cdf",
    ".netcdf",
    ".parquet",
    ".feather",
    ".arrow",
    ".ipc",
    ".zip",
    ".tar",
    ".gz",
    ".bz2",
    ".xz",
}
FITS_EXTENSIONS = (".fits", ".fit", ".fts", ".fz")


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("roots", nargs="+", help="Directories to scan.")
    parser.add_argument("--output-dir", required=True, help="Directory for reports and generated artifacts.")
    parser.add_argument("--max-per-ext", type=int, default=2, help="Max representative files per extension.")
    parser.add_argument(
        "--validation-profile",
        choices=["quick", "deep"],
        default="deep",
        help="quick runs sample inspection plus maintainer regressions and portable core smoke; deep keeps the full historical validation pass.",
    )
    parser.add_argument(
        "--include-system-iwork-samples",
        action="store_true",
        help=(
            "Opt in to supplementing .numbers coverage from the user's system iCloud/Numbers folder "
            "when the provided roots do not contain a .numbers sample. Disabled by default so real-data "
            "audits stay inside the explicit roots."
        ),
    )
    parser.add_argument("--summary-json", help="Optional standard-envelope JSON summary path.")
    parser.add_argument("--manifest-json", help="Optional manifest JSON path.")
    return parser.parse_args()


def emit_payload_safely(payload, summary_json=None, *, full_stdout=False):
    rendered = json.dumps(sanitize_payload(payload), indent=2, ensure_ascii=True) + "\n"
    if summary_json:
        target = Path(summary_json)
        try:
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_text(rendered, encoding="utf-8")
        except OSError as exc:
            raise RuntimeError(
                f"Could not write summary JSON to {public_path(target)}: {exc.__class__.__name__}: {exc}"
            ) from None
    if full_stdout:
        print(rendered, end="")


def validate_roots(roots):
    checked = []
    for raw in roots:
        path = Path(raw).expanduser()
        if not path.exists():
            raise RuntimeError(f"Validation root does not exist: {public_path(path)}")
        if not path.is_dir():
            raise RuntimeError(f"Validation root is not a directory: {public_path(path)}")
        checked.append(path)
    return checked


def ensure_validation_paths(args: argparse.Namespace) -> Path:
    """Reject report paths that could modify any explicit sample root."""

    output_dir = Path(args.output_dir).expanduser()
    input_paths = [operand(f"root[{index}]", raw, tree=True) for index, raw in enumerate(args.roots)]
    if args.include_system_iwork_samples:
        system_numbers = Path.home() / "Library/Mobile Documents/com~apple~Numbers/Documents"
        if system_numbers.exists():
            input_paths.append(operand("system Numbers sample root", system_numbers, tree=True))
    ensure_safe_paths(
        inputs=input_paths,
        outputs=[
            operand("--output-dir", output_dir, tree=True),
            operand("--summary-json", args.summary_json),
            operand("--manifest-json", args.manifest_json),
        ],
    )
    return output_dir


def select_samples(roots, max_per_ext):
    by_ext = defaultdict(list)
    for root in [Path(item).expanduser() for item in roots]:
        for path in root.rglob("*"):
            if not path.is_file():
                continue
            ext = path.suffix.lower() if path.suffix else "[noext]"
            if ext in SUPPORTED_EXTS and len(by_ext[ext]) < max_per_ext:
                by_ext[ext].append(path)
    return by_ext


def supplement_numbers_samples(by_ext, max_per_ext):
    extra_root = Path.home() / "Library/Mobile Documents/com~apple~Numbers/Documents"
    if len(by_ext[".numbers"]) >= max_per_ext or not extra_root.exists():
        return
    for path in sorted(extra_root.glob("*.numbers")):
        if len(by_ext[".numbers"]) >= max_per_ext:
            break
        by_ext[".numbers"].append(path)


def run(cmd, cwd=None, timeout_sec=None):
    if cmd and cmd[0] == "python3":
        cmd = [sys.executable, *cmd[1:]]
    try:
        result = subprocess.run(cmd, cwd=cwd, capture_output=True, text=True, timeout=timeout_sec)
    except subprocess.TimeoutExpired as exc:
        stdout = exc.stdout if isinstance(exc.stdout, str) else (exc.stdout or b"").decode("utf-8", "replace")
        stderr = exc.stderr if isinstance(exc.stderr, str) else (exc.stderr or b"").decode("utf-8", "replace")
        return {
            "command": cmd,
            "returncode": 124,
            "stdout": stdout[-4000:],
            "stderr": stderr[-4000:],
            "timed_out": True,
            "timeout_sec": timeout_sec,
        }
    return {
        "command": cmd,
        "returncode": result.returncode,
        "stdout": result.stdout[-4000:],
        "stderr": result.stderr[-4000:],
    }


def first_image_fits(samples):
    def looks_like_solver_artifact(path):
        lowered = str(path).lower()
        name = path.name.lower()
        if "index files" in lowered or "/astrometry.net-" in lowered:
            return True
        if name.startswith("index-"):
            return True
        if name in {"wcs.fits", "new_fits.fits"}:
            return True
        artifact_tokens = ("rdls", "axy", "corr", "match", "solved", "annotations", "calibration")
        return any(token in name for token in artifact_tokens)

    def candidates():
        for ext in FITS_EXTENSIONS:
            for path in samples.get(ext, []):
                yield path

    try:
        from astropy.io import fits
    except Exception:
        return next(candidates(), None)
    for path in candidates():
        if looks_like_solver_artifact(path):
            continue
        try:
            with fits.open(path, memmap=False) as hdus:
                for hdu in hdus:
                    data = getattr(hdu, "data", None)
                    if data is None or getattr(data, "names", None):
                        continue
                    if getattr(data, "ndim", 0) >= 2 and min(getattr(data, "shape", (0, 0))[:2]) >= 32:
                        return path
        except Exception:
            continue
    return None


def create_synthetic_coordinate_table(output_dir):
    output_dir.mkdir(parents=True, exist_ok=True)
    target = output_dir / "synthetic_coords.csv"
    target.write_text(
        "id,ra_deg,dec_deg\n"
        "1,150.000000,2.000000\n"
        "2,150.000300,2.000200\n"
        "3,150.000600,2.000400\n",
        encoding="utf-8",
    )
    return target, "ra_deg", "dec_deg"


def require_outputs(result, *paths):
    missing = [str(path) for path in paths if not Path(path).exists()]
    if missing:
        result["returncode"] = result["returncode"] or 1
        extra = "Missing expected outputs: " + ", ".join(missing)
        result["stderr"] = (result.get("stderr", "") + "\n" + extra).strip()
    return result


def inspect_zip(path):
    with zipfile.ZipFile(path) as zf:
        return {"entry_count": len(zf.namelist()), "entries": zf.namelist()[:12]}


def inspect_pdf(path):
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    text = ""
    if reader.pages:
        try:
            text = (reader.pages[0].extract_text() or "")[:400]
        except Exception:
            text = ""
    return {"pages": len(reader.pages), "sample_text": text}


def inspect_notebook(path):
    data = json.loads(path.read_text(errors="replace"))
    cells = data.get("cells", [])
    return {"cells": len(cells), "cell_types": [cell.get("cell_type") for cell in cells[:8]]}


def inspect_textual(path):
    with path.open("r", errors="replace") as fh:
        lines = []
        for _ in range(5):
            line = fh.readline()
            if not line:
                break
            lines.append(line.rstrip())
    return {"preview": lines}


def inspect_binary_signature(path):
    with path.open("rb") as fh:
        sig = fh.read(16)
    return {"signature_hex": sig.hex()}


def maybe_find_region_peer(path):
    directory = path.parent
    reg_files = sorted(directory.glob("*.reg"))
    return reg_files[0] if reg_files else None


def choose_filter_expression(path):
    try:
        with path.open("r", errors="replace") as fh:
            header = fh.readline().strip()
        if not header:
            return None
        columns = [item.strip() for item in header.split(",") if item.strip()]
        for name in columns:
            if name.isidentifier():
                return f"({name} == {name})"
    except Exception:
        return None
    return None


def find_coordinate_table(samples):
    try:
        from astropy.table import Table
    except Exception:
        return None

    candidates = []
    for ext in (".ecsv", ".tbl", ".csv", ".fits", ".fit"):
        candidates.extend(samples.get(ext, []))
    pairs = [
        ("ra", "dec"),
        ("RA", "DEC"),
        ("ApertureRA", "ApertureDec"),
        ("CentroidRA", "CentroidDec"),
    ]
    for path in candidates:
        try:
            table = read_table_any(Table, path)
        except Exception:
            continue
        names = set(table.colnames)
        for ra_name, dec_name in pairs:
            if ra_name in names and dec_name in names:
                return path, ra_name, dec_name
    return None


def find_master_bias_and_science(roots):
    master_bias = None
    science = None
    for root in [Path(item) for item in roots]:
        for path in root.rglob("*master_bias*.fits"):
            master_bias = path
            break
        if master_bias is not None:
            break
    if master_bias is None:
        return None
    for candidate in master_bias.parent.glob("ALrd*.fits"):
        if candidate.is_file():
            science = candidate
            break
    if science is None:
        for candidate in master_bias.parent.glob("*.fits"):
            if candidate.is_file() and candidate.name != master_bias.name and "master" not in candidate.name.lower():
                science = candidate
                break
    if science is None:
        return None
    return master_bias, science


def find_teareduce_notebooks(roots, limit=3):
    notebooks = []
    seen = set()
    preferred_names = [
        "P2_01_primeros_pasos.ipynb",
        "P3_04_eliminacion_rayos_cosmicos.ipynb",
        "P3_05_calibracion_longitud_de_onda.ipynb",
    ]
    preferred_found = {}
    for root in [Path(item) for item in roots]:
        for pattern in ("P2_*.ipynb", "P3_*.ipynb"):
            for path in sorted(root.rglob(pattern)):
                if ".ipynb_checkpoints" in path.parts:
                    continue
                resolved = str(path.resolve())
                if resolved in seen:
                    continue
                if path.name in preferred_names and path.name not in preferred_found:
                    preferred_found[path.name] = path
                    seen.add(resolved)
                    continue
                notebooks.append(path)
                seen.add(resolved)
    ordered = [preferred_found[name] for name in preferred_names if name in preferred_found]
    ordered.extend(notebooks)
    return ordered[:limit]


def find_teareduce_fits(roots):
    for root in [Path(item) for item in roots]:
        for path in root.rglob("*.fits"):
            lowered = str(path).lower()
            if "tea_procesado" in lowered or "/n1/" in lowered or "/n2/" in lowered:
                return path
    return None


def create_synthetic_containers(output_dir):
    import numpy as np
    import h5py
    import sqlite3
    import xarray as xr
    from scipy.io import savemat

    output_dir.mkdir(parents=True, exist_ok=True)
    np.save(output_dir / "array.npy", np.arange(12, dtype=float).reshape(3, 4))
    np.savez(output_dir / "bundle.npz", x=np.arange(5), y=np.linspace(0.0, 1.0, 5))
    with h5py.File(output_dir / "sample.h5", "w") as fh:
        fh.create_dataset("group/data", data=np.arange(8).reshape(2, 4))
    savemat(output_dir / "sample.mat", {"signal": np.sin(np.linspace(0, 3.14, 20)), "grid": np.arange(9).reshape(3, 3)})
    sqlite_path = output_dir / "sample.sqlite"
    if sqlite_path.exists():
        sqlite_path.unlink()
    conn = sqlite3.connect(sqlite_path)
    try:
        conn.execute("CREATE TABLE measurements (id INTEGER PRIMARY KEY, value REAL)")
        conn.executemany("INSERT INTO measurements (value) VALUES (?)", [(1.0,), (2.5,), (3.75,)])
        conn.commit()
    finally:
        conn.close()
    ds = xr.Dataset({"flux": ("x", np.linspace(0.0, 1.0, 6))}, coords={"x": np.arange(6)})
    ds.to_netcdf(output_dir / "sample.nc")
    (output_dir / "archive_payload.txt").write_text("synthetic payload\n")
    with zipfile.ZipFile(output_dir / "sample.zip", "w") as zf:
        zf.write(output_dir / "archive_payload.txt", arcname="archive_payload.txt")
    with tarfile.open(output_dir / "sample.tar.gz", "w:gz") as tf:
        tf.add(output_dir / "archive_payload.txt", arcname="archive_payload.txt")
    with gzip.open(output_dir / "sample.txt.gz", "wb") as fh:
        fh.write(b"synthetic gzip payload\n")
    with bz2.open(output_dir / "sample.txt.bz2", "wb") as fh:
        fh.write(b"synthetic bz2 payload\n")
    with lzma.open(output_dir / "sample.txt.xz", "wb") as fh:
        fh.write(b"synthetic xz payload\n")
    ds.to_zarr(output_dir / "sample.zarr", mode="w")
    (output_dir / "sample.geojson").write_text(
        json.dumps(
            {
                "type": "FeatureCollection",
                "features": [
                    {"type": "Feature", "geometry": {"type": "Point", "coordinates": [1.0, 2.0]}, "properties": {"name": "A", "value": 1}},
                    {"type": "Feature", "geometry": {"type": "Point", "coordinates": [3.0, 4.0]}, "properties": {"name": "B", "value": 2}},
                ],
            },
            ensure_ascii=True,
        )
        + "\n"
    )
    (output_dir / "corrupt.txt.gz").write_bytes((output_dir / "sample.txt.gz").read_bytes()[:12])


def create_synthetic_table_variants(output_dir):
    output_dir.mkdir(parents=True, exist_ok=True)
    rows = [
        {"id": 1, "value": 2.5, "label": "a"},
        {"id": 2, "value": 3.5, "label": "b"},
        {"id": 3, "value": 5.0, "label": "c"},
    ]
    (output_dir / "sample.jsonl").write_text("\n".join(json.dumps(row, ensure_ascii=True) for row in rows) + "\n")
    (output_dir / "sample.json").write_text(json.dumps(rows, indent=2, ensure_ascii=True) + "\n")
    (output_dir / "sample.yaml").write_text("rows:\n  - id: 1\n    value: 2.5\n    label: a\n")
    try:
        with suppress_fd_output(True):
            import pandas as pd
    except ImportError:
        return
    dataframe = pd.DataFrame(rows)
    dataframe.to_excel(output_dir / "sample.xlsx", index=False)
    try:
        dataframe.to_csv(output_dir / "sample.csv", index=False)
    except Exception:
        pass
    try:
        with contextlib.redirect_stderr(io.StringIO()), suppress_fd_output(True):
            dataframe.to_parquet(output_dir / "sample.parquet", index=False)
    except Exception:
        pass
    try:
        with contextlib.redirect_stderr(io.StringIO()), suppress_fd_output(True):
            dataframe.to_feather(output_dir / "sample.feather")
    except Exception:
        pass
    try:
        with contextlib.redirect_stderr(io.StringIO()), suppress_fd_output(True):
            import pyarrow as pa
            import pyarrow.ipc as ipc

        table = pa.Table.from_pandas(dataframe)
        with suppress_fd_output(True):
            with pa.OSFile(str(output_dir / "sample.arrow"), "wb") as sink:
                with ipc.new_file(sink, table.schema) as writer:
                    writer.write_table(table)
    except Exception:
        pass
    try:
        from pyxlsbwriter import XlsbWriter

        with XlsbWriter(str(output_dir / "sample.xlsb")) as writer:
            writer.add_sheet("Sheet1")
            writer.write_sheet([list(dataframe.columns)] + dataframe.values.tolist())
    except Exception:
        pass


def create_synthetic_office_docs(output_dir):
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation

    output_dir.mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.add_heading("Validation Memo", level=1)
    doc.add_paragraph("This document includes Placeholder text for round-trip replacement.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Status"
    table.cell(1, 1).text = "Placeholder"
    doc.save(output_dir / "sample.docx")

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Validation Deck"
    slide.placeholders[1].text = "Placeholder bullet for replacement"
    presentation.save(output_dir / "sample.pptx")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"
    worksheet["A1"] = "id"
    worksheet["B1"] = "value"
    worksheet["A2"] = 1
    worksheet["B2"] = "Placeholder"
    workbook.save(output_dir / "sample.xlsx")


def create_synthetic_spectrum(path):
    import numpy as np

    wavelength = np.linspace(6500.0, 6625.0, 400)
    continuum = 1.0 + 0.0003 * (wavelength - 6562.8)
    line = -0.35 * np.exp(-0.5 * ((wavelength - 6562.8) / 2.2) ** 2)
    flux = continuum + line
    path.parent.mkdir(parents=True, exist_ok=True)
    lines = ["wavelength flux"]
    lines.extend(f"{x:.6f} {y:.8f}" for x, y in zip(wavelength, flux))
    path.write_text("\n".join(lines) + "\n")


def create_synthetic_spectrum_2d(path):
    import numpy as np
    from astropy.io import fits

    path.parent.mkdir(parents=True, exist_ok=True)
    wavelength = np.linspace(6500.0, 6625.0, 400)
    continuum = 1.0 + 0.0003 * (wavelength - 6562.8)
    line = -0.35 * np.exp(-0.5 * ((wavelength - 6562.8) / 2.2) ** 2)
    spectrum = continuum + line
    rows = 48
    image = np.random.default_rng(42).normal(0.0, 0.01, size=(rows, wavelength.size))
    for center, scale in [(16.0, 1.0), (33.0, 0.7)]:
        profile = np.exp(-0.5 * ((np.arange(rows) - center) / 2.0) ** 2)
        image += (profile[:, None] * spectrum[None, :]) * scale
    header = fits.Header()
    header["CRVAL1"] = 6500.0
    header["CDELT1"] = (6625.0 - 6500.0) / (wavelength.size - 1)
    header["CRPIX1"] = 1.0
    fits.PrimaryHDU(image.astype("float32"), header=header).writeto(path, overwrite=True)


def create_synthetic_calibration_table(path):
    import numpy as np

    pixels = np.array([0, 100, 200, 300, 399], dtype=float)
    wavelengths = 6500.0 + pixels * ((6625.0 - 6500.0) / 399.0)
    lines = ["pixel wavelength"]
    lines.extend(f"{p:.3f} {w:.6f}" for p, w in zip(pixels, wavelengths))
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n")


def summarize_validation_report(report):
    feature_runs = report.get("feature_runs", {})
    failed_features = sorted(name for name, result in feature_runs.items() if result.get("returncode") != 0)
    sample_errors = []
    sample_count = 0
    for ext, items in report.get("by_extension", {}).items():
        sample_count += len(items)
        for item in items:
            if item.get("status") != "ok":
                sample_errors.append({"extension": ext, "path": item.get("path"), "error": item.get("error")})
    if failed_features or sample_errors:
        status = "fail"
    elif sample_count == 0:
        status = "warning"
    else:
        status = "ok"
    return {
        "status": status,
        "exit_code": 1 if status == "fail" else 0,
        "failed_features": failed_features,
        "sample_errors": sample_errors,
        "metrics": {
            "sample_count": sample_count,
            "extension_count": len(report.get("by_extension", {})),
            "feature_count": len(feature_runs),
            "failed_feature_count": len(failed_features),
            "sample_error_count": len(sample_errors),
        },
    }


def build_summary_payload(report, output_dir, args, report_paths, summary):
    findings = []
    findings.extend(f"feature failed: {name}" for name in summary["failed_features"])
    findings.extend(f"sample error: {item['path']} ({item['error']})" for item in summary["sample_errors"])
    if summary["status"] == "warning":
        findings.append("No supported samples were found under the explicit validation roots.")
    return standard_tool_payload(
        "validate_skill_samples",
        status=summary["status"],
        notes=[
            "Maintainer sample validation with quick and deep profiles for regression and release checks.",
            "The validator is non-destructive and should run on explicit roots or temporary copies.",
        ],
        artifacts={
            "validation_report_json": report_paths["json"],
            "validation_report_md": report_paths["md"],
            "summary_json": getattr(args, "summary_json", None),
            "manifest_json": getattr(args, "manifest_json", None),
            "output_dir": output_dir,
        },
        results={
            "roots": report.get("roots", []),
            "validation_profile": report.get("validation_profile"),
            "overall_status": summary["status"].upper(),
            "failed_features": summary["failed_features"],
            "sample_errors": summary["sample_errors"],
            "metrics": summary["metrics"],
        },
        qa={
            "status": summary["status"],
            "findings": findings,
            "metrics": summary["metrics"],
        },
        include_environment=True,
    )


def build_failure_payload(args, exc, *, blocked=False):
    output_dir = Path(args.output_dir).expanduser() if getattr(args, "output_dir", None) else None
    return standard_tool_payload(
        "validate_skill_samples",
        status="blocked" if blocked else "fail",
        notes=[
            "Maintainer sample validation with quick and deep profiles for regression and release checks.",
            (
                "The validator was blocked before writing because an output could overlap an input."
                if blocked
                else "The validator failed before it could complete cleanly."
            ),
        ],
        artifacts={} if blocked else {
            "summary_json": getattr(args, "summary_json", None),
            "manifest_json": getattr(args, "manifest_json", None),
            "output_dir": output_dir,
        },
        results={
            "roots": getattr(args, "roots", []),
            "validation_profile": getattr(args, "validation_profile", None),
            "overall_status": "BLOCKED_CONTROLADO" if blocked else "FAIL",
            "error_type": exc.__class__.__name__,
            "error": str(exc),
        },
        qa={
            "status": "blocked" if blocked else "fail",
            "findings": [
                {
                    "severity": "high",
                    "title": "validate_skill_samples blocked" if blocked else "validate_skill_samples failed",
                    "detail": str(exc),
                }
            ],
            "metrics": {"failed_feature_count": 0, "sample_error_count": 0},
        },
        include_environment=True,
    )


def write_validation_reports(report, output_dir):
    summary = summarize_validation_report(report)
    report["overall"] = summary
    report_path = output_dir / "validation_report.json"
    report_path.write_text(json.dumps(report, indent=2, ensure_ascii=True) + "\n")

    md_lines = ["# Scientific Data Analysis Validation Report", ""]
    md_lines.append(f"- Validation profile: `{report.get('validation_profile', 'deep')}`")
    md_lines.append(f"- Overall status: `{summary['status'].upper()}`")
    md_lines.append(f"- Feature runs: `{summary['metrics']['feature_count']}` total, `{summary['metrics']['failed_feature_count']}` failed")
    md_lines.append(f"- Samples inspected: `{summary['metrics']['sample_count']}`")
    md_lines.append("")
    md_lines.append("## Sample Coverage")
    for ext, items in sorted(report["by_extension"].items()):
        ok = sum(1 for item in items if item.get("status") == "ok")
        md_lines.append(f"- `{ext}`: {ok}/{len(items)} samples inspected")
    md_lines.append("")
    md_lines.append("## Feature Runs")
    for name, result in sorted(report["feature_runs"].items()):
        status = "PASS" if result["returncode"] == 0 else "FAIL"
        md_lines.append(f"- `{name}`: {status}")
    (output_dir / "validation_report.md").write_text("\n".join(md_lines) + "\n")

    print(f"Validation report written to {report_path.resolve()}")
    print(f"Markdown summary written to {(output_dir / 'validation_report.md').resolve()}")
    return {"json": report_path, "md": output_dir / "validation_report.md"}, summary


def finalize_validation(args, output_dir, report):
    report_paths, summary = write_validation_reports(report, output_dir)
    payload = build_summary_payload(report, output_dir, args, report_paths, summary)
    emit_payload_safely(payload, args.summary_json, full_stdout=bool(args.summary_json))
    if args.manifest_json:
        write_manifest(
            args.manifest_json,
            inputs=[Path(root) for root in args.roots],
            outputs=[
                report_paths["json"],
                report_paths["md"],
                *((Path(args.summary_json),) if args.summary_json else ()),
            ],
            parameters={
                "validation_profile": args.validation_profile,
                "max_per_ext": args.max_per_ext,
                "include_system_iwork_samples": args.include_system_iwork_samples,
            },
            command="validate_skill_samples.py",
            notes=payload["notes"],
            extra={"overall_status": summary["status"].upper()},
        )
    return summary["exit_code"]


def run_validation(args):
    output_dir = ensure_validation_paths(args)
    validate_roots(args.roots)
    output_dir.mkdir(parents=True, exist_ok=True)
    scripts_dir = Path(__file__).resolve().parent
    samples = select_samples(args.roots, args.max_per_ext)
    if args.include_system_iwork_samples:
        supplement_numbers_samples(samples, args.max_per_ext)
    report = {"roots": args.roots, "validation_profile": args.validation_profile, "by_extension": {}, "feature_runs": {}}
    if args.include_system_iwork_samples:
        report["system_iwork_sample_supplement"] = {
            "enabled": True,
            "source": str(Path.home() / "Library/Mobile Documents/com~apple~Numbers/Documents"),
        }
    else:
        report["system_iwork_sample_supplement"] = {
            "enabled": False,
            "note": "Only files under the explicit roots were sampled.",
        }

    for ext, files in sorted(samples.items()):
        ext_report = []
        for path in files:
            entry = {"path": str(path)}
            try:
                if ext in {".docx", ".docm", ".pptx", ".pptm", ".xlsx", ".xlsm", ".pages", ".numbers", ".key", ".odt", ".ods", ".odp", ".epub", ".xlsb"}:
                    entry["inspection"] = inspect_zip(path)
                elif ext == ".pdf":
                    entry["inspection"] = inspect_pdf(path)
                elif ext == ".ipynb":
                    entry["inspection"] = inspect_notebook(path)
                elif ext in {".tex", ".bib", ".txt", ".text", ".dat", ".tbl", ".csv", ".ecsv", ".html", ".htm", ".json", ".jsonl", ".ndjson", ".yaml", ".yml", ".toml", ".rtf", ".reg", ".pref", ".jmars"}:
                    entry["inspection"] = inspect_textual(path)
                elif ext in {".doc", ".xls", ".zip", ".tar", ".gz", ".bz2", ".xz", ".png", ".jpg", ".jpeg", ".fits", ".fit", ".fts", ".fz", ".npy", ".npz", ".h5", ".hdf5", ".mat", ".sqlite", ".sqlite3", ".db", ".nc", ".nc4", ".cdf", ".netcdf", ".parquet", ".feather", ".arrow", ".ipc"}:
                    entry["inspection"] = inspect_binary_signature(path)
                else:
                    entry["inspection"] = {"status": "skipped"}
                entry["status"] = "ok"
            except Exception as exc:
                entry["status"] = "error"
                entry["error"] = f"{exc.__class__.__name__}: {exc}"
            ext_report.append(entry)
        report["by_extension"][ext] = ext_report

    feature_dir = output_dir / "feature_runs"
    feature_dir.mkdir(parents=True, exist_ok=True)

    report["feature_runs"]["audit_step1_regression"] = run(["python3", str(scripts_dir / "audit_step1_regression.py")])
    report["feature_runs"]["audit_step2_contract_regression"] = run(["python3", str(scripts_dir / "audit_step2_contract_regression.py")])
    report["feature_runs"]["audit_physical_qa_regression"] = run(["python3", str(scripts_dir / "audit_physical_qa_regression.py")])
    report["feature_runs"]["audit_coursework_handoff_regression"] = run(["python3", str(scripts_dir / "audit_coursework_handoff_regression.py")])
    report["feature_runs"]["audit_coursework_notebook_fidelity_regression"] = run(
        ["python3", str(scripts_dir / "audit_coursework_notebook_fidelity_regression.py")]
    )
    report["feature_runs"]["audit_document_reporting_regression"] = run(["python3", str(scripts_dir / "audit_document_reporting_regression.py")])
    report["feature_runs"]["audit_practice4_visual_handoff_regression"] = run(
        ["python3", str(scripts_dir / "audit_practice4_visual_handoff_regression.py")]
    )
    report["feature_runs"]["audit_v1_2_release_regression"] = run(
        ["python3", str(scripts_dir / "audit_v1_2_release_regression.py")]
    )
    report["feature_runs"]["audit_v1_3_release_regression"] = run(
        ["python3", str(scripts_dir / "audit_v1_3_release_regression.py")]
    )
    report["feature_runs"]["audit_v1_4_companion_routing_regression"] = run(
        ["python3", str(scripts_dir / "audit_v1_4_companion_routing_regression.py")]
    )
    report["feature_runs"]["audit_fits_rgb_batch_regression"] = run(
        ["python3", str(scripts_dir / "audit_fits_rgb_batch_regression.py")]
    )
    report["feature_runs"]["audit_external_astro_tools_regression"] = run(
        ["python3", str(scripts_dir / "audit_external_astro_tools_regression.py")]
    )

    if args.validation_profile == "quick":
        report["feature_runs"]["portable_smoke_test_core"] = run(
            [
                "python3",
                str(scripts_dir / "portable_smoke_test.py"),
                "--profile",
                "core",
                "--output-dir",
                str(feature_dir / "portable_smoke_core"),
                "--examples-dir",
                str(Path(__file__).resolve().parent.parent / "examples"),
                "--summary-json",
                str(feature_dir / "portable_smoke_core" / "summary.json"),
                "--manifest-json",
                str(feature_dir / "portable_smoke_core" / "manifest.json"),
            ]
        )
        report["feature_runs"]["portable_smoke_test_core"] = require_outputs(
            report["feature_runs"]["portable_smoke_test_core"],
            feature_dir / "portable_smoke_core" / "summary.json",
            feature_dir / "portable_smoke_core" / "manifest.json",
            feature_dir / "portable_smoke_core" / "inspect_fits.json",
            feature_dir / "portable_smoke_core" / "profile_table.json",
        )
        return finalize_validation(args, output_dir, report)

    report["feature_runs"]["env_doctor"] = run(
        [
            "python3",
            str(scripts_dir / "env_doctor.py"),
            "--summary-json",
            str(feature_dir / "env_doctor.json"),
            "--manifest-json",
            str(feature_dir / "env_doctor_manifest.json"),
        ]
    )
    report["feature_runs"]["env_doctor"] = require_outputs(
        report["feature_runs"]["env_doctor"],
        feature_dir / "env_doctor.json",
        feature_dir / "env_doctor_manifest.json",
    )

    def first(ext):
        values = samples.get(ext) or []
        return values[0] if values else None

    fits_sample = first_image_fits(samples)
    if fits_sample:
        fits_out = feature_dir / "inspect_fits"
        fits_out.mkdir(exist_ok=True)
        report["feature_runs"]["inspect_fits"] = run(
            ["python3", str(scripts_dir / "inspect_fits.py"), str(fits_sample), "--summary-json", str(fits_out / "summary.json")],
        )
        quicklook_cmd = [
            "python3",
            str(scripts_dir / "fits_quicklook.py"),
            str(fits_sample),
            "--output",
            str(fits_out / "quicklook.png"),
        ]
        peer_region = maybe_find_region_peer(fits_sample)
        if peer_region is not None:
            quicklook_cmd.extend(["--region-file", str(peer_region)])
        report["feature_runs"]["fits_quicklook"] = run(quicklook_cmd)
        report["feature_runs"]["aperture_photometry"] = run(
            [
                "python3",
                str(scripts_dir / "aperture_photometry.py"),
                str(fits_sample),
                "--auto-brightest",
                "3",
                "--output-json",
                str(fits_out / "apertures.json"),
            ]
        )
        report["feature_runs"]["physical_qa_fits"] = run(
            ["python3", str(scripts_dir / "physical_qa.py"), str(fits_sample), "--summary-json", str(fits_out / "fits_qa.json")]
        )

    csv_sample = first(".csv")
    ecsv_sample = first(".ecsv")
    tbl_sample = first(".tbl")
    table_out = None
    if csv_sample:
        table_out = feature_dir / "tables"
        table_out.mkdir(exist_ok=True)
        report["feature_runs"]["profile_csv"] = run(
            ["python3", str(scripts_dir / "profile_table.py"), str(csv_sample), "--summary-json", str(table_out / "csv_summary.json")]
        )
        report["feature_runs"]["physical_qa_table"] = run(
            ["python3", str(scripts_dir / "physical_qa.py"), str(csv_sample), "--summary-json", str(table_out / "table_qa.json")]
        )
        expression = choose_filter_expression(csv_sample)
        if expression is not None:
            report["feature_runs"]["catalog_filter"] = run(
                [
                    "python3",
                    str(scripts_dir / "catalog_workbench.py"),
                    "filter",
                    str(csv_sample),
                    str(table_out / "filtered.csv"),
                    "--expression",
                    expression,
                ]
            )
    synthetic_table_out = feature_dir / "table_variants"
    create_synthetic_table_variants(synthetic_table_out)
    jsonl_sample = synthetic_table_out / "sample.jsonl"
    if jsonl_sample.exists():
        report["feature_runs"]["profile_jsonl"] = run(
            ["python3", str(scripts_dir / "profile_table.py"), str(jsonl_sample), "--summary-json", str(synthetic_table_out / "jsonl_summary.json")]
        )
        report["feature_runs"]["catalog_filter_jsonl"] = run(
            [
                "python3",
                str(scripts_dir / "catalog_workbench.py"),
                "filter",
                str(jsonl_sample),
                str(synthetic_table_out / "filtered.jsonl"),
                "--expression",
                "(value > 3.0)",
            ]
        )
    synthetic_xlsb = synthetic_table_out / "sample.xlsb"
    if synthetic_xlsb.exists():
        report["feature_runs"]["profile_xlsb_generated"] = run(
            ["python3", str(scripts_dir / "profile_table.py"), str(synthetic_xlsb), "--summary-json", str(synthetic_table_out / "xlsb_summary.json")]
        )
    synthetic_csv = synthetic_table_out / "sample.csv"
    if synthetic_csv.exists() and jsonl_sample.exists():
        report["feature_runs"]["duckdb_query_join"] = run(
            [
                "python3",
                str(scripts_dir / "duckdb_workbench.py"),
                str(synthetic_csv),
                str(jsonl_sample),
                "--sql",
                "SELECT a.id, a.value, b.label FROM source0 a JOIN source1 b USING(id) ORDER BY id",
                "--output",
                str(synthetic_table_out / "duckdb_join.csv"),
                "--summary-json",
                str(synthetic_table_out / "duckdb_join.json"),
            ]
        )
        report["feature_runs"]["semantic_diff_table"] = run(
            [
                "python3",
                str(scripts_dir / "semantic_diff.py"),
                str(synthetic_csv),
                str(synthetic_table_out / "duckdb_join.csv"),
                "--output-json",
                str(synthetic_table_out / "table_diff.json"),
                "--output-md",
                str(synthetic_table_out / "table_diff.md"),
            ]
        )

    office_out = feature_dir / "office_roundtrip"
    create_synthetic_office_docs(office_out)
    report["feature_runs"]["office_docx_replace"] = run(
        [
            "python3",
            str(scripts_dir / "office_roundtrip.py"),
            "docx-replace",
            str(office_out / "sample.docx"),
            str(office_out / "sample_edited.docx"),
            "--find",
            "Placeholder",
            "--replace",
            "Updated",
            "--summary-json",
            str(office_out / "docx_replace.json"),
        ]
    )
    report["feature_runs"]["office_docx_export"] = run(
        [
            "python3",
            str(scripts_dir / "office_roundtrip.py"),
            "docx-export-text",
            str(office_out / "sample_edited.docx"),
            str(office_out / "sample_edited.md"),
        ]
    )
    report["feature_runs"]["office_pptx_replace"] = run(
        [
            "python3",
            str(scripts_dir / "office_roundtrip.py"),
            "pptx-replace",
            str(office_out / "sample.pptx"),
            str(office_out / "sample_edited.pptx"),
            "--find",
            "Placeholder",
            "--replace",
            "Updated",
            "--summary-json",
            str(office_out / "pptx_replace.json"),
        ]
    )
    report["feature_runs"]["office_pptx_backends"] = run(
        [
            "python3",
            str(scripts_dir / "office_roundtrip.py"),
            "pptx-export-backends",
            "--summary-json",
            str(office_out / "pptx_backends.json"),
        ]
    )
    report["feature_runs"]["office_pptx_backends"] = require_outputs(
        report["feature_runs"]["office_pptx_backends"],
        office_out / "pptx_backends.json",
    )
    report["feature_runs"]["office_xlsx_set_cell"] = run(
        [
            "python3",
            str(scripts_dir / "office_roundtrip.py"),
            "xlsx-set-cell",
            str(office_out / "sample.xlsx"),
            str(office_out / "sample_edited.xlsx"),
            "--sheet",
            "Sheet1",
            "--cell",
            "B2",
            "--value",
            "Updated",
            "--summary-json",
            str(office_out / "xlsx_set.json"),
        ]
    )
    report["feature_runs"]["office_xlsx_export_csv"] = run(
        [
            "python3",
            str(scripts_dir / "office_roundtrip.py"),
            "xlsx-export-csv",
            str(office_out / "sample_edited.xlsx"),
            "--output-dir",
            str(office_out / "csv_export"),
        ]
    )
    try:
        backends_summary = json.loads((office_out / "pptx_backends.json").read_text())
    except Exception:
        backends_summary = {}
    if backends_summary.get("soffice"):
        report["feature_runs"]["office_pptx_export_pdf"] = run(
            [
                "python3",
                str(scripts_dir / "office_roundtrip.py"),
                "pptx-export-pdf",
                str(office_out / "sample_edited.pptx"),
                str(office_out / "sample_edited.pdf"),
                "--backend",
                "soffice",
                "--summary-json",
                str(office_out / "pptx_export_pdf.json"),
            ]
        )
        report["feature_runs"]["office_pptx_export_pdf"] = require_outputs(
            report["feature_runs"]["office_pptx_export_pdf"],
            office_out / "sample_edited.pdf",
            office_out / "pptx_export_pdf.json",
        )
    report["feature_runs"]["presentation_workbench_inspect"] = run(
        [
            "python3",
            str(scripts_dir / "presentation_workbench.py"),
            "inspect",
            str(office_out / "sample_edited.pptx"),
            "--summary-json",
            str(office_out / "presentation_inspect.json"),
        ]
    )
    report["feature_runs"]["presentation_workbench_inspect"] = require_outputs(
        report["feature_runs"]["presentation_workbench_inspect"],
        office_out / "presentation_inspect.json",
    )
    presentation_handoff_dir = office_out / "presentation_handoff"
    presentation_handoff_cmd = [
        "python3",
        str(scripts_dir / "presentation_workbench.py"),
        "handoff",
        str(office_out / "sample_edited.pptx"),
        str(presentation_handoff_dir),
        "--summary-json",
        str(office_out / "presentation_handoff.json"),
        "--manifest-json",
        str(office_out / "presentation_handoff_manifest.json"),
    ]
    if backends_summary.get("soffice"):
        presentation_handoff_cmd.extend(["--export-pdf", "--backend", "soffice"])
    report["feature_runs"]["presentation_workbench_handoff"] = run(presentation_handoff_cmd)
    required_presentation_outputs = [
        office_out / "presentation_handoff.json",
        office_out / "presentation_handoff_manifest.json",
        presentation_handoff_dir / "sample_edited.pptx",
        presentation_handoff_dir / "slide_text.md",
        presentation_handoff_dir / "handoff_notes.md",
    ]
    if backends_summary.get("soffice"):
        required_presentation_outputs.append(presentation_handoff_dir / "sample_edited.pdf")
    report["feature_runs"]["presentation_workbench_handoff"] = require_outputs(
        report["feature_runs"]["presentation_workbench_handoff"],
        *required_presentation_outputs,
    )
    report["feature_runs"]["semantic_diff_docx"] = run(
        [
            "python3",
            str(scripts_dir / "semantic_diff.py"),
            str(office_out / "sample.docx"),
            str(office_out / "sample_edited.docx"),
            "--output-json",
            str(office_out / "docx_diff.json"),
            "--output-md",
            str(office_out / "docx_diff.md"),
        ]
    )
    if table_out is None:
        table_out = feature_dir / "tables"
        table_out.mkdir(exist_ok=True)
    coord_path, ra_name, dec_name = create_synthetic_coordinate_table(table_out)
    report["feature_runs"]["catalog_crossmatch"] = run(
        [
            "python3",
            str(scripts_dir / "catalog_workbench.py"),
            "crossmatch-sky",
            str(coord_path),
            str(coord_path),
            str(table_out / "crossmatch.ecsv"),
            "--left-ra",
            ra_name,
            "--left-dec",
            dec_name,
            "--right-ra",
            ra_name,
            "--right-dec",
            dec_name,
            "--radius-arcsec",
            "1.0",
        ]
    )
    if tbl_sample:
        series_out = feature_dir / "series"
        series_out.mkdir(exist_ok=True)
        report["feature_runs"]["analyze_tbl_series"] = run(
            [
                "python3",
                str(scripts_dir / "analyze_series.py"),
                str(tbl_sample),
                "--kind",
                "time",
                "--plot",
                str(series_out / "tbl_series.png"),
                "--summary-json",
                str(series_out / "tbl_series.json"),
            ]
        )

    notebook_sample = first(".ipynb")
    if notebook_sample:
        nb_out = feature_dir / "notebook"
        nb_out.mkdir(exist_ok=True)
        report["feature_runs"]["bootstrap_notebook"] = run(
            [
                "python3",
                str(scripts_dir / "bootstrap_analysis_notebook.py"),
                str(nb_out / "template.ipynb"),
                "--title",
                "Skill Validation Notebook",
                "--domain",
                "astronomy",
                "--language",
                "bilingual",
                "--overwrite",
            ]
        )
    teareduce_notebooks = find_teareduce_notebooks(args.roots)
    teareduce_fits = find_teareduce_fits(args.roots)
    if teareduce_notebooks:
        teareduce_out = feature_dir / "teareduce"
        teareduce_out.mkdir(exist_ok=True)
        report["feature_runs"]["teareduce_healthcheck"] = run(
            [
                "python3",
                str(scripts_dir / "teareduce_healthcheck.py"),
                *[str(path) for path in teareduce_notebooks],
                "--summary-json",
                str(teareduce_out / "healthcheck.json"),
                "--manifest-json",
                str(teareduce_out / "healthcheck_manifest.json"),
            ]
        )
        report["feature_runs"]["teareduce_healthcheck"] = require_outputs(
            report["feature_runs"]["teareduce_healthcheck"],
            teareduce_out / "healthcheck.json",
            teareduce_out / "healthcheck_manifest.json",
        )
        report["feature_runs"]["teareduce_bridge_notebook_scan"] = run(
            [
                "python3",
                str(scripts_dir / "teareduce_bridge.py"),
                "notebook-scan",
                *[str(path) for path in teareduce_notebooks],
                "--summary-json",
                str(teareduce_out / "notebook_scan.json"),
            ]
        )
        report["feature_runs"]["teareduce_router_wavecal"] = run(
            [
                "python3",
                str(scripts_dir / "teareduce_router.py"),
                "--intent",
                "calibracion de longitud de onda de un espectro 2D con notebook de practica",
                str(teareduce_notebooks[-1]),
                "--summary-json",
                str(teareduce_out / "router.json"),
            ]
        )
    if teareduce_fits:
        teareduce_out = feature_dir / "teareduce"
        teareduce_out.mkdir(exist_ok=True)
        report["feature_runs"]["teareduce_bridge_statsummary"] = run(
            [
                "python3",
                str(scripts_dir / "teareduce_bridge.py"),
                "statsummary",
                str(teareduce_fits),
                "--summary-json",
                str(teareduce_out / "statsummary.json"),
                "--manifest-json",
                str(teareduce_out / "statsummary_manifest.json"),
            ]
        )
        report["feature_runs"]["teareduce_bridge_statsummary"] = require_outputs(
            report["feature_runs"]["teareduce_bridge_statsummary"],
            teareduce_out / "statsummary.json",
            teareduce_out / "statsummary_manifest.json",
        )
        report["feature_runs"]["teareduce_bridge_imshow"] = run(
            [
                "python3",
                str(scripts_dir / "teareduce_bridge.py"),
                "imshow-fits",
                str(teareduce_fits),
                "--output",
                str(teareduce_out / "teareduce_quicklook.png"),
                "--summary-json",
                str(teareduce_out / "teareduce_quicklook.json"),
            ]
        )
        report["feature_runs"]["teareduce_bridge_imshow"] = require_outputs(
            report["feature_runs"]["teareduce_bridge_imshow"],
            teareduce_out / "teareduce_quicklook.png",
            teareduce_out / "teareduce_quicklook.json",
        )

    for label, ext in [
        ("document_docx", ".docx"),
        ("document_docm", ".docm"),
        ("document_pptx", ".pptx"),
        ("document_pptm", ".pptm"),
        ("document_xlsx", ".xlsx"),
        ("document_xlsm", ".xlsm"),
        ("document_xls", ".xls"),
        ("document_xlsb", ".xlsb"),
        ("document_doc", ".doc"),
        ("document_pages", ".pages"),
        ("document_numbers", ".numbers"),
        ("document_key", ".key"),
        ("document_odt", ".odt"),
        ("document_ods", ".ods"),
        ("document_odp", ".odp"),
        ("document_epub", ".epub"),
        ("document_rtf", ".rtf"),
        ("document_reg", ".reg"),
        ("document_pref", ".pref"),
        ("document_jmars", ".jmars"),
    ]:
        sample = first(ext)
        if sample:
            report["feature_runs"][label] = run(
                ["python3", str(scripts_dir / "document_semantics.py"), str(sample)]
            )
    pages_sample = first(".pages")
    if pages_sample:
        iwork_out = feature_dir / "iwork_assets"
        iwork_out.mkdir(exist_ok=True)
        report["feature_runs"]["document_pages_assets"] = run(
            [
                "python3",
                str(scripts_dir / "document_semantics.py"),
                str(pages_sample),
                "--output-dir",
                str(iwork_out),
            ]
        )
        report["feature_runs"]["iwork_workbench_pages"] = run(
            [
                "python3",
                str(scripts_dir / "iwork_workbench.py"),
                str(pages_sample),
                "--output-dir",
                str(iwork_out / "pages_workbench"),
            ]
        )
        report["feature_runs"]["iwork_workbench_pages"] = require_outputs(
            report["feature_runs"]["iwork_workbench_pages"],
            iwork_out / "pages_workbench" / "summary.json",
            iwork_out / "pages_workbench" / "report.md",
            iwork_out / "pages_workbench" / "member_manifest.csv",
        )
        report["feature_runs"]["quicklook_bridge_pages"] = run(
            [
                "python3",
                str(scripts_dir / "quicklook_bridge.py"),
                str(pages_sample),
                "--output",
                str(iwork_out / "pages_preview.png"),
                "--summary-json",
                str(iwork_out / "pages_preview.json"),
            ]
        )
        report["feature_runs"]["iwork_roundtrip_convert_pages_md"] = run(
            [
                "python3",
                str(scripts_dir / "iwork_roundtrip.py"),
                "convert",
                str(pages_sample),
                str(iwork_out / "pages_semantic.md"),
                "--enable-ocr",
            ]
        )
    key_sample = first(".key")
    if key_sample:
        key_out = feature_dir / "iwork_assets"
        key_out.mkdir(exist_ok=True)
        report["feature_runs"]["iwork_workbench_key"] = run(
            [
                "python3",
                str(scripts_dir / "iwork_workbench.py"),
                str(key_sample),
                "--output-dir",
                str(key_out / "key_workbench"),
            ]
        )
        report["feature_runs"]["iwork_workbench_key"] = require_outputs(
            report["feature_runs"]["iwork_workbench_key"],
            key_out / "key_workbench" / "summary.json",
            key_out / "key_workbench" / "report.md",
            key_out / "key_workbench" / "member_manifest.csv",
        )
        report["feature_runs"]["iwork_roundtrip_convert_key_png"] = run(
            [
                "python3",
                str(scripts_dir / "iwork_roundtrip.py"),
                "convert",
                str(key_sample),
                str(key_out / "key_preview.png"),
                "--enable-ocr",
            ]
        )
    numbers_sample = first(".numbers")
    if numbers_sample:
        numbers_out = feature_dir / "iwork_numbers"
        numbers_out.mkdir(exist_ok=True)
        report["feature_runs"]["iwork_workbench_numbers"] = run(
            [
                "python3",
                str(scripts_dir / "iwork_workbench.py"),
                str(numbers_sample),
                "--output-dir",
                str(numbers_out / "numbers_workbench"),
                "--html-report",
                str(numbers_out / "numbers_workbench" / "report.html"),
            ]
        )
        report["feature_runs"]["iwork_roundtrip_export_numbers"] = run(
            [
                "python3",
                str(scripts_dir / "iwork_roundtrip.py"),
                "export",
                str(numbers_sample),
                "--output-dir",
                str(numbers_out / "exported"),
                "--fidelity-json",
                str(numbers_out / "exported" / "fidelity.json"),
            ]
        )
        report["feature_runs"]["iwork_roundtrip_edit_numbers"] = run(
            [
                "python3",
                str(scripts_dir / "iwork_roundtrip.py"),
                "set-cell",
                str(numbers_sample),
                str(numbers_out / "edited_copy.numbers"),
                "--cell",
                "B8",
                "--value",
                "12345",
            ]
        )
        report["feature_runs"]["quicklook_bridge_numbers"] = run(
            [
                "python3",
                str(scripts_dir / "quicklook_bridge.py"),
                str(numbers_sample),
                "--output",
                str(numbers_out / "numbers_preview.png"),
                "--timeout-sec",
                "20",
            ],
            timeout_sec=60,
        )

    tex_projects = []
    for root in args.roots:
        for path in Path(root).rglob("main.tex"):
            tex_projects.append(path)
    if tex_projects:
        latex_out = feature_dir / "latex"
        latex_out.mkdir(exist_ok=True)
        tex_main = tex_projects[0]
        report["feature_runs"]["latex_inspect"] = run(
            ["python3", str(scripts_dir / "latex_workbench.py"), "inspect", str(tex_main)]
        )
        report["feature_runs"]["latex_compile"] = run(
            [
                "python3",
                str(scripts_dir / "latex_workbench.py"),
                "compile",
                str(tex_main),
                "--output-dir",
                str(latex_out),
                "--manifest-json",
                str(latex_out / "compile_manifest.json"),
            ]
        )
        report["feature_runs"]["latex_compile"] = require_outputs(
            report["feature_runs"]["latex_compile"],
            latex_out / "main.pdf",
            latex_out / "compile_manifest.json",
        )
        report["feature_runs"]["latex_scaffold"] = run(
            [
                "python3",
                str(scripts_dir / "latex_workbench.py"),
                "scaffold",
                str(latex_out / "scaffold_project"),
                "--title",
                "Validation Report",
                "--kind",
                "article",
            ]
        )

    ccd_pair = find_master_bias_and_science(args.roots)
    if ccd_pair is not None:
        ccd_out = feature_dir / "ccd"
        ccd_out.mkdir(exist_ok=True)
        master_bias, science = ccd_pair
        report["feature_runs"]["reduce_ccd_batch"] = run(
            [
                "python3",
                str(scripts_dir / "reduce_ccd_batch.py"),
                "--master-bias",
                str(master_bias),
                "--science",
                str(science),
                "--output-dir",
                str(ccd_out),
                "--audit-json",
                str(ccd_out / "audit.json"),
                "--manifest-json",
                str(ccd_out / "manifest.json"),
            ]
        )
        report["feature_runs"]["reduce_ccd_batch"] = require_outputs(
            report["feature_runs"]["reduce_ccd_batch"],
            ccd_out / f"reduced_{Path(science).name}",
            ccd_out / "audit.json",
            ccd_out / "manifest.json",
        )

    container_out = feature_dir / "containers"
    create_synthetic_containers(container_out)
    for label, filename in [
        ("container_npy", "array.npy"),
        ("container_npz", "bundle.npz"),
        ("container_h5", "sample.h5"),
        ("container_mat", "sample.mat"),
        ("container_sqlite", "sample.sqlite"),
        ("container_nc", "sample.nc"),
        ("container_zip", "sample.zip"),
        ("container_tar_gz", "sample.tar.gz"),
        ("container_gzip", "sample.txt.gz"),
        ("container_bz2", "sample.txt.bz2"),
        ("container_xz", "sample.txt.xz"),
        ("container_zarr", "sample.zarr"),
        ("container_geojson", "sample.geojson"),
        ("container_corrupt_gzip", "corrupt.txt.gz"),
    ]:
        report["feature_runs"][label] = run(
            ["python3", str(scripts_dir / "inspect_data_container.py"), str(container_out / filename)]
        )
    for label, filename in [("container_parquet", "sample.parquet"), ("container_feather", "sample.feather"), ("container_arrow", "sample.arrow")]:
        candidate = synthetic_table_out / filename
        if candidate.exists():
            report["feature_runs"][label] = run(
                ["python3", str(scripts_dir / "inspect_data_container.py"), str(candidate)]
            )

    spectrum_out = feature_dir / "spectrum"
    synthetic_spectrum = spectrum_out / "synthetic_spectrum.txt"
    create_synthetic_spectrum(synthetic_spectrum)
    report["feature_runs"]["spectral_workbench"] = run(
        [
            "python3",
            str(scripts_dir / "spectral_workbench.py"),
            str(synthetic_spectrum),
            "--line-window",
            "6562.8",
            "20",
            "--normalize",
            "--plot",
            str(spectrum_out / "spectrum.png"),
            "--output-table",
            str(spectrum_out / "spectrum.ecsv"),
            "--summary-json",
            str(spectrum_out / "spectrum.json"),
        ]
    )
    synthetic_spectrum_2d = spectrum_out / "synthetic_spectrum_2d.fits"
    create_synthetic_spectrum_2d(synthetic_spectrum_2d)
    calibration_table = spectrum_out / "calibration_table.txt"
    create_synthetic_calibration_table(calibration_table)
    report["feature_runs"]["spectral_workbench_2d"] = run(
        [
            "python3",
            str(scripts_dir / "spectral_workbench.py"),
            str(synthetic_spectrum_2d),
            "--line-window",
            "6562.8",
            "20",
            "--normalize",
            "--optimal-extraction",
            "--order-count",
            "2",
            "--calibration-table",
            str(calibration_table),
            "--extract-all-orders-dir",
            str(spectrum_out / "orders"),
            "--trace-plot",
            str(spectrum_out / "spectrum_2d_trace.png"),
            "--plot",
            str(spectrum_out / "spectrum_2d.png"),
            "--html-report",
            str(spectrum_out / "spectrum_2d.html"),
            "--summary-json",
            str(spectrum_out / "spectrum_2d.json"),
        ]
    )
    pdf_sample = first(".pdf")
    if pdf_sample:
        pdf_out = feature_dir / "pdf_recovery"
        pdf_out.mkdir(exist_ok=True)
        report["feature_runs"]["pdf_recover_extract"] = run(
            [
                "python3",
                str(scripts_dir / "pdf_recover_extract.py"),
                str(pdf_sample),
                "--output-dir",
                str(pdf_out),
                "--summary-json",
                str(pdf_out / "summary.json"),
            ]
        )

    pipeline_out = feature_dir / "pipeline"
    report["feature_runs"]["pipeline_scaffold"] = run(
        [
            "python3",
            str(scripts_dir / "pipeline_scaffold.py"),
            str(pipeline_out / "project"),
            "--domain",
            "astronomy",
            "--language",
            "bilingual",
            "--profile",
            "spectroscopy",
        ]
    )
    deliverable_out = feature_dir / "deliverables"
    report["feature_runs"]["deliverable_scaffold"] = run(
        [
            "python3",
            str(scripts_dir / "deliverable_factory.py"),
            "scaffold",
            str(deliverable_out / "proposal"),
            "--kind",
            "proposal",
            "--format",
            "markdown",
            "--title",
            "Validation Proposal",
            "--language",
            "bilingual",
        ]
    )
    report["feature_runs"]["deliverable_presentation_scaffold"] = run(
        [
            "python3",
            str(scripts_dir / "deliverable_factory.py"),
            "scaffold",
            str(deliverable_out / "presentation"),
            "--kind",
            "presentation",
            "--format",
            "markdown",
            "--title",
            "Validation Presentation Plan",
            "--language",
            "english",
        ]
    )
    report["feature_runs"]["deliverable_status_report_scaffold"] = run(
        [
            "python3",
            str(scripts_dir / "deliverable_factory.py"),
            "scaffold",
            str(deliverable_out / "status_report"),
            "--kind",
            "status-report",
            "--format",
            "markdown",
            "--title",
            "Validation Status Report",
            "--language",
            "english",
        ]
    )
    report["feature_runs"]["deliverable_project_brief_scaffold"] = run(
        [
            "python3",
            str(scripts_dir / "deliverable_factory.py"),
            "scaffold",
            str(deliverable_out / "project_brief"),
            "--kind",
            "project-brief",
            "--format",
            "markdown",
            "--title",
            "Validation Project Brief",
            "--language",
            "bilingual",
        ]
    )
    report["feature_runs"]["deliverable_qa_summary_scaffold"] = run(
        [
            "python3",
            str(scripts_dir / "deliverable_factory.py"),
            "scaffold",
            str(deliverable_out / "qa_summary"),
            "--kind",
            "qa-summary",
            "--format",
            "markdown",
            "--title",
            "Validation QA Summary",
            "--language",
            "english",
        ]
    )
    report["feature_runs"]["deliverable_package"] = run(
        [
            "python3",
            str(scripts_dir / "deliverable_factory.py"),
            "package",
            str(deliverable_out / "deliverables.zip"),
            str(synthetic_table_out / "duckdb_join.csv"),
            str(office_out / "sample_edited.md"),
            "--title",
            "Validation Bundle",
        ]
    )
    report["feature_runs"]["cross_domain_data_workbench"] = run(
        [
            "python3",
            str(scripts_dir / "cross_domain_data_workbench.py"),
            str(synthetic_table_out),
            "--sql",
            "SELECT * FROM source0 LIMIT 2",
            "--query-output",
            str(feature_dir / "cross_domain" / "query_result.csv"),
            "--output-dir",
            str(feature_dir / "cross_domain"),
            "--summary-json",
            str(feature_dir / "cross_domain" / "summary.json"),
            "--manifest-json",
            str(feature_dir / "cross_domain" / "manifest.json"),
        ]
    )
    report["feature_runs"]["cross_domain_data_workbench"] = require_outputs(
        report["feature_runs"]["cross_domain_data_workbench"],
        feature_dir / "cross_domain" / "inventory.csv",
        feature_dir / "cross_domain" / "report.md",
        feature_dir / "cross_domain" / "summary.json",
        feature_dir / "cross_domain" / "manifest.json",
        feature_dir / "cross_domain" / "query_result.csv",
    )
    intake_extra_note = feature_dir / "document_intake_note.txt"
    intake_extra_note.write_text(
        "Validation note for mixed document intake outside science-specific workflows.\n",
        encoding="utf-8",
    )
    report["feature_runs"]["document_intake_workbench"] = run(
        [
            "python3",
            str(scripts_dir / "document_intake_workbench.py"),
            str(office_out),
            str(intake_extra_note),
            "--output-dir",
            str(feature_dir / "document_intake"),
            "--summary-json",
            str(feature_dir / "document_intake" / "summary.json"),
            "--manifest-json",
            str(feature_dir / "document_intake" / "manifest.json"),
        ]
    )
    report["feature_runs"]["document_intake_workbench"] = require_outputs(
        report["feature_runs"]["document_intake_workbench"],
        feature_dir / "document_intake" / "inventory.csv",
        feature_dir / "document_intake" / "report.md",
        feature_dir / "document_intake" / "summary.json",
        feature_dir / "document_intake" / "manifest.json",
    )
    report["feature_runs"]["portable_smoke_test"] = run(
        [
            "python3",
            str(scripts_dir / "portable_smoke_test.py"),
            "--output-dir",
            str(feature_dir / "portable_smoke"),
            "--examples-dir",
            str(Path(__file__).resolve().parent.parent / "examples"),
            "--summary-json",
            str(feature_dir / "portable_smoke" / "summary.json"),
            "--manifest-json",
            str(feature_dir / "portable_smoke" / "manifest.json"),
        ]
    )
    report["feature_runs"]["portable_smoke_test"] = require_outputs(
        report["feature_runs"]["portable_smoke_test"],
        feature_dir / "portable_smoke" / "summary.json",
        feature_dir / "portable_smoke" / "manifest.json",
        feature_dir / "portable_smoke" / "inspect_fits.json",
        feature_dir / "portable_smoke" / "profile_table.json",
        feature_dir / "portable_smoke" / "portable_smoke_notebook.ipynb",
    )

    return finalize_validation(args, output_dir, report)


def main():
    args = parse_args()
    try:
        # Path safety must run before an optional runtime re-exec or any report
        # setup.  The bootstrap body imports the mother-owned datanalysis_env
        # compatibility script, so expose that sibling root explicitly.
        ensure_validation_paths(args)
        mother_scripts = Path(__file__).resolve().parents[2] / "scientific-data-analysis" / "scripts"
        if mother_scripts.is_dir() and str(mother_scripts) not in sys.path:
            sys.path.insert(0, str(mother_scripts))
        from _internal.datanalysis_bootstrap import ensure_datanalysis_runtime

        ensure_datanalysis_runtime("validate_skill_samples", strict=False)
        configure_runtime("validate_skill_samples")
        exit_code = run_validation(args)
    except Exception as exc:
        blocked = isinstance(exc, PathSafetyError)
        payload = build_failure_payload(args, exc, blocked=blocked)
        summary_target = None if blocked else args.summary_json
        try:
            emit_payload_safely(payload, summary_target, full_stdout=True)
        except Exception as emit_exc:
            payload["notes"].append(str(emit_exc))
            payload["results"]["summary_json_error"] = str(emit_exc)
            emit_payload_safely(payload, None, full_stdout=True)
        raise SystemExit(2 if blocked else 1) from None
    raise SystemExit(exit_code)


if __name__ == "__main__":
    main()
