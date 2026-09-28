#!/usr/bin/env python3
"""Regression checks for exact coursework figure handoff guidance."""

from __future__ import annotations

import base64
import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path


ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "scripts"
REFERENCES = ROOT / "references"


def module_available(module: str) -> bool:
    completed = subprocess.run(
        [sys.executable, "-c", f"import {module}"],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    return completed.returncode == 0


def datanalysis_python() -> str:
    completed = subprocess.run(
        [sys.executable, str(SCRIPTS / "datanalysis_env.py"), "locate"],
        cwd=ROOT,
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    if completed.returncode != 0:
        raise AssertionError(f"Could not locate datanalysis runtime:\n{completed.stderr}")
    payload = json.loads(completed.stdout)
    python_path = payload.get("selected", {}).get("python")
    if not python_path:
        raise AssertionError(f"datanalysis runtime did not report a Python executable: {payload}")
    return str(Path(python_path).expanduser())


def runtime_python() -> str:
    return sys.executable if module_available("pptx") else datanalysis_python()


def run_cmd(args: list[str], *, tmp: Path) -> subprocess.CompletedProcess:
    env = os.environ.copy()
    env["PYTHONDONTWRITEBYTECODE"] = "1"
    env["MPLCONFIGDIR"] = str(tmp / "mplconfig")
    completed = subprocess.run(
        args,
        cwd=ROOT,
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        env=env,
    )
    if completed.returncode != 0:
        raise AssertionError(
            f"Command failed: {' '.join(args)}\nSTDOUT:\n{completed.stdout}\nSTDERR:\n{completed.stderr}"
        )
    return completed


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def check_exact_visual_equivalent_references() -> None:
    equivalent = (REFERENCES / "coursework-figure-equivalents.md").read_text(encoding="utf-8")
    presentation = (REFERENCES / "scientific-presentation-handoff.md").read_text(encoding="utf-8")
    comparison = (REFERENCES / "coursework-notebook-comparison.md").read_text(encoding="utf-8")
    fidelity = (REFERENCES / "coursework-notebook-fidelity.md").read_text(encoding="utf-8")

    required_equivalent = [
        "exact visual equivalent",
        "comparar_fits",
        "detected_sources",
        "sdss_calibrators",
        "final_rgb",
        "diagnostic_crop",
        "viridis",
        "RGB",
        "04_A",
        "If The User Corrects The Figure",
    ]
    missing = [term for term in required_equivalent if term not in equivalent]
    if missing:
        raise AssertionError(f"coursework-figure-equivalents.md is missing terms: {missing}")

    if "coursework-figure-equivalents.md" not in presentation:
        raise AssertionError("scientific-presentation-handoff.md does not route exact figure work to the new reference.")
    if "coursework-figure-equivalents.md" not in comparison:
        raise AssertionError("coursework-notebook-comparison.md does not route candidate equivalence to the new reference.")
    if "coursework-figure-equivalents.md" not in fidelity:
        raise AssertionError("coursework-notebook-fidelity.md does not route rebuilt figures to the new reference.")


def create_science_deck(tmp: Path, python: str) -> Path:
    image_path = tmp / "pixel.png"
    image_path.write_bytes(
        base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAIAAACQd1PeAAAADUlEQVR4nGP4z8DwHwAFgwJ/lc0rWQAAAABJRU5ErkJggg=="
        )
    )
    deck = tmp / "practice4_deck.pptx"
    script = tmp / "make_deck.py"
    script.write_text(
        """
import sys
from pptx import Presentation
from pptx.util import Inches

deck, image_path = sys.argv[1], sys.argv[2]
prs = Presentation()
for index in range(1, 5):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = f"Slide {index}"
    if index == 4:
        slide.shapes.add_picture(image_path, Inches(1.0), Inches(1.4), width=Inches(4.0))
prs.save(deck)
""".lstrip(),
        encoding="utf-8",
    )
    run_cmd([python, str(script), str(deck), str(image_path)], tmp=tmp)
    return deck


def check_presentation_manifest_exact_fields(tmp: Path) -> None:
    python = runtime_python()
    deck = create_science_deck(tmp, python)
    output_dir = tmp / "style_audit"
    summary = output_dir / "summary.json"
    run_cmd(
        [
            python,
            str(SCRIPTS / "presentation_workbench.py"),
            "existing-deck-style-audit",
            str(deck),
            str(output_dir),
            "--reference-slides",
            "4",
            "--summary-json",
            str(summary),
        ],
        tmp=tmp,
    )

    manifest = load_json(output_dir / "scientific_asset_manifest.json")
    required_top = [
        "slide_position_policy",
        "visual_mode_allowed",
        "scientific_role_allowed",
        "coursework_exact_equivalent_fields",
    ]
    missing_top = [field for field in required_top if field not in manifest]
    if missing_top:
        raise AssertionError(f"Scientific asset manifest is missing top-level fields: {missing_top}")

    expected_exact_fields = {
        "final_figure_name",
        "slide_position_code",
        "source_notebook_cell_or_function",
        "source_fits_path",
        "visual_mode",
        "extent",
        "scale_criteria",
        "generation_command",
        "equivalent_target",
    }
    declared = set(manifest["coursework_exact_equivalent_fields"])
    if not expected_exact_fields <= declared:
        raise AssertionError(f"Exact-equivalent fields missing from declaration: {expected_exact_fields - declared}")

    assets = manifest.get("assets") or []
    if len(assets) != 1:
        raise AssertionError(f"Expected one picture asset in synthetic deck, got {len(assets)}: {assets}")
    asset = assets[0]
    if asset.get("slide_position_code") != "04_A":
        raise AssertionError(f"Expected slide-position code 04_A, got {asset.get('slide_position_code')}")
    missing_asset_fields = [field for field in expected_exact_fields if field not in asset]
    if missing_asset_fields:
        raise AssertionError(f"Asset is missing exact-equivalent fields: {missing_asset_fields}")

    payload = load_json(summary)
    status = payload["results"]["scientific_asset_manifest_status"]
    if status.get("asset_count") != 1 or status.get("incomplete_asset_count") != 1:
        raise AssertionError(f"Unexpected manifest status for incomplete synthetic asset: {status}")


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="sda_practice4_visual_handoff_") as raw_tmp:
        tmp = Path(raw_tmp)
        check_exact_visual_equivalent_references()
        print("PASS check_exact_visual_equivalent_references")
        check_presentation_manifest_exact_fields(tmp)
        print("PASS check_presentation_manifest_exact_fields")
    print("All Practice 4 visual handoff regressions passed.")


if __name__ == "__main__":
    main()
