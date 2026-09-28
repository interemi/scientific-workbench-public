#!/usr/bin/env python3
"""Regression tests for office_roundtrip.py docx-styled-replace."""

from __future__ import annotations

import hashlib
import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path

from _internal.public_contract import validate_standard_envelope


ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "scripts"


def run_cmd(args: list[str], *, cwd: Path = ROOT, expect_ok: bool = True) -> subprocess.CompletedProcess:
    env = os.environ.copy()
    env["PYTHONDONTWRITEBYTECODE"] = "1"
    completed = subprocess.run(
        args,
        cwd=cwd,
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        env=env,
    )
    if expect_ok and completed.returncode != 0:
        raise AssertionError(
            f"Command failed: {' '.join(args)}\nSTDOUT:\n{completed.stdout}\nSTDERR:\n{completed.stderr}"
        )
    if not expect_ok and completed.returncode == 0:
        raise AssertionError(f"Command unexpectedly passed: {' '.join(args)}\nSTDOUT:\n{completed.stdout}")
    return completed


def module_available(module: str) -> bool:
    completed = subprocess.run(
        [sys.executable, "-c", f"import {module}"],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    return completed.returncode == 0


def datanalysis_python() -> str:
    completed = run_cmd([sys.executable, str(SCRIPTS / "datanalysis_env.py"), "locate"])
    payload = json.loads(completed.stdout)
    python_path = payload.get("selected", {}).get("python")
    if not python_path:
        raise AssertionError(f"datanalysis runtime did not report a Python executable: {payload}")
    return str(Path(python_path).expanduser())


def runtime_python() -> str:
    return sys.executable if module_available("docx") else datanalysis_python()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def load_stdout_json(completed: subprocess.CompletedProcess) -> dict:
    try:
        return json.loads(completed.stdout)
    except Exception as exc:
        raise AssertionError(f"Expected JSON stdout, got:\n{completed.stdout}\nSTDERR:\n{completed.stderr}") from exc


def create_docx(tmp: Path, python: str, *, mode: str) -> Path:
    script = tmp / f"make_{mode}.py"
    path = tmp / f"{mode}.docx"
    if mode == "marked":
        body = """
from docx import Document
import sys

doc = Document()
p = doc.add_paragraph()
p.add_run("before ")
r = p.add_run("MARKED")
r.bold = True
r.italic = True
r.underline = True
p.add_run(" after")
doc.save(sys.argv[1])
"""
    elif mode == "mixed":
        body = """
from docx import Document
import sys

doc = Document()
p = doc.add_paragraph()
r = p.add_run("TARGET body")
r.bold = True
r.italic = True
r.underline = True
p.add_run(" plain TARGET should stay")
t = doc.add_table(rows=1, cols=1)
r = t.cell(0, 0).paragraphs[0].add_run("TARGET table")
r.bold = True
r.italic = True
r.underline = True
doc.save(sys.argv[1])
"""
    elif mode == "mismatch":
        body = """
from docx import Document
import sys

doc = Document()
p = doc.add_paragraph()
r = p.add_run("TARGET bold-only")
r.bold = True
p.add_run(" plain TARGET")
doc.save(sys.argv[1])
"""
    else:
        raise ValueError(mode)
    script.write_text(body.lstrip(), encoding="utf-8")
    run_cmd([python, str(script), str(path)], cwd=tmp)
    return path


def assert_standard(payload: dict, expected_status: str) -> None:
    issues = validate_standard_envelope(payload)
    if issues:
        raise AssertionError(f"Invalid standard envelope: {issues}\n{payload}")
    if payload["status"] != expected_status:
        raise AssertionError(f"Expected status={expected_status}, got {payload['status']}: {payload}")
    for key in ("replacements", "readback_replacement_match_count"):
        if key not in payload:
            raise AssertionError(f"Legacy key {key!r} is missing from compatibility payload.")


def check_success_and_backup(tmp: Path, python: str) -> None:
    source = create_docx(tmp, python, mode="mixed")
    output = tmp / "edited.docx"
    output.write_bytes(b"previous output")
    summary = tmp / "replace.json"
    manifest = tmp / "manifest.json"
    run_cmd(
        [
            python,
            str(SCRIPTS / "office_roundtrip.py"),
            "docx-styled-replace",
            str(source),
            str(output),
            "--find",
            "TARGET",
            "--replace",
            "REPLACED",
            "--require-bold",
            "--require-italic",
            "--require-underline",
            "--summary-json",
            str(summary),
            "--manifest-json",
            str(manifest),
        ]
    )
    payload = load_json(summary)
    assert_standard(payload, "ok")
    if payload["replacements"] != 2 or payload["readback_replacement_occurrence_count"] != 2:
        raise AssertionError(f"Styled replacement did not verify both replacements: {payload}")
    if not output.exists() or not payload.get("backup_existing_output") or not manifest.exists():
        raise AssertionError(f"Expected edited output, backup, and manifest: {payload}")


def check_zero_replacement_warning(tmp: Path, python: str) -> None:
    source = create_docx(tmp, python, mode="mismatch")
    output = tmp / "zero_replacement.docx"
    summary = tmp / "zero_replacement.json"
    run_cmd(
        [
            python,
            str(SCRIPTS / "office_roundtrip.py"),
            "docx-styled-replace",
            str(source),
            str(output),
            "--find",
            "TARGET",
            "--replace",
            "REPLACED",
            "--require-bold",
            "--require-italic",
            "--require-underline",
            "--summary-json",
            str(summary),
        ]
    )
    payload = load_json(summary)
    assert_standard(payload, "warning")
    if payload["replacements"] != 0 or not payload["qa"]["findings"]:
        raise AssertionError(f"Zero replacement should be a warning with findings: {payload}")


def check_same_input_output_blocked(tmp: Path, python: str) -> None:
    source = create_docx(tmp, python, mode="marked")
    before = sha256(source)
    summary = tmp / "same_path.json"
    completed = run_cmd(
        [
            python,
            str(SCRIPTS / "office_roundtrip.py"),
            "docx-styled-replace",
            str(source),
            str(source),
            "--find",
            "MARKED",
            "--replace",
            "REVISED",
            "--summary-json",
            str(summary),
        ],
        expect_ok=False,
    )
    if "Traceback" in completed.stderr or "Traceback" in completed.stdout:
        raise AssertionError(f"Same-path block leaked a traceback:\nSTDOUT:\n{completed.stdout}\nSTDERR:\n{completed.stderr}")
    if sha256(source) != before:
        raise AssertionError("Same-path replacement modified the input DOCX.")
    payload = load_json(summary)
    assert_standard(payload, "fail")


def check_corrupt_docx_fails_cleanly(tmp: Path, python: str) -> None:
    corrupt = tmp / "corrupt.docx"
    corrupt.write_bytes(b"not a zip-backed docx")
    output = tmp / "corrupt_edited.docx"
    summary = tmp / "corrupt.json"
    completed = run_cmd(
        [
            python,
            str(SCRIPTS / "office_roundtrip.py"),
            "docx-styled-replace",
            str(corrupt),
            str(output),
            "--find",
            "x",
            "--replace",
            "y",
            "--summary-json",
            str(summary),
        ],
        expect_ok=False,
    )
    if "Traceback" in completed.stderr or "Traceback" in completed.stdout:
        raise AssertionError(f"Corrupt DOCX leaked a traceback:\nSTDOUT:\n{completed.stdout}\nSTDERR:\n{completed.stderr}")
    payload = load_json(summary)
    assert_standard(payload, "fail")
    if output.exists():
        raise AssertionError("Corrupt DOCX should not produce an edited output.")


def check_invalid_summary_target_fails_cleanly(tmp: Path, python: str) -> None:
    source = create_docx(tmp, python, mode="marked")
    output = tmp / "invalid_summary_edited.docx"
    summary_dir = tmp / "summary_is_dir"
    summary_dir.mkdir()
    completed = run_cmd(
        [
            python,
            str(SCRIPTS / "office_roundtrip.py"),
            "docx-styled-replace",
            str(source),
            str(output),
            "--find",
            "MARKED",
            "--replace",
            "REVISED",
            "--summary-json",
            str(summary_dir),
        ],
        expect_ok=False,
    )
    if "Traceback" in completed.stderr or "Traceback" in completed.stdout:
        raise AssertionError(f"Invalid summary target leaked a traceback:\nSTDOUT:\n{completed.stdout}\nSTDERR:\n{completed.stderr}")
    payload = load_stdout_json(completed)
    assert_standard(payload, "fail")


def main() -> None:
    python = runtime_python()
    with tempfile.TemporaryDirectory(prefix="sda_office_styled_replace_") as raw_tmp:
        tmp = Path(raw_tmp)
        check_success_and_backup(tmp, python)
        print("PASS check_success_and_backup")
        check_zero_replacement_warning(tmp, python)
        print("PASS check_zero_replacement_warning")
        check_same_input_output_blocked(tmp, python)
        print("PASS check_same_input_output_blocked")
        check_corrupt_docx_fails_cleanly(tmp, python)
        print("PASS check_corrupt_docx_fails_cleanly")
        check_invalid_summary_target_fails_cleanly(tmp, python)
        print("PASS check_invalid_summary_target_fails_cleanly")
    print("All office_roundtrip docx-styled-replace regressions passed.")


if __name__ == "__main__":
    main()
