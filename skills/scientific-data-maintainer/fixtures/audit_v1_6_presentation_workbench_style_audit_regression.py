#!/usr/bin/env python3
"""Regression checks for presentation_workbench.py existing-deck-style-audit."""

from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "presentation_workbench.py"
TMP = Path("/tmp/scientific-data-analysis-v1-6-presentation-style-audit-regression")


def run(cmd: list[str]) -> subprocess.CompletedProcess[str]:
    return subprocess.run(cmd, cwd=ROOT, text=True, encoding="utf-8", errors="replace", capture_output=True, check=False)


def assert_no_traceback(completed: subprocess.CompletedProcess[str]) -> None:
    combined = f"{completed.stdout}\n{completed.stderr}"
    assert "Traceback (most recent call last)" not in combined, combined


def load_payload(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def output_file_count(path: Path) -> int:
    if not path.exists() or not path.is_dir():
        return 0
    return len([item for item in path.rglob("*") if item.is_file()])


def make_image(path: Path, label: str) -> None:
    image = Image.new("RGB", (900, 520), (246, 248, 250))
    draw = ImageDraw.Draw(image)
    draw.rectangle([40, 40, 860, 480], outline=(35, 88, 145), width=6)
    draw.line([90, 400, 810, 130], fill=(225, 170, 62), width=8)
    draw.text((70, 65), label, fill=(20, 20, 20))
    image.save(path)


def add_text(slide, text: str, left: float, top: float, width: float, height: float, size: int) -> None:
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    shape.text_frame.text = text
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(size)


def make_deck(path: Path, first_image: Path, second_image: Path | None = None, *, flattened: bool = False) -> None:
    presentation = Presentation()
    blank = presentation.slide_layouts[6]
    slide = presentation.slides.add_slide(blank)
    if flattened:
        slide.shapes.add_picture(str(first_image), 0, 0, width=presentation.slide_width, height=presentation.slide_height)
    else:
        add_text(slide, "Style Source", 0.5, 0.25, 8.5, 0.55, 28)
        slide.shapes.add_picture(str(first_image), Inches(1.0), Inches(1.15), width=Inches(7.4))
        add_text(slide, "Figure 1. Editable caption.", 0.8, 6.35, 8.0, 0.35, 12)
    if second_image is not None:
        slide = presentation.slides.add_slide(blank)
        add_text(slide, "Comparison", 0.5, 0.25, 8.5, 0.55, 28)
        slide.shapes.add_picture(str(first_image), Inches(0.65), Inches(1.2), width=Inches(4.05))
        slide.shapes.add_picture(str(second_image), Inches(5.0), Inches(1.2), width=Inches(4.05))
        add_text(slide, "Figures A and B share a visual grid.", 0.8, 6.35, 8.0, 0.35, 12)
    presentation.save(path)


def assert_payload(path: Path, *, status: str, qa_status: str) -> dict:
    payload = load_payload(path)
    assert payload["status"] == status, payload
    assert payload["qa"]["status"] == qa_status, payload
    return payload


def main() -> int:
    if TMP.exists():
        shutil.rmtree(TMP)
    TMP.mkdir(parents=True)
    img_a = TMP / "figure_a.png"
    img_b = TMP / "figure_b.png"
    make_image(img_a, "style figure A")
    make_image(img_b, "style figure B")
    happy = TMP / "happy.pptx"
    flattened = TMP / "flattened.pptx"
    corrupt = TMP / "corrupt.pptx"
    make_deck(happy, img_a, img_b)
    make_deck(flattened, img_a, flattened=True)
    corrupt.write_bytes(b"not a real presentation\n")

    happy_summary = TMP / "happy_summary.json"
    happy_manifest = TMP / "happy_manifest.json"
    happy_out = TMP / "happy_out"
    completed = run(
        [
            sys.executable,
            str(SCRIPT),
            "existing-deck-style-audit",
            str(happy),
            str(happy_out),
            "--summary-json",
            str(happy_summary),
            "--manifest-json",
            str(happy_manifest),
            "--reference-slides",
            "1-2",
        ]
    )
    assert completed.returncode == 0, completed.stdout + completed.stderr
    assert_no_traceback(completed)
    assert_payload(happy_summary, status="warning", qa_status="warning")
    for name in ("scientific_asset_manifest.json", "presentation_constraints.json", "slide_content.md", "speaker_script.md", "poster_text.md", "new_slide_templates.json"):
        assert (happy_out / name).exists(), name
    assert happy_manifest.exists(), "manifest should be written for successful audits"

    flattened_summary = TMP / "flattened_summary.json"
    flattened_out = TMP / "flattened_out"
    completed = run(
        [
            sys.executable,
            str(SCRIPT),
            "existing-deck-style-audit",
            str(flattened),
            str(flattened_out),
            "--summary-json",
            str(flattened_summary),
        ]
    )
    assert completed.returncode == 0, completed.stdout + completed.stderr
    assert_no_traceback(completed)
    assert_payload(flattened_summary, status="warning", qa_status="warning")

    corrupt_summary = TMP / "corrupt_summary.json"
    corrupt_manifest = TMP / "corrupt_manifest.json"
    corrupt_out = TMP / "corrupt_out"
    completed = run(
        [
            sys.executable,
            str(SCRIPT),
            "existing-deck-style-audit",
            str(corrupt),
            str(corrupt_out),
            "--summary-json",
            str(corrupt_summary),
            "--manifest-json",
            str(corrupt_manifest),
        ]
    )
    assert completed.returncode == 2, completed.stdout + completed.stderr
    assert_no_traceback(completed)
    assert_payload(corrupt_summary, status="blocked", qa_status="blocked")
    assert output_file_count(corrupt_out) == 0, "blocked corrupt audit must not emit style artifacts"
    assert not corrupt_manifest.exists(), "blocked corrupt audit must not emit manifest"

    bad_selector_summary = TMP / "bad_selector_summary.json"
    bad_selector_out = TMP / "bad_selector_out"
    completed = run(
        [
            sys.executable,
            str(SCRIPT),
            "existing-deck-style-audit",
            str(happy),
            str(bad_selector_out),
            "--reference-slides",
            "1,bad",
            "--summary-json",
            str(bad_selector_summary),
        ]
    )
    assert completed.returncode == 2, completed.stdout + completed.stderr
    assert_no_traceback(completed)
    assert_payload(bad_selector_summary, status="blocked", qa_status="blocked")
    assert output_file_count(bad_selector_out) == 0, "blocked selector audit must not emit style artifacts"

    output_file = TMP / "not_a_directory"
    output_file.write_text("conflict\n", encoding="utf-8")
    output_target_summary = TMP / "output_target_summary.json"
    completed = run(
        [
            sys.executable,
            str(SCRIPT),
            "existing-deck-style-audit",
            str(happy),
            str(output_file),
            "--summary-json",
            str(output_target_summary),
        ]
    )
    assert completed.returncode == 2, completed.stdout + completed.stderr
    assert_no_traceback(completed)
    assert_payload(output_target_summary, status="blocked", qa_status="blocked")

    print("presentation_workbench existing-deck-style-audit regression: PASS")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
