#!/usr/bin/env python3
"""Regression tests for office_roundtrip.py docx-style-inventory."""

from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path

from _internal.public_contract import validate_standard_envelope


ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "scripts"


def run_cmd(args: list[str], *, cwd: Path = ROOT, expect_ok: bool = True) -> subprocess.CompletedProcess:
    env = os.environ.copy()
    env["PYTHONDONTWRITEBYTECODE"] = "1"
    completed = subprocess.run(
        args,
        cwd=cwd,
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        env=env,
    )
    if expect_ok and completed.returncode != 0:
        raise AssertionError(
            f"Command failed: {' '.join(args)}\nSTDOUT:\n{completed.stdout}\nSTDERR:\n{completed.stderr}"
        )
    if not expect_ok and completed.returncode == 0:
        raise AssertionError(f"Command unexpectedly passed: {' '.join(args)}\nSTDOUT:\n{completed.stdout}")
    return completed


def module_available(module: str) -> bool:
    completed = subprocess.run(
        [sys.executable, "-c", f"import {module}"],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    return completed.returncode == 0


def datanalysis_python() -> str:
    completed = run_cmd([sys.executable, str(SCRIPTS / "datanalysis_env.py"), "locate"])
    payload = json.loads(completed.stdout)
    python_path = payload.get("selected", {}).get("python")
    if not python_path:
        raise AssertionError(f"datanalysis runtime did not report a Python executable: {payload}")
    return str(Path(python_path).expanduser())


def runtime_python() -> str:
    return sys.executable if module_available("docx") else datanalysis_python()


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def load_stdout_json(completed: subprocess.CompletedProcess) -> dict:
    try:
        return json.loads(completed.stdout)
    except Exception as exc:
        raise AssertionError(f"Expected JSON stdout, got:\n{completed.stdout}\nSTDERR:\n{completed.stderr}") from exc


def create_styled_docx(tmp: Path, python: str, *, with_marked: bool = True) -> Path:
    script = tmp / ("make_marked_docx.py" if with_marked else "make_plain_docx.py")
    docx_path = tmp / ("marked.docx" if with_marked else "plain.docx")
    if with_marked:
        body = """
from docx import Document
import sys

doc = Document()
doc.add_paragraph("Plain context.")
paragraph = doc.add_paragraph()
paragraph.add_run("before ")
run = paragraph.add_run("MARKED")
run.bold = True
run.italic = True
run.underline = True
paragraph.add_run(" after")
table = doc.add_table(rows=1, cols=1)
cell_run = table.cell(0, 0).paragraphs[0].add_run("TABLE_MARKED")
cell_run.bold = True
cell_run.italic = True
cell_run.underline = True
doc.save(sys.argv[1])
"""
    else:
        body = """
from docx import Document
import sys

doc = Document()
doc.add_paragraph("Plain context without explicit run-level markers.")
doc.add_table(rows=1, cols=1).cell(0, 0).text = "Plain table cell"
doc.save(sys.argv[1])
"""
    script.write_text(body.lstrip(), encoding="utf-8")
    run_cmd([python, str(script), str(docx_path)], cwd=tmp)
    return docx_path


def assert_standard(payload: dict, expected_status: str) -> None:
    issues = validate_standard_envelope(payload)
    if issues:
        raise AssertionError(f"Invalid standard envelope: {issues}\n{payload}")
    if payload["status"] != expected_status:
        raise AssertionError(f"Expected status={expected_status}, got {payload['status']}: {payload}")
    if "match_count" not in payload:
        raise AssertionError("Legacy match_count key is missing from compatibility payload.")


def check_marked_inventory(tmp: Path, python: str) -> None:
    source = create_styled_docx(tmp, python, with_marked=True)
    summary = tmp / "marked_inventory.json"
    csv_path = tmp / "marked_inventory.csv"
    report = tmp / "marked_inventory.md"
    completed = run_cmd(
        [
            python,
            str(SCRIPTS / "office_roundtrip.py"),
            "docx-style-inventory",
            str(source),
            "--require-bold",
            "--require-italic",
            "--require-underline",
            "--output-csv",
            str(csv_path),
            "--report-md",
            str(report),
            "--summary-json",
            str(summary),
        ]
    )
    payload = load_json(summary)
    assert_standard(payload, "ok")
    stdout_payload = load_stdout_json(completed)
    assert_standard(stdout_payload, "ok")
    if payload["match_count"] != 2 or payload["results"]["inventory"]["match_count"] != 2:
        raise AssertionError(f"Unexpected match count for marked inventory: {payload}")
    if not csv_path.exists() or not report.exists():
        raise AssertionError("Inventory CSV/report artifacts were not created.")


def check_zero_match_warning(tmp: Path, python: str) -> None:
    source = create_styled_docx(tmp, python, with_marked=False)
    summary = tmp / "plain_inventory.json"
    run_cmd(
        [
            python,
            str(SCRIPTS / "office_roundtrip.py"),
            "docx-style-inventory",
            str(source),
            "--require-bold",
            "--require-italic",
            "--require-underline",
            "--summary-json",
            str(summary),
        ]
    )
    payload = load_json(summary)
    assert_standard(payload, "warning")
    if payload["match_count"] != 0 or not payload["qa"]["findings"]:
        raise AssertionError(f"Zero-match inventory should warn with findings: {payload}")


def check_corrupt_docx_fails_cleanly(tmp: Path, python: str) -> None:
    corrupt = tmp / "corrupt.docx"
    corrupt.write_bytes(b"not a zip-backed docx")
    summary = tmp / "corrupt_inventory.json"
    completed = run_cmd(
        [
            python,
            str(SCRIPTS / "office_roundtrip.py"),
            "docx-style-inventory",
            str(corrupt),
            "--summary-json",
            str(summary),
        ],
        expect_ok=False,
    )
    if "Traceback" in completed.stderr or "Traceback" in completed.stdout:
        raise AssertionError(f"Corrupt DOCX leaked a traceback:\nSTDOUT:\n{completed.stdout}\nSTDERR:\n{completed.stderr}")
    payload = load_json(summary)
    assert_standard(payload, "fail")
    if payload["match_count"] != 0 or not payload["qa"]["findings"]:
        raise AssertionError(f"Corrupt DOCX failure payload is not useful: {payload}")


def check_invalid_summary_target_fails_cleanly(tmp: Path, python: str) -> None:
    source = create_styled_docx(tmp, python, with_marked=True)
    summary_dir = tmp / "summary_is_dir"
    summary_dir.mkdir()
    completed = run_cmd(
        [
            python,
            str(SCRIPTS / "office_roundtrip.py"),
            "docx-style-inventory",
            str(source),
            "--summary-json",
            str(summary_dir),
        ],
        expect_ok=False,
    )
    if "Traceback" in completed.stderr or "Traceback" in completed.stdout:
        raise AssertionError(f"Invalid summary target leaked a traceback:\nSTDOUT:\n{completed.stdout}\nSTDERR:\n{completed.stderr}")
    payload = load_stdout_json(completed)
    assert_standard(payload, "fail")


def main() -> None:
    python = runtime_python()
    with tempfile.TemporaryDirectory(prefix="sda_office_style_inventory_") as raw_tmp:
        tmp = Path(raw_tmp)
        check_marked_inventory(tmp, python)
        print("PASS check_marked_inventory")
        check_zero_match_warning(tmp, python)
        print("PASS check_zero_match_warning")
        check_corrupt_docx_fails_cleanly(tmp, python)
        print("PASS check_corrupt_docx_fails_cleanly")
        check_invalid_summary_target_fails_cleanly(tmp, python)
        print("PASS check_invalid_summary_target_fails_cleanly")
    print("All office_roundtrip docx-style-inventory regressions passed.")


if __name__ == "__main__":
    main()
