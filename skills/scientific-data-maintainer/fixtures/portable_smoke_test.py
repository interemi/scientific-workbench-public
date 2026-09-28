#!/usr/bin/env python3
"""Run a portable smoke test for the scientific-data-analysis skill."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import subprocess
import sys
from pathlib import Path

from _internal.public_contract import validate_standard_envelope
from _internal.provenance_utils import environment_summary, public_path, sanitize_payload, standard_tool_payload, write_manifest
from _internal.runtime_common import clean_known_stderr, configure_runtime
from maintainer_path_safety import PathSafetyError, ensure_safe_paths, operand


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", required=True, help="Directory for smoke-test artifacts.")
    parser.add_argument(
        "--examples-dir",
        help="Optional examples directory. Defaults to the packaged examples folder in this skill.",
    )
    parser.add_argument(
        "--profile",
        choices=["core", "full"],
        default="core",
        help="Smoke-test depth. 'core' stays lightweight and portable.",
    )
    parser.add_argument("--summary-json", help="Optional summary JSON path.")
    parser.add_argument("--manifest-json", help="Optional manifest JSON path.")
    return parser.parse_args()


def resolve_smoke_context(args: argparse.Namespace) -> tuple[Path, Path, Path, Path]:
    script_dir = Path(__file__).resolve().parent
    root_dir = script_dir.parent
    if root_dir.name == "scientific-data-maintainer":
        mother = root_dir.parent / "scientific-data-analysis"
        if mother.is_dir():
            root_dir = mother
            script_dir = mother / "scripts"
    examples_dir = Path(args.examples_dir).expanduser() if args.examples_dir else root_dir / "examples"
    output_dir = Path(args.output_dir).expanduser()
    return root_dir, script_dir, examples_dir, output_dir


def ensure_portable_smoke_paths(args: argparse.Namespace) -> tuple[Path, Path, Path, Path]:
    """Protect packaged examples before any smoke artifact is created."""

    context = resolve_smoke_context(args)
    _root_dir, _script_dir, examples_dir, output_dir = context
    ensure_safe_paths(
        inputs=[operand("--examples-dir", examples_dir, tree=True)],
        outputs=[
            operand("--output-dir", output_dir, tree=True),
            operand("--summary-json", args.summary_json),
            operand("--manifest-json", args.manifest_json),
        ],
    )
    return context


def emit_payload_safely(payload: dict, summary_json: str | None = None) -> None:
    rendered = json.dumps(sanitize_payload(payload), indent=2, ensure_ascii=True) + "\n"
    if summary_json:
        target = Path(summary_json)
        try:
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_text(rendered, encoding="utf-8")
        except OSError as exc:
            raise RuntimeError(
                f"Could not write summary JSON to {public_path(target)}: {exc.__class__.__name__}: {exc}"
            ) from None
    print(rendered, end="")


def build_failure_payload(args: argparse.Namespace, exc: Exception, *, blocked: bool = False) -> dict:
    output_dir = Path(args.output_dir).expanduser() if getattr(args, "output_dir", None) else None
    return standard_tool_payload(
        "portable_smoke_test",
        status="blocked" if blocked else "fail",
        notes=[
            "This smoke test is designed for fresh machines and portable verification.",
            (
                "The smoke test was blocked before writing because an output could overlap an input."
                if blocked
                else "The smoke test failed before it could complete cleanly."
            ),
        ],
        artifacts={} if blocked else {
            "summary_json": getattr(args, "summary_json", None),
            "output_dir": output_dir,
            "manifest_json": getattr(args, "manifest_json", None),
        },
        results={
            "profile": getattr(args, "profile", None),
            "examples_dir": getattr(args, "examples_dir", None),
            "feature_runs": {},
            "public_surface_coverage": {
                "tiers_enforced": ["core"] if getattr(args, "profile", "core") == "core" else ["core", "full"],
                "missing_coverage": [],
                "coverage_json": None,
            },
            "overall_status": "BLOCKED_CONTROLADO" if blocked else "FAIL",
            "error_type": exc.__class__.__name__,
            "error": str(exc),
        },
        qa={
            "status": "blocked" if blocked else "fail",
            "findings": [
                {
                    "severity": "high",
                    "title": "portable_smoke_test blocked" if blocked else "portable_smoke_test failed",
                    "detail": str(exc),
                }
            ],
            "metrics": {"feature_count": 0, "pass_count": 0, "missing_coverage_count": 0},
        },
        include_environment=True,
    )


def run(cmd: list[str]) -> dict:
    completed = subprocess.run(cmd, text=True, capture_output=True, check=False)
    return {
        "command": cmd,
        "returncode": completed.returncode,
        "stdout": completed.stdout[-4000:],
        "stderr": clean_known_stderr(completed.stderr[-4000:]),
    }


def require_outputs(result: dict, *paths: Path) -> dict:
    missing = [str(path) for path in paths if not path.exists()]
    if missing:
        result["returncode"] = result["returncode"] or 1
        detail = "Missing expected outputs: " + ", ".join(missing)
        result["stderr"] = (result.get("stderr", "") + "\n" + detail).strip()
    return result


def require_standard_envelope(result: dict, summary_path: Path) -> dict:
    if not summary_path.exists():
        return require_outputs(result, summary_path)
    try:
        payload = json.loads(summary_path.read_text(encoding="utf-8"))
    except Exception as exc:
        result["returncode"] = result["returncode"] or 1
        result["stderr"] = (result.get("stderr", "") + f"\nInvalid JSON summary: {exc}").strip()
        return result
    issues = validate_standard_envelope(payload)
    if issues:
        result["returncode"] = result["returncode"] or 1
        detail = "Envelope issues: " + "; ".join(issues)
        result["stderr"] = (result.get("stderr", "") + "\n" + detail).strip()
    return result


def require_unchanged(result: dict, before: dict[Path, str]) -> dict:
    changed = [str(path) for path, digest in before.items()
               if not path.is_file() or hashlib.sha256(path.read_bytes()).hexdigest() != digest]
    result["original_hashes_verified"] = not changed
    if changed:
        result["returncode"] = result["returncode"] or 1
        result["stderr"] = (result.get("stderr", "") + "\nModified synthetic inputs: " + ", ".join(changed)).strip()
    return result


def require_controlled_rejection(result: dict, summary: Path, reason: str) -> dict:
    observed = result["returncode"]
    result["observed_returncode"] = observed
    result["expectation"] = "controlled_rejection"
    try:
        payload = json.loads(summary.read_text())
        valid = (observed != 0 and not validate_standard_envelope(payload)
                 and reason in payload.get("blocked_reason", "")
                 and payload.get("original_modified") is False)
    except (OSError, ValueError):
        valid = False
    result["returncode"] = 0 if valid else 1
    if not valid:
        result["stderr"] = (result.get("stderr", "") + "\nExpected controlled rejection was not demonstrated.").strip()
    return result


def create_generated_office_docs(output_dir: Path) -> list[Path]:
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from PIL import Image, ImageDraw

    output_dir.mkdir(parents=True, exist_ok=True)
    created = []

    doc_path = output_dir / "sample.docx"
    doc = Document()
    doc.add_heading("Portable Smoke Test", level=1)
    doc.add_paragraph("This is a generated DOCX used to verify portable document workflows.")
    doc.save(doc_path)
    created.append(doc_path)

    pptx_path = output_dir / "sample.pptx"
    figure_path = output_dir / "sample_science_figure.png"
    figure = Image.new("RGB", (900, 520), (247, 249, 252))
    draw = ImageDraw.Draw(figure)
    draw.rectangle((80, 80, 820, 420), outline=(45, 92, 156), width=6)
    draw.line((110, 360, 780, 150), fill=(200, 60, 80), width=5)
    draw.line((110, 320, 780, 320), fill=(90, 90, 90), width=3)
    draw.text((120, 95), "Synthetic photometry figure", fill=(30, 45, 60))
    draw.text((120, 435), "Campaign 2026A / reduced light curve", fill=(50, 70, 90))
    figure.save(figure_path)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Portable Smoke Test"
    body = slide.placeholders[1]
    body.text = ""
    tx_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.85), Inches(8.3), Inches(0.5))
    tx_box.text_frame.text = "Reference figure-led slide with editable caption."
    tx_box.text_frame.paragraphs[0].font.size = Pt(18)
    slide.shapes.add_picture(str(figure_path), Inches(1.0), Inches(1.45), width=Inches(8.0), height=Inches(4.45))
    cap_box = slide.shapes.add_textbox(Inches(1.05), Inches(6.0), Inches(7.9), Inches(0.45))
    cap_box.text_frame.text = "Figure 1. Synthetic reduced curve used to exercise editable-caption and safe-area QA."
    cap_box.text_frame.paragraphs[0].font.size = Pt(13)

    slide2 = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide2.shapes.title.text = "Diagnostics"
    slide2.shapes.add_picture(str(figure_path), Inches(0.8), Inches(1.4), width=Inches(4.05), height=Inches(2.6))
    slide2.shapes.add_picture(str(figure_path), Inches(5.05), Inches(1.4), width=Inches(4.05), height=Inches(2.6))
    left_caption = slide2.shapes.add_textbox(Inches(0.85), Inches(4.15), Inches(3.95), Inches(0.55))
    left_caption.text_frame.text = "Figure 2. Control panel."
    left_caption.text_frame.paragraphs[0].font.size = Pt(12)
    right_caption = slide2.shapes.add_textbox(Inches(5.1), Inches(4.15), Inches(3.95), Inches(0.55))
    right_caption.text_frame.text = "Figure 3. Comparison panel."
    right_caption.text_frame.paragraphs[0].font.size = Pt(12)
    presentation.save(pptx_path)
    created.append(pptx_path)
    created.append(figure_path)

    xlsx_path = output_dir / "sample.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"
    worksheet["A1"] = "item"
    worksheet["B1"] = "value"
    worksheet["A2"] = "alpha"
    worksheet["B2"] = 1
    workbook.save(xlsx_path)
    created.append(xlsx_path)

    return created


def create_generated_fits(path: Path) -> Path:
    import numpy as np
    from astropy.io import fits

    path.parent.mkdir(parents=True, exist_ok=True)
    yy, xx = np.mgrid[0:64, 0:64]
    image = 100.0 + 300.0 * np.exp(-((xx - 31.5) ** 2 + (yy - 30.0) ** 2) / (2.0 * 5.0**2))
    header = fits.Header()
    header["OBJECT"] = "SMOKE_TEST"
    header["FILTER"] = "R"
    header["EXPTIME"] = 30.0
    fits.PrimaryHDU(image.astype("float32"), header=header).writeto(path, overwrite=True)
    return path


def create_generated_linear_multispec(path: Path) -> Path:
    """Create a physically supported type-0 MULTISPEC golden-path fixture."""
    import numpy as np
    from astropy.io import fits

    path.parent.mkdir(parents=True, exist_ok=True)
    pixels = 24
    wavelength_zero = 6500
    wavelength_step = 1
    wavelength = wavelength_zero + wavelength_step * np.arange(pixels, dtype=float)
    flux = 1.0 - 0.12 * np.exp(-0.5 * ((wavelength - 6505.0) / 0.8) ** 2)
    hdu = fits.PrimaryHDU(flux.astype("float32").reshape(1, pixels))
    hdu.header["OBJECT"] = "V28_LINEAR_MULTISPEC"
    hdu.header["CTYPE1"] = "MULTISPE"
    hdu.header["WAT0_001"] = "system=equispec"
    hdu.header["WAT1_001"] = "wtype=multispec label=Wavelength units=angstroms"
    hdu.header["WAT2_001"] = (
        f'wtype=multispec spec1="1 1 0 {wavelength_zero} '
        f'{wavelength_step} {pixels} 0 0 0"'
    )
    hdu.writeto(path, overwrite=True)
    return path


def create_generated_existing_wcs_fits(path: Path, header_only: bool = False) -> Path:
    import numpy as np
    from astropy.io import fits
    from astropy.wcs import WCS

    path.parent.mkdir(parents=True, exist_ok=True)
    wcs = WCS(naxis=2)
    wcs.wcs.crpix = [32.0, 32.0]
    wcs.wcs.cdelt = [-0.0001458, 0.0001458]
    wcs.wcs.crval = [250.1234, 36.5678]
    wcs.wcs.ctype = ["RA---TAN", "DEC--TAN"]
    header = wcs.to_header()
    header["OBJECT"] = "SMOKE_WCS"
    header["FILTER"] = "R"
    header["EXPTIME"] = 30.0
    if header_only:
        fits.PrimaryHDU(header=header).writeto(path, overwrite=True)
        return path
    yy, xx = np.mgrid[0:64, 0:64]
    image = 80.0 + 250.0 * np.exp(-((xx - 30.0) ** 2 + (yy - 34.0) ** 2) / (2.0 * 4.5**2))
    fits.PrimaryHDU(image.astype("float32"), header=header).writeto(path, overwrite=True)
    return path


def create_generated_standard_star_table(path: Path) -> Path:
    import numpy as np

    path.parent.mkdir(parents=True, exist_ok=True)
    rng = np.random.default_rng(42)
    airmass = np.array([1.02, 1.08, 1.15, 1.22, 1.31, 1.42, 1.55, 1.68, 1.82, 1.95])
    color = np.array([-0.05, 0.12, 0.25, 0.38, 0.46, 0.58, 0.71, 0.83, 0.94, 1.08])
    zero_point = 24.65
    extinction = -0.18
    color_term = 0.04
    std_mag = np.array([10.4, 10.8, 11.1, 11.4, 11.7, 12.0, 12.2, 12.5, 12.8, 13.0])
    offset = zero_point + extinction * airmass + color_term * color
    noise = rng.normal(0.0, 0.008, size=len(airmass))
    inst_mag = std_mag - offset + noise
    offset_err = np.full(len(airmass), 0.01)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["std_mag", "inst_mag", "airmass", "b_minus_v", "offset_err"])
        for row in zip(std_mag, inst_mag, airmass, color, offset_err):
            writer.writerow([f"{item:.8f}" for item in row])
    return path


def create_shifted_spectrum(source_path: Path, output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    rows = []
    for raw in source_path.read_text(errors="replace").splitlines():
        line = raw.strip()
        if not line or line.startswith("#"):
            continue
        parts = [item for item in line.replace(",", " ").split() if item]
        if len(parts) < 2:
            continue
        try:
            x = float(parts[0])
            y = float(parts[1])
        except Exception:
            continue
        rows.append((x, 0.97 * y + 0.02))
    with output_path.open("w", encoding="utf-8") as handle:
        for x, y in rows:
            handle.write(f"{x:.6f} {y:.6f}\n")
    return output_path


def create_relative_data_notebook(output_dir: Path) -> tuple[Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    notebook_path = output_dir / "relative_data_notebook.ipynb"
    data_path = output_dir / "local_input.txt"
    data_path.write_text("portable smoke relative data\n", encoding="utf-8")
    notebook = {
        "cells": [
            {
                "cell_type": "code",
                "execution_count": None,
                "metadata": {},
                "outputs": [],
                "source": [
                    "from pathlib import Path\n",
                    "text = Path('local_input.txt').read_text(encoding='utf-8').strip()\n",
                    "print(text)\n",
                ],
            }
        ],
        "metadata": {
            "kernelspec": {"display_name": "Python 3", "language": "python", "name": "python3"},
            "language_info": {"name": "python"},
        },
        "nbformat": 4,
        "nbformat_minor": 5,
    }
    notebook_path.write_text(json.dumps(notebook, indent=2), encoding="utf-8")
    return notebook_path, data_path


def create_scientific_writeup_fixtures(output_dir: Path) -> dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    good_md = output_dir / "good_astrophysics_writeup.md"
    good_md.write_text(
        """# Methodology

We used calibrated spectra with wavelength in Angstrom and radial velocity in km/s.
The reduction used a fixed order-selection criterion and rejected orders with low signal.
The uncertainty was estimated from the order-to-order dispersion and is reported as ± 2.0 km/s.

# Results

The measured radial velocity is -12.1 km/s and the equivalent width is 280 mA.
The values are consistent with the literature reference within the stated uncertainty.

# Discussion

The result suggests a physically plausible active late-type star, but the conclusion is limited by
the continuum placement and by the manual order selection. This comparison should be treated as
coursework-level evidence rather than a definitive publication result.

# Conclusions

The analysis is reproducible and the main claim is cautious.
""",
        encoding="utf-8",
    )
    weak_md = output_dir / "weak_astrophysics_writeup.md"
    weak_md.write_text(
        """# Results

The automatic pipeline proves that the star is rotating slowly and confirms the binary solution.
The values are perfect and the conclusion is definitive.
""",
        encoding="utf-8",
    )
    tex_dir = output_dir / "tex_project"
    (tex_dir / "sections").mkdir(parents=True, exist_ok=True)
    main_tex = tex_dir / "main.tex"
    main_tex.write_text(
        r"""\documentclass{article}
\begin{document}
\section{Metodologia}
\input{sections/metodologia}
\section{Resultados}
La velocidad radial medida fue $-11.8\,\mathrm{km/s}$ con incertidumbre $\pm 2.1\,\mathrm{km/s}$.
\section{Discusion}
\input{sections/discusion}
\end{document}
""",
        encoding="utf-8",
    )
    (tex_dir / "sections" / "metodologia.tex").write_text(
        "Se usaron espectros calibrados, seleccion de ordenes y trazabilidad de parametros.\n",
        encoding="utf-8",
    )
    (tex_dir / "sections" / "discusion.tex").write_text(
        "La comparacion con literatura es compatible, pero queda limitada por la seleccion manual y la calidad de la CCF.\n",
        encoding="utf-8",
    )
    return {"good_md": good_md, "weak_md": weak_md, "main_tex": main_tex}


def create_legacy_gate_fixtures(output_dir: Path) -> dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    evidence_dir = output_dir / "evidence"
    evidence_dir.mkdir(parents=True, exist_ok=True)
    (evidence_dir / "rv_summary.csv").write_text("case_id,vhelio_media_kms\npwand,-11.9\n", encoding="utf-8")
    (evidence_dir / "chromosphere_table.csv").write_text("line,ew\nHalpha,0.32\n", encoding="utf-8")
    requirements_csv = output_dir / "requirements_manual.csv"
    rows = [
        {
            "requirement_id": "rv_single_primary",
            "block": "velocidad_radial",
            "priority": "critical",
            "requirement": "RV final de estrella simple",
            "expected_evidence": "tabla RV",
            "evidence_path": str(evidence_dir / "rv_summary.csv"),
            "measurement_state": "measured",
            "adoption_state": "adopted",
            "physical_robustness": "robust",
            "ready_for_report": "true",
            "status": "",
            "notes": "",
        },
        {
            "requirement_id": "rv_sb2_component_a",
            "block": "velocidad_radial",
            "priority": "critical",
            "requirement": "SB2 componente A",
            "expected_evidence": "tabla por orden",
            "evidence_path": str(output_dir / "missing_sb2.csv"),
            "measurement_state": "measured",
            "adoption_state": "measured_not_adopted",
            "physical_robustness": "not_robust",
            "ready_for_report": "false",
            "status": "",
            "notes": "Debe bloquear.",
        },
        {
            "requirement_id": "chromosphere_halpha",
            "block": "actividad_cromosferica",
            "priority": "critical",
            "requirement": "Halpha",
            "expected_evidence": "tabla EW",
            "evidence_path": str(evidence_dir / "chromosphere_table.csv"),
            "measurement_state": "measured",
            "adoption_state": "adopted_with_caution",
            "physical_robustness": "limited",
            "ready_for_report": "true",
            "status": "",
            "notes": "",
        },
    ]
    with requirements_csv.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=[
            "requirement_id",
            "block",
            "priority",
            "requirement",
            "expected_evidence",
            "evidence_path",
            "measurement_state",
            "adoption_state",
            "physical_robustness",
            "ready_for_report",
            "status",
            "notes",
        ])
        writer.writeheader()
        for row in rows:
            writer.writerow(row)
    return {"requirements_csv": requirements_csv}


def create_workspace_audit_fixture(output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "fits_p1").mkdir(parents=True, exist_ok=True)
    (output_dir / "analysis" / "parsed").mkdir(parents=True, exist_ok=True)
    (output_dir / "logs").mkdir(parents=True, exist_ok=True)
    (output_dir / "iSTARMOD").mkdir(parents=True, exist_ok=True)
    (output_dir / "report").mkdir(parents=True, exist_ok=True)
    (output_dir / "fits_p1" / "npwand_n3.fits").write_text("synthetic fits placeholder\n", encoding="utf-8")
    (output_dir / "analysis" / "run_auto_all_summary.json").write_text("{\"status\":\"ok\"}\n", encoding="utf-8")
    (output_dir / "analysis" / "parsed" / "rej_compA_rv_auto.csv").write_text("case_id,vhelio\nrej,12.3\n", encoding="utf-8")
    (output_dir / "logs" / "pipeline.log").write_text("legacy log\n", encoding="utf-8")
    (output_dir / "iSTARMOD" / "rvvalues.dat").write_text("15.200\n", encoding="utf-8")
    (output_dir / "report" / "draft.tex").write_text("\\section{Resultados}\n", encoding="utf-8")
    return output_dir


def create_crossmatch_fixtures(output_dir: Path) -> dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    left = output_dir / "catalog_left.csv"
    right = output_dir / "catalog_right.csv"
    left_rows = [
        ("id", "ra_deg", "dec_deg"),
        ("A", "120.0000", "22.0000"),
        ("B", "120.0050", "22.0020"),
        ("C", "121.0000", "21.5000"),
    ]
    right_rows = [
        ("source", "ra_deg", "dec_deg"),
        ("X", "120.0002", "22.0001"),
        ("Y", "120.0049", "22.0018"),
        ("Z", "130.0000", "10.0000"),
    ]
    for path, rows in ((left, left_rows), (right, right_rows)):
        with path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            writer.writerows(rows)
    return {"left": left, "right": right}


def create_semantic_diff_fixtures(output_dir: Path) -> dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    baseline = output_dir / "baseline.md"
    candidate = output_dir / "candidate.md"
    baseline.write_text(
        "# Status update\n\nThe measured radial velocity is -11.2 km/s.\n",
        encoding="utf-8",
    )
    candidate.write_text(
        "# Status update\n\nThe measured radial velocity is -12.0 km/s and the uncertainty is 2.0 km/s.\n",
        encoding="utf-8",
    )
    return {"baseline": baseline, "candidate": candidate}


def create_container_fixture(output_dir: Path) -> Path:
    import numpy as np

    output_dir.mkdir(parents=True, exist_ok=True)
    container = output_dir / "sample_arrays.npz"
    np.savez(container, flux=np.linspace(0.0, 1.0, 16), mask=np.array([0, 1, 0, 1], dtype="int16"))
    return container


def create_systemic_fixture(output_dir: Path) -> dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    dataset = output_dir / "demo.vels"
    dataset.write_text(
        "# value_units = m/s\n"
        "2450000.0 12.1 1.2\n"
        "2450001.0 10.5 1.0\n"
        "2450002.0 11.4 1.1\n",
        encoding="utf-8",
    )
    system = output_dir / "demo.sys"
    system.write_text(
        "Name\tDemo System\n"
        "Mass\t1.0\n"
        "RV[]\tdemo.vels\n",
        encoding="utf-8",
    )
    return {"system": system, "dataset": dataset}


def run_smoke(args: argparse.Namespace) -> int:
    root_dir, script_dir, examples_dir, output_dir = ensure_portable_smoke_paths(args)
    configure_runtime("portable_smoke_test")
    if not examples_dir.exists():
        raise RuntimeError(f"Examples directory not found: {public_path(examples_dir)}")
    if not examples_dir.is_dir():
        raise RuntimeError(f"Examples path is not a directory: {public_path(examples_dir)}")

    output_dir.mkdir(parents=True, exist_ok=True)
    generated_dir = output_dir / "generated"
    generated_docs_dir = generated_dir / "documents"
    generated_data_dir = generated_dir / "data"
    generated_docs = create_generated_office_docs(generated_docs_dir)
    fits_path = create_generated_fits(generated_data_dir / "smoke_image.fits")
    linear_multispec = create_generated_linear_multispec(generated_data_dir / "smoke_linear_multispec.fits")
    existing_wcs_fits = create_generated_existing_wcs_fits(generated_data_dir / "smoke_existing_wcs.fits")
    header_only_wcs_fits = create_generated_existing_wcs_fits(generated_data_dir / "smoke_header_only_wcs.fits", header_only=True)
    standard_star_table = create_generated_standard_star_table(generated_data_dir / "standard_stars.csv")
    writeup_fixtures = create_scientific_writeup_fixtures(generated_dir / "writeups")
    gate_fixtures = create_legacy_gate_fixtures(generated_dir / "requirements_gate")
    workspace_fixture = create_workspace_audit_fixture(generated_dir / "old_codex_workspace")
    crossmatch_fixtures = create_crossmatch_fixtures(generated_dir / "catalog_crossmatch")
    semantic_diff_fixtures = create_semantic_diff_fixtures(generated_dir / "semantic_diff")
    container_fixture = create_container_fixture(generated_dir / "containers")
    systemic_fixture = create_systemic_fixture(generated_dir / "systemic")

    outputs: list[Path] = []
    feature_runs: dict[str, dict] = {}
    feature_registry_ids: dict[str, str] = {}

    env_doctor_json = output_dir / "env_doctor.json"
    feature_runs["env_doctor"] = run(
        [sys.executable, str(script_dir / "env_doctor.py"), "--summary-json", str(env_doctor_json)]
    )
    feature_runs["env_doctor"] = require_outputs(feature_runs["env_doctor"], env_doctor_json)
    feature_runs["env_doctor"] = require_standard_envelope(feature_runs["env_doctor"], env_doctor_json)
    feature_registry_ids["env_doctor"] = "env_doctor"
    outputs.append(env_doctor_json)

    docs_sync_json = output_dir / "public_surface_docs_check.json"
    feature_runs["sync_public_surface_docs"] = run(
        [
            sys.executable,
            str(script_dir / "sync_public_surface_docs.py"),
            "--check",
            "--summary-json",
            str(docs_sync_json),
        ]
    )
    feature_runs["sync_public_surface_docs"] = require_outputs(feature_runs["sync_public_surface_docs"], docs_sync_json)
    feature_runs["sync_public_surface_docs"] = require_standard_envelope(feature_runs["sync_public_surface_docs"], docs_sync_json)
    outputs.append(docs_sync_json)

    surface_audit_json = output_dir / "skill_surface_audit.json"
    feature_runs["skill_surface_audit"] = run(
        [
            sys.executable,
            str(script_dir / "skill_surface_audit.py"),
            "--skip-hygiene",
            "--summary-json",
            str(surface_audit_json),
        ]
    )
    feature_runs["skill_surface_audit"] = require_outputs(feature_runs["skill_surface_audit"], surface_audit_json)
    feature_runs["skill_surface_audit"] = require_standard_envelope(feature_runs["skill_surface_audit"], surface_audit_json)
    outputs.append(surface_audit_json)

    fits_summary = output_dir / "inspect_fits.json"
    fits_preview = output_dir / "inspect_fits.png"
    feature_runs["inspect_fits"] = run(
        [
            sys.executable,
            str(script_dir / "inspect_fits.py"),
            str(fits_path),
            "--preview",
            str(fits_preview),
            "--summary-json",
            str(fits_summary),
        ]
    )
    feature_runs["inspect_fits"] = require_outputs(feature_runs["inspect_fits"], fits_summary, fits_preview)
    feature_runs["inspect_fits"] = require_standard_envelope(feature_runs["inspect_fits"], fits_summary)
    feature_registry_ids["inspect_fits"] = "inspect_fits"
    outputs.extend([fits_summary, fits_preview])

    legacy_multispec = linear_multispec
    if legacy_multispec.exists():
        legacy_inventory_dir = output_dir / "legacy_multispec_inventory"
        legacy_inventory_json = legacy_inventory_dir / "summary.json"
        legacy_orders_csv = legacy_inventory_dir / "orders_inventory.csv"
        feature_runs["echelle_multispec_inventory"] = run(
            [
                sys.executable,
                str(script_dir / "echelle_multispec_inventory.py"),
                str(legacy_multispec),
                "--output-dir",
                str(legacy_inventory_dir),
            ]
        )
        feature_runs["echelle_multispec_inventory"] = require_outputs(
            feature_runs["echelle_multispec_inventory"],
            legacy_inventory_json,
            legacy_orders_csv,
        )
        feature_runs["echelle_multispec_inventory"] = require_standard_envelope(feature_runs["echelle_multispec_inventory"], legacy_inventory_json)
        feature_registry_ids["echelle_multispec_inventory"] = "echelle_multispec_inventory"
        outputs.extend([legacy_inventory_json, legacy_orders_csv])

        legacy_li_dir = output_dir / "legacy_li_equivalent_width"
        legacy_li_json = legacy_li_dir / "li6708_summary.json"
        legacy_li_plot = legacy_li_dir / "figures" / "li6708_equivalent_width.png"
        feature_runs["li6708_equivalent_width_workbench"] = run(
            [
                sys.executable,
                str(script_dir / "li6708_equivalent_width_workbench.py"),
                "measure",
                str(legacy_multispec),
                "--line-center",
                "6505.0",
                "--line-label",
                "Synthetic line 6505 A",
                "--integration-window",
                "6503.0",
                "6507.0",
                "--continuum-window",
                "6500.0",
                "6502.0",
                "--continuum-window",
                "6508.0",
                "6510.0",
                "--output-dir",
                str(legacy_li_dir),
            ]
        )
        feature_runs["li6708_equivalent_width_workbench"] = require_outputs(
            feature_runs["li6708_equivalent_width_workbench"],
            legacy_li_json,
            legacy_li_plot,
        )
        feature_runs["li6708_equivalent_width_workbench"] = require_standard_envelope(feature_runs["li6708_equivalent_width_workbench"], legacy_li_json)
        feature_registry_ids["li6708_equivalent_width_workbench"] = "li6708_equivalent_width_workbench.measure"
        outputs.extend([legacy_li_json, legacy_li_plot])

    example_ops = examples_dir / "tabular" / "ops.csv"
    profile_json = output_dir / "profile_table.json"
    feature_runs["profile_table"] = run(
        [
            sys.executable,
            str(script_dir / "profile_table.py"),
            str(example_ops),
            "--summary-json",
            str(profile_json),
        ]
    )
    feature_runs["profile_table"] = require_outputs(feature_runs["profile_table"], profile_json)
    feature_runs["profile_table"] = require_standard_envelope(feature_runs["profile_table"], profile_json)
    feature_registry_ids["profile_table"] = "profile_table"
    outputs.append(profile_json)

    photometric_solution_json = output_dir / "photometric_solution.json"
    photometric_solution_coeffs = output_dir / "photometric_solution_coefficients.csv"
    photometric_solution_residuals = output_dir / "photometric_solution_residuals.csv"
    photometric_solution_plot = output_dir / "photometric_solution.png"
    photometric_solution_report = output_dir / "photometric_solution.md"
    feature_runs["photometric_solution"] = run(
        [
            sys.executable,
            str(script_dir / "photometric_solution.py"),
            str(standard_star_table),
            "--color-col",
            "b_minus_v",
            "--include-color-term",
            "--error-col",
            "offset_err",
            "--summary-json",
            str(photometric_solution_json),
            "--coefficients-csv",
            str(photometric_solution_coeffs),
            "--residual-csv",
            str(photometric_solution_residuals),
            "--residual-plot",
            str(photometric_solution_plot),
            "--report-md",
            str(photometric_solution_report),
        ]
    )
    feature_runs["photometric_solution"] = require_outputs(
        feature_runs["photometric_solution"],
        photometric_solution_json,
        photometric_solution_coeffs,
        photometric_solution_residuals,
        photometric_solution_plot,
        photometric_solution_report,
    )
    feature_runs["photometric_solution"] = require_standard_envelope(feature_runs["photometric_solution"], photometric_solution_json)
    feature_registry_ids["photometric_solution"] = "photometric_solution"
    outputs.extend(
        [
            photometric_solution_json,
            photometric_solution_coeffs,
            photometric_solution_residuals,
            photometric_solution_plot,
            photometric_solution_report,
        ]
    )

    noise_budget_json = output_dir / "noise_budget.json"
    noise_budget_report = output_dir / "noise_budget.md"
    feature_runs["photometry_noise_budget"] = run(
        [
            sys.executable,
            str(script_dir / "photometry_noise_budget.py"),
            "--source",
            "25000",
            "--sky-per-pixel",
            "42",
            "--dark-per-pixel",
            "0.2",
            "--read-noise",
            "4.3",
            "--n-pixels",
            "80",
            "--sky-estimate-pixels",
            "600",
            "--n-frames",
            "3",
            "--units",
            "electrons",
            "--summary-json",
            str(noise_budget_json),
            "--report-md",
            str(noise_budget_report),
        ]
    )
    feature_runs["photometry_noise_budget"] = require_outputs(
        feature_runs["photometry_noise_budget"],
        noise_budget_json,
        noise_budget_report,
    )
    feature_runs["photometry_noise_budget"] = require_standard_envelope(feature_runs["photometry_noise_budget"], noise_budget_json)
    feature_registry_ids["photometry_noise_budget"] = "photometry_noise_budget"
    outputs.extend([noise_budget_json, noise_budget_report])

    cross_domain_dir = output_dir / "cross_domain"
    cross_domain_json = cross_domain_dir / "summary.json"
    feature_runs["cross_domain_data_workbench"] = run(
        [
            sys.executable,
            str(script_dir / "cross_domain_data_workbench.py"),
            str(examples_dir / "tabular"),
            "--sql",
            "SELECT * FROM source0 LIMIT 2",
            "--output-dir",
            str(cross_domain_dir),
            "--summary-json",
            str(cross_domain_json),
        ]
    )
    feature_runs["cross_domain_data_workbench"] = require_outputs(
        feature_runs["cross_domain_data_workbench"],
        cross_domain_dir / "inventory.csv",
        cross_domain_dir / "report.md",
        cross_domain_json,
    )
    feature_runs["cross_domain_data_workbench"] = require_standard_envelope(feature_runs["cross_domain_data_workbench"], cross_domain_json)
    feature_registry_ids["cross_domain_data_workbench"] = "cross_domain_data_workbench"
    outputs.extend([cross_domain_dir / "inventory.csv", cross_domain_dir / "report.md", cross_domain_json])

    intake_dir = output_dir / "document_intake"
    intake_json = intake_dir / "summary.json"
    intake_inputs = [str(examples_dir / "documents"), *(str(path) for path in generated_docs)]
    feature_runs["document_intake_workbench"] = run(
        [
            sys.executable,
            str(script_dir / "document_intake_workbench.py"),
            *intake_inputs,
            "--output-dir",
            str(intake_dir),
            "--summary-json",
            str(intake_json),
        ]
    )
    feature_runs["document_intake_workbench"] = require_outputs(
        feature_runs["document_intake_workbench"],
        intake_dir / "inventory.csv",
        intake_dir / "report.md",
        intake_json,
    )
    feature_runs["document_intake_workbench"] = require_standard_envelope(feature_runs["document_intake_workbench"], intake_json)
    feature_registry_ids["document_intake_workbench"] = "document_intake_workbench"
    outputs.extend([intake_dir / "inventory.csv", intake_dir / "report.md", intake_json])

    semantic_diff_json = output_dir / "semantic_diff.json"
    semantic_diff_md = output_dir / "semantic_diff.md"
    feature_runs["semantic_diff"] = run(
        [
            sys.executable,
            str(script_dir / "semantic_diff.py"),
            str(semantic_diff_fixtures["baseline"]),
            str(semantic_diff_fixtures["candidate"]),
            "--summary-json",
            str(semantic_diff_json),
            "--output-md",
            str(semantic_diff_md),
        ]
    )
    feature_runs["semantic_diff"] = require_outputs(
        feature_runs["semantic_diff"],
        semantic_diff_json,
        semantic_diff_md,
    )
    feature_runs["semantic_diff"] = require_standard_envelope(feature_runs["semantic_diff"], semantic_diff_json)
    feature_registry_ids["semantic_diff"] = "semantic_diff"
    outputs.extend([semantic_diff_json, semantic_diff_md])

    container_json = output_dir / "inspect_data_container.json"
    feature_runs["inspect_data_container"] = run(
        [
            sys.executable,
            str(script_dir / "inspect_data_container.py"),
            str(container_fixture),
            "--summary-json",
            str(container_json),
        ]
    )
    feature_runs["inspect_data_container"] = require_outputs(feature_runs["inspect_data_container"], container_json)
    feature_runs["inspect_data_container"] = require_standard_envelope(feature_runs["inspect_data_container"], container_json)
    feature_registry_ids["inspect_data_container"] = "inspect_data_container"
    outputs.append(container_json)

    crossmatch_output = output_dir / "catalog_crossmatch.csv"
    crossmatch_json = output_dir / "catalog_crossmatch.json"
    feature_runs["catalog_workbench_crossmatch"] = run(
        [
            sys.executable,
            str(script_dir / "catalog_workbench.py"),
            "crossmatch-sky",
            str(crossmatch_fixtures["left"]),
            str(crossmatch_fixtures["right"]),
            str(crossmatch_output),
            "--left-ra",
            "ra_deg",
            "--left-dec",
            "dec_deg",
            "--right-ra",
            "ra_deg",
            "--right-dec",
            "dec_deg",
            "--radius-arcsec",
            "2.0",
            "--summary-json",
            str(crossmatch_json),
        ]
    )
    feature_runs["catalog_workbench_crossmatch"] = require_outputs(
        feature_runs["catalog_workbench_crossmatch"],
        crossmatch_output,
        crossmatch_json,
    )
    feature_runs["catalog_workbench_crossmatch"] = require_standard_envelope(feature_runs["catalog_workbench_crossmatch"], crossmatch_json)
    feature_registry_ids["catalog_workbench_crossmatch"] = "catalog_workbench.crossmatch-sky"
    outputs.extend([crossmatch_output, crossmatch_json])

    rv_inspect_json = output_dir / "radial_velocity_inspect.json"
    feature_runs["radial_velocity_workbench_inspect"] = run(
        [
            sys.executable,
            str(script_dir / "radial_velocity_workbench.py"),
            "inspect",
            str(systemic_fixture["system"]),
            "--summary-json",
            str(rv_inspect_json),
        ]
    )
    feature_runs["radial_velocity_workbench_inspect"] = require_outputs(feature_runs["radial_velocity_workbench_inspect"], rv_inspect_json)
    feature_runs["radial_velocity_workbench_inspect"] = require_standard_envelope(feature_runs["radial_velocity_workbench_inspect"], rv_inspect_json)
    feature_registry_ids["radial_velocity_workbench_inspect"] = "radial_velocity_workbench.inspect"
    outputs.append(rv_inspect_json)

    notebook_path = output_dir / "portable_smoke_notebook.ipynb"
    notebook_json = output_dir / "portable_smoke_notebook.json"
    feature_runs["bootstrap_analysis_notebook"] = run(
        [
            sys.executable,
            str(script_dir / "bootstrap_analysis_notebook.py"),
            str(notebook_path),
            "--title",
            "Portable Smoke Notebook",
            "--domain",
            "astronomy",
            "--language",
            "en",
            "--data-path",
            str(standard_star_table),
            "--summary-json",
            str(notebook_json),
            "--overwrite",
        ]
    )
    feature_runs["bootstrap_analysis_notebook"] = require_outputs(
        feature_runs["bootstrap_analysis_notebook"],
        notebook_path,
        notebook_json,
    )
    feature_runs["bootstrap_analysis_notebook"] = require_standard_envelope(feature_runs["bootstrap_analysis_notebook"], notebook_json)
    feature_registry_ids["bootstrap_analysis_notebook"] = "bootstrap_analysis_notebook"
    outputs.extend([notebook_path, notebook_json])

    deliverable_dir = output_dir / "deliverable"
    feature_runs["deliverable_status_report"] = run(
        [
            sys.executable,
            str(script_dir / "deliverable_factory.py"),
            "scaffold",
            str(deliverable_dir),
            "--kind",
            "status-report",
            "--format",
            "markdown",
            "--title",
            "Portable Smoke Status Report",
            "--language",
            "english",
        ]
    )
    feature_runs["deliverable_status_report"] = require_outputs(
        feature_runs["deliverable_status_report"],
        deliverable_dir / "deliverable.md",
    )
    feature_registry_ids["deliverable_status_report"] = "deliverable_factory.scaffold"
    outputs.append(deliverable_dir / "deliverable.md")

    writeup_good_report = output_dir / "scientific_writeup_good.md"
    writeup_good_json = output_dir / "scientific_writeup_good.json"
    feature_runs["scientific_writeup_review_good"] = run(
        [
            sys.executable,
            str(script_dir / "scientific_writeup_review.py"),
            str(writeup_fixtures["good_md"]),
            "--report-md",
            str(writeup_good_report),
            "--summary-json",
            str(writeup_good_json),
        ]
    )
    feature_runs["scientific_writeup_review_good"] = require_outputs(
        feature_runs["scientific_writeup_review_good"],
        writeup_good_report,
        writeup_good_json,
    )
    feature_runs["scientific_writeup_review_good"] = require_standard_envelope(feature_runs["scientific_writeup_review_good"], writeup_good_json)
    feature_registry_ids["scientific_writeup_review_good"] = "scientific_writeup_review"
    outputs.extend([writeup_good_report, writeup_good_json])

    writeup_weak_report = output_dir / "scientific_writeup_weak.md"
    writeup_weak_json = output_dir / "scientific_writeup_weak.json"
    feature_runs["scientific_writeup_review_weak"] = run(
        [
            sys.executable,
            str(script_dir / "scientific_writeup_review.py"),
            str(writeup_fixtures["weak_md"]),
            "--report-md",
            str(writeup_weak_report),
            "--summary-json",
            str(writeup_weak_json),
        ]
    )
    feature_runs["scientific_writeup_review_weak"] = require_outputs(
        feature_runs["scientific_writeup_review_weak"],
        writeup_weak_report,
        writeup_weak_json,
    )
    feature_runs["scientific_writeup_review_weak"] = require_standard_envelope(feature_runs["scientific_writeup_review_weak"], writeup_weak_json)
    outputs.extend([writeup_weak_report, writeup_weak_json])

    writeup_tex_report = output_dir / "scientific_writeup_tex.md"
    writeup_tex_json = output_dir / "scientific_writeup_tex.json"
    feature_runs["scientific_writeup_review_tex_includes"] = run(
        [
            sys.executable,
            str(script_dir / "scientific_writeup_review.py"),
            str(writeup_fixtures["main_tex"]),
            "--report-md",
            str(writeup_tex_report),
            "--summary-json",
            str(writeup_tex_json),
        ]
    )
    feature_runs["scientific_writeup_review_tex_includes"] = require_outputs(
        feature_runs["scientific_writeup_review_tex_includes"],
        writeup_tex_report,
        writeup_tex_json,
    )
    feature_runs["scientific_writeup_review_tex_includes"] = require_standard_envelope(feature_runs["scientific_writeup_review_tex_includes"], writeup_tex_json)
    outputs.extend([writeup_tex_report, writeup_tex_json])

    requirements_template_csv = output_dir / "coursework_requirements_template.csv"
    requirements_template_json = output_dir / "coursework_requirements_template.json"
    feature_runs["coursework_requirements_gate_template"] = run(
        [
            sys.executable,
            str(script_dir / "coursework_requirements_gate.py"),
            "template",
            "--output-csv",
            str(requirements_template_csv),
            "--summary-json",
            str(requirements_template_json),
        ]
    )
    feature_runs["coursework_requirements_gate_template"] = require_outputs(
        feature_runs["coursework_requirements_gate_template"],
        requirements_template_csv,
        requirements_template_json,
    )
    feature_runs["coursework_requirements_gate_template"] = require_standard_envelope(feature_runs["coursework_requirements_gate_template"], requirements_template_json)
    outputs.extend([requirements_template_csv, requirements_template_json])

    requirements_eval_json = output_dir / "coursework_requirements_eval.json"
    requirements_eval_md = output_dir / "coursework_requirements_eval.md"
    requirements_eval_csv = output_dir / "coursework_requirements_eval.csv"
    feature_runs["coursework_requirements_gate_evaluate"] = run(
        [
            sys.executable,
            str(script_dir / "coursework_requirements_gate.py"),
            "evaluate",
            str(gate_fixtures["requirements_csv"]),
            "--output-csv",
            str(requirements_eval_csv),
            "--report-md",
            str(requirements_eval_md),
            "--summary-json",
            str(requirements_eval_json),
        ]
    )
    feature_runs["coursework_requirements_gate_evaluate"] = require_outputs(
        feature_runs["coursework_requirements_gate_evaluate"],
        requirements_eval_csv,
        requirements_eval_md,
        requirements_eval_json,
    )
    feature_runs["coursework_requirements_gate_evaluate"] = require_standard_envelope(feature_runs["coursework_requirements_gate_evaluate"], requirements_eval_json)
    outputs.extend([requirements_eval_csv, requirements_eval_md, requirements_eval_json])

    inheritance_audit_csv = output_dir / "workspace_inheritance_audit.csv"
    inheritance_audit_md = output_dir / "workspace_inheritance_audit.md"
    inheritance_audit_json = output_dir / "workspace_inheritance_audit.json"
    feature_runs["workspace_inheritance_audit"] = run(
        [
            sys.executable,
            str(script_dir / "workspace_inheritance_audit.py"),
            str(workspace_fixture),
            "--output-csv",
            str(inheritance_audit_csv),
            "--report-md",
            str(inheritance_audit_md),
            "--summary-json",
            str(inheritance_audit_json),
        ]
    )
    feature_runs["workspace_inheritance_audit"] = require_outputs(
        feature_runs["workspace_inheritance_audit"],
        inheritance_audit_csv,
        inheritance_audit_md,
        inheritance_audit_json,
    )
    feature_runs["workspace_inheritance_audit"] = require_standard_envelope(feature_runs["workspace_inheritance_audit"], inheritance_audit_json)
    outputs.extend([inheritance_audit_csv, inheritance_audit_md, inheritance_audit_json])

    if args.profile == "full":
        latex_dir = output_dir / "latex_project"
        latex_summary = output_dir / "latex_summary.json"
        feature_runs["latex_scaffold"] = run(
            [
                sys.executable,
                str(script_dir / "latex_workbench.py"),
                "scaffold",
                str(latex_dir),
                "--title",
                "Portable Smoke Report",
                "--kind",
                "academic-report",
                "--language",
                "english",
                "--summary-json",
                str(latex_summary),
            ]
        )
        feature_runs["latex_scaffold"] = require_outputs(
            feature_runs["latex_scaffold"],
            latex_dir / "main.tex",
            latex_summary,
        )
        feature_runs["latex_scaffold"] = require_standard_envelope(feature_runs["latex_scaffold"], latex_summary)
        feature_registry_ids["latex_scaffold"] = "latex_workbench.scaffold"
        outputs.extend([latex_dir / "main.tex", latex_summary])

        latex_compile_dir = output_dir / "latex_compile"
        latex_manifest = output_dir / "latex_compile_manifest.json"
        latex_before = {latex_dir / "main.tex": hashlib.sha256((latex_dir / "main.tex").read_bytes()).hexdigest()}
        feature_runs["latex_compile"] = run([
            sys.executable, str(script_dir / "latex_workbench.py"), "compile",
            str(latex_dir / "main.tex"), "--output-dir", str(latex_compile_dir),
            "--manifest-json", str(latex_manifest),
        ])
        feature_runs["latex_compile"] = require_outputs(feature_runs["latex_compile"], latex_manifest, latex_compile_dir / "main.pdf")
        feature_runs["latex_compile"] = require_unchanged(feature_runs["latex_compile"], latex_before)
        feature_registry_ids["latex_compile"] = "latex_workbench.compile"
        outputs.extend([latex_manifest, latex_compile_dir / "main.pdf"])

        latex_review_md = output_dir / "latex_review.md"
        latex_review_json = output_dir / "latex_review.json"
        feature_runs["latex_review"] = run(
            [
                sys.executable,
                str(script_dir / "latex_workbench.py"),
                "review",
                str(latex_dir / "main.tex"),
                "--report-md",
                str(latex_review_md),
                "--summary-json",
                str(latex_review_json),
            ]
        )
        feature_runs["latex_review"] = require_outputs(
            feature_runs["latex_review"],
            latex_review_md,
            latex_review_json,
        )
        feature_runs["latex_review"] = require_standard_envelope(feature_runs["latex_review"], latex_review_json)
        feature_registry_ids["latex_review"] = "latex_workbench.review"
        outputs.extend([latex_review_md, latex_review_json])

        astrometry_preflight_json = output_dir / "astrometry_preflight.json"
        feature_runs["astrometry_preflight"] = run(
            [
                sys.executable,
                str(script_dir / "astrometry_net_workbench.py"),
                "preflight",
                str(fits_path),
                "--summary-json",
                str(astrometry_preflight_json),
            ]
        )
        feature_runs["astrometry_preflight"] = require_outputs(feature_runs["astrometry_preflight"], astrometry_preflight_json)
        feature_runs["astrometry_preflight"] = require_standard_envelope(feature_runs["astrometry_preflight"], astrometry_preflight_json)
        feature_registry_ids["astrometry_preflight"] = "astrometry_net_workbench.preflight"
        outputs.append(astrometry_preflight_json)

        astrometry_existing_json = output_dir / "astrometry_existing_wcs.json"
        astrometry_existing_dir = output_dir / "astrometry_existing_wcs"
        feature_runs["astrometry_verify_existing_wcs"] = run(
            [
                sys.executable,
                str(script_dir / "astrometry_net_workbench.py"),
                "verify-existing-wcs",
                str(existing_wcs_fits),
                "--output-dir",
                str(astrometry_existing_dir),
                "--summary-json",
                str(astrometry_existing_json),
            ]
        )
        feature_runs["astrometry_verify_existing_wcs"] = require_outputs(
            feature_runs["astrometry_verify_existing_wcs"],
            astrometry_existing_json,
        )
        feature_runs["astrometry_verify_existing_wcs"] = require_standard_envelope(feature_runs["astrometry_verify_existing_wcs"], astrometry_existing_json)
        feature_registry_ids["astrometry_verify_existing_wcs"] = "astrometry_net_workbench.verify-existing-wcs"
        outputs.append(astrometry_existing_json)

        astrometry_header_only_json = output_dir / "astrometry_existing_wcs_header_only.json"
        astrometry_header_only_dir = output_dir / "astrometry_existing_wcs_header_only"
        feature_runs["astrometry_verify_existing_wcs_header_only"] = run(
            [
                sys.executable,
                str(script_dir / "astrometry_net_workbench.py"),
                "verify-existing-wcs",
                str(header_only_wcs_fits),
                "--output-dir",
                str(astrometry_header_only_dir),
                "--summary-json",
                str(astrometry_header_only_json),
            ]
        )
        feature_runs["astrometry_verify_existing_wcs_header_only"] = require_outputs(
            feature_runs["astrometry_verify_existing_wcs_header_only"],
            astrometry_header_only_json,
        )
        feature_runs["astrometry_verify_existing_wcs_header_only"] = require_standard_envelope(feature_runs["astrometry_verify_existing_wcs_header_only"], astrometry_header_only_json)
        outputs.append(astrometry_header_only_json)

        presentation_inspect_json = output_dir / "presentation_inspect.json"
        feature_runs["presentation_workbench_inspect"] = run(
            [
                sys.executable,
                str(script_dir / "presentation_workbench.py"),
                "inspect",
                str(generated_docs_dir / "sample.pptx"),
                "--summary-json",
                str(presentation_inspect_json),
            ]
        )
        feature_runs["presentation_workbench_inspect"] = require_outputs(
            feature_runs["presentation_workbench_inspect"],
            presentation_inspect_json,
        )
        feature_runs["presentation_workbench_inspect"] = require_standard_envelope(feature_runs["presentation_workbench_inspect"], presentation_inspect_json)
        feature_registry_ids["presentation_workbench_inspect"] = "presentation_workbench.inspect"
        outputs.append(presentation_inspect_json)

        presentation_style_dir = output_dir / "presentation_style_audit"
        presentation_style_json = presentation_style_dir / "summary.json"
        feature_runs["presentation_workbench_style_audit"] = run(
            [
                sys.executable,
                str(script_dir / "presentation_workbench.py"),
                "existing-deck-style-audit",
                str(generated_docs_dir / "sample.pptx"),
                str(presentation_style_dir),
                "--reference-slides",
                "1-2",
                "--rule",
                "Do not put notebook screenshots on slides.",
                "--summary-json",
                str(presentation_style_json),
            ]
        )
        feature_runs["presentation_workbench_style_audit"] = require_outputs(
            feature_runs["presentation_workbench_style_audit"],
            presentation_style_json,
            presentation_style_dir / "scientific_asset_manifest.json",
            presentation_style_dir / "presentation_constraints.json",
            presentation_style_dir / "slide_content.md",
            presentation_style_dir / "speaker_script.md",
            presentation_style_dir / "poster_text.md",
            presentation_style_dir / "style_audit_report.md",
            presentation_style_dir / "new_slide_templates.json",
        )
        feature_runs["presentation_workbench_style_audit"] = require_standard_envelope(
            feature_runs["presentation_workbench_style_audit"],
            presentation_style_json,
        )
        feature_registry_ids["presentation_workbench_style_audit"] = "presentation_workbench.existing-deck-style-audit"
        outputs.extend(
            [
                presentation_style_json,
                presentation_style_dir / "scientific_asset_manifest.json",
                presentation_style_dir / "presentation_constraints.json",
                presentation_style_dir / "slide_content.md",
                presentation_style_dir / "speaker_script.md",
                presentation_style_dir / "poster_text.md",
                presentation_style_dir / "style_audit_report.md",
                presentation_style_dir / "new_slide_templates.json",
            ]
        )

        notebook_before = {path: hashlib.sha256(path.read_bytes()).hexdigest()
                           for path in (notebook_path, standard_star_table)}
        blocked_absolute_json = output_dir / "notebook_absolute_blocked.json"
        absolute_result = run([
            sys.executable, str(script_dir / "notebook_workbench.py"), "execute-copy",
            str(notebook_path), "--output-dir", str(output_dir / "notebook_absolute_blocked"),
            "--trust-notebook-code", "--summary-json", str(blocked_absolute_json),
        ])
        feature_runs["notebook_rejects_absolute_input"] = require_unchanged(
            require_controlled_rejection(absolute_result, blocked_absolute_json, "literal absolute paths escape"),
            notebook_before,
        )
        outputs.append(blocked_absolute_json)

        # Adapt a new synthetic fixture copy; retain the bootstrap notebook as
        # the negative case above and stage its data through the supported CLI.
        execution_notebook = output_dir / "portable_execution_notebook.ipynb"
        notebook_copy = json.loads(notebook_path.read_text())
        for cell in notebook_copy["cells"]:
            source = cell.get("source", "")
            if isinstance(source, list):
                cell["source"] = [line.replace(str(standard_star_table), standard_star_table.name) for line in source]
            else:
                cell["source"] = source.replace(str(standard_star_table), standard_star_table.name)
        execution_notebook.write_text(json.dumps(notebook_copy, indent=2) + "\n")
        notebook_before[execution_notebook] = hashlib.sha256(execution_notebook.read_bytes()).hexdigest()

        notebook_run_dir = output_dir / "notebook_workbench"
        notebook_run_json = notebook_run_dir / "summary.json"
        feature_runs["notebook_workbench"] = run(
            [
                sys.executable,
                str(script_dir / "notebook_workbench.py"),
                "execute-copy",
                str(execution_notebook),
                "--output-dir",
                str(notebook_run_dir),
                "--trust-notebook-code",
                "--stage-extra",
                str(standard_star_table),
                "--append-code",
                "smoke_value = 6 * 7\nprint(smoke_value)",
                "--summary-json",
                str(notebook_run_json),
            ]
        )
        feature_runs["notebook_workbench"] = require_outputs(
            feature_runs["notebook_workbench"],
            notebook_run_json,
            notebook_run_dir / execution_notebook.name,
            notebook_run_dir / "_workspace" / standard_star_table.name,
        )
        feature_runs["notebook_workbench"] = require_standard_envelope(feature_runs["notebook_workbench"], notebook_run_json)
        feature_runs["notebook_workbench"] = require_unchanged(feature_runs["notebook_workbench"], notebook_before)
        feature_registry_ids["notebook_workbench"] = "notebook_workbench.execute-copy"
        outputs.extend([notebook_run_json, notebook_run_dir / execution_notebook.name])

        relative_nb_dir = generated_data_dir / "relative_notebook"
        relative_notebook_path, relative_data_path = create_relative_data_notebook(relative_nb_dir)
        relative_before = {path: hashlib.sha256(path.read_bytes()).hexdigest()
                           for path in (relative_notebook_path, relative_data_path)}
        relative_preflight_json = output_dir / "notebook_preflight_execution.json"
        feature_runs["notebook_workbench_preflight_execution"] = run(
            [
                sys.executable,
                str(script_dir / "notebook_workbench.py"),
                "preflight-execution",
                str(relative_notebook_path),
                "--summary-json",
                str(relative_preflight_json),
            ]
        )
        feature_runs["notebook_workbench_preflight_execution"] = require_outputs(
            feature_runs["notebook_workbench_preflight_execution"],
            relative_preflight_json,
        )
        feature_runs["notebook_workbench_preflight_execution"] = require_standard_envelope(feature_runs["notebook_workbench_preflight_execution"], relative_preflight_json)
        outputs.append(relative_preflight_json)

        relative_notebook_run_dir = output_dir / "notebook_workbench_relative"
        relative_notebook_run_json = relative_notebook_run_dir / "summary.json"
        blocked_source_json = output_dir / "notebook_source_dir_blocked.json"
        source_result = run([
            sys.executable, str(script_dir / "notebook_workbench.py"), "execute-copy",
            str(relative_notebook_path), "--output-dir", str(output_dir / "notebook_source_dir_blocked"),
            "--trust-notebook-code", "--run-from-source-dir", "--summary-json", str(blocked_source_json),
        ])
        feature_runs["notebook_rejects_source_dir"] = require_unchanged(
            require_controlled_rejection(source_result, blocked_source_json, "--run-from-source-dir is blocked"),
            relative_before,
        )
        outputs.append(blocked_source_json)

        feature_runs["notebook_workbench_staged_relative_input"] = run(
            [
                sys.executable,
                str(script_dir / "notebook_workbench.py"),
                "execute-copy",
                str(relative_notebook_path),
                "--output-dir",
                str(relative_notebook_run_dir),
                "--trust-notebook-code",
                "--stage-extra",
                str(relative_data_path),
                "--summary-json",
                str(relative_notebook_run_json),
            ]
        )
        feature_runs["notebook_workbench_staged_relative_input"] = require_outputs(
            feature_runs["notebook_workbench_staged_relative_input"],
            relative_notebook_run_json,
            relative_notebook_run_dir / "relative_data_notebook.ipynb",
            relative_notebook_run_dir / "_workspace" / relative_data_path.name,
        )
        feature_runs["notebook_workbench_staged_relative_input"] = require_standard_envelope(feature_runs["notebook_workbench_staged_relative_input"], relative_notebook_run_json)
        feature_runs["notebook_workbench_staged_relative_input"] = require_unchanged(feature_runs["notebook_workbench_staged_relative_input"], relative_before)
        outputs.extend([relative_notebook_run_json, relative_notebook_run_dir / "relative_data_notebook.ipynb"])

        shifted_spectrum = create_shifted_spectrum(
            examples_dir / "science" / "synthetic_spectrum.txt",
            generated_data_dir / "synthetic_spectrum_shifted.txt",
        )
        coursework_dir = output_dir / "spectra_ascii_coursework"
        coursework_json = coursework_dir / "summary.json"
        feature_runs["spectra_ascii_coursework"] = run(
            [
                sys.executable,
                str(script_dir / "spectra_ascii_coursework_workbench.py"),
                str(examples_dir / "science" / "synthetic_spectrum.txt"),
                str(shifted_spectrum),
                "--output-dir",
                str(coursework_dir),
                "--notebook",
                str(execution_notebook),
                "--execute-notebook",
                "--stage-extra",
                str(standard_star_table),
                "--line-window",
                "demo_feature",
                "6560",
                "6570",
                "--summary-json",
                str(coursework_json),
            ]
        )
        feature_runs["spectra_ascii_coursework"] = require_outputs(
            feature_runs["spectra_ascii_coursework"],
            coursework_json,
            coursework_dir / "spectra_analysis" / "spectrum_inventory.csv",
            coursework_dir / "spectra_analysis" / "overlay_raw.png",
            coursework_dir / "spectra_analysis" / "overlay_normalized.png",
            coursework_dir / "coursework_bundle_report.md",
        )
        feature_runs["spectra_ascii_coursework"] = require_standard_envelope(feature_runs["spectra_ascii_coursework"], coursework_json)
        feature_runs["spectra_ascii_coursework"] = require_unchanged(feature_runs["spectra_ascii_coursework"], notebook_before)
        feature_registry_ids["spectra_ascii_coursework"] = "spectra_ascii_coursework_workbench"
        outputs.extend(
            [
                coursework_json,
                coursework_dir / "spectra_analysis" / "spectrum_inventory.csv",
                coursework_dir / "spectra_analysis" / "overlay_raw.png",
                coursework_dir / "spectra_analysis" / "overlay_normalized.png",
                coursework_dir / "coursework_bundle_report.md",
            ]
        )

        if sys.platform == "darwin":
            quicklook_json = output_dir / "quicklook_bridge.json"
            quicklook_png = output_dir / "quicklook_bridge.png"
            feature_runs["quicklook_bridge"] = run(
                [
                    sys.executable,
                    str(script_dir / "quicklook_bridge.py"),
                    str(generated_docs_dir / "sample.pptx"),
                    "--output",
                    str(quicklook_png),
                    "--summary-json",
                    str(quicklook_json),
                ]
            )
            feature_runs["quicklook_bridge"] = require_outputs(feature_runs["quicklook_bridge"], quicklook_json, quicklook_png)
            feature_runs["quicklook_bridge"] = require_standard_envelope(feature_runs["quicklook_bridge"], quicklook_json)
            feature_registry_ids["quicklook_bridge"] = "quicklook_bridge"
            outputs.extend([quicklook_json, quicklook_png])

    from _internal.public_surface_registry import coverage_matrix, load_public_surface_registry

    registry_entries = load_public_surface_registry()
    smoke_matrix = coverage_matrix(registry_entries)
    coverage_rows = []
    run_lookup = {registry_id: run_name for run_name, registry_id in feature_registry_ids.items()}
    tiers_to_enforce = ["core"] if args.profile == "core" else ["core", "full"]
    missing_coverage = []
    for tier, items in smoke_matrix.items():
        for item in items:
            run_name = run_lookup.get(item["id"])
            covered = run_name is not None
            returncode = feature_runs[run_name]["returncode"] if covered else None
            row = dict(item)
            row.update({"covered": covered, "feature_run": run_name, "returncode": returncode})
            coverage_rows.append(row)
            if tier in tiers_to_enforce and not covered:
                missing_coverage.append(item["id"])
    coverage_json = output_dir / "public_surface_coverage.json"
    coverage_json.write_text(json.dumps({"profile": args.profile, "rows": coverage_rows}, indent=2, ensure_ascii=True) + "\n", encoding="utf-8")
    outputs.append(coverage_json)

    overall_ok = all(entry["returncode"] == 0 for entry in feature_runs.values()) and not missing_coverage
    legacy_summary = {
        "tool": "portable_smoke_test",
        "environment": environment_summary(),
        "skill_root": str(root_dir),
        "examples_dir": str(examples_dir),
        "profile": args.profile,
        "generated_inputs": {
            "documents": [public_path(path) for path in generated_docs],
            "fits": public_path(fits_path),
            "linear_multispec": public_path(linear_multispec),
            "existing_wcs_fits": public_path(existing_wcs_fits),
            "header_only_wcs_fits": public_path(header_only_wcs_fits),
            "standard_star_table": public_path(standard_star_table),
        },
        "feature_runs": feature_runs,
        "public_surface_coverage": {
            "tiers_enforced": tiers_to_enforce,
            "missing_coverage": missing_coverage,
            "coverage_json": public_path(coverage_json),
        },
        "overall_status": "PASS" if overall_ok else "FAIL",
        "notes": [
            "This smoke test is designed for fresh machines and portable verification.",
            "It uses packaged examples plus generated core-format samples instead of relying on user datasets.",
        ],
    }
    qa = {
        "status": "ok" if overall_ok else "fail",
        "findings": [
            name for name, entry in feature_runs.items() if entry.get("returncode") != 0
        ]
        + [f"missing coverage: {item}" for item in missing_coverage],
        "metrics": {
            "feature_count": len(feature_runs),
            "pass_count": sum(1 for entry in feature_runs.values() if entry.get("returncode") == 0),
            "coverage_row_count": len(coverage_rows),
            "missing_coverage_count": len(missing_coverage),
        },
    }
    summary = standard_tool_payload(
        "portable_smoke_test",
        status="ok" if overall_ok else "fail",
        notes=legacy_summary["notes"],
        artifacts={"summary_json": args.summary_json, "output_dir": output_dir},
        results=legacy_summary,
        qa=qa,
        legacy=legacy_summary,
    )

    if args.summary_json:
        summary_path = Path(args.summary_json)
        outputs.append(summary_path)
    if args.manifest_json:
        write_manifest(
            args.manifest_json,
            inputs=[examples_dir],
            outputs=outputs,
            parameters={"profile": args.profile},
            command="portable_smoke_test.py",
            notes=summary["notes"],
            extra={"overall_status": summary["overall_status"]},
        )
    emit_payload_safely(summary, args.summary_json)

    return 0 if overall_ok else 1


def main():
    args = parse_args()
    try:
        exit_code = run_smoke(args)
    except Exception as exc:
        blocked = isinstance(exc, PathSafetyError)
        payload = build_failure_payload(args, exc, blocked=blocked)
        summary_target = None if blocked else args.summary_json
        try:
            emit_payload_safely(payload, summary_target)
        except Exception as emit_exc:
            payload["notes"].append(str(emit_exc))
            payload["results"]["summary_json_error"] = str(emit_exc)
            emit_payload_safely(payload, None)
        raise SystemExit(2 if blocked else 1) from None
    raise SystemExit(exit_code)


if __name__ == "__main__":
    main()
