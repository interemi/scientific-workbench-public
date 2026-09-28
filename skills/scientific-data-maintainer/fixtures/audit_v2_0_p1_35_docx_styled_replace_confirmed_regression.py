#!/usr/bin/env python3
"""Regression for v2.0 P1-35 confirmed DOCX styled replacement."""

from __future__ import annotations

import hashlib
import json
import shutil
import subprocess
import sys
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
TMP = ROOT / "tmp" / "v2_0_p1_35_docx_styled_replace_confirmed"
APP = Path(__file__).resolve().parents[3]
OFFICE = ROOT / "scripts" / "office_roundtrip.py"
CONFIRMATION_ID = "p1-35-regression-reviewed"


def require(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def run(args: list[str], *, expect: set[int] = {0}) -> subprocess.CompletedProcess[str]:
    completed = subprocess.run(args, cwd=ROOT, text=True, capture_output=True, check=False)
    require(
        completed.returncode in expect,
        f"Command failed ({completed.returncode}): {' '.join(args)}\nSTDOUT:\n{completed.stdout}\nSTDERR:\n{completed.stderr}",
    )
    require("Traceback" not in completed.stdout + completed.stderr, "Raw traceback leaked")
    return completed


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def make_fixture(python: str, target: Path) -> None:
    script = TMP / "make_fixture.py"
    script.write_text(
        """
from docx import Document
import sys

doc = Document()
doc.add_heading("Reviewed styled replacement", level=1)
p = doc.add_paragraph()
p.add_run("Keep ordinary text. ")
r = p.add_run("APPROVED draft")
r.bold = True
r.italic = True
r.underline = True
p.add_run(" plain draft remains. ")
r = p.add_run("BOLD draft remains")
r.bold = True
table = doc.add_table(rows=1, cols=1)
r = table.cell(0, 0).paragraphs[0].add_run("TABLE draft")
r.bold = True
r.italic = True
r.underline = True
doc.save(sys.argv[1])
""".lstrip(),
        encoding="utf-8",
    )
    run([python, str(script), str(target)])


def read_docx(python: str, path: Path) -> dict:
    script = TMP / "read_docx.py"
    script.write_text(
        """
from docx import Document
import json
import sys

doc = Document(sys.argv[1])
rows = []
for location, paragraphs in (
    ("body", doc.paragraphs),
    ("table", [p for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs]),
):
    for paragraph_index, paragraph in enumerate(paragraphs, start=1):
        for run_index, run in enumerate(paragraph.runs, start=1):
            if run.text:
                rows.append({
                    "location": location,
                    "paragraph_index": paragraph_index,
                    "run_index": run_index,
                    "text": run.text,
                    "bold": run.bold is True,
                    "italic": run.italic is True,
                    "underline": run.underline is not None and run.underline is not False,
                })
print(json.dumps(rows))
""".lstrip(),
        encoding="utf-8",
    )
    return {"runs": json.loads(run([python, str(script), str(path)]).stdout)}


def main() -> int:
    if TMP.exists():
        shutil.rmtree(TMP)
    (TMP / "originals").mkdir(parents=True)
    (TMP / "inputs").mkdir(parents=True)
    (TMP / "runs" / "inventory" / "tables").mkdir(parents=True)
    (TMP / "runs" / "inventory" / "reports").mkdir(parents=True)
    (TMP / "runs" / "blocked_without_confirmation").mkdir(parents=True)
    (TMP / "runs" / "confirmed" / "artifacts").mkdir(parents=True)
    (TMP / "runs" / "confirmed" / "previews").mkdir(parents=True)

    python = sys.executable
    try:
        import docx  # noqa: F401
    except ImportError:
        python = str(Path.home() / "anaconda3/envs/datanalysis/bin/python3")

    original = TMP / "originals" / "reviewed source.docx"
    copied = TMP / "inputs" / "reviewed source copy.docx"
    make_fixture(python, original)
    original_hash = sha256(original)
    shutil.copy2(original, copied)
    copied_hash = sha256(copied)

    inventory_dir = TMP / "runs" / "inventory"
    inventory_summary = inventory_dir / "summary.json"
    run(
        [
            python,
            str(OFFICE),
            "docx-style-inventory",
            str(copied),
            "--require-bold",
            "--require-italic",
            "--require-underline",
            "--output-csv",
            str(inventory_dir / "tables/style_inventory.csv"),
            "--report-md",
            str(inventory_dir / "reports/style_inventory.md"),
            "--summary-json",
            str(inventory_summary),
        ]
    )
    inventory = load(inventory_summary)
    matches = inventory["results"]["inventory"]["matches"]
    require(len(matches) == 2, f"Expected two reviewed runs, got {len(matches)}")

    preview_changes = [
        {
            "location": item["location"],
            "paragraph_index": item["paragraph_index"],
            "run_index": item["run_index"],
            "before_text": item["text"],
            "after_text": item["text"].replace("draft", "final"),
        }
        for item in matches
        if "draft" in item["text"]
    ]
    dry_run_preview = TMP / "runs" / "inventory" / "previews" / "reviewed_diff.json"
    dry_run_preview.parent.mkdir(parents=True)
    dry_run_preview.write_text(
        json.dumps(
            {
                "artifact_type": "app_preview",
                "status": "DRY_RUN",
                "changes": preview_changes,
                "confirmation_required": True,
                "original_modified": False,
            },
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    require(len(preview_changes) == 2, "Dry-run preview did not contain both reviewed changes")

    blocked_dir = TMP / "runs" / "blocked_without_confirmation"
    blocked_output = blocked_dir / "edited_without_confirmation.docx"
    blocked_summary = blocked_dir / "summary.json"
    run(
        [
            python,
            str(OFFICE),
            "docx-styled-replace",
            str(copied),
            str(blocked_output),
            "--find",
            "draft",
            "--replace",
            "final",
            "--require-bold",
            "--require-italic",
            "--require-underline",
            "--require-confirmation",
            "--summary-json",
            str(blocked_summary),
        ],
        expect={2},
    )
    blocked = load(blocked_summary)
    require(blocked["app_status"] == "FAIL", "Missing confirmation did not fail cleanly")
    require(not blocked_output.exists(), "Unconfirmed workflow wrote an edited document")

    confirmed_dir = TMP / "runs" / "confirmed"
    edited = confirmed_dir / "artifacts" / "edited document.docx"
    diff = confirmed_dir / "previews" / "style diff.json"
    summary_path = confirmed_dir / "summary.json"
    manifest_path = confirmed_dir / "manifest.json"
    run(
        [
            python,
            str(OFFICE),
            "docx-styled-replace",
            str(copied),
            str(edited),
            "--find",
            "draft",
            "--replace",
            "final",
            "--require-bold",
            "--require-italic",
            "--require-underline",
            "--require-confirmation",
            "--confirmation-id",
            CONFIRMATION_ID,
            "--diff-json",
            str(diff),
            "--summary-json",
            str(summary_path),
            "--manifest-json",
            str(manifest_path),
        ]
    )

    summary = load(summary_path)
    diff_payload = load(diff)
    manifest = load(manifest_path)
    replacement = summary["results"]["replacement"]
    require(summary["app_status"] == "PASS", f"Confirmed replacement failed: {summary}")
    require(replacement["replacements"] == 2, "Confirmed replacement count is wrong")
    require(replacement["readback_replacement_occurrence_count"] == 2, "Readback did not verify both replacements")
    require(
        replacement["confirmation"]
        == {
            "required": True,
            "enforced_by_backend": True,
            "recorded": True,
            "confirmation_id": CONFIRMATION_ID,
        },
        f"Confirmation trace is incomplete: {replacement['confirmation']}",
    )
    require(diff_payload["review_confirmed"] is True, "Diff does not record confirmation")
    require(diff_payload["confirmation_id"] == CONFIRMATION_ID, "Diff confirmation id is wrong")
    require(diff_payload["changes"] == [
        {
            **preview,
            "table_index": match["table_index"],
            "row_index": match["row_index"],
            "cell_index": match["cell_index"],
            "replacements": 1,
            "bold": True,
            "italic": True,
            "underline": True,
        }
        for preview, match in zip(preview_changes, matches)
    ], "Backend diff diverges from the reviewed preview")

    artifact_types = {item["artifact_type"] for item in summary["typed_artifacts"]}
    require(
        {"edited_document", "app_preview", "summary_json", "manifest_json"} <= artifact_types,
        f"Typed artifacts are incomplete: {artifact_types}",
    )
    manifest_outputs = {Path(item["path"]).name: item for item in manifest["outputs"]}
    require("edited document.docx" in manifest_outputs, "Manifest omits edited document")
    require("style diff.json" in manifest_outputs, "Manifest omits reviewed diff")
    require(manifest_outputs["edited document.docx"]["sha256"], "Edited document lacks manifest hash")
    require(manifest_outputs["style diff.json"]["sha256"], "Diff lacks manifest hash")
    require(manifest["parameters"]["confirmation"]["confirmation_id"] == CONFIRMATION_ID, "Manifest omits confirmation")

    output_runs = read_docx(python, edited)["runs"]
    output_text = " | ".join(item["text"] for item in output_runs)
    require("APPROVED final" in output_text, "Styled body replacement missing")
    require("TABLE final" in output_text, "Styled table replacement missing")
    require("plain draft remains" in output_text, "Plain text was incorrectly replaced")
    require("BOLD draft remains" in output_text, "Non-matching styled text was incorrectly replaced")
    changed_runs = [item for item in output_runs if item["text"] in {"APPROVED final", "TABLE final"}]
    require(len(changed_runs) == 2, "Could not identify both changed runs")
    require(all(item["bold"] and item["italic"] and item["underline"] for item in changed_runs), "Formatting was not preserved")

    require(sha256(original) == original_hash, "Original DOCX was modified")
    require(sha256(copied) == copied_hash, "Staged input copy was modified")
    require(sha256(edited) != copied_hash, "Edited document is byte-identical to input")
    require(summary["original_modified"] is False, "Summary claims original modification")

    app_checks = {
        "Sources/ScientificWorkbench/Services/DocumentWorkflowService.swift": [
            "validateReplacementApproval",
            "confirmationRequired",
            "noReviewedChanges",
        ],
        "Sources/ScientificWorkbench/Stores/WorkbenchStore.swift": [
            "--require-confirmation",
            "--confirmation-id",
            "scientific-workbench:",
        ],
        "Sources/ScientificWorkbench/Views/DocumentRoundtripView.swift": [
            "I reviewed the matches and approve",
            "!store.docxReplacementConfirmed",
            "Replacement Preview",
        ],
        "Tests/ScientificWorkbenchTests/DocumentWorkflowTests.swift": [
            "documentWorkflowRequiresReviewedConfirmationBeforeReplacement",
            "artifactDiscoveryReadsDocumentV2Artifacts",
        ],
    }
    for relative, markers in app_checks.items():
        text = (APP / relative).read_text(encoding="utf-8")
        for marker in markers:
            require(marker in text, f"{relative} lacks {marker}")

    result = {
        "tool": "audit_v2_0_p1_35_docx_styled_replace_confirmed_regression",
        "status": "PASS",
        "original_modified": False,
        "inventory_match_count": len(matches),
        "blocked_without_confirmation": {
            "status": blocked["app_status"],
            "output_exists": blocked_output.exists(),
        },
        "confirmed_replacement": {
            "confirmation_id": CONFIRMATION_ID,
            "replacements": replacement["replacements"],
            "readback_occurrences": replacement["readback_replacement_occurrence_count"],
            "artifact_types": sorted(artifact_types),
        },
        "fingerprints": {
            "original_before": original_hash,
            "original_after": sha256(original),
            "copy_before": copied_hash,
            "copy_after": sha256(copied),
            "edited": sha256(edited),
        },
        "artifacts": {
            "edited_document": str(edited),
            "diff_preview": str(diff),
            "dry_run_preview": str(dry_run_preview),
            "summary_json": str(summary_path),
            "manifest_json": str(manifest_path),
        },
        "warnings": replacement["limitations"],
    }
    output = TMP / "p1_35_summary.json"
    output.write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
