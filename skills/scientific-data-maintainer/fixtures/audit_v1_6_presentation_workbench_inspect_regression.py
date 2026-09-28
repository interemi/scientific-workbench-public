#!/usr/bin/env python3
"""Regression checks for presentation_workbench.py inspect v1.6 hardening."""

from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "presentation_workbench.py"
TMP = Path("/tmp/scientific-data-analysis-v1-6-presentation-inspect-regression")


def run(cmd: list[str]) -> subprocess.CompletedProcess[str]:
    return subprocess.run(cmd, cwd=ROOT, text=True, encoding="utf-8", errors="replace", capture_output=True, check=False)


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def assert_no_traceback(completed: subprocess.CompletedProcess[str]) -> None:
    combined = f"{completed.stdout}\n{completed.stderr}"
    assert "Traceback (most recent call last)" not in combined, combined


def make_image(path: Path) -> None:
    image = Image.new("RGB", (900, 520), (245, 247, 250))
    draw = ImageDraw.Draw(image)
    draw.rectangle([40, 40, 860, 480], outline=(35, 88, 145), width=6)
    draw.line([80, 420, 820, 110], fill=(230, 170, 60), width=8)
    draw.text((70, 65), "regression figure", fill=(20, 20, 20))
    image.save(path)


def add_textbox(slide, text: str, left: float, top: float, width: float, height: float, size: int) -> None:
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    shape.text_frame.text = text
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(size)


def make_deck(path: Path, image_path: Path, *, flattened: bool = False) -> None:
    presentation = Presentation()
    blank = presentation.slide_layouts[6]
    slide = presentation.slides.add_slide(blank)
    if flattened:
        slide.shapes.add_picture(str(image_path), 0, 0, width=presentation.slide_width, height=presentation.slide_height)
    else:
        add_textbox(slide, "Regression Slide", 0.5, 0.25, 8.5, 0.55, 28)
        slide.shapes.add_picture(str(image_path), Inches(1.0), Inches(1.15), width=Inches(7.4))
        add_textbox(slide, "Figure 1. Editable caption.", 0.8, 6.35, 8.0, 0.35, 12)
    presentation.save(path)


def assert_payload(path: Path, *, status: str, qa_status: str) -> dict:
    payload = load_json(path)
    assert payload["status"] == status, payload
    assert payload["qa"]["status"] == qa_status, payload
    return payload


def main() -> int:
    if TMP.exists():
        shutil.rmtree(TMP)
    TMP.mkdir(parents=True)
    image_path = TMP / "figure.png"
    make_image(image_path)
    happy = TMP / "happy.pptx"
    flattened = TMP / "flattened.pptx"
    corrupt = TMP / "corrupt.pptx"
    make_deck(happy, image_path)
    make_deck(flattened, image_path, flattened=True)
    corrupt.write_bytes(b"not a real presentation\n")

    happy_summary = TMP / "happy_summary.json"
    completed = run([sys.executable, str(SCRIPT), "inspect", str(happy), "--summary-json", str(happy_summary)])
    assert completed.returncode == 0, completed.stderr
    assert_no_traceback(completed)
    assert_payload(happy_summary, status="ok", qa_status="ok")

    flattened_summary = TMP / "flattened_summary.json"
    completed = run([sys.executable, str(SCRIPT), "inspect", str(flattened), "--summary-json", str(flattened_summary)])
    assert completed.returncode == 0, completed.stderr
    assert_no_traceback(completed)
    assert_payload(flattened_summary, status="warning", qa_status="warning")

    missing_summary = TMP / "missing_summary.json"
    completed = run([sys.executable, str(SCRIPT), "inspect", str(TMP / "missing.pptx"), "--summary-json", str(missing_summary)])
    assert completed.returncode == 2, completed.stdout + completed.stderr
    assert_no_traceback(completed)
    assert_payload(missing_summary, status="blocked", qa_status="blocked")

    corrupt_summary = TMP / "corrupt_summary.json"
    completed = run([sys.executable, str(SCRIPT), "inspect", str(corrupt), "--summary-json", str(corrupt_summary)])
    assert completed.returncode == 2, completed.stdout + completed.stderr
    assert_no_traceback(completed)
    assert_payload(corrupt_summary, status="blocked", qa_status="blocked")

    conflict_parent = TMP / "preview_parent_is_file"
    conflict_parent.write_text("not a directory\n", encoding="utf-8")
    preview_summary = TMP / "preview_summary.json"
    completed = run(
        [
            sys.executable,
            str(SCRIPT),
            "inspect",
            str(happy),
            "--summary-json",
            str(preview_summary),
            "--export-preview-pdf",
            str(conflict_parent / "preview.pdf"),
        ]
    )
    assert completed.returncode == 0, completed.stdout + completed.stderr
    assert_no_traceback(completed)
    preview_payload = assert_payload(preview_summary, status="warning", qa_status="warning")
    assert preview_payload["results"]["inspection"]["preview_export"]["success"] is False, preview_payload

    sys.path.insert(0, str(ROOT / "scripts"))
    import office_roundtrip

    original_detect = office_roundtrip.detect_presentation_export_backends
    try:
        office_roundtrip.detect_presentation_export_backends = lambda: {"soffice": None, "keynote_available": False}
        result = office_roundtrip.export_pptx_pdf(happy, TMP / "direct_backend_probe.pdf", backend="soffice")
        assert result["success"] is False, result
        assert "not available" in result["stderr"], result
    finally:
        office_roundtrip.detect_presentation_export_backends = original_detect

    print("presentation_workbench inspect regression: PASS")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
