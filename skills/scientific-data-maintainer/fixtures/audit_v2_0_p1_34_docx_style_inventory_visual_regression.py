#!/usr/bin/env python3
"""Regression for the v2.0 P1-34 visual DOCX style inventory workflow."""

from __future__ import annotations

import hashlib
import json
import shutil
import subprocess
import sys
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
TMP = ROOT / "tmp" / "v2_0_p1_34_docx_style_inventory_visual"
APP = Path(__file__).resolve().parents[3]
OFFICE = ROOT / "scripts" / "office_roundtrip.py"


def require(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def run(args: list[str]) -> subprocess.CompletedProcess[str]:
    completed = subprocess.run(args, cwd=ROOT, text=True, capture_output=True, check=False)
    require(
        completed.returncode == 0,
        f"Command failed ({completed.returncode}): {' '.join(args)}\n{completed.stdout}\n{completed.stderr}",
    )
    require("Traceback" not in completed.stdout + completed.stderr, "Raw traceback leaked")
    return completed


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def make_fixture(python: str, target: Path) -> None:
    script = TMP / "make_varied_docx.py"
    script.write_text(
        """
from docx import Document
import sys

doc = Document()
doc.add_heading("Varied explicit styles", level=1)
p = doc.add_paragraph()
p.add_run("Plain body. ")
r = p.add_run("Bold body")
r.bold = True
p.add_run(" | ")
r = p.add_run("Italic body")
r.italic = True
p.add_run(" | ")
r = p.add_run("Underline body")
r.underline = True
p.add_run(" | ")
r = p.add_run("Bold italic body")
r.bold = True
r.italic = True
table = doc.add_table(rows=1, cols=1)
r = table.cell(0, 0).paragraphs[0].add_run("Bold underline table")
r.bold = True
r.underline = True
doc.save(sys.argv[1])
""".lstrip(),
        encoding="utf-8",
    )
    run([python, str(script), str(target)])


def inventory(python: str, copied: Path, name: str, flags: list[str]) -> dict:
    run_dir = TMP / "runs" / name
    (run_dir / "tables").mkdir(parents=True)
    (run_dir / "reports").mkdir(parents=True)
    summary = run_dir / "summary.json"
    run(
        [
            python,
            str(OFFICE),
            "docx-style-inventory",
            str(copied),
            *flags,
            "--output-csv",
            str(run_dir / "tables/style_inventory.csv"),
            "--report-md",
            str(run_dir / "reports/style_inventory.md"),
            "--summary-json",
            str(summary),
        ]
    )
    return load(summary)


def main() -> int:
    if TMP.exists():
        shutil.rmtree(TMP)
    (TMP / "originals").mkdir(parents=True)
    (TMP / "inputs").mkdir(parents=True)

    python = sys.executable
    try:
        import docx  # noqa: F401
    except ImportError:
        python = str(Path.home() / "anaconda3/envs/datanalysis/bin/python3")

    original = TMP / "originals" / "varied explicit styles.docx"
    copied = TMP / "inputs" / "varied explicit styles copy.docx"
    make_fixture(python, original)
    original_hash = sha256(original)
    shutil.copy2(original, copied)

    any_style = inventory(python, copied, "any_special", [])
    bold = inventory(python, copied, "bold", ["--require-bold"])
    bold_italic = inventory(
        python,
        copied,
        "bold_italic",
        ["--require-bold", "--require-italic"],
    )

    any_inventory = any_style["results"]["inventory"]
    require(any_style["app_status"] == "PASS", "Any-style inventory did not pass")
    require(any_inventory["match_count"] == 5, "Expected five explicitly styled runs")
    require(
        any_inventory["style_counts"]
        == {"bold": 3, "italic": 2, "underline": 2, "body": 4, "table": 1},
        f"Unexpected style counts: {any_inventory['style_counts']}",
    )
    require(bold["results"]["inventory"]["match_count"] == 3, "Bold selector count is wrong")
    require(bold_italic["results"]["inventory"]["match_count"] == 1, "Combined selector count is wrong")
    require(original_hash == sha256(original), "Original DOCX was modified")
    require(original_hash == sha256(copied), "Inventory modified the staged copy")
    require(not list(TMP.rglob("edited*.docx")), "Inventory produced an edited DOCX")

    artifact_types = {item["artifact_type"] for item in any_style["typed_artifacts"]}
    require(
        {"summary_json", "table_csv", "report_md"} <= artifact_types,
        f"Inventory artifact typing incomplete: {artifact_types}",
    )

    app_checks = {
        "Sources/ScientificWorkbench/Views/DocumentRoundtripView.swift": [
            "styleCountSummary",
            "Matching Runs",
            "Inspect Copy",
        ],
        "Sources/ScientificWorkbench/Models/DocumentWorkflowModels.swift": [
            "DOCXStyleCounts",
            "styleCounts",
        ],
        "Tests/ScientificWorkbenchTests/DocumentWorkflowTests.swift": [
            "documentWorkflowSummarizesVisibleStyleCounts",
            "artifactDiscoveryReadsDocumentV2Artifacts",
        ],
    }
    for relative, markers in app_checks.items():
        text = (APP / relative).read_text(encoding="utf-8")
        for marker in markers:
            require(marker in text, f"{relative} lacks {marker}")

    summary = {
        "tool": "audit_v2_0_p1_34_docx_style_inventory_visual_regression",
        "status": "PASS",
        "original_modified": False,
        "fixture": str(original),
        "style_table": [
            {"selector": "Any explicit style", "matches": 5, **any_inventory["style_counts"]},
            {
                "selector": "Bold",
                "matches": bold["results"]["inventory"]["match_count"],
                **bold["results"]["inventory"]["style_counts"],
            },
            {
                "selector": "Bold + Italic",
                "matches": bold_italic["results"]["inventory"]["match_count"],
                **bold_italic["results"]["inventory"]["style_counts"],
            },
        ],
        "preview_matches": [
            {
                "text": item["text"],
                "location": item["location"],
                "styles": [
                    name
                    for name in ("bold", "italic", "underline")
                    if item[name]
                ],
            }
            for item in any_inventory["matches"]
        ],
        "artifacts": any_style["typed_artifacts"],
        "warnings": any_inventory["limitations"],
    }
    output = TMP / "p1_34_summary.json"
    output.write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(summary, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
