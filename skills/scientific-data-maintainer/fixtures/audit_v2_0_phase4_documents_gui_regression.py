#!/usr/bin/env python3
"""Regression for v2.0 Phase 4 document, iWork, and GUI workflows."""

from __future__ import annotations

import hashlib
import json
import shutil
import subprocess
import sys
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
TMP = ROOT / "tmp" / "v2_0_phase4_documents_gui"
APP = Path(__file__).resolve().parents[3]
OFFICE = ROOT / "scripts" / "office_roundtrip.py"
KEYNOTE = ROOT / "scripts" / "keynote_export.py"
REFERENCE = ROOT / "references" / "v2-0-phase4-documents-gui.md"


def require(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def run(args: list[str], *, expect: set[int] = {0}) -> subprocess.CompletedProcess[str]:
    completed = subprocess.run(
        args,
        cwd=ROOT,
        text=True,
        capture_output=True,
        check=False,
    )
    require(
        completed.returncode in expect,
        f"Command failed ({completed.returncode}): {' '.join(args)}\nSTDOUT:\n{completed.stdout}\nSTDERR:\n{completed.stderr}",
    )
    require("Traceback" not in completed.stdout + completed.stderr, f"Raw traceback leaked: {' '.join(args)}")
    return completed


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def load(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def make_fixtures(python: str) -> tuple[Path, Path]:
    originals = TMP / "originals"
    originals.mkdir(parents=True, exist_ok=True)
    docx_path = originals / "controlled document.docx"
    pptx_path = originals / "controlled deck.pptx"
    script = TMP / "make_fixtures.py"
    script.write_text(
        """
from docx import Document
from pptx import Presentation
from pptx.util import Inches
import sys

doc = Document()
doc.add_heading("Controlled document", level=1)
p = doc.add_paragraph()
p.add_run("Keep ordinary text. ")
r = p.add_run("APPROVED draft")
r.bold = True
r.italic = True
r.underline = True
p.add_run(" after.")
table = doc.add_table(rows=1, cols=1)
r = table.cell(0, 0).paragraphs[0].add_run("APPROVED table draft")
r.bold = True
r.italic = True
r.underline = True
doc.save(sys.argv[1])

deck = Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[5])
box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
box.text_frame.text = "Controlled Keynote export fixture"
deck.save(sys.argv[2])
""".lstrip(),
        encoding="utf-8",
    )
    run([python, str(script), str(docx_path), str(pptx_path)])
    return docx_path, pptx_path


def artifact_types(payload: dict) -> set[str]:
    return {item.get("artifact_type") for item in payload.get("typed_artifacts") or []}


def validate_app_surface() -> None:
    required = {
        "Sources/ScientificWorkbench/Models/DocumentWorkflowModels.swift": [
            "DOCXStyleReview",
            "KeynotePreflightReview",
        ],
        "Sources/ScientificWorkbench/Services/DocumentWorkflowService.swift": [
            "stageDOCX",
            "stagePresentation",
            "originalIsUnchanged",
        ],
        "Sources/ScientificWorkbench/Views/DocumentRoundtripView.swift": [
            "Replacement Preview",
            "Create Edited Copy",
        ],
        "Sources/ScientificWorkbench/Views/KeynoteExportView.swift": [
            "Run Preflight",
            "Export Copied Deck",
        ],
        "Sources/ScientificWorkbench/Stores/WorkbenchStore.swift": [
            "reviewDOCXStyles",
            "runReviewedDOCXReplacement",
            "runConfirmedKeynoteExport",
        ],
        "Tests/ScientificWorkbenchTests/DocumentWorkflowTests.swift": [
            "documentWorkflowStagesCopyAndDetectsOriginalChanges",
            "artifactDiscoveryReadsDocumentV2Artifacts",
        ],
    }
    for relative, markers in required.items():
        path = APP / relative
        require(path.exists(), f"Missing ScientificWorkbench file: {relative}")
        text = path.read_text(encoding="utf-8")
        for marker in markers:
            require(marker in text, f"{relative} is missing marker: {marker}")


def main() -> int:
    if TMP.exists():
        shutil.rmtree(TMP)
    (TMP / "inputs").mkdir(parents=True)
    (TMP / "runs" / "inventory" / "tables").mkdir(parents=True)
    (TMP / "runs" / "inventory" / "reports").mkdir(parents=True)
    (TMP / "runs" / "replacement" / "artifacts").mkdir(parents=True)
    (TMP / "runs" / "replacement" / "previews").mkdir(parents=True)
    (TMP / "runs" / "keynote_preflight").mkdir(parents=True)

    python = sys.executable
    try:
        import docx  # noqa: F401
        import pptx  # noqa: F401
    except ImportError:
        python = str(Path.home() / "anaconda3/envs/datanalysis/bin/python3")

    original_docx, original_pptx = make_fixtures(python)
    docx_before = sha256(original_docx)
    pptx_before = sha256(original_pptx)
    copied_docx = TMP / "inputs" / "document copy.docx"
    copied_pptx = TMP / "inputs" / "deck copy.pptx"
    shutil.copy2(original_docx, copied_docx)
    shutil.copy2(original_pptx, copied_pptx)

    inventory_run = TMP / "runs" / "inventory"
    inventory_summary = inventory_run / "summary.json"
    run(
        [
            python,
            str(OFFICE),
            "docx-style-inventory",
            str(copied_docx),
            "--require-bold",
            "--require-italic",
            "--require-underline",
            "--output-csv",
            str(inventory_run / "tables/style_inventory.csv"),
            "--report-md",
            str(inventory_run / "reports/style_inventory.md"),
            "--summary-json",
            str(inventory_summary),
        ]
    )
    inventory = load(inventory_summary)
    require(inventory["contract_version"] == "2.0", "Inventory did not emit v2.0 contract")
    require(inventory["app_status"] == "PASS", f"Unexpected inventory status: {inventory}")
    require(inventory["results"]["inventory"]["match_count"] == 2, "Inventory did not find both marked runs")
    require({"summary_json", "table_csv", "report_md"} <= artifact_types(inventory), "Inventory artifact typing incomplete")
    require(inventory["original_modified"] is False, "Inventory claims original modification")

    replacement_run = TMP / "runs" / "replacement"
    edited = replacement_run / "artifacts/edited_copy.docx"
    diff = replacement_run / "previews/style_diff.json"
    replacement_summary = replacement_run / "summary.json"
    replacement_manifest = replacement_run / "manifest.json"
    run(
        [
            python,
            str(OFFICE),
            "docx-styled-replace",
            str(copied_docx),
            str(edited),
            "--find",
            "draft",
            "--replace",
            "final",
            "--require-bold",
            "--require-italic",
            "--require-underline",
            "--diff-json",
            str(diff),
            "--summary-json",
            str(replacement_summary),
            "--manifest-json",
            str(replacement_manifest),
        ]
    )
    replacement = load(replacement_summary)
    diff_payload = load(diff)
    manifest_payload = load(replacement_manifest)
    require(replacement["app_status"] == "PASS", f"Unexpected replacement status: {replacement}")
    require(replacement["replacements"] == 2, "Replacement did not edit both reviewed runs")
    require({"edited_document", "app_preview", "summary_json", "manifest_json"} <= artifact_types(replacement), "Replacement artifact typing incomplete")
    require(replacement["safety"]["requires_confirmation"] is True, "Replacement safety contract lacks confirmation")
    require(diff_payload["artifact_type"] == "app_preview", "Diff sidecar is not app_preview")
    require(len(diff_payload["changes"]) == 2, "Diff sidecar did not record both changes")
    require(edited.exists() and replacement_manifest.exists(), "Edited DOCX or manifest missing")
    manifest_outputs = {Path(item["path"]).name for item in manifest_payload["outputs"]}
    require({"edited_copy.docx", "style_diff.json"} <= manifest_outputs, "Manifest omits edited document or reviewed diff")

    keynote_run = TMP / "runs" / "keynote_preflight"
    keynote_summary = keynote_run / "summary.json"
    keynote = run(
        [
            sys.executable,
            str(KEYNOTE),
            str(copied_pptx),
            str(keynote_run / "preview.pdf"),
            "--preflight-only",
            "--summary-json",
            str(keynote_summary),
        ]
    )
    require(keynote.stdout.strip(), "Keynote preflight emitted no JSON")
    keynote_payload = load(keynote_summary)
    require(keynote_payload["contract_version"] == "2.0", "Keynote preflight did not emit v2.0 contract")
    require(keynote_payload["app_status"] in {"PASS", "BLOCKED_CONTROLADO"}, "Keynote preflight status is dishonest")
    require(keynote_payload["safety"]["requires_confirmation"] is True, "Keynote safety contract lacks confirmation")

    run([sys.executable, str(ROOT / "scripts/audit_v1_6_keynote_export_regression.py")])
    validate_app_surface()
    require(REFERENCE.exists(), "Phase 4 reference is missing")
    reference_text = REFERENCE.read_text(encoding="utf-8")
    for marker in ("P1-34", "P1-35", "P1-37", "edited_document", "app_preview", "BLOCKED_CONTROLADO"):
        require(marker in reference_text, f"Phase 4 reference missing: {marker}")

    require(sha256(original_docx) == docx_before, "Original DOCX was modified")
    require(sha256(original_pptx) == pptx_before, "Original deck was modified")

    summary = {
        "tool": "audit_v2_0_phase4_documents_gui_regression",
        "status": "PASS",
        "targets": {
            "P1-34": {
                "target": "office_roundtrip.py docx-style-inventory",
                "status": inventory["app_status"],
                "match_count": inventory["match_count"],
                "artifact_types": sorted(artifact_types(inventory)),
            },
            "P1-35": {
                "target": "office_roundtrip.py docx-styled-replace",
                "status": replacement["app_status"],
                "replacements": replacement["replacements"],
                "artifact_types": sorted(artifact_types(replacement)),
            },
            "P1-37": {
                "target": "keynote_export.py",
                "status": keynote_payload["app_status"],
                "capabilities": keynote_payload["results"]["capabilities"],
                "recommendation": keynote_payload["results"]["recommendation"],
            },
        },
        "original_modified": False,
        "app_root": str(APP),
        "reference": str(REFERENCE),
        "artifacts": [
            str(inventory_summary),
            str(inventory_run / "tables/style_inventory.csv"),
            str(inventory_run / "reports/style_inventory.md"),
            str(edited),
            str(diff),
            str(replacement_manifest),
            str(keynote_summary),
        ],
        "warnings": [],
        "failures": [],
    }
    output = TMP / "phase4_summary.json"
    output.write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(summary, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
