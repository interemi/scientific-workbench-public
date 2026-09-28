#!/usr/bin/env python3
"""v1.7 phase 1 regression for real-world non-astro intake workflows."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path
from typing import Any


SCRIPT_DIR = Path(__file__).resolve().parent
ROOT = SCRIPT_DIR.parent


def fail(message: str) -> None:
    raise AssertionError(message)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_csv(path: Path, rows: list[list[Any]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerows(rows)
    return path


def write_jsonl(path: Path, rows: list[dict[str, Any]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=True, sort_keys=True) + "\n")
    return path


def write_text(path: Path, text: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    return path


def write_docx(path: Path) -> Path | None:
    try:
        from docx import Document
    except ImportError:
        return None
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading("Resumen administrativo anonimo", level=1)
    doc.add_paragraph("Paquete de operaciones para revision semanal sin datos personales reales.")
    doc.add_paragraph("Riesgos: una incidencia P1 abierta y una metrica financiera con placeholder infinito.")
    doc.save(path)
    return path


def write_notebook(path: Path) -> Path:
    notebook = {
        "cells": [
            {
                "cell_type": "markdown",
                "metadata": {},
                "source": ["# Ops review notebook\n", "Notebook heredado anonimo para Fase 1.\n"],
            },
            {
                "cell_type": "code",
                "execution_count": None,
                "metadata": {},
                "outputs": [],
                "source": [
                    "import pandas as pd\n",
                    "df = pd.read_csv('support_tickets_clean.csv')\n",
                    "summary = df.groupby('team')['resolution_hours'].mean().round(2)\n",
                    "print(summary.to_dict())\n",
                ],
            },
        ],
        "metadata": {
            "kernelspec": {"display_name": "Python 3", "language": "python", "name": "python3"},
            "language_info": {"name": "python", "pygments_lexer": "ipython3"},
        },
        "nbformat": 4,
        "nbformat_minor": 5,
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(notebook, indent=2), encoding="utf-8")
    return path


def create_sqlite(path: Path, tickets_rows: list[list[Any]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(path) as conn:
        conn.execute(
            "CREATE TABLE tickets(ticket_id TEXT, team TEXT, priority TEXT, resolution_hours REAL, sla_breached INTEGER)"
        )
        for row in tickets_rows[1:]:
            conn.execute(
                "INSERT INTO tickets VALUES (?, ?, ?, ?, ?)",
                (row[0], row[4], row[5], float(row[8]), 1 if row[9] == "true" else 0),
            )
        conn.commit()
    return path


def create_npz(path: Path) -> Path | None:
    try:
        import numpy as np
    except ImportError:
        return None
    path.parent.mkdir(parents=True, exist_ok=True)
    np.savez(
        path,
        sensor_counts=np.array([101.0, 102.5, 99.8, 104.2]),
        timestamps=np.array(["2026-05-01T10:00", "2026-05-01T10:05", "2026-05-01T10:10", "2026-05-01T10:15"]),
    )
    return path


def create_fixtures(base: Path) -> dict[str, Any]:
    inputs = base / "inputs"
    tickets_rows = [
        [
            "ticket_id",
            "opened_at",
            "closed_at",
            "channel",
            "team",
            "priority",
            "region",
            "first_response_min",
            "resolution_hours",
            "sla_breached",
            "customer_score",
            "revenue_impact_eur",
            "notes",
        ],
        ["T-1001", "2026-04-01 08:15", "2026-04-01 13:20", "email", "support", "P2", "EU", "35", "5.08", "false", "4.7", "1200", "routine onboarding issue"],
        ["T-1002", "2026-04-02 10:05", "2026-04-04 09:13", "chat", "billing", "P1", "US", "240", "47.13", "true", "", "inf", "bad imported finance placeholder"],
        ["T-1003", "2026-04-03 11:00", "2026-04-03 16:45", "phone", "ops", "P3", "LATAM", "18", "5.75", "false", "4.2", "350", "normal operational record"],
        ["T-1004", "2026-04-04 12:20", "2026-04-05 09:50", "email", "support", "P2", "EU", "55", "21.5", "false", "3.9", "780", "handoff required"],
        ["T-1005", "2026-04-05 09:30", "2026-04-05 11:45", "chat", "ops", "P3", "EU", "12", "2.25", "false", "4.9", "0", "resolved by macro"],
        ["T-1006", "2026-04-06 15:10", "2026-04-07 08:30", "email", "data", "P2", "US", "90", "17.33", "false", "4.0", "430", "data export mismatch"],
    ]
    tickets_clean_rows = [tickets_rows[0], *[row[:] for row in tickets_rows[1:]]]
    tickets_clean_rows[2][11] = "1600"
    team_capacity_rows = [
        ["team", "weekly_capacity_hours", "owner"],
        ["billing", "120", "ops_lead"],
        ["data", "80", "analytics_lead"],
        ["ops", "160", "ops_lead"],
        ["support", "220", "support_lead"],
    ]
    baseline_rows = [
        ["metric", "value", "owner"],
        ["open_p1", "1", "support"],
        ["avg_resolution_hours", "16.5", "ops"],
        ["sla_breaches", "1", "ops"],
    ]
    candidate_rows = [
        ["metric", "value", "owner"],
        ["open_p1", "2", "support"],
        ["avg_resolution_hours", "14.2", "ops"],
        ["sla_breaches", "1", "ops"],
    ]
    events = [
        {"event_id": "E-1", "ticket_id": "T-1001", "event_type": "created", "ts": "2026-04-01T08:15:00"},
        {"event_id": "E-2", "ticket_id": "T-1002", "event_type": "escalated", "ts": "2026-04-02T11:05:00"},
        {"event_id": "E-3", "ticket_id": "T-1006", "event_type": "resolved", "ts": "2026-04-07T08:30:00"},
    ]

    paths: dict[str, Any] = {}
    paths["tickets"] = write_csv(inputs / "tables" / "support_tickets_anon.csv", tickets_rows)
    paths["tickets_clean"] = write_csv(inputs / "tables" / "support_tickets_clean.csv", tickets_clean_rows)
    paths["capacity"] = write_csv(inputs / "tables" / "team_capacity.csv", team_capacity_rows)
    paths["baseline"] = write_csv(inputs / "diff" / "weekly_metrics_baseline.csv", baseline_rows)
    paths["candidate"] = write_csv(inputs / "diff" / "weekly_metrics_candidate.csv", candidate_rows)
    paths["events"] = write_jsonl(inputs / "mixed_project" / "events.jsonl", events)
    shutil.copy2(paths["tickets_clean"], inputs / "mixed_project" / "support_tickets_clean.csv")
    shutil.copy2(paths["capacity"], inputs / "mixed_project" / "team_capacity.csv")
    paths["memo"] = write_text(
        inputs / "documents" / "handoff_memo.md",
        "# Handoff memo\n\nRevision anonima de operaciones. Hay una incidencia P1 y una metrica financiera que debe limpiarse antes de reporting.\n",
    )
    paths["letter"] = write_text(
        inputs / "documents" / "administrative_note.txt",
        "Administrative note\nProject: anonymized operations review\nAction: validate metrics before weekly meeting.\n",
    )
    docx = write_docx(inputs / "documents" / "ops_admin_summary.docx")
    if docx is not None:
        paths["docx"] = docx
    paths["notebook"] = write_notebook(inputs / "notebooks" / "ops_review_legacy.ipynb")
    paths["sqlite"] = create_sqlite(inputs / "containers" / "ops_review.sqlite", tickets_clean_rows)
    npz = create_npz(inputs / "containers" / "sensor_roi_counts.npz")
    if npz is not None:
        paths["npz"] = npz
    package = inputs / "containers" / "ops_package.zip"
    package.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(package, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.write(paths["tickets_clean"], arcname="tables/support_tickets_clean.csv")
        archive.write(paths["events"], arcname="events/events.jsonl")
        archive.writestr("README.txt", "Anonymized operations project package for v1.7 intake regression.\n")
    paths["zip"] = package
    paths["mixed_dir"] = inputs / "mixed_project"
    paths["documents_dir"] = inputs / "documents"
    paths["input_hashes"] = {str(path): sha256(path) for path in inputs.rglob("*") if path.is_file()}
    return paths


def run_command(name: str, command: list[str], logs_dir: Path, timeout: int = 120) -> subprocess.CompletedProcess[str]:
    logs_dir.mkdir(parents=True, exist_ok=True)
    (logs_dir / f"{name}.command.json").write_text(json.dumps(command, indent=2), encoding="utf-8")
    proc = subprocess.run(command, cwd=ROOT, text=True, capture_output=True, timeout=timeout)
    (logs_dir / f"{name}.stdout.txt").write_text(proc.stdout, encoding="utf-8")
    (logs_dir / f"{name}.stderr.txt").write_text(proc.stderr, encoding="utf-8")
    (logs_dir / f"{name}.exitcode").write_text(str(proc.returncode) + "\n", encoding="utf-8")
    if "Traceback (most recent call last)" in proc.stdout or "Traceback (most recent call last)" in proc.stderr:
        fail(f"{name} emitted a traceback")
    return proc


def read_payload(path: Path) -> dict[str, Any]:
    if not path.exists():
        fail(f"expected JSON payload missing: {path}")
    return json.loads(path.read_text(encoding="utf-8"))


def status_label(proc: subprocess.CompletedProcess[str], payload: dict[str, Any] | None = None) -> str:
    if payload and payload.get("status") == "blocked":
        return "BLOCKED_CONTROLADO"
    if proc.returncode != 0:
        return "BLOCKED_CONTROLADO" if payload and payload.get("status") in {"blocked", "fail"} else "FAIL"
    if payload and payload.get("status") == "warning":
        return "WARNING"
    return "PASS"


def require_returncode(proc: subprocess.CompletedProcess[str], name: str) -> None:
    if proc.returncode != 0:
        fail(f"{name} failed: {proc.stderr or proc.stdout}")


def is_sandbox_blocked(payload: dict[str, Any]) -> bool:
    run = payload.get("results", {}).get("run", {})
    assessment = run.get("assessment", {}) if isinstance(run, dict) else {}
    return payload.get("status") == "blocked" and assessment.get("status") == "sandbox_blocked"


def assert_inputs_unchanged(input_hashes: dict[str, str]) -> None:
    for raw_path, before_hash in input_hashes.items():
        path = Path(raw_path)
        if not path.exists():
            fail(f"input disappeared: {path}")
        after_hash = sha256(path)
        if after_hash != before_hash:
            fail(f"input was modified: {path}")


def case_profile(paths: dict[str, Any], base: Path, logs: Path) -> dict[str, Any]:
    out = base / "runs" / "profile_table"
    summary = out / "profile_summary.json"
    manifest = out / "profile_manifest.json"
    proc = run_command(
        "profile_table",
        [
            sys.executable,
            str(SCRIPT_DIR / "profile_table.py"),
            str(paths["tickets"]),
            "--head",
            "5",
            "--max-columns",
            "13",
            "--summary-json",
            str(summary),
            "--manifest-json",
            str(manifest),
        ],
        logs,
    )
    require_returncode(proc, "profile_table")
    payload = read_payload(summary)
    if payload.get("status") != "warning":
        fail("profile_table should warn about the injected infinite numeric value")
    if payload.get("qa", {}).get("metrics", {}).get("infinite_value_count") != 1:
        fail("profile_table should detect one infinite numeric value")
    return {
        "capability": "profile_table.py",
        "status": status_label(proc, payload),
        "input": str(paths["tickets"]),
        "command": "profile_table.py support_tickets_anon.csv --summary-json ...",
        "output": str(summary),
        "warning_error": "1 infinite numeric value detected",
        "diagnostic": "Useful first-pass profile for an anonymized operations table; input preserved.",
        "useful_for_real_person": True,
    }


def case_cross_domain(paths: dict[str, Any], base: Path, logs: Path) -> dict[str, Any]:
    out = base / "runs" / "cross_domain_data_workbench" / "out"
    summary = base / "runs" / "cross_domain_data_workbench" / "summary.json"
    manifest = base / "runs" / "cross_domain_data_workbench" / "manifest.json"
    proc = run_command(
        "cross_domain_data_workbench",
        [
            sys.executable,
            str(SCRIPT_DIR / "cross_domain_data_workbench.py"),
            str(paths["mixed_dir"]),
            "--head",
            "4",
            "--output-dir",
            str(out),
            "--summary-json",
            str(summary),
            "--manifest-json",
            str(manifest),
        ],
        logs,
    )
    require_returncode(proc, "cross_domain_data_workbench")
    payload = read_payload(summary)
    if payload.get("status") not in {"ok", "warning"}:
        fail(f"cross_domain_data_workbench unexpected status: {payload.get('status')}")
    if not (out / "inventory.csv").exists() or not (out / "report.md").exists():
        fail("cross_domain_data_workbench should write inventory.csv and report.md")
    return {
        "capability": "cross_domain_data_workbench.py",
        "status": status_label(proc, payload),
        "input": str(paths["mixed_dir"]),
        "command": "cross_domain_data_workbench.py mixed_project --output-dir ...",
        "output": str(out),
        "warning_error": "none" if payload.get("status") == "ok" else "profile warning propagated",
        "diagnostic": "Useful bundle-level intake for CSV and JSONL operational files.",
        "useful_for_real_person": True,
    }


def case_container(paths: dict[str, Any], base: Path, logs: Path) -> dict[str, Any]:
    out = base / "runs" / "inspect_data_container"
    summary = out / "zip_summary.json"
    manifest = out / "zip_manifest.json"
    proc = run_command(
        "inspect_data_container",
        [
            sys.executable,
            str(SCRIPT_DIR / "inspect_data_container.py"),
            str(paths["zip"]),
            "--summary-json",
            str(summary),
            "--manifest-json",
            str(manifest),
        ],
        logs,
    )
    require_returncode(proc, "inspect_data_container")
    payload = read_payload(summary)
    if payload.get("status") not in {"ok", "warning"}:
        fail(f"inspect_data_container unexpected status: {payload.get('status')}")
    return {
        "capability": "inspect_data_container.py",
        "status": status_label(proc, payload),
        "input": str(paths["zip"]),
        "command": "inspect_data_container.py ops_package.zip --summary-json ...",
        "output": str(summary),
        "warning_error": "none" if payload.get("status") == "ok" else json.dumps(payload.get("qa", {}).get("findings", [])),
        "diagnostic": "Useful safe inspection of a technical project archive before extraction.",
        "useful_for_real_person": True,
    }


def case_document_intake(paths: dict[str, Any], base: Path, logs: Path) -> dict[str, Any]:
    out = base / "runs" / "document_intake_workbench" / "out"
    summary = base / "runs" / "document_intake_workbench" / "summary.json"
    manifest = base / "runs" / "document_intake_workbench" / "manifest.json"
    proc = run_command(
        "document_intake_workbench",
        [
            sys.executable,
            str(SCRIPT_DIR / "document_intake_workbench.py"),
            str(paths["documents_dir"]),
            "--output-dir",
            str(out),
            "--summary-json",
            str(summary),
            "--manifest-json",
            str(manifest),
        ],
        logs,
    )
    require_returncode(proc, "document_intake_workbench")
    payload = read_payload(summary)
    if payload.get("status") not in {"ok", "warning"}:
        fail(f"document_intake_workbench unexpected status: {payload.get('status')}")
    if not (out / "inventory.csv").exists() or not (out / "report.md").exists():
        fail("document_intake_workbench should write inventory.csv and report.md")
    return {
        "capability": "document_intake_workbench.py",
        "status": status_label(proc, payload),
        "input": str(paths["documents_dir"]),
        "command": "document_intake_workbench.py documents --output-dir ...",
        "output": str(out),
        "warning_error": "none" if payload.get("status") == "ok" else json.dumps(payload.get("qa", {}).get("findings", [])),
        "diagnostic": "Useful administrative document intake without opening GUI apps or editing originals.",
        "useful_for_real_person": True,
    }


def case_semantic_diff(paths: dict[str, Any], base: Path, logs: Path) -> dict[str, Any]:
    out = base / "runs" / "semantic_diff"
    summary = out / "summary.json"
    output_json = out / "diff.json"
    output_md = out / "diff.md"
    manifest = out / "manifest.json"
    proc = run_command(
        "semantic_diff",
        [
            sys.executable,
            str(SCRIPT_DIR / "semantic_diff.py"),
            str(paths["baseline"]),
            str(paths["candidate"]),
            "--summary-json",
            str(summary),
            "--output-json",
            str(output_json),
            "--output-md",
            str(output_md),
            "--manifest-json",
            str(manifest),
        ],
        logs,
    )
    require_returncode(proc, "semantic_diff")
    payload = read_payload(summary)
    if payload.get("status") not in {"ok", "warning"}:
        fail(f"semantic_diff unexpected status: {payload.get('status')}")
    if not summary.exists() or not output_md.exists():
        fail("semantic_diff should write summary JSON and Markdown outputs")
    return {
        "capability": "semantic_diff.py",
        "status": status_label(proc, payload),
        "input": f"{paths['baseline']} vs {paths['candidate']}",
        "command": "semantic_diff.py weekly_metrics_baseline.csv weekly_metrics_candidate.csv --output-md ...",
        "output": str(output_md),
        "warning_error": "none" if payload.get("status") == "ok" else json.dumps(payload.get("qa", {}).get("findings", [])),
        "diagnostic": "Useful review of changed operational metrics before a weekly report.",
        "useful_for_real_person": True,
    }


def case_deliverable(paths: dict[str, Any], base: Path, logs: Path) -> dict[str, Any]:
    out = base / "runs" / "deliverable_factory" / "ops_status_pack"
    summary = base / "runs" / "deliverable_factory" / "summary.json"
    manifest = base / "runs" / "deliverable_factory" / "manifest.json"
    proc = run_command(
        "deliverable_factory",
        [
            sys.executable,
            str(SCRIPT_DIR / "deliverable_factory.py"),
            "scaffold",
            str(out),
            "--kind",
            "status-report",
            "--format",
            "markdown",
            "--title",
            "Ops Weekly Status Anonimo",
            "--language",
            "spanish",
            "--summary-json",
            str(summary),
            "--manifest-json",
            str(manifest),
        ],
        logs,
    )
    require_returncode(proc, "deliverable_factory")
    payload = read_payload(summary)
    if payload.get("status") != "ok":
        fail(f"deliverable_factory unexpected status: {payload.get('status')}")
    if not any(out.glob("*.md")):
        fail("deliverable_factory should scaffold a Markdown deliverable")
    return {
        "capability": "deliverable_factory.py scaffold",
        "status": status_label(proc, payload),
        "input": str(out),
        "command": "deliverable_factory.py scaffold ops_status_pack --kind status-report ...",
        "output": str(out),
        "warning_error": "none",
        "diagnostic": "Useful status-report scaffold for handoff after intake.",
        "useful_for_real_person": True,
    }


def case_notebook(paths: dict[str, Any], base: Path, logs: Path) -> dict[str, Any]:
    out = base / "runs" / "notebook_workbench"
    summary = out / "execute_summary.json"
    manifest = out / "execute_manifest.json"
    before_hash = sha256(paths["notebook"])
    proc = run_command(
        "notebook_workbench_execute_copy",
        [
            sys.executable,
            str(SCRIPT_DIR / "notebook_workbench.py"),
            "execute-copy",
            str(paths["notebook"]),
            "--output-dir",
            str(out),
            "--kernel-name",
            "python3",
            "--timeout-sec",
            "90",
            "--stage-extra",
            str(paths["tickets_clean"]),
            "--summary-json",
            str(summary),
            "--manifest-json",
            str(manifest),
        ],
        logs,
        timeout=180,
    )
    require_returncode(proc, "notebook_workbench execute-copy")
    after_hash = sha256(paths["notebook"])
    if before_hash != after_hash:
        fail("notebook_workbench modified the original notebook")
    payload = read_payload(summary)
    if payload.get("status") not in {"ok", "warning"} and not is_sandbox_blocked(payload):
        fail(f"notebook_workbench unexpected status: {payload.get('status')}")
    if not list(out.glob("*.ipynb")):
        fail("notebook_workbench should write an executed copied notebook")
    return {
        "capability": "notebook_workbench.py execute-copy",
        "status": status_label(proc, payload),
        "input": str(paths["notebook"]),
        "command": "notebook_workbench.py execute-copy ops_review_legacy.ipynb --stage-extra support_tickets_clean.csv ...",
        "output": str(out),
        "warning_error": "none" if payload.get("status") == "ok" else json.dumps(payload.get("qa", {}).get("findings", [])),
        "diagnostic": "Useful execution of a copied inherited notebook while preserving the original.",
        "useful_for_real_person": True,
    }


def case_duckdb(paths: dict[str, Any], base: Path, logs: Path) -> dict[str, Any]:
    out = base / "runs" / "duckdb_workbench"
    summary = out / "summary.json"
    output = out / "team_ticket_rollup.csv"
    manifest = out / "manifest.json"
    sql = (
        "SELECT t.team, COUNT(*) AS tickets, ROUND(AVG(t.resolution_hours), 2) AS avg_resolution_hours "
        "FROM source0 t GROUP BY t.team ORDER BY t.team"
    )
    proc = run_command(
        "duckdb_workbench",
        [
            sys.executable,
            str(SCRIPT_DIR / "duckdb_workbench.py"),
            str(paths["tickets_clean"]),
            "--sql",
            sql,
            "--output",
            str(output),
            "--summary-json",
            str(summary),
            "--manifest-json",
            str(manifest),
        ],
        logs,
    )
    payload = read_payload(summary) if summary.exists() else None
    if proc.returncode != 0:
        if payload and payload.get("status") in {"blocked", "fail"}:
            return {
                "capability": "duckdb_workbench.py",
                "status": "BLOCKED_CONTROLADO",
                "input": str(paths["tickets_clean"]),
                "command": "duckdb_workbench.py support_tickets_clean.csv --sql ...",
                "output": str(summary),
                "warning_error": json.dumps(payload.get("qa", {}).get("findings", [])),
                "diagnostic": "DuckDB backend unavailable or blocked cleanly.",
                "useful_for_real_person": False,
            }
        require_returncode(proc, "duckdb_workbench")
    if payload is None or payload.get("status") not in {"ok", "warning"}:
        fail(f"duckdb_workbench unexpected status: {None if payload is None else payload.get('status')}")
    if not output.exists():
        fail("duckdb_workbench should write query output when backend is available")
    return {
        "capability": "duckdb_workbench.py",
        "status": status_label(proc, payload),
        "input": str(paths["tickets_clean"]),
        "command": "duckdb_workbench.py support_tickets_clean.csv --sql rollup --output team_ticket_rollup.csv",
        "output": str(output),
        "warning_error": "none" if payload.get("status") == "ok" else json.dumps(payload.get("qa", {}).get("findings", [])),
        "diagnostic": "Useful local SQL rollup over an operational CSV using explicit source0 naming.",
        "useful_for_real_person": True,
    }


def write_report(base: Path, rows: list[dict[str, Any]]) -> Path:
    report = base / "phase1_real_world_intake_report.md"
    lines = [
        "# AUDITORIA v1.7 - Fase 1: universalizar intake real-world",
        "",
        "## Alcance",
        "",
        "Se ejecuto un paquete anonimo no astrofisico de operaciones/soporte sobre las capabilities target de intake real-world.",
        "",
        "## Resultados",
        "",
        "| Capability | Input | Comando | Estado | Warning/error | Diagnostico | Artefacto principal |",
        "|---|---|---|---|---|---|---|",
    ]
    for row in rows:
        lines.append(
            "| {capability} | `{input}` | `{command}` | {status} | {warning_error} | {diagnostic} | `{output}` |".format(
                **{key: str(value).replace("|", "\\|") for key, value in row.items()}
            )
        )
    lines.extend(
        [
            "",
            "## Decision",
            "",
            "Fase 1 queda cubierta por una regresion integrada y por artefactos temporales reproducibles.",
            "No se modifican originales: el script valida hashes de inputs antes/despues.",
            "No se sincroniza la skill instalada.",
        ]
    )
    report.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return report


def run_phase(output_dir: Path) -> dict[str, Any]:
    output_dir = output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    generated = output_dir / "generated"
    runs = output_dir / "runs"
    logs = output_dir / "logs"
    if generated.exists():
        shutil.rmtree(generated)
    if runs.exists():
        shutil.rmtree(runs)
    if logs.exists():
        shutil.rmtree(logs)
    paths = create_fixtures(generated)
    rows = [
        case_profile(paths, output_dir, logs),
        case_cross_domain(paths, output_dir, logs),
        case_container(paths, output_dir, logs),
        case_document_intake(paths, output_dir, logs),
        case_semantic_diff(paths, output_dir, logs),
        case_deliverable(paths, output_dir, logs),
        case_notebook(paths, output_dir, logs),
        case_duckdb(paths, output_dir, logs),
    ]
    assert_inputs_unchanged(paths["input_hashes"])
    report = write_report(output_dir, rows)
    pass_like = {"PASS", "WARNING", "BLOCKED_CONTROLADO"}
    if any(row["status"] not in pass_like for row in rows):
        fail(f"unexpected failed cases: {rows}")
    payload = {
        "status": "ok",
        "phase": "v1.7 phase 1",
        "capability_count": len(rows),
        "case_status_counts": dict(sorted({status: sum(1 for row in rows if row["status"] == status) for status in {row["status"] for row in rows}}.items())),
        "report_md": str(report),
        "output_dir": str(output_dir),
        "rows": rows,
    }
    (output_dir / "phase1_summary.json").write_text(json.dumps(payload, indent=2, sort_keys=True), encoding="utf-8")
    return payload


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", help="Optional persistent output directory for fixtures and artifacts.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if args.output_dir:
        payload = run_phase(Path(args.output_dir))
        print(json.dumps(payload, indent=2, sort_keys=True))
    else:
        with tempfile.TemporaryDirectory(prefix="sda_v1_7_phase1_intake_") as tmp:
            payload = run_phase(Path(tmp))
            compact = {
                "status": payload["status"],
                "phase": payload["phase"],
                "capability_count": payload["capability_count"],
                "case_status_counts": payload["case_status_counts"],
            }
            print(json.dumps(compact, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
