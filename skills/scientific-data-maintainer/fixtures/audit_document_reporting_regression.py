#!/usr/bin/env python3
"""Regressions for document-reporting handoffs: DOCX marked runs and docs routes."""

from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path


ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "scripts"


def run_cmd(args: list[str], *, tmp: Path) -> subprocess.CompletedProcess:
    env = os.environ.copy()
    env["PYTHONDONTWRITEBYTECODE"] = "1"
    completed = subprocess.run(
        args,
        cwd=ROOT,
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        env=env,
    )
    if completed.returncode != 0:
        raise AssertionError(
            f"Command failed: {' '.join(args)}\nSTDOUT:\n{completed.stdout}\nSTDERR:\n{completed.stderr}"
        )
    return completed


def module_available(module: str) -> bool:
    completed = subprocess.run(
        [sys.executable, "-c", f"import {module}"],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    return completed.returncode == 0


def datanalysis_python() -> str:
    completed = subprocess.run(
        [sys.executable, str(SCRIPTS / "datanalysis_env.py"), "locate"],
        cwd=ROOT,
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    if completed.returncode != 0:
        raise AssertionError(f"Could not locate datanalysis runtime:\n{completed.stderr}")
    payload = json.loads(completed.stdout)
    python_path = payload.get("selected", {}).get("python")
    if not python_path:
        raise AssertionError(f"datanalysis runtime did not report a Python executable: {payload}")
    return str(Path(python_path).expanduser())


def runtime_python() -> str:
    return sys.executable if module_available("docx") else datanalysis_python()


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def create_styled_docx(tmp: Path, python: str) -> Path:
    script = tmp / "make_styled_docx.py"
    docx_path = tmp / "pages_export.docx"
    script.write_text(
        """
from docx import Document
import sys

doc = Document()
doc.add_paragraph("Plain context before the marked fragment.")
paragraph = doc.add_paragraph()
paragraph.add_run("Keep this ordinary text. ")
run = paragraph.add_run("marked fragment")
run.bold = True
run.italic = True
run.underline = True
paragraph.add_run(" after.")
paragraph = doc.add_paragraph()
other = paragraph.add_run("bold only fragment")
other.bold = True
doc.save(sys.argv[1])
""".lstrip(),
        encoding="utf-8",
    )
    run_cmd([python, str(script), str(docx_path)], tmp=tmp)
    return docx_path


def check_docx_marked_runs(tmp: Path) -> None:
    python = runtime_python()
    source = create_styled_docx(tmp, python)
    inventory_json = tmp / "inventory.json"
    inventory_csv = tmp / "inventory.csv"
    report_md = tmp / "inventory.md"
    run_cmd(
        [
            python,
            str(SCRIPTS / "office_roundtrip.py"),
            "docx-style-inventory",
            str(source),
            "--require-bold",
            "--require-italic",
            "--require-underline",
            "--output-csv",
            str(inventory_csv),
            "--report-md",
            str(report_md),
            "--summary-json",
            str(inventory_json),
        ],
        tmp=tmp,
    )
    inventory = load_json(inventory_json)
    if inventory["match_count"] != 1 or inventory["matches"][0]["text"] != "marked fragment":
        raise AssertionError(f"Unexpected style inventory: {inventory}")
    if not inventory_csv.exists() or not report_md.exists():
        raise AssertionError("DOCX style inventory did not write expected artifacts.")

    edited = tmp / "edited.docx"
    replace_json = tmp / "replace.json"
    run_cmd(
        [
            python,
            str(SCRIPTS / "office_roundtrip.py"),
            "docx-styled-replace",
            str(source),
            str(edited),
            "--find",
            "marked fragment",
            "--replace",
            "revised fragment",
            "--require-bold",
            "--require-italic",
            "--require-underline",
            "--summary-json",
            str(replace_json),
        ],
        tmp=tmp,
    )
    replace = load_json(replace_json)
    if replace["replacements"] != 1 or replace["readback_replacement_match_count"] != 1:
        raise AssertionError(f"Styled replacement did not verify readback: {replace}")

    edited_inventory_json = tmp / "edited_inventory.json"
    run_cmd(
        [
            python,
            str(SCRIPTS / "office_roundtrip.py"),
            "docx-style-inventory",
            str(edited),
            "--require-bold",
            "--require-italic",
            "--require-underline",
            "--summary-json",
            str(edited_inventory_json),
        ],
        tmp=tmp,
    )
    edited_inventory = load_json(edited_inventory_json)
    if edited_inventory["matches"][0]["text"] != "revised fragment":
        raise AssertionError(f"Styled readback did not preserve replacement: {edited_inventory}")


def check_reference_routes() -> None:
    formats = (ROOT / "references" / "formats-and-handoffs.md").read_text(encoding="utf-8")
    presentation = (ROOT / "references" / "scientific-presentation-handoff.md").read_text(encoding="utf-8")
    required_formats = [
        "docx-style-inventory",
        "docx-styled-replace",
        "pdftoppm",
        "qlmanage",
        "PDF rendering fallback",
    ]
    missing = [item for item in required_formats if item not in formats]
    if missing:
        raise AssertionError(f"formats-and-handoffs.md is missing route terms: {missing}")
    required_presentation = [
        "Comparative Oral Style",
        "do not copy phrases",
        "speaker voice",
    ]
    missing = [item for item in required_presentation if item not in presentation]
    if missing:
        raise AssertionError(f"scientific-presentation-handoff.md is missing route terms: {missing}")


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="sda_doc_reporting_") as raw_tmp:
        tmp = Path(raw_tmp)
        check_docx_marked_runs(tmp)
        print("PASS check_docx_marked_runs")
        check_reference_routes()
        print("PASS check_reference_routes")
    print("All document/reporting regressions passed.")


if __name__ == "__main__":
    main()
