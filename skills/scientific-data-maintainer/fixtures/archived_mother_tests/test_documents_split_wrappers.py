"""v2.3 document child split checks for historical mother wrappers."""

from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile
import textwrap
import unittest
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
DOCUMENT_CHILD = ROOT.parent / "scientific-data-documents"
DOCUMENT_SCRIPTS = [
    "document_intake_workbench.py",
    "document_semantics.py",
    "pdf_recover_extract.py",
    "office_roundtrip.py",
    "iwork_workbench.py",
    "iwork_roundtrip.py",
    "presentation_workbench.py",
    "quicklook_bridge.py",
    "keynote_export.py",
    "latex_workbench.py",
    "scientific_writeup_review.py",
    "deliverable_factory.py",
    "semantic_diff.py",
]


def python_executable() -> str:
    datanalysis = Path(os.environ.get("DATAANALYSIS_PYTHON") or sys.executable)
    return str(datanalysis) if datanalysis.exists() else sys.executable


class DocumentsSplitWrappersTest(unittest.TestCase):
    def run_command(self, command: list[str], tmp: Path) -> subprocess.CompletedProcess:
        env = os.environ.copy()
        env["PYTHONDONTWRITEBYTECODE"] = "1"
        env["MPLCONFIGDIR"] = str(tmp / "mplconfig")
        return subprocess.run(
            command,
            cwd=ROOT,
            text=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            env=env,
            timeout=120,
        )

    def test_child_has_full_bodies_and_mother_has_wrappers(self) -> None:
        self.assertTrue((DOCUMENT_CHILD / "SKILL.md").exists())
        self.assertTrue((DOCUMENT_CHILD / "scripts" / "datanalysis_env.py").exists())
        for name in DOCUMENT_SCRIPTS:
            mother = SCRIPTS / name
            child = DOCUMENT_CHILD / "scripts" / name
            self.assertIn("child_skill_dispatch", mother.read_text(encoding="utf-8"), name)
            self.assertTrue(child.exists(), name)
            self.assertNotIn("child_skill_dispatch", child.read_text(encoding="utf-8"), name)

    def test_document_workflows_execute_through_mother_wrappers(self) -> None:
        with tempfile.TemporaryDirectory(prefix="sda_docs_split_") as tmp_raw:
            tmp = Path(tmp_raw)
            docs_dir = tmp / "docs"
            docs_dir.mkdir()
            handoff = docs_dir / "handoff.txt"
            handoff.write_text("Anonymous project handoff\nBudget: 1200\nDate: 2026-02-03\n", encoding="utf-8")
            original_handoff = handoff.read_bytes()

            before = tmp / "before.txt"
            after = tmp / "after.txt"
            before.write_text("alpha\nbeta\n", encoding="utf-8")
            after.write_text("alpha\nbeta changed\n", encoding="utf-8")

            pptx = tmp / "deck.pptx"
            create_pptx = textwrap.dedent(
                f"""
                from pptx import Presentation
                prs = Presentation()
                slide = prs.slides.add_slide(prs.slide_layouts[0])
                slide.shapes.title.text = "Quarterly result"
                slide.placeholders[1].text = "One copied presentation fixture."
                prs.save({str(pptx)!r})
                """
            )
            completed = self.run_command([python_executable(), "-c", create_pptx], tmp)
            self.assertEqual(completed.returncode, 0, completed.stderr)

            commands = [
                ([
                    python_executable(),
                    str(SCRIPTS / "document_intake_workbench.py"),
                    str(docs_dir),
                    "--output-dir",
                    str(tmp / "document_intake_out"),
                    "--summary-json",
                    str(tmp / "document_intake.json"),
                    "--max-files",
                    "5",
                ], tmp / "document_intake.json", False),
                ([
                    python_executable(),
                    str(SCRIPTS / "latex_workbench.py"),
                    "scaffold",
                    str(tmp / "latex_project"),
                    "--summary-json",
                    str(tmp / "latex_scaffold.json"),
                    "--title",
                    "Informe",
                ], tmp / "latex_scaffold.json", False),
                ([
                    python_executable(),
                    str(SCRIPTS / "semantic_diff.py"),
                    str(before),
                    str(after),
                    "--summary-json",
                    str(tmp / "semantic_diff.json"),
                    "--output-md",
                    str(tmp / "semantic_diff.md"),
                ], tmp / "semantic_diff.json", False),
                ([
                    python_executable(),
                    str(SCRIPTS / "presentation_workbench.py"),
                    "inspect",
                    str(pptx),
                    "--summary-json",
                    str(tmp / "presentation_inspect.json"),
                ], tmp / "presentation_inspect.json", True),
                ([
                    python_executable(),
                    str(SCRIPTS / "keynote_export.py"),
                    str(pptx),
                    str(tmp / "out.pdf"),
                    "--preflight-only",
                    "--summary-json",
                    str(tmp / "keynote_preflight.json"),
                ], tmp / "keynote_preflight.json", True),
            ]
            for command, summary_path, allow_blocked in commands:
                completed = self.run_command(command, tmp)
                if completed.returncode != 0:
                    self.assertTrue(summary_path.exists(), completed.stdout + completed.stderr)
                    payload = json.loads(summary_path.read_text(encoding="utf-8"))
                    self.assertTrue(allow_blocked, f"unexpected nonzero command: {' '.join(command)}")
                    self.assertEqual(payload.get("app_status"), "BLOCKED_CONTROLADO")
                    self.assertNotIn("Traceback", completed.stdout + completed.stderr)
                else:
                    self.assertEqual(completed.returncode, 0)

            for summary_name in [
                "document_intake.json",
                "latex_scaffold.json",
                "semantic_diff.json",
                "presentation_inspect.json",
                "keynote_preflight.json",
            ]:
                payload = json.loads((tmp / summary_name).read_text(encoding="utf-8"))
                self.assertIn(payload.get("app_status"), {"PASS", "WARNING", "BLOCKED_CONTROLADO", "ok", "warning"})
                self.assertFalse(payload.get("original_modified", False), summary_name)
            self.assertEqual(handoff.read_bytes(), original_handoff)


if __name__ == "__main__":
    unittest.main()
