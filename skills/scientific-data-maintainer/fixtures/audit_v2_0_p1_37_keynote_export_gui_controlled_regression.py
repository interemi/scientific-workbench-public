#!/usr/bin/env python3
"""Regression for v2.0 P1-37 controlled Keynote GUI export."""

from __future__ import annotations

import contextlib
import hashlib
import io
import json
import logging
import shutil
import subprocess
import sys
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
TMP = ROOT / "tmp" / "v2_0_p1_37_keynote_export_gui_controlled"
APP = Path(__file__).resolve().parents[3]
SCRIPT = ROOT / "scripts" / "keynote_export.py"
CONFIRMATION_ID = "p1-37-regression-reviewed"


def rerun_in_datanalysis_if_needed() -> int | None:
    try:
        import pptx  # noqa: F401
        import pypdf  # noqa: F401

        return None
    except ImportError:
        sys.path.insert(0, str(ROOT / "scripts"))
        from datanalysis_env import locate_env_real

        selected = (locate_env_real().get("selected") or {}).get("python")
        if not selected or Path(selected).resolve() == Path(sys.executable).resolve():
            raise RuntimeError("P1-37 regression needs python-pptx and pypdf to create and inspect its synthetic fixture.")
        completed = subprocess.run(
            [str(selected), str(Path(__file__).resolve())],
            cwd=ROOT,
            text=True,
            capture_output=True,
            check=False,
        )
        print(completed.stdout, end="")
        print(completed.stderr, end="", file=sys.stderr)
        return completed.returncode


def require(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def run(args: list[str], *, expect: set[int] = {0}, timeout: int = 180) -> subprocess.CompletedProcess[str]:
    completed = subprocess.run(
        args,
        cwd=ROOT,
        text=True,
        capture_output=True,
        check=False,
        timeout=timeout,
    )
    require(
        completed.returncode in expect,
        f"Command failed ({completed.returncode}): {' '.join(args)}\nSTDOUT:\n{completed.stdout}\nSTDERR:\n{completed.stderr}",
    )
    require("Traceback" not in completed.stdout + completed.stderr, "Raw traceback leaked")
    return completed


def make_deck(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.0), Inches(11.5), Inches(3.0))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = "Controlled Keynote export P1-37"
    paragraph.font.size = Pt(30)
    paragraph.font.bold = True
    second = box.text_frame.add_paragraph()
    second.text = "Copied deck, confirmed GUI automation, separate PDF."
    second.font.size = Pt(18)
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(path)


def ready_report() -> dict:
    return {
        "status": "ok",
        "blocking_findings": [],
        "warning_findings": [],
        "recommendation": "Synthetic Keynote backend is ready.",
        "capabilities": {"platform": "darwin", "osascript": True, "keynote_app": True},
    }


def blocked_report() -> dict:
    return {
        "status": "blocked",
        "blocking_findings": ["Keynote.app is not installed or not discoverable by AppleScript."],
        "warning_findings": [],
        "recommendation": "Use a copy-based fallback.",
        "capabilities": {"platform": "darwin", "osascript": True, "keynote_app": False},
    }


def invoke_main_with_fakes(
    argv: list[str],
    *,
    report: dict,
    exporter,
) -> tuple[int, str, str]:
    sys.path.insert(0, str(ROOT / "scripts"))
    import keynote_export

    old_argv = sys.argv[:]
    # The public wrapper re-exports ``main`` from the documents child. Python
    # functions retain the child module globals where they were defined, so
    # patch those globals directly; patching only wrapper attributes would let
    # the simulated cases invoke the real osascript exporter.
    runtime_globals = keynote_export.main.__globals__
    old_report = runtime_globals["build_capability_report"]
    old_export = runtime_globals["export_with_keynote"]
    stdout = io.StringIO()
    stderr = io.StringIO()
    exit_code = 0
    try:
        runtime_globals["build_capability_report"] = lambda: report
        runtime_globals["export_with_keynote"] = exporter
        sys.argv = [str(SCRIPT), *argv]
        with contextlib.redirect_stdout(stdout), contextlib.redirect_stderr(stderr):
            try:
                keynote_export.main()
            except SystemExit as exc:
                exit_code = int(exc.code or 0)
    finally:
        runtime_globals["build_capability_report"] = old_report
        runtime_globals["export_with_keynote"] = old_export
        sys.argv = old_argv
    require("Traceback" not in stdout.getvalue() + stderr.getvalue(), "Fake backend leaked traceback")
    return exit_code, stdout.getvalue(), stderr.getvalue()


def artifact_types(payload: dict) -> set[str]:
    return {item["artifact_type"] for item in payload.get("typed_artifacts", [])}


def main() -> int:
    rerun_code = rerun_in_datanalysis_if_needed()
    if rerun_code is not None:
        return rerun_code
    if TMP.exists():
        shutil.rmtree(TMP)
    original = TMP / "originals" / "Controlled deck P1-37.pptx"
    copied = TMP / "inputs" / "reviewer's copied deck P1-37.pptx"
    make_deck(original)
    copied.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(original, copied)
    original_hash = sha256(original)
    copied_hash = sha256(copied)

    preflight_dir = TMP / "runs" / "preflight"
    preflight_summary = preflight_dir / "summary.json"
    preflight_log = preflight_dir / "logs" / "keynote_preflight.txt"
    preflight_pdf = preflight_dir / "previews" / "keynote_export.pdf"
    run(
        [
            sys.executable,
            str(SCRIPT),
            str(copied),
            str(preflight_pdf),
            "--preflight-only",
            "--log-txt",
            str(preflight_log),
            "--summary-json",
            str(preflight_summary),
        ]
    )
    preflight = load(preflight_summary)
    require(preflight["app_status"] in {"PASS", "BLOCKED_CONTROLADO"}, "Dishonest real preflight status")
    require(preflight_log.exists(), "Preflight log was not captured")
    require("log_txt" in artifact_types(preflight), "Preflight log is not typed")
    require(not preflight_pdf.exists(), "Preflight unexpectedly created a PDF")

    fake_dir = TMP / "runs" / "simulated"
    fake_deck = fake_dir / "copied.pptx"
    fake_deck.parent.mkdir(parents=True, exist_ok=True)
    fake_deck.write_text("synthetic copied deck", encoding="utf-8")

    export_calls: list[str] = []

    def forbidden_export(_input: Path, _output: Path):
        export_calls.append("called")
        return subprocess.CompletedProcess(["osascript"], 0, stdout="", stderr="")

    unconfirmed_pdf = fake_dir / "unconfirmed.pdf"
    unconfirmed_summary = fake_dir / "unconfirmed_summary.json"
    unconfirmed_log = fake_dir / "unconfirmed.log"
    code, _, _ = invoke_main_with_fakes(
        [
            str(fake_deck),
            str(unconfirmed_pdf),
            "--require-confirmation",
            "--log-txt",
            str(unconfirmed_log),
            "--summary-json",
            str(unconfirmed_summary),
        ],
        report=ready_report(),
        exporter=forbidden_export,
    )
    require(code == 2, "Missing confirmation did not block")
    require(not export_calls, "Keynote exporter ran without confirmation")
    require(not unconfirmed_pdf.exists(), "Unconfirmed flow created a PDF")
    require(load(unconfirmed_summary)["app_status"] == "BLOCKED_CONTROLADO", "Unconfirmed status is not controlled")

    missing_pdf = fake_dir / "missing_backend.pdf"
    missing_summary = fake_dir / "missing_backend_summary.json"
    missing_log = fake_dir / "missing_backend.log"
    code, _, _ = invoke_main_with_fakes(
        [
            str(fake_deck),
            str(missing_pdf),
            "--require-confirmation",
            "--confirmation-id",
            CONFIRMATION_ID,
            "--log-txt",
            str(missing_log),
            "--summary-json",
            str(missing_summary),
        ],
        report=blocked_report(),
        exporter=forbidden_export,
    )
    require(code == 2, "Missing Keynote backend did not block")
    require(not missing_pdf.exists(), "Missing backend flow created a PDF")
    require(load(missing_summary)["app_status"] == "BLOCKED_CONTROLADO", "Missing backend status is not controlled")
    require("copy-based" in " ".join(load(missing_summary)["next_actions"]), "Fallback action is missing")

    partial_pdf = fake_dir / "partial_native_failure.pdf"
    partial_summary = fake_dir / "partial_native_failure_summary.json"
    partial_log = fake_dir / "partial_native_failure.log"

    def partial_failure(_input: Path, output: Path):
        output.write_bytes(b"%PDF-1.4\npartial")
        return subprocess.CompletedProcess(["osascript"], 1, stdout="", stderr="Synthetic AppleScript denial")

    code, _, _ = invoke_main_with_fakes(
        [
            str(fake_deck),
            str(partial_pdf),
            "--require-confirmation",
            "--confirmation-id",
            CONFIRMATION_ID,
            "--log-txt",
            str(partial_log),
            "--summary-json",
            str(partial_summary),
        ],
        report=ready_report(),
        exporter=partial_failure,
    )
    require(code == 2, "Native AppleScript failure did not block")
    require(not partial_pdf.exists(), "Partial failed PDF was not removed")
    require("Synthetic AppleScript denial" in partial_log.read_text(encoding="utf-8"), "Native failure log is incomplete")

    timeout_pdf = fake_dir / "partial_timeout.pdf"
    timeout_summary = fake_dir / "partial_timeout_summary.json"
    timeout_log = fake_dir / "partial_timeout.log"

    def timed_out_export(_input: Path, output: Path):
        output.write_bytes(b"%PDF-1.4\npartial")
        raise subprocess.TimeoutExpired(["osascript"], timeout=300)

    code, _, _ = invoke_main_with_fakes(
        [
            str(fake_deck),
            str(timeout_pdf),
            "--require-confirmation",
            "--confirmation-id",
            CONFIRMATION_ID,
            "--log-txt",
            str(timeout_log),
            "--summary-json",
            str(timeout_summary),
        ],
        report=ready_report(),
        exporter=timed_out_export,
    )
    require(code == 2, "Native AppleScript timeout did not block")
    require(not timeout_pdf.exists(), "Partial timed-out PDF was not removed")
    require(load(timeout_summary)["app_status"] == "BLOCKED_CONTROLADO", "Timeout status is not controlled")
    require("safety timeout" in timeout_log.read_text(encoding="utf-8"), "Timeout log is incomplete")

    real_status = preflight["app_status"]
    real_artifacts: dict[str, str] = {
        "preflight_summary": str(preflight_summary),
        "preflight_log": str(preflight_log),
    }
    pdf_text = None
    if real_status == "PASS":
        confirmed_dir = TMP / "runs" / "confirmed"
        pdf = confirmed_dir / "previews" / "Controlled Keynote export.pdf"
        log = confirmed_dir / "logs" / "keynote_export.txt"
        summary_path = confirmed_dir / "summary.json"
        manifest_path = confirmed_dir / "manifest.json"
        run(
            [
                sys.executable,
                str(SCRIPT),
                str(copied),
                str(pdf),
                "--require-confirmation",
                "--confirmation-id",
                CONFIRMATION_ID,
                "--log-txt",
                str(log),
                "--summary-json",
                str(summary_path),
                "--manifest-json",
                str(manifest_path),
            ]
        )
        summary = load(summary_path)
        manifest = load(manifest_path)
        require(summary["app_status"] == "PASS", "Confirmed real export did not pass")
        require(pdf.exists() and pdf.stat().st_size > 0, "Confirmed Keynote PDF is missing or empty")
        require(log.exists() and "status=PASS" in log.read_text(encoding="utf-8"), "Confirmed execution log is missing")
        require(
            {"preview_pdf", "log_txt", "summary_json", "manifest_json"} <= artifact_types(summary),
            f"Confirmed artifact typing is incomplete: {artifact_types(summary)}",
        )
        confirmation = summary["results"]["confirmation"]
        require(confirmation["enforced_by_backend"] is True, "Backend confirmation was not enforced")
        require(confirmation["confirmation_id"] == CONFIRMATION_ID, "Confirmation identifier is missing")
        require(manifest["parameters"]["confirmation"]["confirmation_id"] == CONFIRMATION_ID, "Manifest lacks confirmation")
        output_names = {Path(item["path"]).name for item in manifest["outputs"]}
        require({"Controlled Keynote export.pdf", "keynote_export.txt"} <= output_names, "Manifest lacks PDF or log")
        logging.getLogger("pypdf").setLevel(logging.ERROR)
        from pypdf import PdfReader

        reader = PdfReader(str(pdf))
        require(len(reader.pages) == 1, "Expected one exported slide/page")
        pdf_text = "\n".join(page.extract_text() or "" for page in reader.pages)
        require("Controlled Keynote export P1-37" in pdf_text, "Exported PDF text is not traceable to the fixture")
        real_artifacts.update(
            {
                "pdf": str(pdf),
                "log": str(log),
                "summary": str(summary_path),
                "manifest": str(manifest_path),
            }
        )

    require(sha256(original) == original_hash, "Original deck was modified")
    require(sha256(copied) == copied_hash, "Staged copied deck was modified")

    app_checks = {
        "Sources/ScientificWorkbench/Services/DocumentWorkflowService.swift": [
            "validateKeynoteExportApproval",
            "keynoteConfirmationRequired",
            "keynotePreflightRequired",
        ],
        "Sources/ScientificWorkbench/Stores/WorkbenchStore.swift": [
            "--require-confirmation",
            "--confirmation-id",
            "logs/keynote_export.txt",
            "isExportingKeynote",
        ],
        "Sources/ScientificWorkbench/Views/KeynoteExportView.swift": [
            "Keynote automation is running on the copied deck",
            "I approve opening Keynote",
        ],
        "Tests/ScientificWorkbenchTests/DocumentWorkflowTests.swift": [
            "documentWorkflowRequiresReadyPreflightAndKeynoteConfirmation",
            "artifactDiscoveryReadsControlledKeynoteArtifacts",
        ],
    }
    for relative, markers in app_checks.items():
        text = (APP / relative).read_text(encoding="utf-8")
        for marker in markers:
            require(marker in text, f"{relative} lacks {marker}")

    result = {
        "tool": "audit_v2_0_p1_37_keynote_export_gui_controlled_regression",
        "status": "PASS",
        "real_backend_status": real_status,
        "real_capabilities": preflight["results"]["capabilities"],
        "original_modified": False,
        "simulated_cases": {
            "unconfirmed": "BLOCKED_CONTROLADO",
            "backend_absent": "BLOCKED_CONTROLADO",
            "native_failure_with_partial_pdf": "BLOCKED_CONTROLADO",
            "native_timeout_with_partial_pdf": "BLOCKED_CONTROLADO",
        },
        "fingerprints": {
            "original_before": original_hash,
            "original_after": sha256(original),
            "copy_before": copied_hash,
            "copy_after": sha256(copied),
        },
        "pdf_text": pdf_text,
        "artifacts": real_artifacts,
        "warnings": [] if real_status == "PASS" else preflight["results"]["blocking_findings"],
    }
    output = TMP / "p1_37_summary.json"
    output.write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
