#!/usr/bin/env python3
"""v1.7 phase 3 regression for transferable astro-born patterns."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path


SCRIPT_DIR = Path(__file__).resolve().parent
SKILL_ROOT = SCRIPT_DIR.parent
DEFAULT_OUTPUT_DIR = Path(tempfile.gettempdir()) / "sda_v1_7_phase3_astro_transfer_patterns"

TARGETS = {
    "catalog_workbench.py crossmatch-sky": SCRIPT_DIR / "catalog_workbench.py",
    "fits_rgb_batch.py": SCRIPT_DIR / "fits_rgb_batch.py",
    "rgb_visual_fits_export.py": SCRIPT_DIR / "rgb_visual_fits_export.py",
    "astrometry_net_workbench.py preflight": SCRIPT_DIR / "astrometry_net_workbench.py",
    "radial_velocity_workbench.py validate-manifest": SCRIPT_DIR / "radial_velocity_workbench.py",
    "presentation_workbench.py existing-deck-style-audit": SCRIPT_DIR / "presentation_workbench.py",
}


def fail(message: str) -> None:
    raise AssertionError(message)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output-dir",
        default=str(DEFAULT_OUTPUT_DIR),
        help="Directory for generated fixtures, command logs, and reports.",
    )
    return parser.parse_args()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_json(path: Path) -> dict:
    return json.loads(
        path.read_text(encoding="utf-8"),
        parse_constant=lambda token: fail(f"non-strict JSON constant emitted: {token}"),
    )


def write_csv(path: Path, fieldnames: list[str], rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def copy_to_temp(source: Path, destination: Path) -> Path:
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, destination)
    return destination


def run_command(case_dir: Path, command: list[str], timeout: int = 120) -> subprocess.CompletedProcess[str]:
    case_dir.mkdir(parents=True, exist_ok=True)
    proc = subprocess.run(
        command,
        cwd=SKILL_ROOT,
        text=True,
        capture_output=True,
        timeout=timeout,
    )
    (case_dir / "command.txt").write_text(" ".join(command) + "\n", encoding="utf-8")
    (case_dir / "stdout.txt").write_text(proc.stdout, encoding="utf-8")
    (case_dir / "stderr.txt").write_text(proc.stderr, encoding="utf-8")
    (case_dir / "returncode.txt").write_text(str(proc.returncode) + "\n", encoding="utf-8")
    combined = proc.stdout + "\n" + proc.stderr
    if "Traceback (most recent call last)" in combined:
        fail(f"traceback leaked for command: {' '.join(command)}")
    return proc


def require_status(payload: dict, expected: set[str], label: str) -> str:
    status = payload.get("status")
    qa_status = (payload.get("qa") or {}).get("status")
    if status not in expected:
        fail(f"{label}: expected status in {sorted(expected)}, got {status}")
    if qa_status != status:
        fail(f"{label}: qa.status {qa_status} did not match status {status}")
    return str(status)


def make_gaussian(shape: tuple[int, int], y0: float, x0: float):
    import numpy as np

    yy, xx = np.indices(shape)
    return (25.0 + 1800.0 * np.exp(-(((yy - y0) ** 2) + ((xx - x0) ** 2)) / (2.0 * 3.0**2))).astype("float32")


def write_generic_channel_fits(path: Path, filt: str, y0: float, x0: float) -> None:
    from astropy.io import fits

    data = make_gaussian((64, 64), y0, x0)
    header = fits.Header()
    header["OBJECT"] = "QC_PANEL"
    header["FILTER"] = filt
    header["EXPTIME"] = 1.0
    header["BUNIT"] = "adu"
    header["ORIGIN"] = "v1.7 non-astro synthetic sensor channel"
    fits.PrimaryHDU(data=data, header=header).writeto(path, overwrite=True)


def write_calibration_like_fits(path: Path) -> None:
    import numpy as np
    from astropy.io import fits

    data = np.ones((10, 190), dtype="float32")
    data[:, 25:28] = 150.0
    data[:, 80:83] = 220.0
    data[:, 140:143] = 180.0
    header = fits.Header()
    header["OBJECT"] = "barcode_calibration_panel"
    header["FILTER"] = "CLEAR"
    header["EXPTIME"] = 0.1
    fits.PrimaryHDU(data=data, header=header).writeto(path, overwrite=True)


def write_rgb_png(path: Path, *, red_offset: int = 0, size: tuple[int, int] = (48, 36)) -> None:
    import numpy as np
    from PIL import Image

    width, height = size
    yy, xx = np.indices((height, width))
    red = np.clip(xx * 255 / max(1, width - 1) + red_offset, 0, 255)
    green = np.clip(yy * 255 / max(1, height - 1), 0, 255)
    blue = np.full((height, width), 80)
    Image.fromarray(np.dstack([red, green, blue]).astype("uint8"), mode="RGB").save(path)


def write_deck_figure(path: Path, label: str) -> None:
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (900, 520), (246, 248, 250))
    draw = ImageDraw.Draw(image)
    draw.rectangle([40, 40, 860, 480], outline=(45, 96, 145), width=6)
    draw.line([90, 410, 810, 135], fill=(225, 170, 62), width=8)
    draw.text((70, 65), label, fill=(20, 20, 20))
    image.save(path)


def make_business_deck(path: Path, figure_a: Path, figure_b: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    presentation = Presentation()
    blank = presentation.slide_layouts[6]

    def add_text(slide, text: str, left: float, top: float, width: float, height: float, size: int) -> None:
        shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        shape.text_frame.text = text
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(size)

    slide = presentation.slides.add_slide(blank)
    add_text(slide, "Operations QA", 0.55, 0.25, 8.5, 0.55, 28)
    slide.shapes.add_picture(str(figure_a), Inches(1.0), Inches(1.15), width=Inches(7.4))
    add_text(slide, "Figure 1. Weekly sensor health overview.", 0.8, 6.35, 8.0, 0.35, 12)

    slide = presentation.slides.add_slide(blank)
    add_text(slide, "Before / After", 0.55, 0.25, 8.5, 0.55, 28)
    slide.shapes.add_picture(str(figure_a), Inches(0.65), Inches(1.2), width=Inches(4.05))
    slide.shapes.add_picture(str(figure_b), Inches(5.0), Inches(1.2), width=Inches(4.05))
    add_text(slide, "Figures A and B share layout for visual QA.", 0.8, 6.35, 8.0, 0.35, 12)
    presentation.save(path)


def make_source_fixtures(source_dir: Path) -> dict[str, Path]:
    source_dir.mkdir(parents=True, exist_ok=True)

    left_sites = source_dir / "field_assets_lonlat.csv"
    right_sites = source_dir / "maintenance_observations_lonlat.csv"
    write_csv(
        left_sites,
        ["asset_id", "lon_deg", "lat_deg", "asset_type"],
        [
            {"asset_id": "A-001", "lon_deg": "10.000000", "lat_deg": "20.000000", "asset_type": "sensor"},
            {"asset_id": "A-002", "lon_deg": "10.001000", "lat_deg": "20.001000", "asset_type": "sensor"},
            {"asset_id": "A-003", "lon_deg": "11.500000", "lat_deg": "21.500000", "asset_type": "camera"},
        ],
    )
    write_csv(
        right_sites,
        ["observation_id", "lon_deg", "lat_deg", "status"],
        [
            {"observation_id": "M-101", "lon_deg": "10.000120", "lat_deg": "20.000080", "status": "ok"},
            {"observation_id": "M-102", "lon_deg": "10.001090", "lat_deg": "20.001070", "status": "review"},
            {"observation_id": "M-999", "lon_deg": "42.000000", "lat_deg": "2.000000", "status": "unrelated"},
        ],
    )

    fits_root = source_dir / "generic_sensor_fits_channels"
    fits_root.mkdir()
    write_generic_channel_fits(fits_root / "qc_panel_B.fits", "B", 32.0, 32.0)
    write_generic_channel_fits(fits_root / "qc_panel_G.fits", "G", 32.2, 31.9)
    write_generic_channel_fits(fits_root / "qc_panel_R.fits", "R", 31.8, 32.1)

    calibration_like = source_dir / "barcode_calibration_panel.fits"
    write_calibration_like_fits(calibration_like)

    png_dir = source_dir / "rgb_png"
    png_dir.mkdir()
    previous_png = png_dir / "previous_dashboard_rgb.png"
    new_png = png_dir / "new_dashboard_rgb.png"
    write_rgb_png(previous_png, red_offset=0)
    write_rgb_png(new_png, red_offset=18)

    deck_dir = source_dir / "deck"
    deck_dir.mkdir()
    fig_a = deck_dir / "ops_figure_a.png"
    fig_b = deck_dir / "ops_figure_b.png"
    deck = deck_dir / "operations_qa_deck.pptx"
    write_deck_figure(fig_a, "operations figure A")
    write_deck_figure(fig_b, "operations figure B")
    make_business_deck(deck, fig_a, fig_b)

    rv_manifest_dir = source_dir / "rv_manifest_pattern"
    figures = rv_manifest_dir / "figures"
    figures.mkdir(parents=True)
    for name in ("raw_rv.png", "initial_periodogram.png", "fitted_rv.png", "statistics.png", "residual_periodogram.png", "dynamics.png"):
        write_rgb_png(figures / name, red_offset=5, size=(24, 18))
    manifest = {
        "schema_version": "1.0",
        "tool": "radial_velocity_workbench",
        "language": "en",
        "title": "Generic measurement-session manifest pattern",
        "backend": "generic-rv",
        "system_name": "QC Probe Session",
        "source_files": ["production_probe_measurements.vels"],
        "star_metadata": {"name": "QC Probe Session"},
        "rv_datasets": [{"source": "production_probe_measurements.vels", "value_units": "micrometers", "sample_count": 12}],
        "objective": "Validate that a manually curated measurement-session manifest is complete before a handoff.",
        "initial_search": {
            "raw_rv_figure": "figures/raw_rv.png",
            "initial_periodogram_figure": "figures/initial_periodogram.png",
            "dominant_periods_days": [],
            "false_alarm_probabilities": [],
            "notes": "Non-astro pattern fixture; no planet claim is made.",
        },
        "fitted_model": {
            "planet_count": 0,
            "fitted_rv_figure": "figures/fitted_rv.png",
            "statistics_figure": "figures/statistics.png",
            "residual_periodogram_figure": "figures/residual_periodogram.png",
            "chi2_before": 18.2,
            "chi2_after": 7.1,
            "rms_before": 0.82,
            "rms_after": 0.31,
            "optimization_notes": "Manifest-completeness pattern only.",
        },
        "dynamics": {"integration_horizon": "not applicable", "dynamics_figure": "figures/dynamics.png", "notes": "No dynamical inference."},
        "conclusions": "Manifest contains the expected figures and scalar diagnostics for handoff.",
        "caveats": "This is not a general measurement-report schema; it remains an RV workbench contract.",
        "workflow_notes": ["v1.7 pattern fixture"],
        "report_readiness_notes": [],
        "validation": {},
        "planets": [],
    }
    manifest_path = rv_manifest_dir / "session_manifest.json"
    manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=True) + "\n", encoding="utf-8")

    return {
        "left_sites": left_sites,
        "right_sites": right_sites,
        "fits_root": fits_root,
        "calibration_like": calibration_like,
        "previous_png": previous_png,
        "new_png": new_png,
        "deck": deck,
        "rv_manifest": manifest_path,
    }


def case_catalog_crossmatch(output_dir: Path, left_source: Path, right_source: Path) -> dict:
    case_dir = output_dir / "runs" / "catalog_crossmatch_sky"
    left = copy_to_temp(left_source, case_dir / "input_copy" / left_source.name)
    right = copy_to_temp(right_source, case_dir / "input_copy" / right_source.name)
    before = {left: sha256(left), right: sha256(right)}
    output = case_dir / "matched_assets.csv"
    summary = case_dir / "crossmatch_summary.json"
    manifest = case_dir / "crossmatch_manifest.json"
    command = [
        sys.executable,
        str(TARGETS["catalog_workbench.py crossmatch-sky"]),
        "crossmatch-sky",
        str(left),
        str(right),
        str(output),
        "--left-ra",
        "lon_deg",
        "--left-dec",
        "lat_deg",
        "--right-ra",
        "lon_deg",
        "--right-dec",
        "lat_deg",
        "--radius-arcsec",
        "0.6",
        "--summary-json",
        str(summary),
        "--manifest-json",
        str(manifest),
    ]
    proc = run_command(case_dir, command)
    if proc.returncode != 0:
        fail(f"catalog crossmatch failed: {proc.stderr or proc.stdout}")
    if {left: sha256(left), right: sha256(right)} != before:
        fail("catalog crossmatch modified copied inputs")
    payload = read_json(summary)
    status = require_status(payload, {"ok"}, "catalog_workbench.py crossmatch-sky")
    matched_rows = payload["results"].get("matched_rows")
    if matched_rows != 2:
        fail(f"expected 2 coordinate matches, got {matched_rows}")
    if not output.exists() or not manifest.exists():
        fail("catalog crossmatch missing output table or manifest")
    return {
        "capability": "catalog_workbench.py crossmatch-sky",
        "astro_domain": "sky-coordinate nearest-neighbour matching with RA/Dec ranges and angular radius in arcsec",
        "reusable_pattern": "tolerance-based spherical coordinate crossmatch with explicit manifest",
        "decision": "adapt_with_limits",
        "input": f"{left} + {right}",
        "command": " ".join(command),
        "status": status.upper(),
        "warning_error": "none",
        "diagnostic": f"matched_rows={matched_rows}; max_sep_arcsec={payload.get('qa', {}).get('metrics', {}).get('max_sep_arcsec'):.4f}",
        "modifies_originals": "NO: copied input hashes preserved",
        "useful_for_real_person": "YES: useful for lon/lat-like spherical coordinate dedupe when arcsec tolerance is meaningful",
        "limits": "Not a generic geospatial engine: no CRS, projected distances, polygons, addresses, or road/network distance.",
        "artifact": str(summary),
    }


def case_fits_rgb_batch(output_dir: Path, source_root: Path) -> dict:
    case_dir = output_dir / "runs" / "fits_rgb_batch"
    input_root = case_dir / "input_copy" / source_root.name
    shutil.copytree(source_root, input_root)
    before = {path: sha256(path) for path in sorted(input_root.glob("*.fits"))}
    out_dir = case_dir / "out"
    summary = case_dir / "fits_rgb_summary.json"
    command = [
        sys.executable,
        str(TARGETS["fits_rgb_batch.py"]),
        "--input-root",
        str(input_root),
        "--output-dir",
        str(out_dir),
        "--alignment-mode",
        "none",
        "--no-stacks",
        "--clean-derived",
        "--summary-json",
        str(summary),
    ]
    proc = run_command(case_dir, command, timeout=180)
    if proc.returncode != 0:
        fail(f"fits_rgb_batch failed: {proc.stderr or proc.stdout}")
    after = {path: sha256(path) for path in sorted(input_root.glob("*.fits"))}
    if after != before:
        fail("fits_rgb_batch modified copied FITS inputs")
    payload = read_json(summary)
    status = require_status(payload, {"ok", "warning"}, "fits_rgb_batch.py")
    results = payload["results"]
    if results["counts"].get("groups_written") != 1 or results["modes"].get("true_rgb") != 1:
        fail(f"fits_rgb_batch did not create one true_rgb group: {results}")
    group = results["group_reports"][0]
    qa_path = Path(str(group["alignment_qa"]).replace("~", str(Path.home())))
    png_path = Path(str(group["png"]).replace("~", str(Path.home())))
    if not qa_path.exists() or not png_path.exists():
        fail("fits_rgb_batch missing QA or PNG output")
    return {
        "capability": "fits_rgb_batch.py",
        "astro_domain": "FITS image grouping by OBJECT/FILTER, visual RGB/pseudo-RGB rendering, WCS/phase alignment, and calibration-state warnings",
        "reusable_pattern": "multichannel image product with channel mapping, alignment QA, and reproducible run summary",
        "decision": "document_pattern_only",
        "input": str(input_root),
        "command": " ".join(command),
        "status": status.upper(),
        "warning_error": "; ".join(payload.get("qa", {}).get("findings", [])) or "none",
        "diagnostic": f"groups_written=1; mode=true_rgb; png={png_path.name}",
        "modifies_originals": "NO: copied FITS hashes preserved",
        "useful_for_real_person": "YES, if a non-astro instrument already stores channels in FITS with compatible metadata",
        "limits": "Still FITS/OBJECT/FILTER oriented; not a general image-batch tool for arbitrary PNG/JPEG folders.",
        "artifact": str(summary),
    }


def case_rgb_visual_export(output_dir: Path, source_png: Path, previous_source: Path) -> dict:
    case_dir = output_dir / "runs" / "rgb_visual_fits_export"
    input_png = copy_to_temp(source_png, case_dir / "input_copy" / source_png.name)
    previous_png = copy_to_temp(previous_source, case_dir / "input_copy" / previous_source.name)
    before = {input_png: sha256(input_png), previous_png: sha256(previous_png)}
    output_fits = case_dir / "dashboard_visual_rgb.fits"
    comparison = case_dir / "before_after.png"
    summary = case_dir / "rgb_visual_summary.json"
    manifest = case_dir / "rgb_visual_manifest.json"
    command = [
        sys.executable,
        str(TARGETS["rgb_visual_fits_export.py"]),
        str(input_png),
        str(output_fits),
        "--previous-png",
        str(previous_png),
        "--comparison-png",
        str(comparison),
        "--summary-json",
        str(summary),
        "--manifest-json",
        str(manifest),
        "--object-name",
        "Operations dashboard RGB",
        "--reason",
        "v1.7 non-astro display-only RGB export pattern",
    ]
    proc = run_command(case_dir, command, timeout=120)
    if proc.returncode != 0:
        fail(f"rgb_visual_fits_export failed: {proc.stderr or proc.stdout}")
    if {input_png: sha256(input_png), previous_png: sha256(previous_png)} != before:
        fail("rgb_visual_fits_export modified copied PNG inputs")
    payload = read_json(summary)
    status = require_status(payload, {"ok"}, "rgb_visual_fits_export.py")
    if payload["results"].get("rgb_cube_shape") != [3, 36, 48]:
        fail(f"unexpected RGB cube shape: {payload['results'].get('rgb_cube_shape')}")
    for path in (output_fits, comparison, manifest):
        if not path.exists():
            fail(f"rgb_visual_fits_export missing artifact: {path}")
    return {
        "capability": "rgb_visual_fits_export.py",
        "astro_domain": "display-only RGB FITS convention for rendered astronomy RGB products",
        "reusable_pattern": "reproducible display-only RGB container with channel extensions, comparison image, and manifest",
        "decision": "adapt_with_limits",
        "input": str(input_png),
        "command": " ".join(command),
        "status": status.upper(),
        "warning_error": "none",
        "diagnostic": f"rgb_cube_shape={payload['results'].get('rgb_cube_shape')}; comparison={comparison.name}",
        "modifies_originals": "NO: copied PNG hashes preserved",
        "useful_for_real_person": "YES: useful when a rendered RGB needs traceable archive/container handoff",
        "limits": "Display-only FITS; not calibrated data and not a generic image editor.",
        "artifact": str(summary),
    }


def case_astrometry_preflight(output_dir: Path, source_fits: Path) -> dict:
    case_dir = output_dir / "runs" / "astrometry_preflight"
    input_fits = copy_to_temp(source_fits, case_dir / "input_copy" / source_fits.name)
    before = sha256(input_fits)
    summary = case_dir / "preflight_summary.json"
    report = case_dir / "preflight_report.md"
    manifest = case_dir / "preflight_manifest.json"
    command = [
        sys.executable,
        str(TARGETS["astrometry_net_workbench.py preflight"]),
        "preflight",
        str(input_fits),
        "--summary-json",
        str(summary),
        "--report-md",
        str(report),
        "--manifest-json",
        str(manifest),
    ]
    proc = run_command(case_dir, command, timeout=120)
    if proc.returncode != 0:
        fail(f"astrometry preflight should warn, not fail: {proc.stderr or proc.stdout}")
    if sha256(input_fits) != before:
        fail("astrometry preflight modified copied FITS input")
    payload = read_json(summary)
    status = require_status(payload, {"warning"}, "astrometry_net_workbench.py preflight")
    findings = " ".join(payload.get("qa", {}).get("findings", []))
    if "spectroscopy" not in findings and "calibration" not in findings:
        fail(f"astrometry preflight did not warn against non-imaging solve: {findings}")
    return {
        "capability": "astrometry_net_workbench.py preflight",
        "astro_domain": "Astrometry.net solve-readiness, celestial WCS hints, and imaging-vs-calibration classification",
        "reusable_pattern": "preflight before an expensive or optional backend path",
        "decision": "do_not_generalize",
        "input": str(input_fits),
        "command": " ".join(command),
        "status": status.upper(),
        "warning_error": findings,
        "diagnostic": "correctly warned that the non-astro barcode/calibration-like FITS should not be blindly solved",
        "modifies_originals": "NO: copied FITS hash preserved",
        "useful_for_real_person": "YES as a negative gate: it prevents sending unsuitable calibration-like images to an astro solver",
        "limits": "The backend and hints are astrometry-specific; use the preflight pattern elsewhere, not this solver contract.",
        "artifact": str(summary),
    }


def case_rv_validate_manifest(output_dir: Path, source_manifest: Path) -> dict:
    case_dir = output_dir / "runs" / "rv_validate_manifest"
    input_root = case_dir / "input_copy"
    shutil.copytree(source_manifest.parent, input_root)
    manifest_path = input_root / source_manifest.name
    before = {path: sha256(path) for path in sorted(input_root.rglob("*")) if path.is_file()}
    summary = case_dir / "validation_summary.json"
    manifest = case_dir / "validation_manifest.json"
    command = [
        sys.executable,
        str(TARGETS["radial_velocity_workbench.py validate-manifest"]),
        "validate-manifest",
        str(manifest_path),
        "--summary-json",
        str(summary),
        "--manifest-json",
        str(manifest),
        "--language",
        "en",
    ]
    proc = run_command(case_dir, command)
    if proc.returncode != 0:
        fail(f"rv validate-manifest failed: {proc.stderr or proc.stdout}")
    after = {path: sha256(path) for path in sorted(input_root.rglob("*")) if path.is_file()}
    if after != before:
        fail("rv validate-manifest modified copied manifest tree")
    payload = read_json(summary)
    status = require_status(payload, {"ok"}, "radial_velocity_workbench.py validate-manifest")
    if payload["qa"]["metrics"].get("warning_count") != 0:
        fail(f"rv validate-manifest emitted unexpected warnings: {payload['qa']}")
    return {
        "capability": "radial_velocity_workbench.py validate-manifest",
        "astro_domain": "Systemic/RV session-manifest completeness before report generation",
        "reusable_pattern": "manual-analysis manifest validation with required figures, metrics, conclusions, and provenance",
        "decision": "document_pattern_only",
        "input": str(manifest_path),
        "command": " ".join(command),
        "status": status.upper(),
        "warning_error": "none",
        "diagnostic": "manifest schema completeness passed for a copied generic measurement-session fixture",
        "modifies_originals": "NO: copied manifest tree hashes preserved",
        "useful_for_real_person": "YES as a manifest/checklist pattern for manual handoff, but only through the RV-shaped schema",
        "limits": "Do not advertise as a general manifest validator; fields remain RV/Systemic specific.",
        "artifact": str(summary),
    }


def case_presentation_style_audit(output_dir: Path, source_deck: Path) -> dict:
    case_dir = output_dir / "runs" / "presentation_style_audit"
    input_deck = copy_to_temp(source_deck, case_dir / "input_copy" / source_deck.name)
    before = sha256(input_deck)
    out_dir = case_dir / "style_audit_out"
    summary = case_dir / "style_audit_summary.json"
    manifest = case_dir / "style_audit_manifest.json"
    command = [
        sys.executable,
        str(TARGETS["presentation_workbench.py existing-deck-style-audit"]),
        "existing-deck-style-audit",
        str(input_deck),
        str(out_dir),
        "--reference-slides",
        "1-2",
        "--summary-json",
        str(summary),
        "--manifest-json",
        str(manifest),
        "--rule",
        "Keep operational figures editable and avoid flattening charts into screenshots.",
    ]
    proc = run_command(case_dir, command, timeout=180)
    if proc.returncode != 0:
        fail(f"presentation style audit failed: {proc.stderr or proc.stdout}")
    if sha256(input_deck) != before:
        fail("presentation style audit modified copied input deck")
    payload = read_json(summary)
    status = require_status(payload, {"ok", "warning"}, "presentation_workbench.py existing-deck-style-audit")
    expected_files = ["scientific_asset_manifest.json", "presentation_constraints.json", "slide_content.md", "new_slide_templates.json"]
    missing = [name for name in expected_files if not (out_dir / name).exists()]
    if missing:
        fail(f"presentation style audit missing expected handoff artifacts: {missing}")
    return {
        "capability": "presentation_workbench.py existing-deck-style-audit",
        "astro_domain": "scientific deck geometry/provenance and visual QA for figure handoff",
        "reusable_pattern": "visual style audit, editability-risk scan, constraints file, and asset manifest for a copied deck",
        "decision": "adapt_with_limits",
        "input": str(input_deck),
        "command": " ".join(command),
        "status": status.upper(),
        "warning_error": "; ".join(str(item) for item in payload.get("qa", {}).get("findings", [])) or "none",
        "diagnostic": f"handoff artifacts={len(expected_files)}; output_dir={out_dir.name}",
        "modifies_originals": "NO: copied deck hash preserved",
        "useful_for_real_person": "YES: useful for professional deck handoff QA and visual/style provenance",
        "limits": "Does not replace a full presentation editor or visual sign-off; scientific field names may remain in artifacts.",
        "artifact": str(summary),
    }


def write_report(output_dir: Path, cases: list[dict]) -> Path:
    report = output_dir / "phase3_astro_transfer_patterns_report.md"
    lines = [
        "# AUDITORIA v1.7 - Fase 3: patrones astro transferibles sin romper astro",
        "",
        "## Alcance",
        "",
        "Se probaron o delimitaron seis capabilities astro-born con fixtures no astro copiados a temporal cuando el contrato lo soporta.",
        "",
        "## Resultados",
        "",
        "| Capability | Decision | Patron reusable | Input | Comando | Estado | Warning/error | Diagnostico | Modifica originales | Utilidad real | Artefacto |",
        "|---|---|---|---|---|---|---|---|---|---|---|",
    ]
    for case in cases:
        lines.append(
            "| {capability} | {decision} | {pattern} | `{input}` | `{command}` | {status} | {warning} | {diagnostic} | {modifies} | {useful} | `{artifact}` |".format(
                capability=case["capability"],
                decision=case["decision"],
                pattern=case["reusable_pattern"],
                input=case["input"],
                command=case["command"].replace("|", "\\|"),
                status=case["status"],
                warning=case["warning_error"].replace("|", "\\|"),
                diagnostic=case["diagnostic"].replace("|", "\\|"),
                modifies=case["modifies_originals"],
                useful=case["useful_for_real_person"],
                artifact=case["artifact"],
            )
        )
    lines.extend(["", "## Dominio estricto y limites", ""])
    for case in cases:
        lines.extend(
            [
                f"### {case['capability']}",
                "",
                f"- Parte estrictamente astro/cientifica: {case['astro_domain']}",
                f"- Patron reusable: {case['reusable_pattern']}",
                f"- Decision: {case['decision']}",
                f"- Limite honesto: {case['limits']}",
                "",
            ]
        )
    lines.extend(
        [
            "## Decision",
            "",
            "No se cambiaron contratos astro ni se crearon nuevas capabilities. La fase protege patrones transferibles mediante docs y regresion estrecha.",
        ]
    )
    report.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")
    return report


def run_phase(output_dir: Path) -> dict:
    output_dir = output_dir.expanduser().resolve()
    if output_dir.exists():
        shutil.rmtree(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    fixtures = make_source_fixtures(output_dir / "generated" / "source_inputs")
    cases = [
        case_catalog_crossmatch(output_dir, fixtures["left_sites"], fixtures["right_sites"]),
        case_fits_rgb_batch(output_dir, fixtures["fits_root"]),
        case_rgb_visual_export(output_dir, fixtures["new_png"], fixtures["previous_png"]),
        case_astrometry_preflight(output_dir, fixtures["calibration_like"]),
        case_rv_validate_manifest(output_dir, fixtures["rv_manifest"]),
        case_presentation_style_audit(output_dir, fixtures["deck"]),
    ]
    report = write_report(output_dir, cases)
    summary = {
        "phase": "v1.7 phase 3",
        "status": "ok",
        "capability_count": len(cases),
        "case_status_counts": {status: sum(1 for case in cases if case["status"] == status) for status in sorted({case["status"] for case in cases})},
        "output_dir": str(output_dir),
        "report_md": str(report),
        "cases": cases,
    }
    (output_dir / "phase3_summary.json").write_text(json.dumps(summary, indent=2, ensure_ascii=True) + "\n", encoding="utf-8")
    return summary


def main() -> int:
    args = parse_args()
    summary = run_phase(Path(args.output_dir))
    print(json.dumps({key: summary[key] for key in ("phase", "status", "capability_count", "case_status_counts", "report_md")}, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
