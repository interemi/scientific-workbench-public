#!/usr/bin/env python3
"""Regressions for coursework branch comparison, FITS fallback, and deck style inspection."""

from __future__ import annotations

import base64
import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path


ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "scripts"


def module_available(module: str) -> bool:
    completed = subprocess.run(
        [sys.executable, "-c", f"import {module}"],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    return completed.returncode == 0


def datanalysis_python() -> str:
    completed = subprocess.run(
        [sys.executable, str(SCRIPTS / "datanalysis_env.py"), "locate"],
        cwd=ROOT,
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    if completed.returncode != 0:
        raise AssertionError(f"Could not locate datanalysis runtime:\n{completed.stderr}")
    payload = json.loads(completed.stdout)
    python_path = payload.get("selected", {}).get("python")
    if not python_path:
        raise AssertionError(f"datanalysis runtime did not report a Python executable: {payload}")
    return str(Path(python_path).expanduser())


def run_cmd(args: list[str], *, tmp: Path) -> subprocess.CompletedProcess:
    env = os.environ.copy()
    env["PYTHONDONTWRITEBYTECODE"] = "1"
    env["MPLCONFIGDIR"] = str(tmp / "mplconfig")
    completed = subprocess.run(
        args,
        cwd=ROOT,
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        env=env,
    )
    if completed.returncode != 0:
        raise AssertionError(
            f"Command failed: {' '.join(args)}\nSTDOUT:\n{completed.stdout}\nSTDERR:\n{completed.stderr}"
        )
    return completed


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def write_minimal_notebook(path: Path) -> None:
    path.write_text(
        json.dumps(
            {
                "cells": [],
                "metadata": {"kernelspec": {"name": "python3", "display_name": "Python 3"}},
                "nbformat": 4,
                "nbformat_minor": 5,
            }
        )
        + "\n",
        encoding="utf-8",
    )


def check_notebook_branch_compare(tmp: Path) -> None:
    root = tmp / "branches"
    for branch in ["A", "C", "EDU"]:
        (root / branch).mkdir(parents=True)
    for suffix in [".png", ".fits", ".html", ".ipynb"]:
        target = root / "A" / f"m51_R_2024_photometry{suffix}"
        if suffix == ".ipynb":
            write_minimal_notebook(target)
        else:
            target.write_bytes(b"synthetic")
    for suffix in [".png", ".fits", ".ipynb"]:
        target = root / "C" / f"m51_R_2024_photometry{suffix}"
        if suffix == ".ipynb":
            write_minimal_notebook(target)
        else:
            target.write_bytes(b"synthetic")
    (root / "EDU" / "m51_R_2024_photometry.png").write_bytes(b"synthetic")
    (root / "EDU" / "m51_R_2024_photometry_copy.png").write_bytes(b"synthetic")

    output_dir = tmp / "branch_compare"
    summary = output_dir / "summary.json"
    run_cmd(
        [
            sys.executable,
            str(SCRIPTS / "notebook_branch_compare.py"),
            str(root),
            "--output-dir",
            str(output_dir),
            "--summary-json",
            str(summary),
        ],
        tmp=tmp,
    )
    payload = load_json(summary)
    results = payload["results"]
    if results["best_branch"]["branch"] != "A":
        raise AssertionError(f"Expected branch A as most consistent, got {results['best_branch']}")
    if results["missing_product_type_group_count"] < 1 or results["duplicate_group_count"] < 1:
        raise AssertionError(f"Expected missing and duplicate findings: {results}")
    for artifact in ["branch_product_inventory.csv", "equivalent_products.csv", "candidate_figures.md"]:
        if not (output_dir / artifact).exists():
            raise AssertionError(f"Missing branch compare artifact: {artifact}")


def fits_card(keyword: str, value) -> bytes:
    if isinstance(value, bool):
        rendered = "T" if value else "F"
    elif isinstance(value, str):
        rendered = f"'{value}'"
    else:
        rendered = str(value)
    return f"{keyword:<8}= {rendered:>20}".ljust(80).encode("ascii")


def write_simple_fits(path: Path) -> None:
    import numpy as np

    data = np.linspace(0, 1, 20 * 16, dtype=">f4").reshape(16, 20)
    cards = [
        fits_card("SIMPLE", True),
        fits_card("BITPIX", -32),
        fits_card("NAXIS", 2),
        fits_card("NAXIS1", 20),
        fits_card("NAXIS2", 16),
        fits_card("OBJECT", "SIMPLE_FALLBACK"),
        b"END".ljust(80),
    ]
    header = b"".join(cards)
    header += b" " * ((2880 - len(header) % 2880) % 2880)
    payload = data.tobytes()
    payload += b"\0" * ((2880 - len(payload) % 2880) % 2880)
    path.write_bytes(header + payload)


def check_inspect_fits_simple_fallback(tmp: Path) -> None:
    fits_path = tmp / "simple.fits"
    preview = tmp / "simple.png"
    summary = tmp / "simple.json"
    write_simple_fits(fits_path)
    run_cmd(
        [
            sys.executable,
            str(SCRIPTS / "inspect_fits.py"),
            str(fits_path),
            "--force-simple-fallback",
            "--preview",
            str(preview),
            "--crop",
            "2,2,10,8",
            "--summary-json",
            str(summary),
        ],
        tmp=tmp,
    )
    payload = load_json(summary)
    if payload["results"].get("fallback_mode") != "simple_primary_image":
        raise AssertionError(f"Simple fallback was not used: {payload['results']}")
    if not preview.exists() or preview.stat().st_size == 0:
        raise AssertionError("Simple fallback preview PNG was not written.")


def check_presentation_style_inspection(tmp: Path) -> None:
    runtime_python = sys.executable if module_available("pptx") else datanalysis_python()
    image_path = tmp / "pixel.png"
    image_path.write_bytes(
        base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAIAAACQd1PeAAAADUlEQVR4nGP4z8DwHwAFgwJ/lc0rWQAAAABJRU5ErkJggg=="
        )
    )
    deck = tmp / "science_deck.pptx"
    make_deck = tmp / "make_deck.py"
    make_deck.write_text(
        """
import sys
from pptx import Presentation
from pptx.util import Inches

deck, image_path = sys.argv[1], sys.argv[2]
prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[5])
slide.shapes.title.text = "Result comparison"
slide.shapes.add_picture(image_path, Inches(1.0), Inches(1.5), width=Inches(4.0))
slide = prs.slides.add_slide(prs.slide_layouts[5])
slide.shapes.title.text = "Dense notes"
box = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(8), Inches(4))
box.text_frame.text = " ".join(["method"] * 110)
prs.save(deck)
""".lstrip(),
        encoding="utf-8",
    )
    run_cmd([runtime_python, str(make_deck), str(deck), str(image_path)], tmp=tmp)

    summary = tmp / "deck_summary.json"
    run_cmd([runtime_python, str(SCRIPTS / "presentation_workbench.py"), "inspect", str(deck), "--summary-json", str(summary)], tmp=tmp)
    payload = load_json(summary)
    inspection = payload["results"]["inspection"]["inspection"]
    if inspection["slide_count"] != 2 or inspection["image_count"] < 1:
        raise AssertionError(f"Unexpected deck inspection: {inspection}")
    if "style_summary" not in inspection or not inspection["style_summary"]["patterns"]:
        raise AssertionError(f"Missing style summary: {inspection}")


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="sda_coursework_handoff_") as raw_tmp:
        tmp = Path(raw_tmp)
        checks = [
            check_notebook_branch_compare,
            check_inspect_fits_simple_fallback,
            check_presentation_style_inspection,
        ]
        for check in checks:
            check(tmp)
            print(f"PASS {check.__name__}")
    print("All coursework handoff regressions passed.")


if __name__ == "__main__":
    main()
