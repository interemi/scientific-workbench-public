#!/usr/bin/env python3
"""Inspect, audit, and package copied presentation decks for scientific handoff."""

from __future__ import annotations

import argparse
import json
import shutil
from pathlib import Path
from statistics import median

from office_roundtrip import export_pptx_pdf, export_pptx_text
from quicklook_bridge import generate_preview
from _internal.path_safety import find_output_input_collisions
from _internal.pdfium_backend import render_pdf_pages
from _internal.public_contract import build_tool_payload, emit_payload
from _internal.provenance_utils import public_path, sanitize_payload, write_manifest
from _internal.runtime_common import clean_known_stderr, configure_runtime, detect_presentation_export_backends


PPTX_EXTS = {".pptx", ".pptm"}
PRESENTATION_EXTS = PPTX_EXTS | {".key", ".odp"}
TITLE_TOP_RATIO = 0.26
CAPTION_BOTTOM_RATIO = 0.72
EDGE_CLIP_EPS = 0.003
SAFE_MARGIN_DEFAULT = 0.035


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)

    inspect = subparsers.add_parser("inspect", help="Inspect a presentation deck or package.")
    inspect.add_argument("input")
    inspect.add_argument("--summary-json")
    inspect.add_argument("--export-preview-pdf", help="Optional PDF preview path for PPTX/PPTM inputs.")
    inspect.add_argument("--backend", choices=["auto", "soffice", "keynote"], default="auto")

    audit = subparsers.add_parser(
        "existing-deck-style-audit",
        help="Audit a reference deck and derive reusable title/figure/caption style guidance.",
    )
    audit.add_argument("input")
    audit.add_argument("output_dir")
    audit.add_argument("--reference-slides", help="Optional 1-based slide selection like 1,2,5-7.")
    audit.add_argument("--backend", choices=["auto", "soffice", "keynote"], default="auto")
    audit.add_argument("--summary-json")
    audit.add_argument("--manifest-json")
    audit.add_argument(
        "--rule",
        action="append",
        default=[],
        help="Persist a user constraint, for example: 'Do not put notebook screenshots on slides'.",
    )

    handoff = subparsers.add_parser("handoff", help="Create a copied presentation handoff bundle.")
    handoff.add_argument("input")
    handoff.add_argument("output_dir")
    handoff.add_argument("--export-pdf", action="store_true")
    handoff.add_argument("--backend", choices=["auto", "soffice", "keynote"], default="auto")
    handoff.add_argument("--summary-json")
    handoff.add_argument("--manifest-json")
    handoff.add_argument(
        "--rule",
        action="append",
        default=[],
        help="Persist a user constraint, for example: 'Scope slides to photometric reduction only'.",
    )

    return parser.parse_args()


def _requested_outputs(args, input_path: Path) -> list[tuple[str, Path | str | None]]:
    outputs: list[tuple[str, Path | str | None]] = []
    if args.command == "inspect":
        return [
            ("--summary-json", args.summary_json),
            ("--export-preview-pdf", args.export_preview_pdf),
        ]
    output_dir = Path(args.output_dir)
    outputs.extend(
        [
            ("output_dir", output_dir),
            ("--summary-json", args.summary_json),
            ("--manifest-json", args.manifest_json),
        ]
    )
    derived_names = [
        "scientific_asset_manifest.json",
        "presentation_constraints.json",
        "slide_content.md",
        "speaker_script.md",
        "poster_text.md",
        "new_slide_templates.json",
        "style_audit_report.md",
        "reference_preview.pdf",
        "reference_preview.png",
    ]
    outputs.extend((f"derived:{name}", output_dir / name) for name in derived_names)
    if output_dir.is_dir():
        outputs.extend(
            ("existing-output-member", item)
            for item in output_dir.rglob("*")
            if item.is_file()
        )
    if args.command == "handoff":
        outputs.extend(
            [
                ("derived:copied_input", output_dir / input_path.name),
                ("derived:bundle_summary.json", output_dir / "bundle_summary.json"),
                ("derived:handoff_notes.md", output_dir / "handoff_notes.md"),
                ("derived:slide_text.md", output_dir / "slide_text.md"),
                ("derived:exported_pdf", output_dir / f"{input_path.stem}.pdf"),
            ]
        )
    return outputs


def _emit_collision_only(args, input_path: Path, collisions: list[dict[str, str]]) -> int:
    message = "Refusing presentation outputs that overlap the input deck or package."
    payload = build_tool_payload(
        f"presentation_workbench.{args.command}",
        status="blocked",
        notes=[message, "No output files were written and the input presentation was not modified."],
        artifacts={},
        results={
            "blocked_reason": message,
            "error_type": "output_input_collision",
            "input": public_path(input_path),
            "collisions": collisions,
        },
        qa={
            "status": "blocked",
            "findings": [{"severity": "high", "title": "Input/output collision", "detail": message}],
            "metrics": {"blocking_count": len(collisions)},
        },
        legacy={"blocked_reason": message, "collisions": collisions},
    )
    print(json.dumps(sanitize_payload(payload), indent=2, ensure_ascii=True))
    return 2


def _preflight_output_safety(args) -> int | None:
    input_path = Path(args.input).expanduser()
    collisions = find_output_input_collisions(
        [("input", input_path)],
        _requested_outputs(args, input_path),
    )
    return _emit_collision_only(args, input_path, collisions) if collisions else None


def copy_input(source: Path, target: Path) -> Path:
    if source.is_dir():
        shutil.copytree(source, target, dirs_exist_ok=True)
    else:
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, target)
    return target


def parse_slide_selector(selector: str | None, slide_count: int) -> list[int]:
    if not selector:
        return list(range(1, slide_count + 1))
    selected: set[int] = set()
    for token in selector.split(","):
        token = token.strip()
        if not token:
            continue
        if "-" in token:
            start_text, end_text = token.split("-", 1)
            try:
                start = int(start_text)
                end = int(end_text)
            except ValueError as exc:
                raise ValueError(
                    f"Invalid slide selector token {token!r}; use comma-separated slide numbers or ranges like 1,2,5-7."
                ) from exc
            for value in range(min(start, end), max(start, end) + 1):
                if 1 <= value <= slide_count:
                    selected.add(value)
        else:
            try:
                value = int(token)
            except ValueError as exc:
                raise ValueError(
                    f"Invalid slide selector token {token!r}; use comma-separated slide numbers or ranges like 1,2,5-7."
                ) from exc
            if 1 <= value <= slide_count:
                selected.add(value)
    if not selected:
        raise ValueError(f"Reference slide selector {selector!r} did not match any slide in a {slide_count}-slide deck.")
    return sorted(selected)


def _safe_float(value) -> float:
    try:
        return float(value)
    except Exception:
        return 0.0


def _median(values: list[float], fallback: float | None = None) -> float | None:
    cleaned = [float(value) for value in values if value is not None]
    if not cleaned:
        return fallback
    return float(median(cleaned))


def _box_from_shape(shape, slide_width: float, slide_height: float) -> dict:
    left = _safe_float(getattr(shape, "left", 0.0)) / slide_width
    top = _safe_float(getattr(shape, "top", 0.0)) / slide_height
    width = _safe_float(getattr(shape, "width", 0.0)) / slide_width
    height = _safe_float(getattr(shape, "height", 0.0)) / slide_height
    return {
        "left": round(left, 5),
        "top": round(top, 5),
        "width": round(width, 5),
        "height": round(height, 5),
        "right": round(left + width, 5),
        "bottom": round(top + height, 5),
        "center_x": round(left + width / 2.0, 5),
        "center_y": round(top + height / 2.0, 5),
        "area_ratio": round(width * height, 5),
    }


def _shape_kind(shape) -> str:
    shape_type = str(getattr(shape, "shape_type", "")).lower()
    if getattr(shape, "has_chart", False) or "chart" in shape_type:
        return "chart"
    if getattr(shape, "has_table", False):
        return "table"
    try:
        if hasattr(shape, "image") or "picture" in shape_type:
            return "picture"
    except Exception:
        pass
    if getattr(shape, "has_text_frame", False):
        return "text"
    return "other"


def _shape_text(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    lines = []
    for paragraph in shape.text_frame.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)
    return "\n".join(lines)


def _shape_font_points(shape) -> list[float]:
    points = []
    if not getattr(shape, "has_text_frame", False):
        return points
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            size = getattr(getattr(run, "font", None), "size", None)
            if size is not None:
                try:
                    points.append(float(size.pt))
                except Exception:
                    continue
    return points


def _placeholder_label(shape) -> str | None:
    try:
        if getattr(shape, "is_placeholder", False):
            return str(shape.placeholder_format.type)
    except Exception:
        return None
    return None


def _box_overlap_ratio(a: dict, b: dict) -> float:
    left = max(a["left"], b["left"])
    top = max(a["top"], b["top"])
    right = min(a["right"], b["right"])
    bottom = min(a["bottom"], b["bottom"])
    if right <= left or bottom <= top:
        return 0.0
    overlap = (right - left) * (bottom - top)
    denom = min(a["area_ratio"], b["area_ratio"]) or 1e-6
    return overlap / denom


def _pick_title_shape(text_shapes: list[dict]) -> dict | None:
    if not text_shapes:
        return None
    candidates = []
    for shape in text_shapes:
        box = shape["box"]
        if box["top"] > TITLE_TOP_RATIO:
            continue
        placeholder = (shape.get("placeholder") or "").lower()
        placeholder_bonus = 100.0 if "title" in placeholder else 0.0
        font_score = shape.get("font_pt_median") or 0.0
        score = placeholder_bonus + font_score - box["top"] * 40.0 - shape["word_count"] * 0.1
        candidates.append((score, shape))
    if not candidates:
        return min(text_shapes, key=lambda item: (item["box"]["top"], -item.get("font_pt_median", 0.0)))
    candidates.sort(key=lambda item: item[0], reverse=True)
    return candidates[0][1]


def _pick_caption_shapes(text_shapes: list[dict], figure_shapes: list[dict], title_shape: dict | None) -> list[dict]:
    captions = []
    title_id = title_shape.get("shape_id") if title_shape else None
    for shape in text_shapes:
        if shape["shape_id"] == title_id:
            continue
        box = shape["box"]
        word_count = shape["word_count"]
        if word_count == 0:
            continue
        if word_count > 45:
            continue
        if box["height"] > 0.22:
            continue
        likely_caption = box["top"] >= CAPTION_BOTTOM_RATIO or (shape.get("font_pt_median") or 0.0) <= 18.0
        if not likely_caption and figure_shapes:
            for figure in figure_shapes:
                if box["top"] >= figure["box"]["bottom"] - 0.01 and _box_overlap_ratio(box, figure["box"]) > 0.15:
                    likely_caption = True
                    break
        if likely_caption:
            captions.append(shape)
    captions.sort(key=lambda item: (item["box"]["top"], item["box"]["left"]))
    return captions


def _content_bounds(shapes: list[dict]) -> dict | None:
    if not shapes:
        return None
    return {
        "left": min(shape["box"]["left"] for shape in shapes),
        "top": min(shape["box"]["top"] for shape in shapes),
        "right": max(shape["box"]["right"] for shape in shapes),
        "bottom": max(shape["box"]["bottom"] for shape in shapes),
    }


def inspect_pptx(path: Path) -> dict:
    from pptx import Presentation

    presentation = Presentation(path)
    slide_width = float(presentation.slide_width)
    slide_height = float(presentation.slide_height)
    slide_titles = []
    text_blocks = 0
    image_count = 0
    table_count = 0
    chart_count = 0
    slide_summaries = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        shapes = []
        text_shapes = []
        figure_shapes = []
        chart_shapes = []
        table_shapes = []
        editable_nonfigure_count = 0
        slide_text_blocks = 0
        slide_words = 0
        for shape_index, shape in enumerate(slide.shapes, start=1):
            kind = _shape_kind(shape)
            box = _box_from_shape(shape, slide_width, slide_height)
            text = _shape_text(shape)
            font_points = _shape_font_points(shape)
            placeholder = _placeholder_label(shape)
            row = {
                "shape_id": shape_index,
                "name": getattr(shape, "name", f"Shape {shape_index}"),
                "kind": kind,
                "placeholder": placeholder,
                "box": box,
                "text_preview": text.splitlines()[0][:140] if text else None,
                "word_count": len(text.split()) if text else 0,
                "font_pt_median": round(_median(font_points, 0.0) or 0.0, 2) if font_points else None,
            }
            shapes.append(row)
            if kind == "picture":
                image_count += 1
                figure_shapes.append(row)
            elif kind == "chart":
                chart_count += 1
                chart_shapes.append(row)
                editable_nonfigure_count += 1
            elif kind == "table":
                table_count += 1
                table_shapes.append(row)
                editable_nonfigure_count += 1
            elif kind == "text":
                if text:
                    text_blocks += 1
                    slide_text_blocks += 1
                    slide_words += row["word_count"]
                    text_shapes.append(row)
                    editable_nonfigure_count += 1
            elif kind == "other":
                editable_nonfigure_count += 1
        title_shape = _pick_title_shape(text_shapes)
        title = title_shape["text_preview"] if title_shape and title_shape.get("text_preview") else f"Slide {slide_index}"
        slide_titles.append(title)
        captions = _pick_caption_shapes(text_shapes, figure_shapes + chart_shapes, title_shape)
        main_figure = None
        figure_like = figure_shapes + chart_shapes
        if figure_like:
            main_figure = max(figure_like, key=lambda item: item["box"]["area_ratio"])
        content_bounds = _content_bounds(shapes)
        rasterized_slide_risk = bool(
            main_figure
            and main_figure["box"]["area_ratio"] >= 0.82
            and editable_nonfigure_count <= 2
            and len(captions) <= 1
        )
        slide_summaries.append(
            {
                "index": slide_index,
                "title": title,
                "text_blocks": slide_text_blocks,
                "word_count": slide_words,
                "image_count": len(figure_shapes),
                "table_count": len(table_shapes),
                "chart_count": len(chart_shapes),
                "editable_nonfigure_count": editable_nonfigure_count,
                "shape_count": len(shapes),
                "title_box": title_shape["box"] if title_shape else None,
                "title_shape_id": title_shape["shape_id"] if title_shape else None,
                "caption_count": len(captions),
                "caption_boxes": [item["box"] for item in captions[:3]],
                "main_figure_box": main_figure["box"] if main_figure else None,
                "figure_boxes": [item["box"] for item in sorted(figure_like, key=lambda row: (row["box"]["left"], row["box"]["top"]))[:4]],
                "content_bounds": content_bounds,
                "rasterized_slide_risk": rasterized_slide_risk,
                "shape_inventory": shapes,
            }
        )
    style = classify_presentation_style(slide_summaries)
    return {
        "slide_count": len(presentation.slides),
        "slide_size": {
            "width_emu": int(slide_width),
            "height_emu": int(slide_height),
            "aspect_ratio": round(slide_width / slide_height, 4) if slide_height else None,
        },
        "slide_titles": slide_titles[:20],
        "text_blocks": text_blocks,
        "image_count": image_count,
        "table_count": table_count,
        "chart_count": chart_count,
        "slides": slide_summaries,
        "style_summary": style,
    }


def classify_presentation_style(slides: list[dict]) -> dict:
    if not slides:
        return {"dominant_style": "empty", "patterns": [], "acceptance_findings": ["No slides detected."]}
    slide_count = len(slides)
    total_images = sum(slide["image_count"] + slide["chart_count"] for slide in slides)
    total_text_blocks = sum(slide["text_blocks"] for slide in slides)
    total_words = sum(slide["word_count"] for slide in slides)
    comparative_slides = sum(1 for slide in slides if (slide["image_count"] + slide["chart_count"]) >= 2)
    dense_slides = sum(1 for slide in slides if slide["word_count"] > 90 or slide["text_blocks"] > 5)
    rasterized_risk_count = sum(1 for slide in slides if slide.get("rasterized_slide_risk"))
    avg_words = total_words / slide_count
    avg_images = total_images / slide_count
    avg_text_blocks = total_text_blocks / slide_count
    patterns = []
    if avg_images >= 1.0 and avg_text_blocks <= 3.0:
        patterns.append("figure-led")
    if avg_text_blocks >= 3.0 or avg_words > 70:
        patterns.append("text-led")
    if comparative_slides:
        patterns.append("comparative")
    if dense_slides:
        patterns.append("dense")
    if rasterized_risk_count:
        patterns.append("editable-risk")
    if not patterns:
        patterns.append("mixed-light")
    findings = []
    if "text-led" in patterns:
        findings.append("Slides look text-led; consider replacing prose with one protagonist figure plus short editable captions.")
    if "dense" in patterns:
        findings.append("At least one slide looks dense by text-block/word-count heuristics.")
    if total_images == 0:
        findings.append("No image-like scientific figures were detected.")
    if rasterized_risk_count:
        findings.append("At least one slide may be functioning as a flattened image instead of editable native objects.")
    return {
        "dominant_style": patterns[0],
        "patterns": patterns,
        "metrics": {
            "avg_words_per_slide": round(avg_words, 2),
            "avg_images_per_slide": round(avg_images, 2),
            "avg_text_blocks_per_slide": round(avg_text_blocks, 2),
            "comparative_slide_count": comparative_slides,
            "dense_slide_count": dense_slides,
            "editable_risk_slide_count": rasterized_risk_count,
        },
        "acceptance_findings": findings,
    }


def inspect_presentation(path: Path) -> dict:
    path = Path(path)
    summary = {
        "path": public_path(path),
        "exists": path.exists(),
        "suffix": path.suffix.lower(),
        "is_dir": path.is_dir(),
        "backends": detect_presentation_export_backends(),
    }
    if not path.exists():
        return summary
    if path.suffix.lower() in PPTX_EXTS and path.is_file():
        try:
            summary["inspection"] = inspect_pptx(path)
        except Exception as exc:
            message = clean_known_stderr(f"{exc.__class__.__name__}: {exc}") or exc.__class__.__name__
            summary["inspection_error"] = message
            summary["inspection"] = {
                "error": f"Could not inspect PPTX/PPTM deck: {message}",
                "error_type": exc.__class__.__name__,
            }
    elif path.suffix.lower() == ".key" and path.is_dir():
        members = sorted(str(item.relative_to(path)) for item in path.rglob("*") if item.is_file())
        summary["inspection"] = {
            "bundle_member_count": len(members),
            "bundle_members_preview": members[:25],
            "best_effort_note": "Keynote bundle inspection is best-effort. Use a backup, export previews, review visually, and keep a rollback path.",
        }
    else:
        summary["inspection"] = {"note": "No specialized inspector available for this format."}
    return summary


def build_dependency_preflight(backends: dict, suffix: str | None = None) -> dict:
    missing = []
    warnings = []
    if not backends.get("soffice") and not backends.get("keynote_available"):
        missing.append("No PDF export backend is currently available for copied presentation handoffs.")
    if suffix == ".key":
        warnings.append("Keynote support is best-effort and should always be validated visually on a backup copy.")
    recommendation = (
        "Install LibreOffice or use macOS Keynote automation if you need copied-deck PDF exports."
        if missing
        else "Presentation inspection is ready, and at least one export backend is available."
    )
    if warnings and not missing:
        recommendation += " For .key files, keep backup, preview export, and rollback discipline."
    return {
        "status": "warning" if missing or warnings else "ok",
        "optional_backend_missing": missing,
        "warning_findings": warnings,
        "recommendation": recommendation,
    }


def _median_box(boxes: list[dict]) -> dict | None:
    if not boxes:
        return None
    return {
        key: round(_median([box[key] for box in boxes], 0.0) or 0.0, 5)
        for key in ("left", "top", "width", "height", "right", "bottom", "center_x", "center_y", "area_ratio")
    }


def _compute_style_grid(reference_slides: list[dict]) -> dict:
    title_boxes = [slide["title_box"] for slide in reference_slides if slide.get("title_box")]
    primary_caption_boxes = [slide["caption_boxes"][0] for slide in reference_slides if slide.get("caption_boxes")]
    content_bounds = [slide["content_bounds"] for slide in reference_slides if slide.get("content_bounds")]
    single_figure_boxes = [
        slide["main_figure_box"]
        for slide in reference_slides
        if slide.get("main_figure_box") and (slide["image_count"] + slide["chart_count"]) == 1
    ]
    comparative_left = []
    comparative_right = []
    for slide in reference_slides:
        figure_boxes = list(slide.get("figure_boxes") or [])
        if len(figure_boxes) >= 2:
            ordered = sorted(figure_boxes, key=lambda box: (box["left"], box["top"]))
            comparative_left.append(ordered[0])
            comparative_right.append(ordered[1])
    safe_left = _median([bounds["left"] for bounds in content_bounds], SAFE_MARGIN_DEFAULT) or SAFE_MARGIN_DEFAULT
    safe_top = _median([bounds["top"] for bounds in content_bounds], SAFE_MARGIN_DEFAULT) or SAFE_MARGIN_DEFAULT
    safe_right_margin = _median([1.0 - bounds["right"] for bounds in content_bounds], SAFE_MARGIN_DEFAULT) or SAFE_MARGIN_DEFAULT
    safe_bottom_margin = _median([1.0 - bounds["bottom"] for bounds in content_bounds], SAFE_MARGIN_DEFAULT) or SAFE_MARGIN_DEFAULT
    return {
        "safe_area": {
            "left": round(safe_left, 5),
            "top": round(safe_top, 5),
            "right_margin": round(safe_right_margin, 5),
            "bottom_margin": round(safe_bottom_margin, 5),
        },
        "title_box": _median_box(title_boxes),
        "primary_caption_box": _median_box(primary_caption_boxes),
        "single_figure_box": _median_box(single_figure_boxes),
        "comparative_pair_boxes": {
            "left": _median_box(comparative_left),
            "right": _median_box(comparative_right),
        },
        "visual_density": {
            "median_words_per_slide": round(_median([slide["word_count"] for slide in reference_slides], 0.0) or 0.0, 2),
            "median_text_blocks_per_slide": round(_median([slide["text_blocks"] for slide in reference_slides], 0.0) or 0.0, 2),
            "median_figures_per_slide": round(
                _median([slide["image_count"] + slide["chart_count"] for slide in reference_slides], 0.0) or 0.0,
                2,
            ),
        },
    }


def _slide_visual_findings(slide: dict, style_grid: dict) -> list[dict]:
    findings = []
    safe_area = style_grid.get("safe_area") or {}
    bounds = slide.get("content_bounds")
    if bounds:
        if bounds["left"] < max(0.0, safe_area.get("left", SAFE_MARGIN_DEFAULT) - 0.01):
            findings.append({"severity": "medium", "title": "Left margin drift", "detail": "Some content sits noticeably left of the inferred safe area."})
        if bounds["top"] < max(0.0, safe_area.get("top", SAFE_MARGIN_DEFAULT) - 0.01):
            findings.append({"severity": "medium", "title": "Top margin drift", "detail": "Some content sits noticeably above the inferred safe area."})
        if (1.0 - bounds["right"]) < max(0.0, safe_area.get("right_margin", SAFE_MARGIN_DEFAULT) - 0.01):
            findings.append({"severity": "medium", "title": "Right margin drift", "detail": "Some content sits too close to the right edge relative to neighboring slides."})
        if (1.0 - bounds["bottom"]) < max(0.0, safe_area.get("bottom_margin", SAFE_MARGIN_DEFAULT) - 0.01):
            findings.append({"severity": "medium", "title": "Bottom margin drift", "detail": "Some content sits too close to the bottom edge relative to neighboring slides."})
    for shape in slide.get("shape_inventory", []):
        box = shape["box"]
        if box["left"] < -EDGE_CLIP_EPS or box["top"] < -EDGE_CLIP_EPS or box["right"] > 1.0 + EDGE_CLIP_EPS or box["bottom"] > 1.0 + EDGE_CLIP_EPS:
            findings.append({"severity": "high", "title": "Possible clipping", "detail": f"Shape `{shape['name']}` appears to fall outside slide bounds."})
    title_box = slide.get("title_box")
    template_title = style_grid.get("title_box")
    if title_box and template_title:
        if abs(title_box["left"] - template_title["left"]) > 0.08 or abs(title_box["top"] - template_title["top"]) > 0.06:
            findings.append({"severity": "low", "title": "Title alignment drift", "detail": "The title box deviates from the dominant title position in the reference slides."})
    if (slide["image_count"] + slide["chart_count"]) > 0 and slide.get("caption_count", 0) == 0:
        findings.append({"severity": "medium", "title": "Missing visible caption", "detail": "This slide contains a figure-like object but no likely editable caption was detected."})
    if slide.get("rasterized_slide_risk"):
        findings.append({"severity": "high", "title": "Editable-native-first risk", "detail": "The slide may be functioning as one large raster image rather than editable native objects."})
    if slide["word_count"] > 95 or slide["text_blocks"] > 5:
        findings.append({"severity": "low", "title": "Dense slide", "detail": "The slide looks dense by word-count or text-block heuristics."})
    main_figure_box = slide.get("main_figure_box")
    template_figure = style_grid.get("single_figure_box")
    if main_figure_box and template_figure and (slide["image_count"] + slide["chart_count"]) == 1:
        if abs(main_figure_box["center_x"] - template_figure["center_x"]) > 0.08:
            findings.append({"severity": "low", "title": "Main figure off-center", "detail": "The protagonist figure is noticeably off the typical horizontal alignment."})
    return findings


SEVERITY_RANK = {"high": 0, "medium": 1, "low": 2}
MAX_QA_FINDINGS = 18


def _annotated_visual_findings(visual_audit: dict) -> list[dict]:
    rows = []
    for finding in visual_audit.get("deck_findings") or []:
        item = dict(finding)
        item.setdefault("scope", "deck")
        rows.append(item)
    for slide in visual_audit.get("slide_qa", []):
        for finding in slide.get("findings") or []:
            item = dict(finding)
            item["scope"] = "slide"
            item["slide_index"] = slide.get("index")
            item["slide_title"] = slide.get("title")
            rows.append(item)
    return rows


def _summarize_visual_findings(visual_audit: dict, *, limit: int = MAX_QA_FINDINGS) -> dict:
    findings = _annotated_visual_findings(visual_audit)
    severity_counts = {severity: 0 for severity in ("high", "medium", "low")}
    title_counts: dict[str, int] = {}
    slide_counts: dict[str, dict] = {}
    for finding in findings:
        severity = finding.get("severity", "low")
        severity_counts[severity] = severity_counts.get(severity, 0) + 1
        title = finding.get("title", "Finding")
        title_counts[title] = title_counts.get(title, 0) + 1
        if finding.get("slide_index") is not None:
            key = str(finding["slide_index"])
            slide_counts.setdefault(key, {"slide_index": finding["slide_index"], "slide_title": finding.get("slide_title"), "finding_count": 0})
            slide_counts[key]["finding_count"] += 1
    ordered = sorted(
        findings,
        key=lambda item: (
            SEVERITY_RANK.get(item.get("severity", "low"), 3),
            item.get("slide_index") if item.get("slide_index") is not None else -1,
            item.get("title", ""),
        ),
    )
    return {
        "finding_count": len(findings),
        "severity_counts": severity_counts,
        "top_titles": [
            {"title": title, "count": count}
            for title, count in sorted(title_counts.items(), key=lambda item: (-item[1], item[0]))[:8]
        ],
        "top_slides": sorted(slide_counts.values(), key=lambda item: (-item["finding_count"], item["slide_index"]))[:8],
        "top_findings": ordered[:limit],
        "truncated": len(ordered) > limit,
        "limit": limit,
    }


def _build_visual_audit(inspection: dict, reference_slide_selector: str | None = None) -> dict:
    slides = inspection.get("slides") or []
    selected_indices = parse_slide_selector(reference_slide_selector, len(slides))
    reference_slides = [slide for slide in slides if slide["index"] in selected_indices]
    style_grid = _compute_style_grid(reference_slides)
    slide_qa = []
    all_findings = []
    for slide in reference_slides:
        findings = _slide_visual_findings(slide, style_grid)
        slide_qa.append(
            {
                "index": slide["index"],
                "title": slide["title"],
                "status": "warning" if findings else "ok",
                "findings": findings,
                "metrics": {
                    "word_count": slide["word_count"],
                    "text_blocks": slide["text_blocks"],
                    "figure_count": slide["image_count"] + slide["chart_count"],
                    "caption_count": slide["caption_count"],
                },
            }
        )
        all_findings.extend(findings)
    deck_findings = []
    if style_grid.get("title_box") is None:
        deck_findings.append({"severity": "medium", "title": "No stable title pattern", "detail": "The audit could not derive a robust title region from the selected reference slides."})
    if style_grid.get("single_figure_box") is None and not style_grid.get("comparative_pair_boxes", {}).get("left"):
        deck_findings.append({"severity": "medium", "title": "No stable figure grid", "detail": "The audit could not derive a reusable figure region from the selected reference slides."})
    if any(finding["severity"] == "high" for finding in all_findings):
        status = "warning"
    elif deck_findings or all_findings:
        status = "warning"
    else:
        status = "ok"
    new_slide_templates = {
        "single_figure": {
            "title_box": style_grid.get("title_box"),
            "figure_box": style_grid.get("single_figure_box"),
            "caption_box": style_grid.get("primary_caption_box"),
            "safe_area": style_grid.get("safe_area"),
        },
        "comparative_pair": {
            "title_box": style_grid.get("title_box"),
            "left_figure_box": (style_grid.get("comparative_pair_boxes") or {}).get("left"),
            "right_figure_box": (style_grid.get("comparative_pair_boxes") or {}).get("right"),
            "caption_box": style_grid.get("primary_caption_box"),
            "safe_area": style_grid.get("safe_area"),
        },
    }
    visual_audit = {
        "status": status,
        "reference_slide_indices": selected_indices,
        "style_grid": style_grid,
        "new_slide_templates": new_slide_templates,
        "deck_findings": deck_findings,
        "slide_qa": slide_qa,
    }
    visual_audit["finding_summary"] = _summarize_visual_findings(visual_audit)
    return visual_audit


def _render_pdf_pages(pdf_path: Path, output_dir: Path, selected_pages: list[int]) -> list[str]:
    try:
        import pypdfium2
    except Exception:
        return []
    requests = [(number, output_dir / f"slide_{number:02d}.png") for number in selected_pages]
    return [str(path.resolve()) for path in render_pdf_pages(pdf_path, requests)]


def _export_reference_previews(input_path: Path, output_dir: Path, backend: str, selected_slides: list[int]) -> dict:
    suffix = input_path.suffix.lower()
    preview = {"success": False, "artifacts": {}, "notes": []}
    if suffix in PPTX_EXTS and input_path.is_file():
        preview_pdf = output_dir / "reference_preview.pdf"
        export_result = export_pptx_pdf(input_path, preview_pdf, backend=backend)
        preview["export_result"] = export_result
        if export_result.get("success") and preview_pdf.exists():
            preview["success"] = True
            preview["artifacts"]["preview_pdf"] = str(preview_pdf.resolve())
            preview_pages = _render_pdf_pages(preview_pdf, output_dir / "reference_preview_pages", selected_slides)
            if preview_pages:
                preview["artifacts"]["preview_pages"] = preview_pages
            else:
                preview["notes"].append("Preview PDF rendered, but page PNG export was unavailable.")
        else:
            preview["notes"].append(export_result.get("stderr") or "Reference preview PDF export did not succeed.")
        return preview
    if suffix == ".key" and input_path.exists():
        preview_png = output_dir / "reference_preview.png"
        try:
            preview_result = generate_preview(input_path, preview_png)
        except Exception as exc:
            preview_result = {"ok": False, "native_error": str(exc), "method": "best_effort_keynote_preview"}
        preview["export_result"] = preview_result
        if preview_result.get("ok") and preview_png.exists():
            preview["success"] = True
            preview["artifacts"]["preview_png"] = str(preview_png.resolve())
        else:
            preview["notes"].append(preview_result.get("native_error") or "Best-effort Keynote preview failed.")
        preview["notes"].append("`.key` audits stay best-effort: require backup, preview review, and rollback discipline.")
        return preview
    preview["notes"].append("No reference preview exporter is available for this presentation format.")
    return preview


def _slot_letter(index: int) -> str:
    index = max(1, int(index))
    letters = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(ord("A") + remainder) + letters
    return letters


def _build_scientific_asset_manifest(deck_path: Path, inspection: dict) -> dict:
    slides = inspection.get("slides") or []
    assets = []
    for slide in slides:
        figure_position = 0
        for shape_position, shape in enumerate(slide.get("shape_inventory", []), start=1):
            if shape["kind"] not in {"picture", "chart"}:
                continue
            figure_position += 1
            slide_position_code = f"{slide['index']:02d}_{_slot_letter(figure_position)}"
            assets.append(
                {
                    "asset_id": f"slide{slide['index']:02d}_asset{shape_position:02d}",
                    "slide_index": slide["index"],
                    "slide_title": slide["title"],
                    "slide_position_code": slide_position_code,
                    "shape_inventory_position": shape_position,
                    "shape_name": shape["name"],
                    "asset_kind": shape["kind"],
                    "bbox_ratio": shape["box"],
                    "final_figure_name": None,
                    "figure_source_path": None,
                    "source_original_path": None,
                    "source_notebook": None,
                    "source_notebook_cell_or_function": None,
                    "source_fits_path": None,
                    "source_caption_or_reference_figure": None,
                    "semantic_intent": None,
                    "object_or_target": None,
                    "astronomical_object": None,
                    "filter_or_channel": None,
                    "campaign_or_year": None,
                    "dataset_or_run": None,
                    "notebook_or_product_origin": None,
                    "visual_mode": None,
                    "extent": None,
                    "crop_or_region": None,
                    "scale_criteria": None,
                    "colormap": None,
                    "vmin_vmax_or_percentiles": None,
                    "origin": None,
                    "overlays": None,
                    "annotations": None,
                    "generation_command": None,
                    "equivalent_target": None,
                    "scientific_role": None,
                    "justification": None,
                    "status": "needs_completion",
                    "editable_native_exception": "scientific_figure",
                }
            )
    return {
        "deck_path": public_path(deck_path),
        "editable_native_first": True,
        "slide_position_policy": "Use slide-position codes such as 04_A or 04_B for final presentation assets.",
        "visual_mode_allowed": ["rgb", "viridis", "grayscale", "false_color", "diagnostic_crop", "plot", "chart", "other"],
        "scientific_role_allowed": ["protagonist", "control", "oral_support", "diagnostic_crop", "reject_or_backup"],
        "required_fields": [
            "figure_source_path",
            "notebook_or_product_origin",
            "semantic_intent",
            "visual_mode",
            "scale_criteria",
            "justification",
        ],
        "coursework_exact_equivalent_fields": [
            "final_figure_name",
            "slide_position_code",
            "source_original_path",
            "source_notebook",
            "source_notebook_cell_or_function",
            "source_fits_path",
            "source_caption_or_reference_figure",
            "astronomical_object",
            "filter_or_channel",
            "visual_mode",
            "extent",
            "crop_or_region",
            "scale_criteria",
            "colormap",
            "vmin_vmax_or_percentiles",
            "origin",
            "overlays",
            "annotations",
            "generation_command",
            "equivalent_target",
        ],
        "assets": assets,
    }


def _asset_manifest_status(asset_manifest: dict) -> dict:
    assets = asset_manifest.get("assets") or []
    required_fields = asset_manifest.get("required_fields") or []
    incomplete = 0
    for asset in assets:
        if any(not asset.get(field) for field in required_fields):
            incomplete += 1
    return {
        "asset_count": len(assets),
        "incomplete_asset_count": incomplete,
        "complete": incomplete == 0 and bool(assets),
    }


def _build_constraints_record(deck_path: Path, suffix: str, user_rules: list[str]) -> dict:
    return {
        "deck_path": public_path(deck_path),
        "working_policy": {
            "work_on_local_copy_only": True,
            "touch_original_downloaded_deck": False,
            "editable_native_first": True,
            "rasterized_full_slide_prohibited": True,
            "keynote_support": "best_effort" if suffix == ".key" else "native_object_first",
        },
        "layer_separation": {
            "slide_content": "Only visible slide content: titles, short bullets, editable captions, and scientific figures.",
            "speaker_script": "Talk track, methodological nuance, oral transitions, and detail that should not crowd the slide.",
            "poster_text": "Longer narrative prose or poster-style explanatory text kept outside the slide layer.",
        },
        "user_rules": list(user_rules or []),
        "suggested_rules_if_relevant": [
            "Do not put notebook screenshots or notebook prose on slides.",
            "Keep scope limited to the requested scientific thread only.",
            "Do not touch the downloaded deck from Drive; work only on a local copy.",
        ],
    }


def _write_json(path: str | Path, payload: dict) -> Path:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(sanitize_payload(payload), indent=2, ensure_ascii=True) + "\n", encoding="utf-8")
    return target


def _payload_status_from_qa(qa: dict) -> str:
    status = qa.get("status")
    if status in {"blocked", "fail"}:
        return status
    return "ok" if status == "ok" else "warning"


def _emit_style_audit_blocked(args, summary: dict, qa: dict) -> int:
    payload = build_tool_payload(
        "presentation_workbench.existing-deck-style-audit",
        status=_payload_status_from_qa(qa),
        notes=[
            "Style audit was blocked before emitting handoff artifacts.",
            "Use a readable copied deck and a valid --reference-slides selector before deriving reusable geometry.",
        ],
        artifacts={"summary_json": args.summary_json},
        results={"inspection": summary, "style_audit": None, "reference_preview": None, "scientific_asset_manifest_status": None},
        qa=qa,
        legacy=summary,
    )
    emit_payload(payload, args.summary_json)
    return 2


def _write_layer_stubs(output_dir: Path, constraints: dict) -> dict[str, Path]:
    slide_content = output_dir / "slide_content.md"
    speaker_script = output_dir / "speaker_script.md"
    poster_text = output_dir / "poster_text.md"
    slide_content.write_text(
        "\n".join(
            [
                "# Slide Content Layer",
                "",
                "- Keep only slide-visible content here.",
                "- Use editable titles, short captions, and scientific figures.",
                "- Do not paste notebook prose, long reductions, or speaker-only detail into the slide layer.",
                "",
                "## Constraints carried into slide editing",
            ]
            + [f"- {rule}" for rule in constraints.get("user_rules", [])]
        )
        + "\n",
        encoding="utf-8",
    )
    speaker_script.write_text(
        "\n".join(
            [
                "# Speaker Script Layer",
                "",
                "- Use this layer for oral transitions, detail, caveats, and explanation that should not crowd the slide.",
                "- If a slide needs more context than fits cleanly, put it here instead of forcing it into the deck.",
                "",
                "## Constraints carried into the oral layer",
            ]
            + [f"- {rule}" for rule in constraints.get("user_rules", [])]
        )
        + "\n",
        encoding="utf-8",
    )
    poster_text.write_text(
        "\n".join(
            [
                "# Poster Text Layer",
                "",
                "- Use this layer for poster-like prose or longer explanatory text that should remain separate from the slides.",
                "- Do not backflow this text into slide captions unless the deck style truly supports it.",
                "",
                "## Constraints carried into the poster layer",
            ]
            + [f"- {rule}" for rule in constraints.get("user_rules", [])]
        )
        + "\n",
        encoding="utf-8",
    )
    return {
        "slide_content": slide_content,
        "speaker_script": speaker_script,
        "poster_text": poster_text,
    }


def _write_style_audit_report(
    path: Path,
    deck_path: Path,
    inspection: dict,
    visual_audit: dict,
    asset_manifest_path: Path,
    constraints_path: Path,
    layer_paths: dict[str, Path],
    preview_summary: dict,
) -> Path:
    lines = [
        "# Existing Deck Style Audit",
        "",
        f"- Deck: `{public_path(deck_path)}`",
        f"- Slide count: `{inspection.get('slide_count')}`",
        f"- Reference slides: `{', '.join(str(value) for value in visual_audit.get('reference_slide_indices', [])) or '[none]'}`",
        f"- Scientific asset manifest: `{public_path(asset_manifest_path)}`",
        f"- Persisted constraints: `{public_path(constraints_path)}`",
        f"- Layer stubs: slide=`{public_path(layer_paths['slide_content'])}`, script=`{public_path(layer_paths['speaker_script'])}`, poster=`{public_path(layer_paths['poster_text'])}`",
        "",
        "## Derived Grid",
        f"- Safe area: `{json.dumps(visual_audit.get('style_grid', {}).get('safe_area', {}), ensure_ascii=True)}`",
        f"- Title box: `{json.dumps(visual_audit.get('style_grid', {}).get('title_box', {}), ensure_ascii=True)}`",
        f"- Single-figure box: `{json.dumps(visual_audit.get('style_grid', {}).get('single_figure_box', {}), ensure_ascii=True)}`",
        f"- Caption box: `{json.dumps(visual_audit.get('style_grid', {}).get('primary_caption_box', {}), ensure_ascii=True)}`",
        "",
        "## Deck Findings",
    ]
    deck_findings = visual_audit.get("deck_findings") or []
    if deck_findings:
        for finding in deck_findings:
            lines.append(f"- `{finding['severity']}` {finding['title']}: {finding['detail']}")
    else:
        lines.append("- No deck-level style blockers detected automatically.")
    lines.extend(["", "## Slide QA"])
    for slide in visual_audit.get("slide_qa", []):
        lines.append(f"### Slide {slide['index']} — {slide['title']}")
        lines.append(f"- Status: `{slide['status']}`")
        if slide["findings"]:
            for finding in slide["findings"]:
                lines.append(f"- `{finding['severity']}` {finding['title']}: {finding['detail']}")
        else:
            lines.append("- No automatic visual issues detected.")
        lines.append(
            f"- Metrics: figures={slide['metrics']['figure_count']}, captions={slide['metrics']['caption_count']}, text_blocks={slide['metrics']['text_blocks']}, words={slide['metrics']['word_count']}"
        )
        lines.append("")
    lines.extend(["## Preview export", f"- Success: `{preview_summary.get('success')}`"])
    for key, value in (preview_summary.get("artifacts") or {}).items():
        lines.append(f"- {key}: `{value}`")
    for note in preview_summary.get("notes") or []:
        lines.append(f"- Note: {note}")
    path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")
    return path


def build_presentation_qa(summary: dict) -> dict:
    inspection = summary.get("inspection") or {}
    findings = []
    if not summary.get("exists"):
        return {
            "status": "blocked",
            "findings": [{"severity": "high", "title": "Input missing", "detail": "Presentation input does not exist."}],
            "metrics": {"blocking_count": 1},
        }
    if summary.get("inspection_error"):
        return {
            "status": "blocked",
            "findings": [
                {
                    "severity": "high",
                    "title": "Unreadable presentation",
                    "detail": f"Could not inspect the presentation: {summary['inspection_error']}",
                }
            ],
            "metrics": {"blocking_count": 1},
        }
    metrics = {
        "slide_count": inspection.get("slide_count"),
        "text_blocks": inspection.get("text_blocks"),
        "image_count": inspection.get("image_count"),
        "dense_slide_count": (inspection.get("style_summary") or {}).get("metrics", {}).get("dense_slide_count"),
        "bundle_member_count": inspection.get("bundle_member_count"),
        "visual_warning_slide_count": sum(1 for slide in (summary.get("visual_audit") or {}).get("slide_qa", []) if slide.get("status") != "ok"),
    }
    status = "ok"
    if summary.get("suffix") in PPTX_EXTS and not inspection.get("slide_count"):
        status = "warning"
        findings.append({"severity": "medium", "title": "Empty deck detection", "detail": "Presentation exists but no slides were detected."})
    style_findings = (inspection.get("style_summary") or {}).get("acceptance_findings") or []
    findings.extend({"severity": "low", "title": "Style heuristic", "detail": item} for item in style_findings)
    if style_findings:
        status = "warning"
    visual_audit = summary.get("visual_audit") or {}
    visual_finding_summary = visual_audit.get("finding_summary")
    if visual_audit and visual_finding_summary is None:
        visual_finding_summary = _summarize_visual_findings(visual_audit)
    findings.extend(visual_audit.get("deck_findings") or [])
    for slide in visual_audit.get("slide_qa", []):
        for finding in slide.get("findings") or []:
            item = dict(finding)
            item["slide_index"] = slide.get("index")
            item["slide_title"] = slide.get("title")
            findings.append(item)
    if visual_audit.get("deck_findings") or any(slide.get("findings") for slide in visual_audit.get("slide_qa", [])):
        status = "warning"
    if summary.get("preview_export", {}).get("success") is False:
        status = "warning"
        findings.append({"severity": "low", "title": "Preview export missing", "detail": "Requested preview PDF export did not succeed."})
    if summary.get("suffix") == ".key":
        status = "warning"
        findings.append(
            {
                "severity": "medium",
                "title": "Best-effort Keynote support",
                "detail": "`.key` handling should be treated as best-effort: keep a backup, export previews, review visually, and be ready to roll back.",
            }
        )
    manifest_status = summary.get("scientific_asset_manifest_status") or {}
    if manifest_status.get("asset_count", 0) and manifest_status.get("incomplete_asset_count", 0):
        status = "warning"
        findings.append(
            {
                "severity": "medium",
                "title": "Scientific asset manifest incomplete",
                "detail": "One or more scientific figures still lack source/origin/intent/visual-mode/scale/justification fields in scientific_asset_manifest.json.",
            }
        )
    severity_counts = {}
    for finding in findings:
        severity = finding.get("severity", "low") if isinstance(finding, dict) else "low"
        severity_counts[severity] = severity_counts.get(severity, 0) + 1
    ordered_findings = sorted(
        findings,
        key=lambda item: (
            SEVERITY_RANK.get(item.get("severity", "low"), 3) if isinstance(item, dict) else 3,
            item.get("slide_index") if isinstance(item, dict) and item.get("slide_index") is not None else -1,
            item.get("title", "") if isinstance(item, dict) else str(item),
        ),
    )
    metrics["finding_count"] = len(findings)
    metrics["finding_counts_by_severity"] = severity_counts
    metrics["findings_truncated"] = len(ordered_findings) > MAX_QA_FINDINGS
    return {
        "status": status,
        "findings": ordered_findings[:MAX_QA_FINDINGS],
        "metrics": metrics,
        "finding_summary": visual_finding_summary,
    }


def write_handoff_notes(
    path: Path,
    copied_input: Path,
    summary: dict,
    exported_pdf: Path | None,
    text_export: Path | None,
    asset_manifest_path: Path,
    constraints_path: Path,
    layer_paths: dict[str, Path],
) -> None:
    lines = [
        "# Presentation Handoff",
        "",
        f"- Source copy: `{public_path(copied_input)}`",
        f"- Format: `{summary['suffix']}`",
        f"- Available backends: `{summary['backends']}`",
        f"- Scientific asset manifest: `{public_path(asset_manifest_path)}`",
        f"- Persisted constraints: `{public_path(constraints_path)}`",
        f"- Slide content layer: `{public_path(layer_paths['slide_content'])}`",
        f"- Speaker script layer: `{public_path(layer_paths['speaker_script'])}`",
        f"- Poster text layer: `{public_path(layer_paths['poster_text'])}`",
    ]
    inspection = summary.get("inspection", {})
    if inspection.get("slide_count") is not None:
        lines.append(f"- Slide count: {inspection['slide_count']}")
    if text_export is not None and text_export.exists():
        lines.append(f"- Extracted slide text: `{public_path(text_export)}`")
    if exported_pdf is not None:
        lines.append(f"- Exported PDF: `{public_path(exported_pdf)}`")
    lines.extend(
        [
            "",
            "## Notes",
            "- The original presentation is untouched; this bundle works on a copied deck or package.",
            "- Editable-native-first applies: slide text, captions, arrows, labels, and shapes should remain editable objects.",
            "- Full-slide rasterization is discouraged unless the element is itself a scientific figure and no native-object path exists.",
            "- `.key` handling is best-effort and should always keep backup, preview review, and rollback discipline.",
        ]
    )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _prepare_style_artifacts(
    input_path: Path,
    output_dir: Path,
    summary: dict,
    user_rules: list[str],
    reference_slide_selector: str | None,
    backend: str,
) -> dict:
    output_dir.mkdir(parents=True, exist_ok=True)
    inspection = summary.get("inspection") or {}
    visual_audit = None
    preview_summary = {"success": False, "artifacts": {}, "notes": []}
    new_slide_templates_path = None
    if summary.get("suffix") in PPTX_EXTS and inspection.get("slides"):
        visual_audit = _build_visual_audit(inspection, reference_slide_selector)
        preview_summary = _export_reference_previews(
            input_path,
            output_dir,
            backend=backend,
            selected_slides=visual_audit["reference_slide_indices"],
        )
        new_slide_templates_path = _write_json(output_dir / "new_slide_templates.json", visual_audit["new_slide_templates"])
    elif summary.get("suffix") == ".key":
        preview_summary = _export_reference_previews(input_path, output_dir, backend=backend, selected_slides=[1])
    asset_manifest = _build_scientific_asset_manifest(input_path, inspection)
    asset_manifest_path = _write_json(output_dir / "scientific_asset_manifest.json", asset_manifest)
    asset_manifest_status = _asset_manifest_status(asset_manifest)
    constraints = _build_constraints_record(input_path, summary.get("suffix"), user_rules)
    constraints_path = _write_json(output_dir / "presentation_constraints.json", constraints)
    layer_paths = _write_layer_stubs(output_dir, constraints)
    report_path = None
    if visual_audit:
        report_path = _write_style_audit_report(
            output_dir / "style_audit_report.md",
            input_path,
            inspection,
            visual_audit,
            asset_manifest_path,
            constraints_path,
            layer_paths,
            preview_summary,
        )
    return {
        "visual_audit": visual_audit,
        "preview_summary": preview_summary,
        "asset_manifest_path": asset_manifest_path,
        "asset_manifest_status": asset_manifest_status,
        "constraints_path": constraints_path,
        "layer_paths": layer_paths,
        "report_path": report_path,
        "new_slide_templates_path": new_slide_templates_path,
    }


def cmd_inspect(args):
    configure_runtime("presentation_workbench")
    input_path = Path(args.input)
    summary = inspect_presentation(input_path)
    summary["dependency_preflight"] = build_dependency_preflight(summary.get("backends") or {}, suffix=summary.get("suffix"))
    if summary.get("suffix") in PPTX_EXTS and (summary.get("inspection") or {}).get("slides"):
        summary["visual_audit"] = _build_visual_audit(summary["inspection"])
    preview_pdf = None
    if args.export_preview_pdf:
        preview_pdf = Path(args.export_preview_pdf)
        if input_path.suffix.lower() in PPTX_EXTS and input_path.is_file():
            try:
                summary["preview_export"] = export_pptx_pdf(input_path, preview_pdf, backend=args.backend)
            except Exception as exc:
                message = clean_known_stderr(f"{exc.__class__.__name__}: {exc}") or exc.__class__.__name__
                summary["preview_export"] = {
                    "success": False,
                    "stderr": f"Preview PDF export failed cleanly: {message}",
                    "error_type": exc.__class__.__name__,
                }
        else:
            summary["preview_export"] = {
                "success": False,
                "stderr": "Preview PDF export from inspect currently supports PPTX/PPTM inputs only.",
            }
    qa = build_presentation_qa(summary)
    payload_status = qa["status"] if qa["status"] in {"blocked", "fail"} else ("ok" if qa["status"] == "ok" else "warning")
    payload = build_tool_payload(
        "presentation_workbench.inspect",
        status=payload_status,
        notes=["Use inspect first when you want a copy-safe read on a deck before building a handoff bundle."],
        artifacts={"summary_json": args.summary_json, "preview_pdf": preview_pdf if preview_pdf and preview_pdf.exists() else None},
        results={"inspection": summary},
        qa=qa,
        legacy=summary,
    )
    emit_payload(payload, args.summary_json)
    if qa["status"] in {"blocked", "fail"}:
        raise SystemExit(2)


def _selector_qa(message: str) -> dict:
    return {
        "status": "blocked",
        "findings": [{"severity": "high", "title": "Invalid reference slide selector", "detail": message}],
        "metrics": {"blocking_count": 1},
    }


def _output_target_qa(message: str) -> dict:
    return {
        "status": "blocked",
        "findings": [{"severity": "high", "title": "Invalid style audit output target", "detail": message}],
        "metrics": {"blocking_count": 1},
    }


def cmd_existing_deck_style_audit(args):
    configure_runtime("presentation_workbench")
    input_path = Path(args.input)
    output_dir = Path(args.output_dir)
    summary = inspect_presentation(input_path)
    summary["dependency_preflight"] = build_dependency_preflight(summary.get("backends") or {}, suffix=summary.get("suffix"))
    preflight_qa = build_presentation_qa(summary)
    if preflight_qa["status"] in {"blocked", "fail"}:
        raise SystemExit(_emit_style_audit_blocked(args, summary, preflight_qa))
    if output_dir.exists() and not output_dir.is_dir():
        qa = _output_target_qa(f"Output target exists and is not a directory: {output_dir}")
        raise SystemExit(_emit_style_audit_blocked(args, summary, qa))
    if args.reference_slides and summary.get("suffix") in PPTX_EXTS and (summary.get("inspection") or {}).get("slides"):
        try:
            parse_slide_selector(args.reference_slides, len(summary["inspection"]["slides"]))
        except ValueError as exc:
            qa = _selector_qa(str(exc))
            raise SystemExit(_emit_style_audit_blocked(args, summary, qa))
    style_outputs = _prepare_style_artifacts(
        input_path,
        output_dir,
        summary,
        user_rules=args.rule,
        reference_slide_selector=args.reference_slides,
        backend=args.backend,
    )
    if style_outputs["visual_audit"] is not None:
        summary["visual_audit"] = style_outputs["visual_audit"]
    summary["reference_preview"] = style_outputs["preview_summary"]
    summary["scientific_asset_manifest_status"] = style_outputs["asset_manifest_status"]
    qa = build_presentation_qa(summary)
    payload = build_tool_payload(
        "presentation_workbench.existing-deck-style-audit",
        status=_payload_status_from_qa(qa),
        notes=[
            "This audit derives deck style from an existing reference deck without editing it.",
            "Use the generated templates, scientific asset manifest, and layer files before touching slides.",
        ],
        artifacts={
            "summary_json": args.summary_json,
            "output_dir": output_dir,
            "style_audit_report": style_outputs["report_path"],
            "scientific_asset_manifest": style_outputs["asset_manifest_path"],
            "presentation_constraints": style_outputs["constraints_path"],
            "slide_content": style_outputs["layer_paths"]["slide_content"],
            "speaker_script": style_outputs["layer_paths"]["speaker_script"],
            "poster_text": style_outputs["layer_paths"]["poster_text"],
            "new_slide_templates": style_outputs["new_slide_templates_path"],
        },
        results={
            "inspection": summary,
            "style_audit": style_outputs["visual_audit"],
            "reference_preview": style_outputs["preview_summary"],
            "scientific_asset_manifest_status": style_outputs["asset_manifest_status"],
        },
        qa=qa,
        legacy=summary,
    )
    if args.summary_json:
        _write_json(args.summary_json, payload)
    print(json.dumps(sanitize_payload(payload), indent=2, ensure_ascii=True))
    if args.manifest_json:
        outputs = [
            style_outputs["asset_manifest_path"],
            style_outputs["constraints_path"],
            style_outputs["layer_paths"]["slide_content"],
            style_outputs["layer_paths"]["speaker_script"],
            style_outputs["layer_paths"]["poster_text"],
        ]
        if style_outputs["report_path"] is not None:
            outputs.append(style_outputs["report_path"])
        if style_outputs["new_slide_templates_path"] is not None:
            outputs.append(style_outputs["new_slide_templates_path"])
        preview_artifacts = style_outputs["preview_summary"].get("artifacts") or {}
        for value in preview_artifacts.values():
            if isinstance(value, list):
                outputs.extend(Path(item) for item in value)
            elif value:
                outputs.append(Path(value))
        if args.summary_json:
            outputs.append(Path(args.summary_json))
        write_manifest(
            args.manifest_json,
            inputs=[input_path],
            outputs=outputs,
            parameters={"backend": args.backend, "reference_slides": args.reference_slides, "rules": args.rule},
            command="presentation_workbench.py existing-deck-style-audit",
            notes=payload["notes"],
        )


def cmd_handoff(args):
    configure_runtime("presentation_workbench")
    source = Path(args.input)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    copied = copy_input(source, output_dir / source.name)
    summary = inspect_presentation(copied)
    summary["dependency_preflight"] = build_dependency_preflight(summary.get("backends") or {}, suffix=summary.get("suffix"))
    bundle_summary_path = output_dir / "bundle_summary.json"
    style_outputs = _prepare_style_artifacts(
        copied,
        output_dir,
        summary,
        user_rules=args.rule,
        reference_slide_selector=None,
        backend=args.backend,
    )
    if style_outputs["visual_audit"] is not None:
        summary["visual_audit"] = style_outputs["visual_audit"]
    summary["reference_preview"] = style_outputs["preview_summary"]
    summary["scientific_asset_manifest_status"] = style_outputs["asset_manifest_status"]
    text_export = None
    if copied.suffix.lower() in PPTX_EXTS and copied.is_file():
        text_export = output_dir / "slide_text.md"
        export_pptx_text(copied, text_export)
    exported_pdf = None
    if args.export_pdf:
        exported_pdf = output_dir / f"{copied.stem}.pdf"
        summary["pdf_export"] = export_pptx_pdf(copied, exported_pdf, backend=args.backend)
    notes_path = output_dir / "handoff_notes.md"
    write_handoff_notes(
        notes_path,
        copied,
        summary,
        exported_pdf if exported_pdf and exported_pdf.exists() else None,
        text_export,
        style_outputs["asset_manifest_path"],
        style_outputs["constraints_path"],
        style_outputs["layer_paths"],
    )
    qa = build_presentation_qa(summary)
    payload = build_tool_payload(
        "presentation_workbench.handoff",
        status="ok" if qa["status"] == "ok" and summary.get("pdf_export", {}).get("success", True) else "warning",
        notes=["Presentation handoff bundles stay copy-based, preserve editability, and keep the original deck untouched."],
        artifacts={
            "bundle_summary_json": bundle_summary_path,
            "handoff_notes": notes_path,
            "copied_input": copied,
            "text_export": text_export if text_export and text_export.exists() else None,
            "exported_pdf": exported_pdf if exported_pdf and exported_pdf.exists() else None,
            "scientific_asset_manifest": style_outputs["asset_manifest_path"],
            "presentation_constraints": style_outputs["constraints_path"],
            "slide_content": style_outputs["layer_paths"]["slide_content"],
            "speaker_script": style_outputs["layer_paths"]["speaker_script"],
            "poster_text": style_outputs["layer_paths"]["poster_text"],
            "new_slide_templates": style_outputs["new_slide_templates_path"],
        },
        results={
            "handoff": summary,
            "style_audit": style_outputs["visual_audit"],
            "reference_preview": style_outputs["preview_summary"],
            "scientific_asset_manifest_status": style_outputs["asset_manifest_status"],
        },
        qa=qa,
        legacy=summary,
    )
    _write_json(bundle_summary_path, payload)
    if args.summary_json:
        _write_json(args.summary_json, payload)
    if args.manifest_json:
        outputs = [
            copied,
            notes_path,
            bundle_summary_path,
            style_outputs["asset_manifest_path"],
            style_outputs["constraints_path"],
            style_outputs["layer_paths"]["slide_content"],
            style_outputs["layer_paths"]["speaker_script"],
            style_outputs["layer_paths"]["poster_text"],
        ]
        if text_export is not None and text_export.exists():
            outputs.append(text_export)
        if exported_pdf is not None and exported_pdf.exists():
            outputs.append(exported_pdf)
        if style_outputs["report_path"] is not None:
            outputs.append(style_outputs["report_path"])
        if style_outputs["new_slide_templates_path"] is not None:
            outputs.append(style_outputs["new_slide_templates_path"])
        preview_artifacts = style_outputs["preview_summary"].get("artifacts") or {}
        for value in preview_artifacts.values():
            if isinstance(value, list):
                outputs.extend(Path(item) for item in value)
            elif value:
                outputs.append(Path(value))
        write_manifest(
            args.manifest_json,
            inputs=[source],
            outputs=outputs,
            parameters={"export_pdf": bool(args.export_pdf), "backend": args.backend, "rules": args.rule},
            command="presentation_workbench.py handoff",
            notes=["Presentation handoff bundles are copy-based and keep the original deck untouched."],
            extra={"summary": summary},
        )
    if args.export_pdf and summary.get("pdf_export", {}).get("success") is False:
        raise SystemExit(summary["pdf_export"].get("stderr") or "Presentation PDF export failed.")
    print(f"Saved presentation handoff bundle: {public_path(output_dir)}")


def main():
    args = parse_args()
    collision_status = _preflight_output_safety(args)
    if collision_status is not None:
        raise SystemExit(collision_status)
    if args.command == "inspect":
        cmd_inspect(args)
    elif args.command == "existing-deck-style-audit":
        cmd_existing_deck_style_audit(args)
    else:
        cmd_handoff(args)


if __name__ == "__main__":
    main()
