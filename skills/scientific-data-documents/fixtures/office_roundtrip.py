#!/usr/bin/env python3
"""Safe copy-based round-trip helpers for common Office formats."""

import argparse
import csv
import json
import subprocess
from collections.abc import Iterable
from pathlib import Path

from _internal.path_safety import find_output_input_collisions, paths_alias
from _internal.provenance_utils import public_path, sanitize_payload, standard_tool_payload, write_manifest
from _internal.runtime_common import detect_presentation_export_backends


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)

    docx_replace = subparsers.add_parser("docx-replace", help="Replace text in a DOCX copy.")
    docx_replace.add_argument("input")
    docx_replace.add_argument("output")
    docx_replace.add_argument("--find", required=True)
    docx_replace.add_argument("--replace", required=True)
    docx_replace.add_argument("--summary-json")
    docx_replace.add_argument("--manifest-json")

    docx_export = subparsers.add_parser("docx-export-text", help="Export DOCX text to TXT or Markdown.")
    docx_export.add_argument("input")
    docx_export.add_argument("output")
    docx_export.add_argument("--manifest-json")

    docx_style = subparsers.add_parser("docx-style-inventory", help="List DOCX runs matching explicit character-style flags.")
    docx_style.add_argument("input")
    docx_style.add_argument("--require-bold", action="store_true")
    docx_style.add_argument("--require-italic", action="store_true")
    docx_style.add_argument("--require-underline", action="store_true")
    docx_style.add_argument("--output-csv")
    docx_style.add_argument("--report-md")
    docx_style.add_argument("--summary-json")
    docx_style.add_argument("--manifest-json")

    docx_style_replace = subparsers.add_parser("docx-styled-replace", help="Replace text only inside DOCX runs matching explicit character-style flags.")
    docx_style_replace.add_argument("input")
    docx_style_replace.add_argument("output")
    docx_style_replace.add_argument("--find", required=True)
    docx_style_replace.add_argument("--replace", required=True)
    docx_style_replace.add_argument("--require-bold", action="store_true")
    docx_style_replace.add_argument("--require-italic", action="store_true")
    docx_style_replace.add_argument("--require-underline", action="store_true")
    docx_style_replace.add_argument(
        "--require-confirmation",
        action="store_true",
        help="Block before writing unless --confirmation-id records an external reviewed approval.",
    )
    docx_style_replace.add_argument("--confirmation-id", help="Opaque identifier for the reviewed external confirmation.")
    docx_style_replace.add_argument("--diff-json")
    docx_style_replace.add_argument("--summary-json")
    docx_style_replace.add_argument("--manifest-json")

    pptx_replace = subparsers.add_parser("pptx-replace", help="Replace text in a PPTX copy.")
    pptx_replace.add_argument("input")
    pptx_replace.add_argument("output")
    pptx_replace.add_argument("--find", required=True)
    pptx_replace.add_argument("--replace", required=True)
    pptx_replace.add_argument("--summary-json")
    pptx_replace.add_argument("--manifest-json")

    pptx_export = subparsers.add_parser("pptx-export-text", help="Export slide text from a PPTX.")
    pptx_export.add_argument("input")
    pptx_export.add_argument("output")
    pptx_export.add_argument("--manifest-json")

    pptx_export_pdf = subparsers.add_parser("pptx-export-pdf", help="Export a PPTX copy to PDF when a supported backend is available.")
    pptx_export_pdf.add_argument("input")
    pptx_export_pdf.add_argument("output")
    pptx_export_pdf.add_argument("--backend", choices=["auto", "soffice", "keynote"], default="auto")
    pptx_export_pdf.add_argument("--summary-json")
    pptx_export_pdf.add_argument("--manifest-json")

    pptx_backends = subparsers.add_parser("pptx-export-backends", help="Report available PPTX to PDF export backends.")
    pptx_backends.add_argument("--summary-json")

    xlsx_set = subparsers.add_parser("xlsx-set-cell", help="Set a cell value in an XLSX copy.")
    xlsx_set.add_argument("input")
    xlsx_set.add_argument("output")
    xlsx_set.add_argument("--sheet", required=True)
    xlsx_set.add_argument("--cell", required=True)
    xlsx_set.add_argument("--value", required=True)
    xlsx_set.add_argument("--summary-json")
    xlsx_set.add_argument("--manifest-json")

    xlsx_export = subparsers.add_parser("xlsx-export-csv", help="Export workbook sheets to CSV.")
    xlsx_export.add_argument("input")
    xlsx_export.add_argument("--output-dir", required=True)
    xlsx_export.add_argument("--manifest-json")

    return parser.parse_args()


def parse_scalar(value):
    lowered = value.strip().lower()
    if lowered in {"true", "false"}:
        return lowered == "true"
    try:
        if "." in value:
            return float(value)
        return int(value)
    except Exception:
        return value


def replace_in_paragraph(paragraph, find_text, replace_text):
    original_text = paragraph.text
    total = original_text.count(find_text)
    if total == 0:
        return 0
    replaced = 0
    for run in paragraph.runs:
        count = run.text.count(find_text)
        if count:
            run.text = run.text.replace(find_text, replace_text)
            replaced += count
    if replaced == total:
        return replaced
    element = paragraph._element
    for child in list(element):
        element.remove(child)
    paragraph.add_run(original_text.replace(find_text, replace_text))
    return total


def docx_replace_text(input_path, output_path, find_text, replace_text):
    input_path = Path(input_path)
    output_path = Path(output_path)
    if paths_refer_to_same_file(input_path, output_path):
        raise ValueError("Refusing to overwrite the input DOCX; choose a distinct output path.")
    from docx import Document

    doc = Document(input_path)
    replacements = 0
    for paragraph in doc.paragraphs:
        replacements += replace_in_paragraph(paragraph, find_text, replace_text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replacements += replace_in_paragraph(paragraph, find_text, replace_text)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return {"replacements": replacements}


def export_docx_text(input_path, output_path):
    from docx import Document

    doc = Document(input_path)
    pieces = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            pieces.append(text)
    for table in doc.tables:
        pieces.append("")
        for row in table.rows:
            pieces.append(" | ".join(cell.text.strip() for cell in row.cells))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prefix = "# Exported DOCX Text\n\n" if output_path.suffix.lower() == ".md" else ""
    output_path.write_text(prefix + "\n".join(pieces).strip() + "\n")


def style_requirements(args):
    requirements = {
        "bold": bool(getattr(args, "require_bold", False)),
        "italic": bool(getattr(args, "require_italic", False)),
        "underline": bool(getattr(args, "require_underline", False)),
    }
    if not any(requirements.values()):
        requirements["any_special_style"] = True
    return requirements


def run_style_flags(run):
    underline = run.underline
    return {
        "bold": run.bold is True,
        "italic": run.italic is True,
        "underline": underline is not None and underline is not False,
    }


def style_matches(flags, requirements):
    if requirements.get("any_special_style"):
        return any(flags.values())
    return all(flags[key] for key, required in requirements.items() if required)


def summarize_style_counts(rows):
    return {
        "bold": sum(1 for row in rows if row["bold"]),
        "italic": sum(1 for row in rows if row["italic"]),
        "underline": sum(1 for row in rows if row["underline"]),
        "body": sum(1 for row in rows if row["location"] == "body"),
        "table": sum(1 for row in rows if row["location"] == "table"),
    }


def iter_docx_runs(doc) -> Iterable[dict]:
    for paragraph_index, paragraph in enumerate(doc.paragraphs, start=1):
        for run_index, run in enumerate(paragraph.runs, start=1):
            yield {
                "location": "body",
                "paragraph_index": paragraph_index,
                "table_index": None,
                "row_index": None,
                "cell_index": None,
                "run_index": run_index,
                "paragraph_style": getattr(paragraph.style, "name", None),
                "run": run,
            }
    for table_index, table in enumerate(doc.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            for cell_index, cell in enumerate(row.cells, start=1):
                for paragraph_index, paragraph in enumerate(cell.paragraphs, start=1):
                    for run_index, run in enumerate(paragraph.runs, start=1):
                        yield {
                            "location": "table",
                            "paragraph_index": paragraph_index,
                            "table_index": table_index,
                            "row_index": row_index,
                            "cell_index": cell_index,
                            "run_index": run_index,
                            "paragraph_style": getattr(paragraph.style, "name", None),
                            "run": run,
                        }


def docx_style_inventory(input_path, requirements):
    from docx import Document

    doc = Document(input_path)
    rows = []
    for item in iter_docx_runs(doc):
        run = item["run"]
        text = run.text
        if not text.strip():
            continue
        flags = run_style_flags(run)
        if not style_matches(flags, requirements):
            continue
        rows.append(
            {
                "location": item["location"],
                "paragraph_index": item["paragraph_index"],
                "table_index": item["table_index"],
                "row_index": item["row_index"],
                "cell_index": item["cell_index"],
                "run_index": item["run_index"],
                "text": text,
                "bold": flags["bold"],
                "italic": flags["italic"],
                "underline": flags["underline"],
                "paragraph_style": item["paragraph_style"],
                "character_style": getattr(run.style, "name", None),
            }
        )
    return {
        "input": public_path(input_path),
        "requirements": requirements,
        "match_count": len(rows),
        "style_counts": summarize_style_counts(rows),
        "matches": rows,
        "limitations": [
            "This detects explicit DOCX run-level formatting. Styles inherited from a paragraph or theme may require visual readback.",
            "Use this on a copy or on a DOCX exported from Pages, not on the original .pages package.",
        ],
    }


def write_style_csv(path, rows):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "location",
        "paragraph_index",
        "table_index",
        "row_index",
        "cell_index",
        "run_index",
        "text",
        "bold",
        "italic",
        "underline",
        "paragraph_style",
        "character_style",
    ]
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def write_style_report(path, summary):
    lines = [
        "# DOCX Styled Run Inventory",
        "",
        f"- Input: `{summary['input']}`",
        f"- Matches: `{summary['match_count']}`",
        f"- Requirements: `{json.dumps(summary['requirements'], ensure_ascii=True)}`",
        "",
        "## Style Counts",
        "",
        f"- Bold: `{summary['style_counts']['bold']}`",
        f"- Italic: `{summary['style_counts']['italic']}`",
        f"- Underline: `{summary['style_counts']['underline']}`",
        f"- Body: `{summary['style_counts']['body']}`",
        f"- Table: `{summary['style_counts']['table']}`",
        "",
        "## Matches",
    ]
    if not summary["matches"]:
        lines.append("- No matching runs found.")
    for row in summary["matches"][:80]:
        flags = ",".join(flag for flag in ("bold", "italic", "underline") if row[flag])
        location = row["location"]
        if location == "table":
            where = f"table {row['table_index']}, row {row['row_index']}, cell {row['cell_index']}, paragraph {row['paragraph_index']}, run {row['run_index']}"
        else:
            where = f"paragraph {row['paragraph_index']}, run {row['run_index']}"
        lines.append(f"- `{row['text']}` ({where}; {flags or 'unstyled'})")
    lines.extend(["", "## Limitations"])
    for item in summary["limitations"]:
        lines.append(f"- {item}")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def apply_document_v2_contract(payload, *, requires_confirmation):
    payload["contract_version"] = "2.0"
    payload["job_metadata"] = {
        "safe_to_retry": True,
        "supports_cancel": False,
        "recommended_exposure": "normal",
    }
    payload["safety"] = {
        "input_policy": "copied_input_only",
        "output_policy": "separate_run_directory",
        "secrets_redacted": True,
        "requires_confirmation": bool(requires_confirmation),
    }
    payload["provenance"]["skill_contract_version"] = "v2.0"
    payload["provenance"]["originals_policy"] = "copied_inputs_only"
    return payload


def build_style_inventory_payload(summary, *, artifacts=None, status=None, findings=None):
    findings = list(findings or [])
    status = status or ("warning" if findings else "ok")
    payload = standard_tool_payload(
        "office_roundtrip.docx-style-inventory",
        status=status,
        notes=[
            "Lists explicit DOCX run-level formatting for copy-based Pages/DOCX marked-fragment workflows.",
            "Use on a DOCX copy or on a DOCX exported from Pages; verify visually before carrying edits back to the source document.",
        ],
        artifacts=artifacts or {},
        results={"inventory": summary},
        qa={
            "status": "warning" if findings else "ok",
            "findings": findings,
            "metrics": {
                "match_count": summary.get("match_count", 0),
                "style_counts": summary.get("style_counts", {}),
            },
        },
        inputs=[summary.get("input")],
        legacy=summary,
    )
    return apply_document_v2_contract(payload, requires_confirmation=False)


def build_style_inventory_failure(input_path, requirements, error, *, artifacts=None, extra_findings=None):
    findings = list(extra_findings or [])
    findings.append(f"Could not complete DOCX style inventory: {type(error).__name__}: {error}")
    legacy = {
        "input": public_path(input_path),
        "requirements": requirements,
        "match_count": 0,
        "style_counts": {"bold": 0, "italic": 0, "underline": 0, "body": 0, "table": 0},
        "matches": [],
        "limitations": [
            "The DOCX could not be opened. Work on a fresh export/copy and verify the file is a valid .docx package.",
        ],
        "error": str(error),
    }
    payload = standard_tool_payload(
        "office_roundtrip.docx-style-inventory",
        status="fail",
        notes=[
            "DOCX style inventory failed before producing trustworthy matches.",
        ],
        artifacts=artifacts or {},
        results={"inventory": legacy},
        qa={
            "status": "fail",
            "findings": findings,
            "metrics": {"match_count": 0},
        },
        inputs=[input_path],
        legacy=legacy,
    )
    return apply_document_v2_contract(payload, requires_confirmation=False)


def paths_refer_to_same_file(input_path, output_path):
    input_path = Path(input_path)
    output_path = Path(output_path)
    try:
        if input_path.exists() and output_path.exists():
            return input_path.samefile(output_path)
    except OSError:
        pass
    try:
        return input_path.resolve(strict=False) == output_path.resolve(strict=False)
    except OSError:
        return str(input_path.absolute()) == str(output_path.absolute())


def office_output_collisions(command, args):
    if not hasattr(args, "input"):
        return []
    input_path = Path(args.input)
    requested = []
    for attribute, flag in (
        ("output", "output"),
        ("output_dir", "--output-dir"),
        ("output_csv", "--output-csv"),
        ("report_md", "--report-md"),
        ("diff_json", "--diff-json"),
        ("summary_json", "--summary-json"),
        ("manifest_json", "--manifest-json"),
    ):
        value = getattr(args, attribute, None)
        if value:
            requested.append((flag, value))
    output_dir = getattr(args, "output_dir", None)
    if output_dir and Path(output_dir).expanduser().is_dir():
        requested.extend(
            ("existing-output-member", item)
            for item in Path(output_dir).expanduser().rglob("*")
            if item.is_file()
        )
    collisions = [
        {
            "flag": item["flag"],
            "protected_path": "input",
            "path": public_path(input_path),
        }
        for item in find_output_input_collisions([("input", input_path)], requested)
        if item.get("collision_kind") != "output_output_alias"
    ]
    output_path = getattr(args, "output", None)
    if output_path:
        for flag, pathlike in requested:
            if flag == "output":
                continue
            if paths_alias(pathlike, output_path):
                collisions.append({"flag": flag, "protected_path": "output", "path": public_path(output_path)})
    present = [(flag, pathlike) for flag, pathlike in requested if flag not in {"output", "--output-dir"}]
    for index, (left_flag, left_path) in enumerate(present):
        for right_flag, right_path in present[index + 1 :]:
            if paths_alias(left_path, right_path):
                collisions.append(
                    {
                        "flag": f"{left_flag},{right_flag}",
                        "protected_path": "auxiliary_output",
                        "path": public_path(left_path),
                    }
                )
    return collisions


def emit_office_collision_only(command, args, collisions):
    rendered = ", ".join(
        f"{item['flag']}->{item['protected_path']}" for item in collisions
    )
    message = f"Refusing Office output paths that overlap protected inputs or outputs: {rendered}."
    payload = standard_tool_payload(
        f"office_roundtrip.{command}",
        status="blocked",
        notes=[message, "No output files were written."],
        artifacts={},
        results={
            "blocked_reason": message,
            "error_type": "output_input_collision",
            "collisions": collisions,
        },
        qa={
            "status": "blocked",
            "findings": [message],
            "metrics": {"blocking_count": len(collisions)},
        },
        inputs=[args.input],
        legacy={"blocked_reason": message},
    )
    apply_document_v2_contract(
        payload,
        requires_confirmation=command == "docx-styled-replace",
    )
    print(json.dumps(sanitize_payload(payload), indent=2, ensure_ascii=True))
    return 2


def build_styled_replace_payload(summary, *, artifacts=None, status=None, findings=None):
    findings = list(findings or [])
    status = status or ("warning" if findings else "ok")
    payload = standard_tool_payload(
        "office_roundtrip.docx-styled-replace",
        status=status,
        notes=[
            "Replaces text only inside DOCX runs matching explicit style requirements.",
            "The command writes a separate edited DOCX copy and verifies the replacement through a styled-run readback.",
        ],
        artifacts=artifacts or {},
        results={"replacement": summary},
        qa={
            "status": status if status in {"ok", "warning", "fail"} else "warning",
            "findings": findings,
            "metrics": {
                "replacements": summary.get("replacements", 0),
                "touched_run_count": len(summary.get("touched_runs", [])),
                "readback_replacement_match_count": summary.get("readback_replacement_match_count", 0),
                "readback_replacement_occurrence_count": summary.get("readback_replacement_occurrence_count", 0),
            },
        },
        inputs=[summary.get("input")],
        legacy=summary,
    )
    return apply_document_v2_contract(payload, requires_confirmation=True)


def classify_styled_replace_summary(summary):
    findings = []
    status = "ok"
    replacements = int(summary.get("replacements") or 0)
    readback_occurrences = int(summary.get("readback_replacement_occurrence_count") or 0)
    if summary.get("find") == summary.get("replace"):
        findings.append("Find and replace text are identical; the output copy is not a meaningful edit.")
        status = "warning"
    if replacements == 0:
        findings.append(
            "No matching styled runs were edited. Check the exact text, split DOCX runs, and explicit style requirements."
        )
        status = "warning"
    elif readback_occurrences < replacements:
        findings.append(
            "Styled-run readback found fewer replacement occurrences than the number of replacements reported."
        )
        status = "fail"
    return status, findings


def build_styled_replace_failure(input_path, output_path, find_text, replace_text, requirements, error, *, artifacts=None, extra_findings=None):
    findings = list(extra_findings or [])
    findings.append(f"Could not complete DOCX styled replacement: {type(error).__name__}: {error}")
    legacy = {
        "input": public_path(input_path),
        "output": public_path(output_path),
        "backup_existing_output": None,
        "requirements": requirements,
        "find": find_text,
        "replace": replace_text,
        "replacements": 0,
        "touched_runs": [],
        "readback_replacement_match_count": 0,
        "readback_replacement_occurrence_count": 0,
        "limitations": [
            "The replacement was not completed. Work on a valid DOCX copy and keep the output path separate from the input.",
        ],
        "error": str(error),
    }
    payload = standard_tool_payload(
        "office_roundtrip.docx-styled-replace",
        status="fail",
        notes=[
            "DOCX styled replacement failed before producing a trustworthy edited copy.",
        ],
        artifacts=artifacts or {},
        results={"replacement": legacy},
        qa={
            "status": "fail",
            "findings": findings,
            "metrics": {
                "replacements": 0,
                "touched_run_count": 0,
                "readback_replacement_match_count": 0,
                "readback_replacement_occurrence_count": 0,
            },
        },
        inputs=[input_path],
        legacy=legacy,
    )
    return apply_document_v2_contract(payload, requires_confirmation=True)


def backup_existing_output(output_path):
    output_path = Path(output_path)
    if not output_path.exists():
        return None
    for index in range(1, 1000):
        candidate = output_path.with_name(f"{output_path.name}.bak{index}")
        if not candidate.exists():
            candidate.write_bytes(output_path.read_bytes())
            return candidate
    raise RuntimeError(f"Could not create backup for existing output: {output_path}")


def confirmation_record(*, enforced, confirmation_id):
    return {
        "required": True,
        "enforced_by_backend": bool(enforced),
        "recorded": bool(confirmation_id),
        "confirmation_id": confirmation_id or None,
    }


def docx_styled_replace(
    input_path,
    output_path,
    find_text,
    replace_text,
    requirements,
    *,
    confirmation_id=None,
    enforce_confirmation=False,
):
    from docx import Document

    input_path = Path(input_path)
    output_path = Path(output_path)
    if enforce_confirmation and not confirmation_id:
        raise ValueError("--confirmation-id is required when --require-confirmation is set.")
    if paths_refer_to_same_file(input_path, output_path):
        raise ValueError("Output path must be different from input path for copy-based DOCX edits.")
    doc = Document(input_path)
    replacements = 0
    touched_runs = []
    for item in iter_docx_runs(doc):
        run = item["run"]
        if not run.text or find_text not in run.text:
            continue
        flags = run_style_flags(run)
        if not style_matches(flags, requirements):
            continue
        count = run.text.count(find_text)
        before_text = run.text
        run.text = run.text.replace(find_text, replace_text)
        after_text = run.text
        replacements += count
        touched_runs.append(
            {
                "location": item["location"],
                "paragraph_index": item["paragraph_index"],
                "table_index": item["table_index"],
                "row_index": item["row_index"],
                "cell_index": item["cell_index"],
                "run_index": item["run_index"],
                "replacements": count,
                "bold": flags["bold"],
                "italic": flags["italic"],
                "underline": flags["underline"],
                "before_text": before_text,
                "after_text": after_text,
            }
        )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    backup = backup_existing_output(output_path)
    doc.save(output_path)
    readback = docx_style_inventory(output_path, requirements)
    readback_matches = [row for row in readback["matches"] if replace_text in row["text"]]
    readback_occurrences = sum(row["text"].count(replace_text) for row in readback_matches) if replace_text else 0
    return {
        "input": public_path(input_path),
        "output": public_path(output_path),
        "backup_existing_output": public_path(backup) if backup else None,
        "requirements": requirements,
        "find": find_text,
        "replace": replace_text,
        "confirmation": confirmation_record(
            enforced=enforce_confirmation,
            confirmation_id=confirmation_id,
        ),
        "replacements": replacements,
        "touched_runs": touched_runs,
        "readback_replacement_match_count": len(readback_matches),
        "readback_replacement_occurrence_count": readback_occurrences,
        "limitations": [
            "Replacement is deliberately run-local to preserve formatting; text split across multiple runs is reported as zero replacements.",
            "For .pages sources, verify the DOCX export visually before importing or copying edits back into Pages.",
        ],
    }


def write_styled_diff(path, summary):
    diff = {
        "contract_version": "2.0",
        "artifact_type": "app_preview",
        "input": summary.get("input"),
        "output": summary.get("output"),
        "find": summary.get("find"),
        "replace": summary.get("replace"),
        "replacement_count": summary.get("replacements", 0),
        "review_confirmed": bool((summary.get("confirmation") or {}).get("recorded")),
        "confirmation_id": (summary.get("confirmation") or {}).get("confirmation_id"),
        "changes": [
            {
                key: row.get(key)
                for key in (
                    "location",
                    "paragraph_index",
                    "table_index",
                    "row_index",
                    "cell_index",
                    "run_index",
                    "before_text",
                    "after_text",
                    "replacements",
                    "bold",
                    "italic",
                    "underline",
                )
            }
            for row in summary.get("touched_runs", [])
        ],
        "original_modified": False,
        "review_required": True,
    }
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(sanitize_payload(diff), indent=2, ensure_ascii=True) + "\n", encoding="utf-8")


def replace_in_shape(shape, find_text, replace_text):
    replacements = 0
    if getattr(shape, "has_text_frame", False):
        for paragraph in shape.text_frame.paragraphs:
            count = paragraph.text.count(find_text)
            if count:
                paragraph.text = paragraph.text.replace(find_text, replace_text)
                replacements += count
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            for cell in row.cells:
                for paragraph in cell.text_frame.paragraphs:
                    count = paragraph.text.count(find_text)
                    if count:
                        paragraph.text = paragraph.text.replace(find_text, replace_text)
                        replacements += count
    return replacements


def pptx_replace_text(input_path, output_path, find_text, replace_text):
    input_path = Path(input_path)
    output_path = Path(output_path)
    if paths_refer_to_same_file(input_path, output_path):
        raise ValueError("Refusing to overwrite the input PPTX; choose a distinct output path.")
    from pptx import Presentation

    presentation = Presentation(input_path)
    replacements = 0
    for slide in presentation.slides:
        for shape in slide.shapes:
            replacements += replace_in_shape(shape, find_text, replace_text)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output_path)
    return {"replacements": replacements, "slide_count": len(presentation.slides)}


def export_pptx_text(input_path, output_path):
    from pptx import Presentation

    presentation = Presentation(input_path)
    lines = []
    for index, slide in enumerate(presentation.slides, start=1):
        lines.append(f"## Slide {index}")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = "\n".join(paragraph.text.strip() for paragraph in shape.text_frame.paragraphs if paragraph.text.strip())
                if text:
                    lines.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        lines.append("")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prefix = "# Exported PPTX Text\n\n" if output_path.suffix.lower() == ".md" else ""
    output_path.write_text(prefix + "\n".join(lines).strip() + "\n")


def export_pptx_pdf(input_path, output_path, backend="auto"):
    backends = detect_presentation_export_backends()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    selected = backend
    if backend == "auto":
        if backends.get("soffice"):
            selected = "soffice"
        elif backends.get("keynote_available"):
            selected = "keynote"
        else:
            selected = "none"
    result = {
        "backend_requested": backend,
        "backend_used": selected,
        "backends": backends,
        "success": False,
        "stdout": "",
        "stderr": "",
    }
    if selected == "soffice":
        if not backends.get("soffice"):
            result["stderr"] = "LibreOffice/soffice PDF export backend was requested but is not available."
            return result
        cmd = [
            backends["soffice"],
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output_path.parent),
            str(input_path),
        ]
        completed = subprocess.run(cmd, text=True, capture_output=True, check=False)
        result["stdout"] = completed.stdout
        result["stderr"] = completed.stderr
        result["success"] = completed.returncode == 0 and output_path.exists()
        return result
    if selected == "keynote":
        if not backends.get("keynote_available"):
            result["stderr"] = "Keynote PDF export backend was requested but is not available."
            return result
        script = [
            'tell application "Keynote"',
            "activate",
            f'set theDoc to open POSIX file "{Path(input_path).resolve()}"',
            f'export theDoc to POSIX file "{Path(output_path).resolve()}" as PDF',
            "close theDoc saving no",
            "end tell",
        ]
        cmd = ["osascript"]
        for line in script:
            cmd.extend(["-e", line])
        completed = subprocess.run(cmd, text=True, capture_output=True, check=False)
        result["stdout"] = completed.stdout
        result["stderr"] = completed.stderr
        result["success"] = completed.returncode == 0 and output_path.exists()
        return result
    result["stderr"] = "No supported PPTX export backend is available."
    return result


def xlsx_set_cell(input_path, output_path, sheet_name, cell_ref, value):
    input_path = Path(input_path)
    output_path = Path(output_path)
    if paths_refer_to_same_file(input_path, output_path):
        raise ValueError("Refusing to overwrite the input XLSX; choose a distinct output path.")
    from openpyxl import load_workbook

    workbook = load_workbook(input_path)
    worksheet = workbook[sheet_name]
    worksheet[cell_ref] = parse_scalar(value)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return {"sheet": sheet_name, "cell": cell_ref, "value": parse_scalar(value)}


def export_xlsx_csv(input_path, output_dir):
    from openpyxl import load_workbook

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(input_path, data_only=True)
    outputs = []
    for worksheet in workbook.worksheets:
        target = output_dir / f"{worksheet.title}.csv"
        with target.open("w", newline="") as handle:
            writer = csv.writer(handle)
            for row in worksheet.iter_rows(values_only=True):
                writer.writerow(list(row))
        outputs.append(public_path(target))
    return outputs


def emit_summary(path, payload):
    if not path:
        return None
    summary_path = Path(path)
    summary_path.parent.mkdir(parents=True, exist_ok=True)
    if summary_path.exists() and summary_path.is_dir():
        raise RuntimeError(f"--summary-json points to a directory, not a file: {public_path(summary_path)}")
    rendered = json.dumps(sanitize_payload(payload), indent=2, ensure_ascii=True) + "\n"
    summary_path.write_text(rendered, encoding="utf-8")
    return summary_path


def main():
    args = parse_args()
    command = args.command

    output_collisions = office_output_collisions(command, args)
    if output_collisions:
        raise SystemExit(emit_office_collision_only(command, args, output_collisions))

    mutating_commands = {
        "docx-replace": "DOCX",
        "pptx-replace": "PPTX",
        "xlsx-set-cell": "XLSX",
    }
    if command in mutating_commands and paths_refer_to_same_file(args.input, args.output):
        message = (
            f"Refusing to overwrite the input {mutating_commands[command]}; "
            "choose a distinct output path."
        )
        payload = standard_tool_payload(
            f"office_roundtrip.{command}",
            status="blocked",
            notes=[message],
            artifacts={},
            results={"blocked_reason": message},
            qa={"status": "blocked", "findings": [message], "metrics": {"blocking_count": 1}},
            inputs=[args.input],
            legacy={"blocked_reason": message},
        )
        print(json.dumps(sanitize_payload(payload), indent=2, ensure_ascii=True))
        raise SystemExit(2)

    if command == "docx-replace":
        output = Path(args.output)
        summary = docx_replace_text(args.input, output, args.find, args.replace)
        if args.summary_json:
            emit_summary(args.summary_json, summary)
        if args.manifest_json:
            write_manifest(args.manifest_json, inputs=[args.input], outputs=[output], parameters={"find": args.find, "replace": args.replace})
        print(f"Saved DOCX copy: {public_path(output)}")
    elif command == "docx-export-text":
        output = Path(args.output)
        export_docx_text(args.input, output)
        if args.manifest_json:
            write_manifest(args.manifest_json, inputs=[args.input], outputs=[output], parameters={"command": command})
        print(f"Saved DOCX text export: {public_path(output)}")
    elif command == "docx-style-inventory":
        requirements = style_requirements(args)
        artifacts = {
            "output_csv": args.output_csv,
            "report_md": args.report_md,
            "summary_json": args.summary_json,
            "manifest_json": args.manifest_json,
        }
        try:
            summary = docx_style_inventory(Path(args.input), requirements)
            findings = []
            if summary["match_count"] == 0:
                findings.append(
                    "No matching styled runs found. For Pages-exported marked fragments, verify the export retained explicit run-level styles or relax the style requirements."
                )
            payload = build_style_inventory_payload(
                summary,
                artifacts=artifacts,
                status="warning" if findings else "ok",
                findings=findings,
            )
            outputs = []
            if args.output_csv:
                write_style_csv(args.output_csv, summary["matches"])
                outputs.append(Path(args.output_csv))
            if args.report_md:
                write_style_report(args.report_md, summary)
                outputs.append(Path(args.report_md))
            if args.summary_json:
                emit_summary(args.summary_json, payload)
                outputs.append(Path(args.summary_json))
            if args.manifest_json:
                write_manifest(
                    args.manifest_json,
                    inputs=[args.input],
                    outputs=outputs,
                    parameters={"requirements": requirements},
                    command="office_roundtrip.py docx-style-inventory",
                )
            print(json.dumps(sanitize_payload(payload), indent=2, ensure_ascii=True))
        except Exception as error:
            payload = build_style_inventory_failure(Path(args.input), requirements, error, artifacts=artifacts)
            if args.summary_json:
                try:
                    emit_summary(args.summary_json, payload)
                except Exception as summary_error:
                    payload = build_style_inventory_failure(
                        Path(args.input),
                        requirements,
                        error,
                        artifacts=artifacts,
                        extra_findings=[f"Could not write --summary-json: {type(summary_error).__name__}: {summary_error}"],
                    )
            print(json.dumps(sanitize_payload(payload), indent=2, ensure_ascii=True))
            raise SystemExit(2)
    elif command == "docx-styled-replace":
        output = Path(args.output)
        requirements = style_requirements(args)
        artifacts = {
            "output_docx": output,
            "diff_json": args.diff_json,
            "summary_json": args.summary_json,
            "manifest_json": args.manifest_json,
        }
        try:
            summary = docx_styled_replace(
                Path(args.input),
                output,
                args.find,
                args.replace,
                requirements,
                confirmation_id=args.confirmation_id,
                enforce_confirmation=args.require_confirmation,
            )
            if summary.get("backup_existing_output"):
                artifacts["backup_existing_output"] = summary["backup_existing_output"]
            status, findings = classify_styled_replace_summary(summary)
            payload = build_styled_replace_payload(summary, artifacts=artifacts, status=status, findings=findings)
            if args.diff_json:
                write_styled_diff(args.diff_json, summary)
            if args.summary_json:
                emit_summary(args.summary_json, payload)
            if args.manifest_json:
                outputs = [output] if output.exists() else []
                if args.diff_json and Path(args.diff_json).exists():
                    outputs.append(args.diff_json)
                if summary.get("backup_existing_output"):
                    outputs.append(summary["backup_existing_output"])
                write_manifest(
                    args.manifest_json,
                    inputs=[args.input],
                    outputs=outputs,
                    parameters={
                        "find": args.find,
                        "replace": args.replace,
                        "requirements": summary["requirements"],
                        "confirmation": summary["confirmation"],
                    },
                )
            print(json.dumps(sanitize_payload(payload), indent=2, ensure_ascii=True))
            if status == "fail":
                raise SystemExit(2)
        except Exception as error:
            payload = build_styled_replace_failure(Path(args.input), output, args.find, args.replace, requirements, error, artifacts=artifacts)
            if args.summary_json:
                try:
                    emit_summary(args.summary_json, payload)
                except Exception as summary_error:
                    payload = build_styled_replace_failure(
                        Path(args.input),
                        output,
                        args.find,
                        args.replace,
                        requirements,
                        error,
                        artifacts=artifacts,
                        extra_findings=[f"Could not write --summary-json: {type(summary_error).__name__}: {summary_error}"],
                    )
            print(json.dumps(sanitize_payload(payload), indent=2, ensure_ascii=True))
            raise SystemExit(2)
    elif command == "pptx-replace":
        output = Path(args.output)
        summary = pptx_replace_text(args.input, output, args.find, args.replace)
        if args.summary_json:
            emit_summary(args.summary_json, summary)
        if args.manifest_json:
            write_manifest(args.manifest_json, inputs=[args.input], outputs=[output], parameters={"find": args.find, "replace": args.replace})
        print(f"Saved PPTX copy: {public_path(output)}")
    elif command == "pptx-export-text":
        output = Path(args.output)
        export_pptx_text(args.input, output)
        if args.manifest_json:
            write_manifest(args.manifest_json, inputs=[args.input], outputs=[output], parameters={"command": command})
        print(f"Saved PPTX text export: {public_path(output)}")
    elif command == "pptx-export-pdf":
        output = Path(args.output)
        summary = export_pptx_pdf(args.input, output, backend=args.backend)
        if args.summary_json:
            emit_summary(args.summary_json, summary)
        if args.manifest_json:
            write_manifest(args.manifest_json, inputs=[args.input], outputs=[output] if output.exists() else [], parameters=summary)
        if not summary["success"]:
            raise SystemExit(summary["stderr"] or "PPTX to PDF export failed.")
        print(f"Saved PPTX PDF export: {public_path(output)}")
    elif command == "pptx-export-backends":
        summary = detect_presentation_export_backends()
        if args.summary_json:
            emit_summary(args.summary_json, summary)
        print(json.dumps(summary, indent=2, ensure_ascii=True))
    elif command == "xlsx-set-cell":
        output = Path(args.output)
        summary = xlsx_set_cell(args.input, output, args.sheet, args.cell, args.value)
        if args.summary_json:
            emit_summary(args.summary_json, summary)
        if args.manifest_json:
            write_manifest(args.manifest_json, inputs=[args.input], outputs=[output], parameters=summary)
        print(f"Saved XLSX copy: {public_path(output)}")
    elif command == "xlsx-export-csv":
        outputs = export_xlsx_csv(args.input, args.output_dir)
        if args.manifest_json:
            write_manifest(args.manifest_json, inputs=[args.input], outputs=outputs, parameters={"command": command})
        print(json.dumps({"exports": outputs}, indent=2, ensure_ascii=True))


if __name__ == "__main__":
    main()
