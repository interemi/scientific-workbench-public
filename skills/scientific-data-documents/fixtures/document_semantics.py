#!/usr/bin/env python3
"""Extract semantic summaries from documents, Office files, and iWork packages."""

import argparse
import html
import json
import os
import plistlib
import re
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET

from _internal.iwork_iwa import inspect_iwa_member
from _internal.ocr_utils import ocr_image, ocr_pdf
from _internal.path_safety import find_output_input_collisions
from _internal.public_contract import build_tool_payload
from _internal.provenance_utils import public_path, sanitize_payload


TEXTUAL_EXTS = {
    ".txt",
    ".md",
    ".tex",
    ".bib",
    ".json",
    ".jsonl",
    ".ndjson",
    ".yaml",
    ".yml",
    ".toml",
    ".xml",
    ".html",
    ".htm",
    ".reg",
    ".pref",
    ".jmars",
    ".ini",
    ".cfg",
    ".log",
    ".rtf",
}


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("path", help="Input document or package path.")
    parser.add_argument("--max-chars", type=int, default=2500, help="Maximum characters for extracted text excerpts.")
    parser.add_argument("--output-json", help="Optional output JSON path.")
    parser.add_argument("--output-dir", help="Optional directory for extracted preview assets or helper files.")
    parser.add_argument("--enable-ocr", action="store_true", help="Try OCR on iWork preview images when available.")
    return parser.parse_args()


def _emit_collision_only(path: Path, collisions: list[dict[str, str]]) -> int:
    message = "Refusing document-semantic outputs that overlap the input document or package."
    payload = build_tool_payload(
        "document_semantics",
        status="blocked",
        notes=[message, "No extracted asset or JSON summary was written."],
        artifacts={},
        results={
            "blocked_reason": message,
            "error_type": "output_input_collision",
            "input": public_path(path),
            "collisions": collisions,
        },
        qa={
            "status": "blocked",
            "findings": [{"severity": "high", "title": "Input/output collision", "detail": message}],
            "metrics": {"blocking_count": len(collisions)},
        },
        legacy={"blocked_reason": message, "collisions": collisions},
    )
    print(json.dumps(sanitize_payload(payload), indent=2, ensure_ascii=True))
    return 2


def _preflight_output_safety(args, path: Path) -> int | None:
    outputs: list[tuple[str, Path | str | None]] = [
        ("--output-json", args.output_json),
        ("--output-dir", args.output_dir),
    ]
    output_dir = Path(args.output_dir).expanduser() if args.output_dir else None
    if output_dir and output_dir.is_dir():
        outputs.extend(
            ("existing-output-member", item)
            for item in output_dir.rglob("*")
            if item.is_file()
        )
    collisions = find_output_input_collisions([("input", path)], outputs)
    return _emit_collision_only(path, collisions) if collisions else None


def clean_excerpt(text, max_chars):
    text = text.replace("\r", "\n")
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = text.strip()
    return text[:max_chars]


def run_subprocess(cmd, env=None, timeout=None):
    try:
        result = subprocess.run(cmd, capture_output=True, text=True, env=env, timeout=timeout)
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError(f"Command timed out: {' '.join(cmd)}") from exc
    if result.returncode != 0:
        raise RuntimeError(result.stderr.strip() or result.stdout.strip() or f"Command failed: {' '.join(cmd)}")
    return result.stdout


def safe_output_name(name):
    return re.sub(r"[^A-Za-z0-9._-]+", "_", Path(name).name)


def extract_plain_text(path, max_chars):
    return {"method": "plain_text", "text_excerpt": clean_excerpt(path.read_text(errors="replace"), max_chars)}


def extract_html_text(path, max_chars):
    raw = path.read_text(errors="replace")
    text = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", raw)
    text = re.sub(r"(?s)<[^>]+>", " ", text)
    text = html.unescape(text)
    return {"method": "html_regex", "text_excerpt": clean_excerpt(text, max_chars)}


def extract_textutil(path, max_chars):
    text = run_subprocess(["/usr/bin/textutil", "-stdout", "-convert", "txt", str(path)])
    return {"method": "textutil", "text_excerpt": clean_excerpt(text, max_chars)}


def extract_pandoc(path, max_chars):
    pandoc = shutil_which("pandoc")
    if not pandoc:
        raise RuntimeError("pandoc is not available")
    text = run_subprocess([pandoc, "-t", "plain", str(path)])
    return {"method": "pandoc", "text_excerpt": clean_excerpt(text, max_chars)}


def shutil_which(name):
    from shutil import which

    return which(name)


def xml_text_from_bytes(xml_bytes, namespaces=None, tags=None):
    root = ET.fromstring(xml_bytes)
    texts = []
    for elem in root.iter():
        tag = elem.tag
        if "}" in tag:
            local = tag.split("}", 1)[1]
        else:
            local = tag
        if tags is None or local in tags:
            if elem.text and elem.text.strip():
                texts.append(elem.text.strip())
    return texts


def summarize_ooxml_members(zf):
    names = zf.namelist()
    embedded = [
        name
        for name in names
        if "/embeddings/" in name or name.lower().endswith((".bin", ".ole", ".emf", ".wmf"))
    ]
    media = [name for name in names if "/media/" in name][:40]
    custom_xml = [name for name in names if name.startswith("customXml/")][:20]
    vba = [name for name in names if "vbaProject.bin" in name]
    active_x = [name for name in names if "activeX" in name][:20]
    return {
        "embedded_objects": embedded[:20],
        "embedded_object_count": len(embedded),
        "media_assets": media[:20],
        "media_asset_count": len(media),
        "custom_xml_parts": custom_xml,
        "activex_members": active_x,
        "has_vba": bool(vba),
    }


def extract_docx(path, max_chars):
    with zipfile.ZipFile(path) as zf:
        text = "\n".join(xml_text_from_bytes(zf.read("word/document.xml"), tags={"t"}))
        summary = {
            "method": "docx_xml",
            "text_excerpt": clean_excerpt(text, max_chars),
            "members": [name for name in zf.namelist()[:20]],
        }
        summary.update(summarize_ooxml_members(zf))
        return summary


def extract_pptx(path, max_chars):
    with zipfile.ZipFile(path) as zf:
        slide_names = sorted(name for name in zf.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml"))
        slides = []
        combined = []
        for name in slide_names[:20]:
            text = " ".join(xml_text_from_bytes(zf.read(name), tags={"t"}))
            slides.append({"slide": name, "text": clean_excerpt(text, 300)})
            combined.append(text)
        summary = {
            "method": "pptx_xml",
            "slide_count": len(slide_names),
            "slides": slides[:8],
            "text_excerpt": clean_excerpt("\n".join(combined), max_chars),
        }
        summary.update(summarize_ooxml_members(zf))
        return summary


def extract_xlsx(path):
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for sheet_name in workbook.sheetnames[:10]:
        ws = workbook[sheet_name]
        rows = []
        for row in ws.iter_rows(min_row=1, max_row=5, values_only=True):
            rows.append([cell for cell in row[:8]])
        sheets.append({"sheet": sheet_name, "preview_rows": rows})
    summary = {"method": "openpyxl", "sheet_count": len(workbook.sheetnames), "sheets": sheets}
    with zipfile.ZipFile(path) as zf:
        summary.update(summarize_ooxml_members(zf))
    return summary


def extract_xlsb(path):
    try:
        import pandas as pd
    except Exception:
        return extract_zip_members(path)
    try:
        workbook = pd.ExcelFile(path, engine="pyxlsb")
    except Exception:
        return extract_zip_members(path)
    sheets = []
    for sheet_name in workbook.sheet_names[:10]:
        frame = pd.read_excel(path, sheet_name=sheet_name, engine="pyxlsb").head(5)
        sheets.append({"sheet": sheet_name, "preview_rows": frame.iloc[:, :8].where(frame.notna(), None).values.tolist()})
    return {"method": "pyxlsb", "sheet_count": len(workbook.sheet_names), "sheets": sheets, "editable_backend": "pyxlsbwriter"}


def extract_zip_members(path):
    with zipfile.ZipFile(path) as zf:
        members = zf.namelist()
        return {"method": "zip_members", "member_count": len(members), "members": members[:30]}


def summarize_extension_counts(names, limit=12):
    counts = Counter()
    for name in names:
        suffix = Path(name).suffix.lower() or "<none>"
        counts[suffix] += 1
    return dict(counts.most_common(limit))


def summarize_top_level_entries(names, limit=8):
    counts = Counter()
    for name in names:
        top = name.split("/", 1)[0] if "/" in name else name
        counts[top] += 1
    return dict(counts.most_common(limit))


def sanitize_plist_value(value):
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    if isinstance(value, bytes):
        try:
            return value.decode("utf-8", "replace")
        except Exception:
            return repr(value)
    if isinstance(value, (list, tuple)):
        return [sanitize_plist_value(item) for item in value]
    if isinstance(value, dict):
        return {str(key): sanitize_plist_value(item) for key, item in value.items()}
    return str(value)


def read_embedded_plist(zf, member_name):
    if member_name not in zf.namelist():
        return None
    try:
        return plistlib.loads(zf.read(member_name))
    except Exception:
        return None


def extract_iwork_bundle_metadata(zf):
    properties = read_embedded_plist(zf, "Metadata/Properties.plist")
    build_history = read_embedded_plist(zf, "Metadata/BuildVersionHistory.plist")
    bundle_metadata = {}
    if isinstance(properties, dict):
        for key in [
            "documentUUID",
            "versionUUID",
            "stableDocumentUUID",
            "shareUUID",
            "privateUUID",
            "revision",
            "fileFormatVersion",
            "isMultiPage",
            "hasExternalReferenceOrMissingOrUnmaterializedRemoteData",
        ]:
            if key in properties:
                bundle_metadata[key] = sanitize_plist_value(properties[key])
    return bundle_metadata, sanitize_plist_value(build_history) if build_history is not None else []


def infer_iwork_structure(path, names, asset_hints, iwa_hints):
    counts = {
        "slide_members": 0,
        "template_slide_members": 0,
        "sheet_members": 0,
        "table_members": 0,
        "equation_members": 0,
        "screenshot_assets": 0,
        "preview_members": 0,
    }
    for name in names:
        base = Path(name).name.lower()
        if base.startswith("slide") and base.endswith(".iwa"):
            counts["slide_members"] += 1
        if base.startswith("templateslide") and base.endswith(".iwa"):
            counts["template_slide_members"] += 1
        if base.startswith("sheet") and base.endswith(".iwa"):
            counts["sheet_members"] += 1
        if base.startswith("table") and base.endswith(".iwa"):
            counts["table_members"] += 1
        if re.fullmatch(r"equation-\d+\.pdf", base):
            counts["equation_members"] += 1
        if "screenshot" in base:
            counts["screenshot_assets"] += 1
        if "preview" in base or "thumbnail" in base:
            counts["preview_members"] += 1
    suffix = path.suffix.lower()
    roles = []
    if suffix == ".key":
        if counts["slide_members"]:
            roles.append(f"slide deck with {counts['slide_members']} slide-like members")
        if counts["template_slide_members"]:
            roles.append(f"{counts['template_slide_members']} template slides")
    elif suffix == ".numbers":
        if counts["sheet_members"]:
            roles.append(f"spreadsheet-like package with {counts['sheet_members']} sheet members")
        if counts["table_members"]:
            roles.append(f"{counts['table_members']} table members detected")
    elif suffix == ".pages":
        if counts["equation_members"]:
            roles.append(f"document with {counts['equation_members']} equation PDFs")
        if counts["screenshot_assets"]:
            roles.append(f"{counts['screenshot_assets']} screenshot assets")
    if asset_hints:
        roles.append(f"{len(asset_hints)} named assets with semantic hints")
    if iwa_hints:
        roles.append(f"{min(len(iwa_hints), 20)} textual hints recovered from IWA streams")
    return {"counts": counts, "roles": roles}


def export_numbers_tables(document, output_dir):
    exports = []
    if output_dir is None:
        return exports
    from csv import writer

    output_dir.mkdir(parents=True, exist_ok=True)
    for sheet_idx, sheet in enumerate(document.sheets, start=1):
        for table_idx, table in enumerate(sheet.tables, start=1):
            safe_name = safe_output_name(f"{sheet_idx:02d}_{sheet.name}_{table_idx:02d}_{table.name}.csv")
            target = output_dir / safe_name
            with target.open("w", newline="") as handle:
                csv_writer = writer(handle)
                for row in table.rows(values_only=True):
                    csv_writer.writerow(row)
            exports.append(str(target.resolve()))
    return exports


def summarize_numbers_document(path, output_dir=None):
    from numbers_parser import Document

    document = Document(path)
    sheets = []
    relation_hints = []
    for sheet in document.sheets:
        tables = []
        for table in sheet.tables:
            preview_rows = []
            formula_count = 0
            string_hints = []
            for row_idx, row in enumerate(table.rows()):
                if row_idx < 5:
                    preview_rows.append([getattr(cell, "formatted_value", getattr(cell, "value", None)) for cell in row[:8]])
                for cell in row:
                    if getattr(cell, "is_formula", False) and getattr(cell, "formula", None):
                        formula_count += 1
                        formula_text = str(cell.formula)
                        if formula_text not in relation_hints:
                            relation_hints.append(formula_text)
                    value = getattr(cell, "formatted_value", getattr(cell, "value", None))
                    if isinstance(value, str) and len(value.strip()) >= 4 and value not in string_hints:
                        string_hints.append(value.strip())
            tables.append(
                {
                    "name": table.name,
                    "rows": int(table.num_rows),
                    "cols": int(table.num_cols),
                    "header_rows": int(table.num_header_rows),
                    "header_cols": int(table.num_header_cols),
                    "formula_count": formula_count,
                    "preview_rows": preview_rows,
                    "string_hints": string_hints[:15],
                }
            )
        sheets.append({"sheet": sheet.name, "tables": tables})
    exported_tables = export_numbers_tables(document, output_dir / "numbers_tables" if output_dir else None)
    return {
        "sheet_count": len(document.sheets),
        "sheets": sheets,
        "numbers_relation_hints": relation_hints[:20],
        "exported_tables": exported_tables,
    }


def extract_pdf(path, max_chars):
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    texts = []
    for page in reader.pages[:5]:
        try:
            texts.append(page.extract_text() or "")
        except Exception:
            continue
    excerpt = clean_excerpt("\n".join(texts), max_chars)
    summary = {"method": "pypdf", "pages": len(reader.pages), "text_excerpt": excerpt}
    if not excerpt.strip():
        try:
            with tempfile.TemporaryDirectory(prefix="scientific-data-analysis-pdfocr-") as tmpdir:
                ocr_summary = ocr_pdf(path, tmpdir, max_pages=min(3, len(reader.pages)))
            summary["ocr_excerpt"] = clean_excerpt(ocr_summary.get("text", ""), max_chars)
            summary["ocr_backend"] = "rapidocr"
            summary["ocr_rendered_images"] = ocr_summary.get("rendered_images", [])
            if summary["ocr_excerpt"]:
                summary["text_excerpt"] = summary["ocr_excerpt"]
                summary["method"] = "pypdf+ocr"
        except Exception:
            pass
    return summary


def extract_xls(path, max_chars):
    import olefile

    summary = {
        "method": "olefile",
        "vba_assessment": "not_assessed",
        "vba_assessment_note": (
            "This legacy .xls inspection does not detect VBA macros; "
            "the absence of has_vba is not evidence that the file is macro-free."
        ),
    }
    ole = olefile.OleFileIO(str(path))
    try:
        summary["streams"] = ["/".join(item) for item in ole.listdir()[:20]]
        try:
            meta = ole.get_metadata()
            def safe_meta(value):
                if isinstance(value, bytes):
                    try:
                        return value.decode("latin1", "replace")
                    except Exception:
                        return repr(value)
                return str(value) if value is not None else None
            summary["metadata"] = {
                "title": safe_meta(getattr(meta, "title", None)),
                "author": safe_meta(getattr(meta, "author", None)),
                "last_saved_by": safe_meta(getattr(meta, "last_saved_by", None)),
                "create_time": safe_meta(getattr(meta, "create_time", None)),
            }
        except Exception:
            pass
    finally:
        ole.close()
    strings_bin = shutil_which("strings")
    if strings_bin:
        try:
            text = run_subprocess([strings_bin, "-n", "6", str(path)])
            summary["text_excerpt"] = clean_excerpt(text, max_chars)
        except Exception:
            pass
    return summary


def interesting_asset_names(names):
    hints = []
    for name in names:
        base = Path(name).name
        if base.startswith(("preview", "thumbnail", "preview-")):
            continue
        if re.fullmatch(r"(mt|st)-[A-F0-9-]+-\d+\.(jpg|jpeg|png)", base, flags=re.IGNORECASE):
            continue
        if re.fullmatch(r"equation-\d+\.pdf", base, flags=re.IGNORECASE):
            continue
        if re.search(r"[A-Za-z]{3,}", base):
            hints.append(base)
    deduped = []
    seen = set()
    for item in hints:
        if item not in seen:
            deduped.append(item)
            seen.add(item)
    return deduped[:20]


def iwa_string_hints(data):
    strings = re.findall(rb"[ -~]{6,}", data)
    hints = []
    for raw in strings:
        text = raw.decode("ascii", "ignore").strip()
        if not text:
            continue
        if text in {"January", "February", "March", "April", "August", "September", "October", "November", "December"}:
            continue
        if text not in hints:
            hints.append(text)
    return hints[:20]


def unique_preserve(items, limit=None):
    out = []
    seen = set()
    for item in items:
        if not item or item in seen:
            continue
        out.append(item)
        seen.add(item)
        if limit is not None and len(out) >= limit:
            break
    return out


def preprocess_images_for_ocr(image_paths):
    try:
        from PIL import Image, ImageOps
    except Exception:
        return list(image_paths)
    processed = []
    for path in image_paths:
        try:
            image = Image.open(path).convert("L")
            image = ImageOps.autocontrast(image)
            scale = 2 if max(image.size) < 1800 else 1
            if scale != 1:
                image = image.resize((image.width * scale, image.height * scale))
            target = path.with_name(path.stem + "_ocr.png")
            image.save(target)
            processed.append(target)
        except Exception:
            processed.append(path)
    return processed


def run_apple_vision_ocr(image_paths):
    swift = shutil_which("swift")
    script_path = Path(__file__).resolve().parent / "_internal" / "apple_vision_ocr.swift"
    if not swift or not script_path.exists() or not image_paths:
        return []
    cache_root = Path(tempfile.gettempdir()) / "scientific-data-analysis-swift-cache"
    module_cache = cache_root / "module-cache"
    module_cache.mkdir(parents=True, exist_ok=True)
    env = os.environ.copy()
    env["CLANG_MODULE_CACHE_PATH"] = str(module_cache)
    env["SWIFT_MODULECACHE_PATH"] = str(module_cache)
    try:
        prepared_paths = preprocess_images_for_ocr(image_paths)
        raw = run_subprocess([swift, str(script_path), *[str(path) for path in prepared_paths]], env=env, timeout=45)
        payload = json.loads(raw)
    except Exception:
        return []
    if isinstance(payload, list):
        return payload
    return []


def maybe_extract_assets(zf, member_names, output_dir):
    if output_dir is None:
        return []
    output_dir.mkdir(parents=True, exist_ok=True)
    exported = []
    for idx, name in enumerate(member_names[:8], start=1):
        target = output_dir / f"{idx:02d}_{safe_output_name(name)}"
        try:
            target.write_bytes(zf.read(name))
            exported.append(str(target.resolve()))
        except Exception:
            continue
    return exported


def extract_iwork_package(path, max_chars, output_dir=None, enable_ocr=False):
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        preview_files = [name for name in names if "preview" in name.lower() or "thumbnail" in name.lower()][:20]
        image_assets = [name for name in names if Path(name).suffix.lower() in {".jpg", ".jpeg", ".png"}]
        equation_assets = [name for name in names if re.fullmatch(r".*equation-\d+\.pdf", name, flags=re.IGNORECASE)]
        iwa_files = [name for name in names if name.lower().endswith(".iwa")]
        iwa_member_summaries = []
        iwa_hints = []
        for name in iwa_files[:20]:
            try:
                raw_member = zf.read(name)
                iwa_hints.extend(iwa_string_hints(raw_member))
                iwa_summary = inspect_iwa_member(raw_member, name)
                iwa_member_summaries.append(iwa_summary)
                iwa_hints.extend(iwa_summary.get("string_hints", []))
            except Exception:
                continue
        iwa_hints = unique_preserve(iwa_hints, limit=40)
        ocr_text = ""
        exported_assets = []
        ocr_sources = []
        candidate_images = unique_preserve(preview_files + image_assets, limit=4)
        if output_dir is not None:
            exported_assets = maybe_extract_assets(zf, unique_preserve(preview_files + equation_assets + image_assets, limit=8), output_dir)
        if enable_ocr and candidate_images:
            with tempfile.TemporaryDirectory(prefix="scientific-data-analysis-iwork-") as tmpdir:
                temp_paths = []
                for name in candidate_images:
                    target = Path(tmpdir) / safe_output_name(name)
                    try:
                        target.write_bytes(zf.read(name))
                    except Exception:
                        continue
                    temp_paths.append(target)
                ocr_results = []
                for item in temp_paths:
                    payload = ocr_image(item)
                    if payload.get("text"):
                        ocr_results.append({"image": str(item), "text": payload["text"], "backend": payload.get("backend")})
                ocr_sources = [item.get("image", "") for item in ocr_results if item.get("text")]
                ocr_text = clean_excerpt("\n".join(item.get("text", "") for item in ocr_results if item.get("text")), max_chars)
        bundle_metadata, build_history = extract_iwork_bundle_metadata(zf)
        semantic_blocks = []
        if ocr_text:
            semantic_blocks.append(ocr_text)
        if iwa_hints:
            semantic_blocks.append("IWA hints: " + "; ".join(iwa_hints[:20]))
        asset_hints = interesting_asset_names(names)
        structure_hints = infer_iwork_structure(path, names, asset_hints, iwa_hints)
        numbers_summary = None
        if path.suffix.lower() == ".numbers":
            try:
                numbers_summary = summarize_numbers_document(path, output_dir=output_dir)
            except Exception as exc:
                numbers_summary = {"error": str(exc)}
        if asset_hints:
            semantic_blocks.append("Asset hints: " + "; ".join(asset_hints[:15]))
        if build_history:
            semantic_blocks.append("Build history: " + "; ".join(str(item) for item in build_history[:3]))
        if bundle_metadata:
            concise_metadata = []
            for key in ["fileFormatVersion", "isMultiPage", "hasExternalReferenceOrMissingOrUnmaterializedRemoteData"]:
                if key in bundle_metadata:
                    concise_metadata.append(f"{key}={bundle_metadata[key]}")
            if concise_metadata:
                semantic_blocks.append("Bundle metadata: " + "; ".join(concise_metadata))
        if structure_hints.get("roles"):
            semantic_blocks.append("Structure hints: " + "; ".join(structure_hints["roles"][:5]))
        if numbers_summary and numbers_summary.get("sheets"):
            blocks = []
            for sheet in numbers_summary["sheets"][:3]:
                table_names = ", ".join(table["name"] for table in sheet["tables"][:4])
                blocks.append(f"{sheet['sheet']}: {table_names}")
            semantic_blocks.append("Numbers sheets: " + "; ".join(blocks))
        return {
            "method": "iwork_zip",
            "package_type": path.suffix.lower(),
            "sheet_count": numbers_summary.get("sheet_count") if isinstance(numbers_summary, dict) and numbers_summary.get("sheet_count") is not None else None,
            "preview_files": preview_files,
            "image_asset_count": len(image_assets),
            "equation_pdf_count": len(equation_assets),
            "iwa_file_count": len(iwa_files),
            "asset_name_hints": asset_hints,
            "document_iwa_hints": iwa_hints,
            "iwa_member_summaries": iwa_member_summaries[:10],
            "bundle_metadata": bundle_metadata,
            "build_history": build_history[:5] if isinstance(build_history, list) else build_history,
            "member_extension_counts": summarize_extension_counts(names, limit=12),
            "top_level_entries": summarize_top_level_entries(names, limit=8),
            "structure_hints": structure_hints,
            "numbers_summary": numbers_summary,
            "ocr_excerpt": ocr_text,
            "ocr_sources": ocr_sources,
            "exported_assets": exported_assets,
            "text_excerpt": clean_excerpt("\n\n".join(semantic_blocks), max_chars) if semantic_blocks else "",
            "member_count": len(names),
        }


def extract_odf_package(path, max_chars):
    with zipfile.ZipFile(path) as zf:
        members = zf.namelist()
        text = ""
        meta_text = ""
        if "content.xml" in members:
            text = "\n".join(xml_text_from_bytes(zf.read("content.xml")))
        if "meta.xml" in members:
            meta_text = "\n".join(xml_text_from_bytes(zf.read("meta.xml")))
        return {
            "method": "odf_zip",
            "member_count": len(members),
            "members": members[:20],
            "text_excerpt": clean_excerpt("\n".join([meta_text, text]), max_chars),
        }


def extract_epub(path, max_chars):
    with zipfile.ZipFile(path) as zf:
        members = zf.namelist()
        pieces = []
        for name in members:
            lower = name.lower()
            if lower.endswith((".xhtml", ".html", ".htm", ".ncx", ".opf")):
                try:
                    pieces.extend(xml_text_from_bytes(zf.read(name)))
                except Exception:
                    continue
        return {
            "method": "epub_zip",
            "member_count": len(members),
            "members": members[:20],
            "text_excerpt": clean_excerpt("\n".join(pieces), max_chars),
        }


def extract_reg(path, max_chars):
    lines = path.read_text(errors="replace").splitlines()
    shapes = []
    for line in lines:
        stripped = line.strip()
        if not stripped or stripped.startswith("#") or stripped.startswith("global") or stripped == "physical":
            continue
        shape = stripped.split("(", 1)[0]
        shapes.append(shape)
    return {
        "method": "ds9_region",
        "shape_counts": {shape: shapes.count(shape) for shape in sorted(set(shapes))},
        "text_excerpt": clean_excerpt("\n".join(lines[:20]), max_chars),
    }


def extract_properties(path, max_chars):
    lines = path.read_text(errors="replace").splitlines()
    props = {}
    for line in lines:
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        props[key.strip()] = value.strip()
        if len(props) >= 30:
            break
    return {"method": "properties", "properties": props, "text_excerpt": clean_excerpt("\n".join(lines[:40]), max_chars)}


def detect_and_extract(path, max_chars, output_dir=None, enable_ocr=False):
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".tex", ".bib", ".json", ".jsonl", ".ndjson", ".yaml", ".yml", ".toml", ".xml", ".ini", ".cfg", ".log", ".csv", ".tsv"}:
        return extract_plain_text(path, max_chars)
    if suffix == ".reg":
        return extract_reg(path, max_chars)
    if suffix in {".pref", ".jmars"}:
        return extract_properties(path, max_chars)
    if suffix in {".html", ".htm"}:
        for func in (extract_html_text, extract_textutil, extract_pandoc):
            try:
                return func(path, max_chars)
            except Exception:
                continue
        return extract_plain_text(path, max_chars)
    if suffix in {".rtf", ".doc"}:
        for func in (extract_textutil, extract_pandoc):
            try:
                return func(path, max_chars)
            except Exception:
                continue
        return extract_plain_text(path, max_chars)
    if suffix in {".docx", ".docm"}:
        try:
            return extract_docx(path, max_chars)
        except Exception:
            for func in (extract_textutil, extract_pandoc):
                try:
                    return func(path, max_chars)
                except Exception:
                    continue
            raise
    if suffix in {".pptx", ".pptm"}:
        return extract_pptx(path, max_chars)
    if suffix in {".xlsx", ".xlsm"}:
        return extract_xlsx(path)
    if suffix == ".xlsb":
        return extract_xlsb(path)
    if suffix == ".xls":
        return extract_xls(path, max_chars)
    if suffix in {".pages", ".key", ".numbers"}:
        return extract_iwork_package(path, max_chars, output_dir=output_dir, enable_ocr=enable_ocr)
    if suffix == ".pdf":
        return extract_pdf(path, max_chars)
    if suffix in {".odt", ".ods", ".odp"}:
        for func in (extract_odf_package, extract_pandoc, extract_textutil):
            try:
                return func(path, max_chars)
            except Exception:
                continue
        return extract_zip_members(path)
    if suffix == ".epub":
        return extract_epub(path, max_chars)
    return {"method": "fallback", "text_excerpt": clean_excerpt(path.read_text(errors="replace"), max_chars)}


def print_summary(path, summary):
    numbers_summary = summary.get("numbers_summary") or {}
    print(f"Document: {path.resolve()}")
    for key in ["method", "package_type", "pages", "slide_count", "sheet_count", "member_count", "equation_pdf_count", "image_asset_count"]:
        if key in summary:
            print(f"{key}: {summary[key]}")
    if "metadata" in summary:
        print(f"metadata: {json.dumps(summary['metadata'], ensure_ascii=True)}")
    if "shape_counts" in summary:
        print(f"shape_counts: {json.dumps(summary['shape_counts'], ensure_ascii=True)}")
    if "properties" in summary:
        keys = list(summary["properties"].keys())[:12]
        print("properties: " + ", ".join(keys))
    if "preview_files" in summary and summary["preview_files"]:
        print("preview_files: " + ", ".join(summary["preview_files"][:5]))
    if "embedded_objects" in summary and summary["embedded_objects"]:
        print("embedded_objects: " + ", ".join(summary["embedded_objects"][:8]))
    if "media_assets" in summary and summary["media_assets"]:
        print("media_assets: " + ", ".join(summary["media_assets"][:8]))
    if "asset_name_hints" in summary and summary["asset_name_hints"]:
        print("asset_hints: " + ", ".join(summary["asset_name_hints"][:8]))
    if "document_iwa_hints" in summary and summary["document_iwa_hints"]:
        print("iwa_hints: " + ", ".join(summary["document_iwa_hints"][:8]))
    if "ocr_sources" in summary and summary["ocr_sources"]:
        print("ocr_sources: " + ", ".join(summary["ocr_sources"][:4]))
    if "ocr_rendered_images" in summary and summary["ocr_rendered_images"]:
        print("ocr_rendered_images: " + ", ".join(summary["ocr_rendered_images"][:3]))
    if "exported_assets" in summary and summary["exported_assets"]:
        print("exported_assets:")
        for asset in summary["exported_assets"][:6]:
            print(f"- {asset}")
    if "slides" in summary and summary["slides"]:
        print("slide_preview:")
        for slide in summary["slides"][:3]:
            print(f"- {slide['slide']}: {slide['text'][:120]}")
    if "sheets" in summary and summary["sheets"]:
        print("sheet_preview:")
        for sheet in summary["sheets"][:3]:
            print(f"- {sheet['sheet']}: {sheet['preview_rows'][:2]}")
    if numbers_summary.get("sheets"):
        print("numbers_sheet_preview:")
        for sheet in numbers_summary["sheets"][:3]:
            table_names = ", ".join(table["name"] for table in sheet["tables"][:4])
            print(f"- {sheet['sheet']}: {table_names}")
    if "text_excerpt" in summary and summary["text_excerpt"]:
        print("")
        print(summary["text_excerpt"])


def main():
    args = parse_args()
    path = Path(args.path).expanduser()
    collision_status = _preflight_output_safety(args, path)
    if collision_status is not None:
        raise SystemExit(collision_status)
    if not path.exists():
        raise SystemExit(f"File not found: {path}")
    output_dir = Path(args.output_dir) if args.output_dir else None
    summary = detect_and_extract(path, args.max_chars, output_dir=output_dir, enable_ocr=args.enable_ocr)
    summary["path"] = str(path.resolve())
    print_summary(path, summary)
    if args.output_json:
        out = Path(args.output_json)
        out.parent.mkdir(parents=True, exist_ok=True)
        payload = build_tool_payload(
            "document_semantics",
            status="ok",
            notes=[
                "Document semantics are an intake summary, not an authoritative conversion of layout or scientific content.",
            ],
            artifacts={"output_json": str(out)},
            results=summary,
            qa={
                "status": "ok",
                "findings": [],
                "metrics": {
                    "text_excerpt_chars": len(summary.get("text_excerpt") or ""),
                    "image_asset_count": int(summary.get("image_asset_count") or 0),
                    "member_count": int(summary.get("member_count") or 0),
                },
            },
            legacy=summary,
        )
        out.write_text(json.dumps(payload, indent=2, ensure_ascii=True) + "\n")
        print("")
        print(f"Saved summary: {out.resolve()}")


if __name__ == "__main__":
    main()
