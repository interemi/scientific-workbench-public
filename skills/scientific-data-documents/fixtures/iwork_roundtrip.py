#!/usr/bin/env python3
"""Safe copy-based edits and conversions for iWork documents."""

import argparse
import csv
import json
import subprocess
from pathlib import Path

from document_semantics import detect_and_extract
from _internal.path_safety import find_output_input_collisions
from _internal.provenance_utils import public_path, sanitize_payload, standard_tool_payload, write_manifest
from quicklook_bridge import generate_preview


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)

    export_cmd = subparsers.add_parser("export", help="Export tables or text from an iWork file.")
    export_cmd.add_argument("input")
    export_cmd.add_argument("--output-dir", required=True)
    export_cmd.add_argument("--fidelity-json")
    export_cmd.add_argument("--manifest-json")
    export_cmd.add_argument("--enable-ocr", action="store_true")

    set_cmd = subparsers.add_parser("set-cell", help="Edit a Numbers cell in a copied document.")
    set_cmd.add_argument("input")
    set_cmd.add_argument("output")
    set_cmd.add_argument("--sheet", default=None)
    set_cmd.add_argument("--table", default=None)
    set_cmd.add_argument("--cell", required=True, help="A1 or row/col notation accepted by numbers-parser.")
    set_cmd.add_argument("--value", required=True)
    set_cmd.add_argument("--manifest-json")

    replace_cmd = subparsers.add_parser("replace-table-csv", help="Replace a Numbers table from CSV in a copied document.")
    replace_cmd.add_argument("input")
    replace_cmd.add_argument("csv_path")
    replace_cmd.add_argument("output")
    replace_cmd.add_argument("--sheet", default=None)
    replace_cmd.add_argument("--table", default=None)
    replace_cmd.add_argument("--manifest-json")

    convert_cmd = subparsers.add_parser("convert", help="Convert an iWork file into a more interoperable derivative.")
    convert_cmd.add_argument("input")
    convert_cmd.add_argument("output")
    convert_cmd.add_argument("--manifest-json")
    convert_cmd.add_argument("--fidelity-json")
    convert_cmd.add_argument("--enable-ocr", action="store_true")
    return parser.parse_args()


def _collision_request(args) -> tuple[list[tuple[str, str | Path | None]], list[tuple[str, str | Path | None]]]:
    inputs: list[tuple[str, str | Path | None]] = [("input", args.input)]
    outputs: list[tuple[str, str | Path | None]] = []
    if args.command == "export":
        output_dir = Path(args.output_dir).expanduser()
        outputs.extend(
            [
                ("--output-dir", output_dir),
                ("--fidelity-json", args.fidelity_json),
                ("--manifest-json", args.manifest_json),
                ("derived:semantic.txt", output_dir / f"{Path(args.input).stem}_semantic.txt"),
                ("derived:semantic.md", output_dir / f"{Path(args.input).stem}_semantic.md"),
                ("derived:semantic.json", output_dir / f"{Path(args.input).stem}_semantic.json"),
                ("derived:preview.png", output_dir / f"{Path(args.input).stem}_preview.png"),
                ("derived:xlsx", output_dir / f"{Path(args.input).stem}.xlsx"),
            ]
        )
        if output_dir.is_dir():
            outputs.extend(
                ("existing-output-member", item)
                for item in output_dir.rglob("*")
                if item.is_file()
            )
    else:
        output = Path(args.output).expanduser()
        outputs.extend(
            [
                ("output", output),
                ("--manifest-json", args.manifest_json),
            ]
        )
        if args.command == "replace-table-csv":
            inputs.append(("csv_path", args.csv_path))
        if args.command == "convert":
            outputs.extend(
                [
                    ("--fidelity-json", args.fidelity_json),
                    ("derived:temporary_csv", output.parent / f"{Path(args.input).stem}_temp_export.csv"),
                    ("derived:intermediate_xlsx", output.parent / f"{Path(args.input).stem}.xlsx"),
                ]
            )
    return inputs, outputs


def _emit_collision_only(args, collisions: list[dict[str, str]]) -> int:
    message = "Refusing iWork round-trip outputs that overlap protected inputs."
    payload = standard_tool_payload(
        "iwork_roundtrip",
        status="blocked",
        notes=[message, "No copy, export, fidelity report, or manifest was written."],
        artifacts={},
        results={
            "blocked_reason": message,
            "error_type": "output_input_collision",
            "command": args.command,
            "collisions": collisions,
        },
        qa={
            "status": "blocked",
            "findings": [{"severity": "high", "title": "Input/output collision", "detail": message}],
            "metrics": {"blocking_count": len(collisions)},
        },
        legacy={"blocked_reason": message, "collisions": collisions},
    )
    print(json.dumps(sanitize_payload(payload), indent=2, ensure_ascii=True))
    return 2


def _preflight_output_safety(args) -> int | None:
    inputs, outputs = _collision_request(args)
    collisions = find_output_input_collisions(inputs, outputs)
    return _emit_collision_only(args, collisions) if collisions else None


def parse_scalar(value):
    lowered = value.strip().lower()
    if lowered in {"true", "false"}:
        return lowered == "true"
    try:
        if "." in value:
            return float(value)
        return int(value)
    except Exception:
        return value


def load_numbers_document(path):
    from numbers_parser import Document

    return Document(path)


def choose_table(document, sheet_name=None, table_name=None):
    sheets = document.sheets
    sheet = sheets[0] if sheet_name is None else next(item for item in sheets if item.name == sheet_name)
    tables = sheet.tables
    table = tables[0] if table_name is None else next(item for item in tables if item.name == table_name)
    return sheet, table


def export_numbers(path, output_dir):
    from openpyxl import Workbook

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    document = load_numbers_document(path)
    exports = []
    workbook = Workbook()
    workbook.remove(workbook.active)
    fidelity = {"sheet_count": len(document.sheets), "tables": []}
    for sheet in document.sheets:
        for table in sheet.tables:
            csv_name = f"{sheet.name}__{table.name}.csv"
            csv_path = output_dir / csv_name.replace("/", "_")
            with csv_path.open("w", newline="") as handle:
                csv_writer = csv.writer(handle)
                for row in table.rows(values_only=True):
                    csv_writer.writerow(row)
            exports.append(str(csv_path.resolve()))
            ws = workbook.create_sheet((f"{sheet.name}__{table.name}")[:31])
            for row in table.rows(values_only=True):
                ws.append(row)
            fidelity["tables"].append(
                {
                    "sheet": sheet.name,
                    "table": table.name,
                    "rows": int(table.num_rows),
                    "cols": int(table.num_cols),
                    "formula_count": sum(
                        1
                        for table_row in table.rows()
                        for cell in table_row
                        if getattr(cell, "is_formula", False)
                    ),
                }
            )
    xlsx_path = output_dir / f"{Path(path).stem}.xlsx"
    workbook.save(xlsx_path)
    exports.append(str(xlsx_path.resolve()))
    return exports, fidelity


def write_semantic_exports(path, output_dir, enable_ocr=False):
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    summary = detect_and_extract(Path(path), 8000, output_dir=output_dir, enable_ocr=enable_ocr)
    exports = []
    stem = Path(path).stem
    text_excerpt = (summary.get("text_excerpt") or "").strip()
    if text_excerpt:
        txt_path = output_dir / f"{stem}_semantic.txt"
        txt_path.write_text(text_excerpt + "\n")
        exports.append(str(txt_path.resolve()))
        md_path = output_dir / f"{stem}_semantic.md"
        md_path.write_text(f"# Semantic Export\n\n{text_excerpt}\n")
        exports.append(str(md_path.resolve()))
    json_path = output_dir / f"{stem}_semantic.json"
    json_path.write_text(json.dumps(summary, indent=2, ensure_ascii=True) + "\n")
    exports.append(str(json_path.resolve()))
    preview_path = output_dir / f"{stem}_preview.png"
    preview_result = generate_preview(path, preview_path, size=1800)
    if preview_result.get("ok"):
        exports.append(str(preview_path.resolve()))
    fidelity = {
        "package_type": Path(path).suffix.lower(),
        "preview_method": preview_result.get("method"),
        "preview_native_error": preview_result.get("native_error"),
        "text_available": bool(text_excerpt),
        "member_count": summary.get("member_count"),
    }
    if summary.get("structure_hints"):
        fidelity["structure_hints"] = summary["structure_hints"]
    return exports, fidelity


def export_pages_or_key(path, output_dir, enable_ocr=False):
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    exports = []
    suffix = Path(path).suffix.lower()
    if suffix == ".pages":
        for ext in ("txt", "html", "rtf"):
            target = output_dir / f"{Path(path).stem}.{ext}"
            result = subprocess.run(
                ["/usr/bin/textutil", "-convert", ext, "-output", str(target), str(path)],
                capture_output=True,
                text=True,
            )
            if result.returncode == 0 and target.exists():
                exports.append(str(target.resolve()))
                break
    semantic_exports, fidelity = write_semantic_exports(path, output_dir, enable_ocr=enable_ocr)
    exports.extend(item for item in semantic_exports if item not in exports)
    fidelity["export_count"] = len(exports)
    return exports, fidelity


def write_xlsb_from_csv(csv_path, xlsb_path):
    from pyxlsbwriter import XlsbWriter

    with open(csv_path, newline="") as handle:
        rows = list(csv.reader(handle))
    with XlsbWriter(str(xlsb_path)) as writer:
        writer.add_sheet("Sheet1")
        writer.write_sheet(rows)


def cmd_export(args):
    input_path = Path(args.input)
    if input_path.suffix.lower() == ".numbers":
        exports, fidelity = export_numbers(input_path, args.output_dir)
    else:
        exports, fidelity = export_pages_or_key(input_path, args.output_dir, enable_ocr=args.enable_ocr)
    if args.fidelity_json:
        Path(args.fidelity_json).write_text(json.dumps(fidelity, indent=2, ensure_ascii=True) + "\n")
    if args.manifest_json:
        write_manifest(args.manifest_json, inputs=[input_path], outputs=exports, parameters={"command": "export"})
    print(json.dumps({"exports": exports, "fidelity": fidelity}, indent=2, ensure_ascii=True))


def cmd_set_cell(args):
    doc = load_numbers_document(args.input)
    _, table = choose_table(doc, args.sheet, args.table)
    table.write(args.cell, parse_scalar(args.value))
    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    if args.manifest_json:
        write_manifest(args.manifest_json, inputs=[args.input], outputs=[output_path], parameters={"cell": args.cell, "value": args.value})
    print(f"Saved edited Numbers copy: {output_path.resolve()}")


def cmd_replace_table_csv(args):
    doc = load_numbers_document(args.input)
    _, table = choose_table(doc, args.sheet, args.table)
    with open(args.csv_path, newline="") as handle:
        rows = list(csv.reader(handle))
    while table.num_rows < len(rows):
        table.add_row()
    max_cols = max((len(row) for row in rows), default=0)
    while table.num_cols < max_cols:
        table.add_column()
    for row_idx, row in enumerate(rows):
        for col_idx, value in enumerate(row):
            table.write(row_idx, col_idx, parse_scalar(value))
    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    if args.manifest_json:
        write_manifest(args.manifest_json, inputs=[args.input, args.csv_path], outputs=[output_path], parameters={"command": "replace-table-csv"})
    print(f"Saved CSV-updated Numbers copy: {output_path.resolve()}")


def cmd_convert(args):
    input_path = Path(args.input)
    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fidelity = {}
    if input_path.suffix.lower() == ".numbers" and output_path.suffix.lower() == ".xlsx":
        exports, fidelity = export_numbers(input_path, output_path.parent)
        generated = next((Path(item) for item in exports if item.endswith(".xlsx")), None)
        if generated is None:
            raise SystemExit("Could not generate XLSX export from Numbers.")
        generated.replace(output_path)
    elif output_path.suffix.lower() == ".xlsb":
        with tempfile_csv(input_path, output_path.parent) as csv_path:
            write_xlsb_from_csv(csv_path, output_path)
        fidelity = {"converted_via": "csv_to_xlsb"}
    else:
        if input_path.suffix.lower() in {".pages", ".key"} and output_path.suffix.lower() in {".txt", ".md", ".html", ".png", ".json"}:
            summary = detect_and_extract(input_path, 12000, output_dir=output_path.parent, enable_ocr=args.enable_ocr)
            text_excerpt = (summary.get("text_excerpt") or "").strip()
            if output_path.suffix.lower() == ".png":
                result = generate_preview(input_path, output_path, size=1800)
                if not result.get("ok"):
                    raise SystemExit(result.get("native_error") or "Preview generation failed")
                fidelity = {"converted_via": result.get("method"), "native_error": result.get("native_error")}
            elif output_path.suffix.lower() == ".json":
                output_path.write_text(json.dumps(summary, indent=2, ensure_ascii=True) + "\n")
                fidelity = {"converted_via": "semantic_json"}
            elif output_path.suffix.lower() == ".html":
                html = (
                    "<!doctype html><html><head><meta charset='utf-8'><title>Semantic Export</title></head><body>"
                    f"<h1>{input_path.name}</h1><pre>{text_excerpt}</pre></body></html>"
                )
                output_path.write_text(html + "\n")
                fidelity = {"converted_via": "semantic_html"}
            else:
                prefix = "# Semantic Export\n\n" if output_path.suffix.lower() == ".md" else ""
                output_path.write_text(prefix + text_excerpt + "\n")
                fidelity = {"converted_via": "semantic_text"}
        else:
            if output_path.suffix.lower() in {".txt", ".html", ".rtf"}:
                result = subprocess.run(
                    ["/usr/bin/textutil", "-convert", output_path.suffix.lower().lstrip("."), "-output", str(output_path), str(input_path)],
                    capture_output=True,
                    text=True,
                )
                if result.returncode != 0:
                    raise SystemExit(result.stderr.strip() or result.stdout.strip() or "textutil conversion failed")
                fidelity = {"converted_via": "textutil"}
            elif output_path.suffix.lower() == ".png":
                result = generate_preview(input_path, output_path, size=1800)
                if not result.get("ok"):
                    raise SystemExit(result.get("native_error") or "Preview generation failed")
                fidelity = {"converted_via": result.get("method"), "native_error": result.get("native_error")}
            else:
                raise SystemExit(f"Unsupported conversion target: {output_path.suffix}")
    if args.fidelity_json:
        Path(args.fidelity_json).write_text(json.dumps(fidelity, indent=2, ensure_ascii=True) + "\n")
    if args.manifest_json:
        write_manifest(args.manifest_json, inputs=[input_path], outputs=[output_path], parameters={"command": "convert"})
    print(f"Saved converted file: {output_path.resolve()}")


class tempfile_csv:
    def __init__(self, input_path, output_dir):
        self.input_path = Path(input_path)
        self.output_dir = Path(output_dir)
        self.path = self.output_dir / f"{self.input_path.stem}_temp_export.csv"

    def __enter__(self):
        exports, _ = export_numbers(self.input_path, self.output_dir)
        csv_candidates = [Path(item) for item in exports if item.endswith(".csv")]
        if not csv_candidates:
            raise SystemExit("Could not create a temporary CSV export from Numbers.")
        csv_candidates[0].replace(self.path)
        return self.path

    def __exit__(self, exc_type, exc_val, exc_tb):
        if self.path.exists():
            self.path.unlink()


def main():
    args = parse_args()
    collision_status = _preflight_output_safety(args)
    if collision_status is not None:
        raise SystemExit(collision_status)
    if args.command == "export":
        cmd_export(args)
    elif args.command == "set-cell":
        cmd_set_cell(args)
    elif args.command == "replace-table-csv":
        cmd_replace_table_csv(args)
    elif args.command == "convert":
        cmd_convert(args)


if __name__ == "__main__":
    main()
