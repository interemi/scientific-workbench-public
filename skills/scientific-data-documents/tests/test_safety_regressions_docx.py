"""Focused regression tests split from a larger safety module."""

from __future__ import annotations

from document_safety_support import *  # noqa: F403


class DocumentSafetyRegressionTest(unittest.TestCase):
    def test_docx_style_inventory_output_input_collisions_are_stdout_only(self) -> None:
        for flag in ("--output-csv", "--report-md", "--summary-json", "--manifest-json"):
            with self.subTest(flag=flag), tempfile.TemporaryDirectory() as raw:
                input_docx = Path(raw) / "invalid.docx"
                input_docx.write_bytes(b"not a docx")
                before = sha256(input_docx)

                completed = run_script("office_roundtrip.py", "docx-style-inventory", input_docx, flag, input_docx)

                self.assertEqual(completed.returncode, 2, completed.stderr or completed.stdout)
                self.assertEqual(completed.stderr, "")
                self.assertEqual(sha256(input_docx), before)
                payload = strict_json(completed.stdout)
                self.assertEqual(payload["app_status"], "BLOCKED_CONTROLADO")
                self.assertEqual(payload["results"]["collisions"][0]["flag"], flag)

    def test_docx_style_inventory_success_writes_manifest_without_touching_input(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            input_docx = root / "input.docx"
            document = Document()
            run = document.add_paragraph().add_run("marked text")
            run.bold = True
            document.save(input_docx)
            before = sha256(input_docx)
            output_csv = root / "run" / "styles.csv"
            report = root / "run" / "styles.md"
            summary = root / "run" / "summary.json"
            manifest = root / "run" / "manifest.json"

            completed = run_script(
                "office_roundtrip.py",
                "docx-style-inventory",
                input_docx,
                "--require-bold",
                "--output-csv",
                output_csv,
                "--report-md",
                report,
                "--summary-json",
                summary,
                "--manifest-json",
                manifest,
            )

            self.assertEqual(completed.returncode, 0, completed.stderr or completed.stdout)
            self.assertEqual(sha256(input_docx), before)
            self.assertTrue(output_csv.is_file())
            self.assertTrue(report.is_file())
            strict_json(summary.read_text(encoding="utf-8"))
            strict_json(manifest.read_text(encoding="utf-8"))

    def test_docx_styled_replace_auxiliary_collisions_preserve_input_and_output(self) -> None:
        for flag in ("--diff-json", "--summary-json", "--manifest-json"):
            for protected in ("input", "output"):
                with self.subTest(flag=flag, protected=protected), tempfile.TemporaryDirectory() as raw:
                    root = Path(raw)
                    input_docx = root / "input.docx"
                    output_docx = root / "output.docx"
                    input_docx.write_bytes(b"input bytes")
                    output_docx.write_bytes(b"output bytes")
                    before_input = sha256(input_docx)
                    before_output = sha256(output_docx)
                    colliding = input_docx if protected == "input" else output_docx

                    completed = run_script(
                        "office_roundtrip.py",
                        "docx-styled-replace",
                        input_docx,
                        output_docx,
                        "--find",
                        "old",
                        "--replace",
                        "new",
                        flag,
                        colliding,
                    )

                    self.assertEqual(completed.returncode, 2, completed.stderr or completed.stdout)
                    self.assertEqual(completed.stderr, "")
                    self.assertEqual(sha256(input_docx), before_input)
                    self.assertEqual(sha256(output_docx), before_output)
                    self.assertFalse((root / "output.docx.bak1").exists())
                    payload = strict_json(completed.stdout)
                    self.assertEqual(payload["app_status"], "BLOCKED_CONTROLADO")
                    collision = payload["results"]["collisions"][0]
                    self.assertEqual(collision["flag"], flag)
                    self.assertEqual(collision["protected_path"], protected)

    def test_docx_styled_replace_main_output_cannot_be_input(self) -> None:
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            input_docx = root / "input.docx"
            input_docx.write_bytes(b"input bytes")
            before = sha256(input_docx)
            summary = root / "summary.json"

            completed = run_script(
                "office_roundtrip.py",
                "docx-styled-replace",
                input_docx,
                input_docx,
                "--find",
                "old",
                "--replace",
                "new",
                "--summary-json",
                summary,
            )

            self.assertEqual(completed.returncode, 2, completed.stderr or completed.stdout)
            self.assertEqual(completed.stderr, "")
            self.assertEqual(sha256(input_docx), before)
            self.assertFalse(summary.exists())
            payload = strict_json(completed.stdout)
            self.assertEqual(payload["app_status"], "BLOCKED_CONTROLADO")
            self.assertEqual(payload["results"]["collisions"][0]["flag"], "output")

    def test_docx_styled_replace_success_keeps_input_and_emits_artifacts(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            input_docx = root / "input.docx"
            document = Document()
            run = document.add_paragraph().add_run("old value")
            run.bold = True
            document.save(input_docx)
            before = sha256(input_docx)
            output_docx = root / "run" / "output.docx"
            diff = root / "run" / "diff.json"
            summary = root / "run" / "summary.json"
            manifest = root / "run" / "manifest.json"

            completed = run_script(
                "office_roundtrip.py",
                "docx-styled-replace",
                input_docx,
                output_docx,
                "--find",
                "old",
                "--replace",
                "new",
                "--require-bold",
                "--diff-json",
                diff,
                "--summary-json",
                summary,
                "--manifest-json",
                manifest,
            )

            self.assertEqual(completed.returncode, 0, completed.stderr or completed.stdout)
            self.assertEqual(sha256(input_docx), before)
            self.assertTrue(output_docx.is_file())
            strict_json(diff.read_text(encoding="utf-8"))
            strict_json(summary.read_text(encoding="utf-8"))
            strict_json(manifest.read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
